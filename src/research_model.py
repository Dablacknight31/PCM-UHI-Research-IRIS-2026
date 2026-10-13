"""Delhi PCM rooftop and neighbourhood atmospheric research model.

Import does not acquire data or run a simulation. Fresh execution requires
validated inputs and explicit environment configuration. Public saved-result
inspection uses src/validation/verify_publication.py. See REPRODUCIBILITY.md.
"""

from pathlib import Path
from dataclasses import dataclass
from copy import deepcopy
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from types import SimpleNamespace
import os, sys, re, json, io, uuid, hashlib, math, shutil, warnings, inspect, time, platform, operator
import importlib.util, importlib.metadata
import numpy as np
import pandas as pd
import openpyxl
import requests
import rasterio
from affine import Affine
from pyproj import CRS, Transformer, Geod
from rasterio.features import geometry_mask
from rasterio.io import MemoryFile
from shapely.geometry import shape, mapping, box
from shapely.ops import unary_union, transform as geometry_transform
from shapely.ops import transform as transform_geometry
from scipy.optimize import brentq
from numba import njit
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from IPython.display import display
IN_COLAB = False
C3_IN_COLAB = False
C1_NUMBA_AVAILABLE = True
REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PCM_FULL_PACKAGE = REPOSITORY_ROOT
PCM_SELECTED_WORKBOOK = REPOSITORY_ROOT / 'data/inputs/materials/PCM_Details_As_Executed.xlsx'
ENVIRONMENT_VERSIONS = {'python': sys.version.split()[0], 'colab': False,
    **{name: importlib.metadata.version(name) for name in ('numpy', 'pandas', 'openpyxl', 'rasterio', 'pyproj', 'affine')}}


PACKAGE_ROOT = PCM_FULL_PACKAGE.resolve()

PCM_WORKBOOK_PATH = PCM_SELECTED_WORKBOOK

CACHE_FIRST = False

RUN_FRESH_UPSTREAM = True

USE_GOOGLE_DRIVE = False

C3_STORAGE_ROOT = Path(os.environ.get('PCM_UHI_OUTPUT_ROOT', str(REPOSITORY_ROOT / 'runtime')))

C3_UPSTREAM_ROOT = None

C3_CACHE_ZIP = os.environ.get('PCM_UHI_PRIVATE_CACHE_ARCHIVE')

C3_EXECUTION = {'mode': 'VALIDATE_ONLY', 'upstream_resume_dir': None, 'batch_size': 1, 'chunk_hours': 168, 'resume_verified_chunks': True, 'cfl': 0.45, 'dt_cap_s': 120.0, 'effect_threshold_K': 0.001, 'building_confidence_min': 0.75, 'building_page_size': 500, 'building_max_features': 100000, 'coverage_fractions': [0.1, 0.25, 0.5, 0.75, 1.0], 'arrangements': ['dispersed', 'clustered', 'roof_area_priority'], 'sensitivity_patterns': [(1.0, 'dispersed'), (0.5, 'dispersed'), (0.5, 'clustered')], 'domain_candidates_m': [3000, 6000, 10000, 20000, 30000], 'pilot_buffer_m': 6000, 'baseline_buffer_m': 6000, 'design_seed': 20251008, 'feedback_iterations': 2}

CATEGORY3_READY = False

CATEGORY3_HANDOFF = None

C3_ALLOWED_MODES = {'VALIDATE_ONLY', 'REAL_48H_PILOT', 'REAL_BASELINE_ANNUAL', 'REAL_FULL_EXPERIMENT'}

FRESH_PCM_WORKBOOK_PATH = None

FRESH_PCM_WORKBOOK_SHA256 = '0dda90b01987f99819c24ff1f7b18d51bfee4d3ea73358f73c7e616977e76aed'

CATEGORY4_READY = False

CATEGORY4_HANDOFF = None

CATEGORY4_PARTIAL_STATUS = 'NOT_RUN'

PCM_WORKBOOK_LOCAL_PATH = str(PCM_WORKBOOK_PATH)

PCM_WORKBOOK_PORTABLE_PATH = str(PCM_WORKBOOK_PATH)

PCM_WORKBOOK_EXPECTED_SHA256 = '0dda90b01987f99819c24ff1f7b18d51bfee4d3ea73358f73c7e616977e76aed'

PCM_WORKBOOK_REVIEW = {'status': 'USER_VERIFIED', 'workbook_sha256': '0dda90b01987f99819c24ff1f7b18d51bfee4d3ea73358f73c7e616977e76aed', 'basis': 'Author-reported property checking; independent physical verification not established', 'scope': 'PCM material values; no assertion of completed simulations or roof assembly safety'}

CONFIG = {'city': 'Delhi', 'country': 'India', 'study_year': 2025, 'timezone': 'Asia/Kolkata', 'study_start_local': '2025-01-01', 'study_end_local_exclusive': '2026-01-01', 'timestamp_label': 'start', 'hourly_interval_seconds': 3600, 'data_mode': 'REAL', 'project_root': os.environ.get('PCM_UHI_ROOT', '.'), 'input_dir': '.', 'output_dir': 'runtime', 'cache_dir': '.pcm_uhi_cache', 'pcm_workbook_path': os.environ.get('PCM_WORKBOOK', 'data/inputs/materials/PCM_Details_As_Executed.xlsx'), 'pcm_sheet_name': 'pcm_properties', 'pcm_expected_sha256': '0dda90b01987f99819c24ff1f7b18d51bfee4d3ea73358f73c7e616977e76aed', 'spatial': {'crs': 'EPSG:32643', 'resolution_m': (150.0, 150.0), 'anchor_xy_m': (0.0, 0.0), 'alignment_tolerance_m': 1e-06, 'bounds_xy_m': None}}

CONFIG['project_root'] = REPOSITORY_ROOT
CONFIG['output_dir'] = Path(os.environ.get('PCM_UHI_OUTPUT_ROOT', str(REPOSITORY_ROOT / 'runtime')))
CONFIG['pcm_workbook_path'] = str(PCM_SELECTED_WORKBOOK)
CONFIG = {**CONFIG, 'project_root': Path(CONFIG['project_root'])}

class InputValidationError(ValueError):
    """Required scientific input or metadata violate the declared contract."""

def resolve_path(value, base, *, must_exist=False, within_base=False, label='path'):
    """Resolve a declared path; optionally forbid escape from its base."""
    if value is None or not str(value).strip():
        raise InputValidationError(f'{label} must be supplied.')
    base = Path(base).expanduser().resolve()
    candidate = Path(value).expanduser()
    result = (candidate if candidate.is_absolute() else base / candidate).resolve()
    if within_base and (not result.is_relative_to(base)):
        raise InputValidationError(f'{label} must stay inside project_root: {result}')
    if must_exist and (not result.exists()):
        raise FileNotFoundError(f'Missing {label}: {result}')
    return result

def sha256_file(path):
    """Stream file contents so provenance hashes do not load a large file at once."""
    digest = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()

def require_file(path, *, label, expected_sha256=None):
    """Require a nonempty file and, when declared, the exact scientific artifact."""
    path = Path(path)
    if not path.is_file():
        remedy = "Provide the permission-cleared compiled PCM workbook at the configured input path, or set CONFIG['input_dir'] and CONFIG['pcm_workbook_path']/PCM_WORKBOOK to its location. No substitute or demo workbook will be generated." if 'pcm workbook' in label.casefold() else 'Provide this file at its declared input path.'
        raise FileNotFoundError(f'{label} is required, but was not found at: {path}. ' + remedy)
    if path.stat().st_size == 0:
        raise InputValidationError(f'{label} is empty: {path}')
    actual = sha256_file(path)
    if expected_sha256 is not None and actual != expected_sha256.lower():
        raise InputValidationError(f"{label} SHA-256 mismatch. Expected {expected_sha256}; received {actual}. Supply the configured current workbook, or explicitly review the replacement and update CONFIG['pcm_expected_sha256']. No legacy copy is selected automatically.")
    stat = path.stat()
    return {'file_path': str(path.resolve()), 'sha256': actual, 'file_size_bytes': stat.st_size, 'file_modified_at_UTC': datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()}

def validate_config(config):
    """Resolve paths and reject inconsistent calendar, grid, or real-data policy."""
    required = {'city', 'country', 'study_year', 'timezone', 'study_start_local', 'study_end_local_exclusive', 'timestamp_label', 'hourly_interval_seconds', 'data_mode', 'project_root', 'input_dir', 'output_dir', 'cache_dir', 'pcm_workbook_path', 'pcm_sheet_name', 'pcm_expected_sha256', 'spatial'}
    if not isinstance(config, dict) or set(config) != required:
        raise InputValidationError('Configuration must contain exactly the declared master keys; check missing keys or typos.')
    cfg = deepcopy(config)
    if cfg['data_mode'] != 'REAL':
        raise InputValidationError("This model requires data_mode='REAL'.")
    if any((not isinstance(cfg[key], str) or not cfg[key].strip() for key in ('city', 'country', 'pcm_sheet_name'))):
        raise InputValidationError('City, country and PCM sheet must be nonempty.')
    year = cfg['study_year']
    if isinstance(year, bool) or not isinstance(year, (int, np.integer)) or (not 1900 <= year <= 2200):
        raise InputValidationError('study_year must be an integer between 1900 and 2200.')
    ZoneInfo(cfg['timezone'])
    for key in ('study_start_local', 'study_end_local_exclusive'):
        stamp = pd.Timestamp(cfg[key])
        if pd.isna(stamp) or stamp.tzinfo is not None or stamp != stamp.normalize():
            raise InputValidationError(f'{key} must be a timezone-naive local calendar date.')
    expected = (pd.Timestamp(f'{year}-01-01'), pd.Timestamp(f'{year + 1}-01-01'))
    supplied = tuple((pd.Timestamp(cfg[key]) for key in ('study_start_local', 'study_end_local_exclusive')))
    if supplied != expected:
        raise InputValidationError('Study dates must span the complete selected local calendar year.')
    if cfg['timestamp_label'] not in {'start', 'midpoint', 'end'}:
        raise InputValidationError('timestamp_label must be start, midpoint or end.')
    if cfg['hourly_interval_seconds'] != 3600:
        raise InputValidationError('Section 0 defines exactly hourly (3600-second) intervals.')
    if not re.fullmatch('[0-9a-fA-F]{64}', str(cfg['pcm_expected_sha256'])):
        raise InputValidationError('Supply an explicit 64-character reviewed PCM SHA-256.')
    cfg['pcm_expected_sha256'] = cfg['pcm_expected_sha256'].lower()
    root = resolve_path(cfg['project_root'], Path.cwd(), must_exist=True, label='project_root')
    if not root.is_dir():
        raise InputValidationError('project_root must be a directory.')
    cfg['project_root'] = root
    cfg['input_dir'] = resolve_path(cfg['input_dir'], root, must_exist=True, label='input_dir')
    if not cfg['input_dir'].is_dir():
        raise InputValidationError('input_dir must be a directory.')
    for key in ('output_dir', 'cache_dir'):
        cfg[key] = resolve_path(cfg[key], root, within_base=True, label=key)
        if cfg[key].exists() and (not cfg[key].is_dir()):
            raise InputValidationError(f'{key} names a file, not a directory.')
        if cfg[key] in (root, cfg['input_dir']):
            raise InputValidationError(f'{key} cannot alias project_root or input_dir.')
    if cfg['output_dir'] == cfg['cache_dir']:
        raise InputValidationError('output_dir and cache_dir must be distinct.')
    cfg['pcm_workbook_path'] = resolve_path(cfg['pcm_workbook_path'], cfg['input_dir'], label='PCM workbook')
    spatial = cfg['spatial']
    if not isinstance(spatial, dict) or set(spatial) != {'crs', 'resolution_m', 'anchor_xy_m', 'alignment_tolerance_m', 'bounds_xy_m'}:
        raise InputValidationError('Spatial configuration must contain exactly the declared grid-policy keys.')
    crs = CRS.from_user_input(spatial['crs'])
    if not crs.is_projected or len(crs.axis_info) < 2 or any((not np.isclose(axis.unit_conversion_factor, 1.0) for axis in crs.axis_info[:2])):
        raise InputValidationError('The common spatial CRS must be projected in metres.')
    spatial['crs'] = crs.to_string()
    res = np.asarray(spatial['resolution_m'], dtype=float)
    anchor = np.asarray(spatial['anchor_xy_m'], dtype=float)
    if res.shape != (2,) or not np.isfinite(res).all() or (res <= 0).any():
        raise InputValidationError('resolution_m must contain two positive finite metre values.')
    if anchor.shape != (2,) or not np.isfinite(anchor).all():
        raise InputValidationError('anchor_xy_m must contain two finite coordinates.')
    tol = spatial['alignment_tolerance_m']
    if not np.isfinite(tol) or tol <= 0 or tol >= min(res) / 100:
        raise InputValidationError('alignment_tolerance_m must be positive and much smaller than a cell.')
    if spatial['bounds_xy_m'] is not None:
        bounds = np.asarray(spatial['bounds_xy_m'], dtype=float)
        if bounds.shape != (4,) or not np.isfinite(bounds).all() or (not (bounds[0] < bounds[2] and bounds[1] < bounds[3])):
            raise InputValidationError('bounds_xy_m must be (xmin, ymin, xmax, ymax).')
    return cfg

def prepare_output_directories(config):
    """Explicit later-stage operation; Section 0 itself does not create directories."""
    cfg = validate_config(config)
    for key in ('output_dir', 'cache_dir'):
        cfg[key].mkdir(parents=True, exist_ok=True)
    return {key: cfg[key] for key in ('output_dir', 'cache_dir')}

CONFIG = validate_config(CONFIG)

CUMULATIVE_SCREEN_METHOD = 'MELTING_INTERVAL_OVERLAP_INCLUSIVE_V1'

CUMULATIVE_SCREEN_EQUATION = 'Tm_high_C >= lower_C AND Tm_low_C <= upper_C; PCM_RUNNABLE only'

CUMULATIVE_RUN_ID = 'through_category4_' + datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ') + '_' + __import__('uuid').uuid4().hex[:8]

C1_RECORDED_SOURCE_SHA256 = '5a44293120e3b2ee58f12718fcc86772f91212c530ca6a8c39d4eeec3f0010bc'

CUMULATIVE_SOURCE_IDENTITY = {'source_content_sha256': '39f3a8d6707020cd2d67e7d97ba50fd3345e7e7f19ce208c523227e26ec91ca5', 'executable_source_sha256': C1_RECORDED_SOURCE_SHA256, 'fingerprint_policy': 'Recorded computational source fingerprint; exact source retained privately.'}

CATEGORY1_HANDOFF = None

CATEGORY1_READY = False

def is_unknown(value):
    """Recognize unknown scalar metadata without treating a sourced zero as absent."""
    if value is None:
        return True
    if isinstance(value, (dict, list, tuple, set)):
        return len(value) == 0
    if isinstance(value, str):
        return value.strip().upper() in {'', 'UNKNOWN', 'NOT CHECKED', 'TODO', 'TBD', 'NOT AVAILABLE', 'N/A', 'NA', 'NAN', 'NONE', 'NULL', 'MISSING', 'NOT REPORTED', 'UNSPECIFIED', '-'}
    if isinstance(value, (float, np.floating)):
        return not np.isfinite(value)
    return value is pd.NA or value is pd.NaT

def require_fields(mapping, names, *, context):
    """Stop on absent/unknown required metadata; preserve explicit numerical zero."""
    missing = [name for name in names if name not in mapping or is_unknown(mapping[name])]
    if missing:
        raise InputValidationError(f'{context}: required metadata missing or unknown: {missing}')

def require_columns(frame, names, *, context):
    """Reject duplicate headers, empty tables and missing required columns."""
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        raise InputValidationError(f'{context}: require a nonempty DataFrame.')
    if frame.columns.duplicated().any():
        raise InputValidationError(f'{context}: duplicate column headers.')
    missing = sorted(set(names) - set(frame.columns))
    if missing:
        raise InputValidationError(f'{context}: missing required columns: {missing}')

def require_finite_numeric(frame, names, *, context):
    """Reject nonnumeric, missing, infinite and boolean scientific values."""
    require_columns(frame, names, context=context)
    for name in names:
        if frame[name].map(lambda value: isinstance(value, (bool, np.bool_))).any():
            raise InputValidationError(f'{context}: {name} contains boolean physical values.')
        numbers = pd.to_numeric(frame[name], errors='coerce')
        bad = ~np.isfinite(numbers.to_numpy(dtype=float, na_value=np.nan))
        if bad.any():
            raise InputValidationError(f'{context}: {name} is missing/nonfinite/nonnumeric at row indices {frame.index[bad].tolist()[:10]}.')

class InputProvenanceRegistry:
    """In-memory dataset records; accessors return copies to avoid silent mutation."""
    QUALITY_STATUSES = {'UNKNOWN', 'NOT CHECKED', 'PARTIAL', 'PASS', 'FAIL'}
    AVAILABILITY_STATUSES = {'PENDING_FUTURE_STAGE', 'PRESENT', 'MISSING', 'INVALID'}
    NATURES = {'UNKNOWN', 'REAL', 'SYNTHETIC', 'DEMO'}

    def __init__(self):
        self._records = {}

    def register(self, dataset_id, *, source=None, planned_source=None, availability='PENDING_FUTURE_STAGE', data_nature='UNKNOWN', scope_kind='study', observed_city=None, expected_city=None, required_in_section0=False, data_start=None, data_end_exclusive=None, expected_start=None, expected_end_exclusive=None, acquired_at_UTC=None, units=None, expected_units=None, native_crs=None, native_resolution=None, target_crs=None, target_resolution=None, file_info=None, quality_status=None, source_quality_status=None, checks=None, notes='', replace=False):
        """Register supplied metadata; never infer quality from availability or file existence."""
        if not isinstance(dataset_id, str) or is_unknown(dataset_id):
            raise InputValidationError('dataset_id must be a nonempty identifier.')
        if dataset_id in self._records and (not replace):
            raise InputValidationError(f'Already registered: {dataset_id}; use explicit replace=True.')
        quality = 'UNKNOWN' if is_unknown(quality_status) and quality_status != 'NOT CHECKED' else quality_status
        source_quality = 'UNKNOWN' if is_unknown(source_quality_status) and source_quality_status != 'NOT CHECKED' else source_quality_status
        if quality not in self.QUALITY_STATUSES or source_quality not in self.QUALITY_STATUSES:
            raise InputValidationError('Quality must be UNKNOWN, NOT CHECKED, PARTIAL, PASS or FAIL.')
        if availability not in self.AVAILABILITY_STATUSES or data_nature not in self.NATURES:
            raise InputValidationError('Invalid availability or data_nature.')
        if scope_kind not in {'study', 'material', 'general'}:
            raise InputValidationError('scope_kind must be study, material or general.')
        checks = deepcopy(checks or [])
        for check in checks:
            require_fields(check, ('name', 'detail'), context=f'{dataset_id} quality check')
            if check.get('result') not in {'PASS', 'FAIL', 'PARTIAL', 'NOT CHECKED'}:
                raise InputValidationError(f'{dataset_id}: invalid check result.')
        if quality == 'PASS' and (not checks or any((c['result'] != 'PASS' for c in checks))):
            raise InputValidationError(f'{dataset_id}: PASS requires explicit all-passing checks.')
        file_info = deepcopy(file_info or {})
        allowed_file_keys = {'file_path', 'sha256', 'file_size_bytes', 'file_modified_at_UTC'}
        if set(file_info) - allowed_file_keys:
            raise InputValidationError('Unexpected file metadata keys.')
        record = dict(dataset_id=dataset_id, source=source or 'UNKNOWN', planned_source=planned_source or 'UNKNOWN', availability=availability, data_nature=data_nature, scope_kind=scope_kind, observed_city=observed_city or 'UNKNOWN', expected_city=expected_city or 'UNKNOWN', required_in_section0=bool(required_in_section0), data_start=data_start or 'UNKNOWN', data_end_exclusive=data_end_exclusive or 'UNKNOWN', expected_start=expected_start or 'UNKNOWN', expected_end_exclusive=expected_end_exclusive or 'UNKNOWN', acquired_at_UTC=acquired_at_UTC or 'UNKNOWN', units=deepcopy(units) if units is not None else 'UNKNOWN', expected_units=deepcopy(expected_units) if expected_units is not None else 'UNKNOWN', native_crs=native_crs or 'UNKNOWN', native_resolution=deepcopy(native_resolution) if native_resolution is not None else 'UNKNOWN', target_crs=target_crs or 'UNKNOWN', target_resolution=deepcopy(target_resolution) if target_resolution is not None else 'UNKNOWN', quality_status=quality, source_quality_status=source_quality, checks=checks, file_path=None, sha256=None, file_size_bytes=None, file_modified_at_UTC=None, registered_at_UTC=datetime.now(timezone.utc).isoformat(), notes=notes)
        record.update(file_info)
        self._records[dataset_id] = record
        return self.get(dataset_id)

    def get(self, dataset_id):
        if dataset_id not in self._records:
            raise InputValidationError(f'Unregistered required input: {dataset_id}')
        return deepcopy(self._records[dataset_id])

    def to_frame(self):
        return pd.DataFrame(deepcopy(list(self._records.values())))

def require_registered_input(dataset_id, registry, config):
    """Permit a later real-data section only after explicit source/quality validation."""
    row = registry.get(dataset_id)
    if row['availability'] != 'PRESENT':
        raise InputValidationError(f"{dataset_id}: input is {row['availability']}; provide and validate it first.")
    if config['data_mode'] != 'REAL' or row['data_nature'] != 'REAL':
        raise InputValidationError(f"{dataset_id}: synthetic/demo/unknown inputs cannot enter the real {config['city']} pipeline.")
    require_fields(row, ('source', 'units'), context=dataset_id)
    if row['quality_status'] != 'PASS' or row['source_quality_status'] != 'PASS':
        raise InputValidationError(f"{dataset_id}: structural quality={row['quality_status']}; source quality={row['source_quality_status']}. Explicit passing checks are required.")
    if not row['checks'] or any((check['result'] != 'PASS' for check in row['checks'])):
        raise InputValidationError(f'{dataset_id}: quality-check evidence is incomplete.')
    if row['scope_kind'] == 'study':
        require_fields(row, ('observed_city', 'data_start', 'data_end_exclusive'), context=dataset_id)
        if row['observed_city'].casefold() != config['city'].casefold():
            raise InputValidationError(f"{dataset_id}: source geography is not {config['city']}.")
    if row['file_path'] is not None:
        require_fields(row, ('sha256',), context=f'{dataset_id} file provenance')
        require_file(row['file_path'], label=dataset_id, expected_sha256=row['sha256'])
    return row

PROVENANCE_REGISTRY = InputProvenanceRegistry()

def to_utc_index(values, *, source_timezone=None):
    """Normalize actual instants to UTC; naive timestamps require an explicit timezone."""
    if isinstance(values, (str, pd.Timestamp, datetime)):
        raise InputValidationError('Supply a timestamp sequence, not a scalar.')
    try:
        values = list(values)
        if any((isinstance(value, (bool, int, float, np.number)) for value in values)):
            raise InputValidationError('Numeric/boolean timestamps are ambiguous; provide explicit dated timestamps.')
        stamps = [pd.Timestamp(value) for value in values]
    except (ValueError, TypeError) as error:
        raise InputValidationError(f'Unparseable timestamp: {error}') from error
    if not stamps or any((pd.isna(stamp) for stamp in stamps)):
        raise InputValidationError('Timestamps must be nonempty and nonmissing.')
    awareness = [stamp.tzinfo is not None for stamp in stamps]
    if any(awareness) and (not all(awareness)):
        raise InputValidationError('Mixed naive and timezone-aware timestamps are forbidden.')
    if all(awareness):
        return pd.DatetimeIndex([stamp.tz_convert('UTC') for stamp in stamps])
    if source_timezone is None:
        raise InputValidationError('Naive timestamps require explicit source_timezone; UTC is never assumed.')
    ZoneInfo(source_timezone)
    try:
        return pd.DatetimeIndex(stamps).tz_localize(source_timezone, ambiguous='raise', nonexistent='raise').tz_convert('UTC')
    except Exception as error:
        raise InputValidationError(f'Ambiguous/nonexistent source-local time: {error}') from error

def to_local_index(values, *, timezone_name, source_timezone=None):
    """Display the same UTC instants in a named local timezone without changing them."""
    ZoneInfo(timezone_name)
    return to_utc_index(values, source_timezone=source_timezone).tz_convert(timezone_name)

def expected_hourly_intervals(start_local, end_local_exclusive, *, timezone_name, timestamp_label='start'):
    """Create time coordinates only, with explicit boundaries and label convention."""
    if timestamp_label not in {'start', 'midpoint', 'end'}:
        raise InputValidationError('Use timestamp_label=start, midpoint or end.')
    edges = to_utc_index([start_local, end_local_exclusive], source_timezone=timezone_name)
    local_edges = edges.tz_convert(timezone_name)
    if any((stamp.minute or stamp.second or stamp.microsecond or stamp.nanosecond for stamp in local_edges)):
        raise InputValidationError('Study boundaries must be exact local hours.')
    hours = (edges[1] - edges[0]).total_seconds() / 3600
    if hours <= 0 or hours != int(hours):
        raise InputValidationError('The interval window must span a positive whole number of elapsed hours.')
    starts = pd.date_range(edges[0], edges[1], freq='h', inclusive='left')
    ends = starts + np.timedelta64(1, 'h')
    labels = starts + np.timedelta64({'start': 0, 'midpoint': 30, 'end': 60}[timestamp_label], 'm')
    result = pd.DataFrame({'interval_start_UTC': starts, 'interval_end_UTC': ends, 'interval_start_local': starts.tz_convert(timezone_name), 'interval_end_local': ends.tz_convert(timezone_name), 'timestamp_UTC': labels, 'timestamp_local': labels.tz_convert(timezone_name)})
    result.attrs.update(timezone=timezone_name, timestamp_label=timestamp_label, window_convention='[start, end)', kind='CALENDAR_EXPECTATION_ONLY')
    return result

def annual_hourly_intervals(year, *, timezone_name, timestamp_label='start'):
    """Define expected elapsed-hour coverage for a complete local calendar year, including leap days."""
    if isinstance(year, bool) or not isinstance(year, (int, np.integer)) or (not 1900 <= year <= 2200):
        raise InputValidationError('Annual coverage year must be an integer between 1900 and 2200.')
    return expected_hourly_intervals(f'{year}-01-01', f'{year + 1}-01-01', timezone_name=timezone_name, timestamp_label=timestamp_label)

def validate_hourly_coverage(timestamps, expected_intervals, *, timestamp_label, source_timezone=None):
    """Require exact expected labels, order and hourly coverage; never sort, fill or truncate."""
    require_columns(expected_intervals, ('timestamp_UTC',), context='expected calendar')
    if timestamp_label != expected_intervals.attrs.get('timestamp_label'):
        raise InputValidationError('Timestamp-label convention differs from the declared expected calendar.')
    actual = to_utc_index(timestamps, source_timezone=source_timezone)
    expected = to_utc_index(expected_intervals['timestamp_UTC'])
    if expected.has_duplicates or not expected.is_monotonic_increasing or (len(expected) > 1 and (not np.all((expected[1:] - expected[:-1]).total_seconds() == 3600))):
        raise InputValidationError('Expected calendar itself must contain unique consecutive elapsed-hour labels.')
    if actual.has_duplicates:
        raise InputValidationError('Duplicate hourly timestamps.')
    if not actual.is_monotonic_increasing:
        raise InputValidationError('Hourly timestamps are unsorted; order must be resolved explicitly.')
    missing, extra = (expected.difference(actual), actual.difference(expected))
    if len(missing) or len(extra):
        raise InputValidationError(f'Hourly coverage mismatch: expected={len(expected)}, actual={len(actual)}, missing={len(missing)}, extra={len(extra)}; first missing={list(missing[:3])}; first extra={list(extra[:3])}.')
    if not actual.equals(expected):
        raise InputValidationError('Hourly timestamps do not match the expected sequence.')
    return {'coverage_status': 'PASS', 'expected_hours': len(expected), 'observed_hours': len(actual), 'timestamp_label': timestamp_label, 'timezone': expected_intervals.attrs['timezone']}

ANNUAL_HOURLY_INTERVALS = annual_hourly_intervals(CONFIG['study_year'], timezone_name=CONFIG['timezone'], timestamp_label=CONFIG['timestamp_label'])

EXPECTED_ANNUAL_HOURS = len(ANNUAL_HOURLY_INTERVALS)

@dataclass(frozen=True)
class GridSpec:
    """Explicit grid metadata; a NumPy array alone has no CRS or spatial alignment."""
    shape: tuple
    crs: object
    transform: Affine

def validate_grid_spec(spec, config=None):
    """Check north-up metre-grid geometry; optionally enforce the master grid policy."""
    if not isinstance(spec, GridSpec):
        raise InputValidationError('Provide GridSpec(shape, crs, Affine transform); array shape alone is insufficient.')
    if not isinstance(spec.shape, (tuple, list)) or len(spec.shape) != 2 or any((isinstance(v, (bool, np.bool_)) or not isinstance(v, (int, np.integer)) or v < 1 for v in spec.shape)):
        raise InputValidationError('Raster shape must be positive integer (rows, columns).')
    if spec.crs is None:
        raise InputValidationError('Raster CRS is missing.')
    crs = CRS.from_user_input(spec.crs)
    if not crs.is_projected or len(crs.axis_info) < 2 or any((not np.isclose(axis.unit_conversion_factor, 1.0) for axis in crs.axis_info[:2])):
        raise InputValidationError('Common grids must use a projected CRS in metres.')
    if not isinstance(spec.transform, Affine):
        raise InputValidationError('Raster transform must be an Affine, never inferred from image shape.')
    transform = spec.transform
    if not np.isfinite(tuple(transform)[:6]).all():
        raise InputValidationError('Raster transform contains nonfinite coefficients.')
    if abs(transform.b) > 1e-12 or abs(transform.d) > 1e-12 or transform.a <= 0 or (transform.e >= 0):
        raise InputValidationError('Require an unrotated north-up grid: a>0, e<0, b=d=0.')
    resolution = (transform.a, -transform.e)
    if config is not None:
        policy = config['spatial']
        tol = policy['alignment_tolerance_m']
        if not crs.equals(CRS.from_user_input(policy['crs'])):
            raise InputValidationError('Raster CRS differs from the configured common CRS.')
        if not np.allclose(resolution, policy['resolution_m'], rtol=0, atol=tol):
            raise InputValidationError(f"Raster resolution {resolution} differs from common {policy['resolution_m']} metres.")
        for origin, anchor, step in zip((transform.c, transform.f), policy['anchor_xy_m'], resolution):
            nearest = anchor + round((origin - anchor) / step) * step
            if abs(origin - nearest) > tol:
                raise InputValidationError('Raster origin is not aligned to the configured grid anchor.')
    return {'grid_status': 'PASS', 'shape': tuple(spec.shape), 'crs': crs.to_string(), 'resolution_m': resolution, 'transform': tuple(transform)[:6]}

def define_common_grid(bounds_xy_m, config):
    """Snap supplied projected boundary bounds outward to the common anchor; no boundary is fabricated."""
    if bounds_xy_m is None:
        raise InputValidationError('A sourced study boundary is required before defining the common grid extent.')
    bounds = np.asarray(bounds_xy_m, dtype=float)
    if bounds.shape != (4,) or not np.isfinite(bounds).all() or (not (bounds[0] < bounds[2] and bounds[1] < bounds[3])):
        raise InputValidationError('Require finite (xmin, ymin, xmax, ymax) projected bounds.')
    dx, dy = config['spatial']['resolution_m']
    ax, ay = config['spatial']['anchor_xy_m']
    left = ax + np.floor((bounds[0] - ax) / dx) * dx
    right = ax + np.ceil((bounds[2] - ax) / dx) * dx
    bottom = ay + np.floor((bounds[1] - ay) / dy) * dy
    top = ay + np.ceil((bounds[3] - ay) / dy) * dy
    grid = GridSpec((int(round((top - bottom) / dy)), int(round((right - left) / dx))), config['spatial']['crs'], Affine(dx, 0, left, 0, -dy, top))
    validate_grid_spec(grid, config)
    return grid

def raster_grid_from_file(path):
    """Read actual raster shape/CRS/transform only; do not retrieve or alter raster data."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f'Required raster file is missing: {path}')
    with rasterio.open(path) as dataset:
        if dataset.crs is None:
            raise InputValidationError(f'Raster has no CRS: {path}')
        return GridSpec((dataset.height, dataset.width), dataset.crs.to_wkt(), dataset.transform)

def validate_grid_alignment(spec, reference, *, config=None, tolerance_m=1e-06):
    """Require full coincident grids, including origins; there is no automatic warp or resize."""
    validate_grid_spec(spec, config)
    validate_grid_spec(reference, config)
    if config is not None:
        tolerance_m = config['spatial']['alignment_tolerance_m']
    if not np.isfinite(tolerance_m) or tolerance_m <= 0 or tolerance_m >= min(spec.transform.a, -spec.transform.e, reference.transform.a, -reference.transform.e) / 100:
        raise InputValidationError('Alignment tolerance must be positive and much smaller than a cell.')
    if tuple(spec.shape) != tuple(reference.shape):
        raise InputValidationError(f'Raster shape mismatch: {spec.shape} vs {reference.shape}.')
    if not CRS.from_user_input(spec.crs).equals(CRS.from_user_input(reference.crs)):
        raise InputValidationError('Raster CRS mismatch.')
    if not np.allclose(tuple(spec.transform)[:6], tuple(reference.transform)[:6], rtol=0, atol=tolerance_m):
        raise InputValidationError('Raster transform/resolution/origin mismatch: the grids are not aligned.')
    return {'alignment_status': 'PASS', 'shape': spec.shape, 'tolerance_m': tolerance_m}

def validate_raster_grid(array, spec, reference, *, config=None):
    """Match a 2D or band-first 3D array to its declared metadata and the common grid."""
    array = np.asanyarray(array)
    if array.ndim not in (2, 3) or (array.ndim == 3 and array.shape[0] < 1):
        raise InputValidationError('Raster must be 2D or a nonempty band-first 3D array.')
    validate_grid_spec(spec, config)
    if tuple(array.shape[-2:]) != tuple(spec.shape):
        raise InputValidationError('Raster array shape differs from its declared grid metadata.')
    return validate_grid_alignment(spec, reference, config=config)

COMMON_GRID = None

PCM_COLUMNS = ['pcm_id', 'material_name', 'composition', 'Tm_low_C', 'Tm_high_C', 'Tf_low_C', 'Tf_high_C', 'L_Jkg', 'k_solid_WmK', 'k_liquid_WmK', 'cp_solid_JkgK', 'cp_liquid_JkgK', 'rho_ref_kgm3', 'supercooling_C', 'phase_model', 'evidence_tier', 'source', 'source_locator', 'notes']

PCM_NUMERIC_COLUMNS = ['Tm_low_C', 'Tm_high_C', 'Tf_low_C', 'Tf_high_C', 'L_Jkg', 'k_solid_WmK', 'k_liquid_WmK', 'cp_solid_JkgK', 'cp_liquid_JkgK', 'rho_ref_kgm3', 'supercooling_C']

PCM_REQUIRED_NUMBERS = ['Tm_low_C', 'Tm_high_C', 'L_Jkg', 'k_solid_WmK', 'k_liquid_WmK', 'cp_solid_JkgK', 'cp_liquid_JkgK', 'rho_ref_kgm3']

def clean_pcm_properties(frame):
    """Keep all rows. Missing/ambiguous numbers are NaN and receive audit reasons."""
    raw = frame.copy().reset_index(drop=True)
    if raw.columns.duplicated().any():
        raise ValueError('PCM column headers must be unique.')
    out = pd.DataFrame(index=raw.index)
    for name in PCM_COLUMNS:
        out[name] = raw[name] if name in raw else np.nan
    problems = [[] for _ in range(len(out))]
    for name in PCM_NUMERIC_COLUMNS:
        original = out[name]
        numbers = pd.to_numeric(original, errors='coerce')
        bad = original.notna() & numbers.isna()
        for i in out.index[bad]:
            problems[i].append(f'{name}: nonnumeric text')
        for i in out.index[numbers.notna() & ~np.isfinite(numbers)]:
            problems[i].append(f'{name}: nonfinite value')
        out[name] = numbers
    for name in ('pcm_id', 'material_name', 'phase_model'):
        out[name] = out[name].map(lambda v: str(v).strip() if pd.notna(v) else v)
    duplicate = out['pcm_id'].notna() & out['pcm_id'].duplicated(keep=False)
    statuses = []
    for i, row in out.iterrows():
        missing = [c for c in PCM_REQUIRED_NUMBERS if pd.isna(row[c])]
        for name in ('pcm_id', 'material_name', 'phase_model'):
            if pd.isna(row[name]) or row[name] == '':
                missing.append(name)
        if missing:
            problems[i].append('missing: ' + ', '.join(missing))
        if duplicate[i]:
            problems[i].append('duplicate pcm_id')
        if pd.notna(row.Tm_low_C) and pd.notna(row.Tm_high_C):
            if not -273.15 < row.Tm_low_C <= row.Tm_high_C:
                problems[i].append('require -273.15 < Tm_low_C <= Tm_high_C')
        for name in PCM_REQUIRED_NUMBERS[2:]:
            if pd.notna(row[name]) and (not np.isfinite(row[name]) or row[name] <= 0):
                problems[i].append(f'{name}: require a finite positive value')
        if pd.notna(row.supercooling_C) and row.supercooling_C < 0:
            problems[i].append('supercooling_C must be nonnegative')
        if pd.notna(row.Tf_low_C) and pd.notna(row.Tf_high_C) and (row.Tf_low_C > row.Tf_high_C):
            problems[i].append('Tf_low_C exceeds Tf_high_C')
        phase = row.phase_model
        if phase == 'excluded':
            status = 'EXCLUDED'
        elif problems[i]:
            status = 'NEEDS_SOURCE_DATA'
        elif phase not in ('equilibrium', 'equilibrium_screen'):
            status = 'UNSUPPORTED_PHASE_MODEL'
            problems[i].append('history_aware/measured curves need an agreed constitutive model')
        else:
            status = 'READY_EQUILIBRIUM_SCREEN'
        statuses.append(status)
    out['excel_row'] = out.index + 2
    out['model_readiness'] = statuses
    out['input_issues'] = ['; '.join(p) for p in problems]
    return pd.concat([out, raw.add_prefix('raw_')], axis=1)

def load_pcm_properties(path=None, sheet_name='pcm_properties', source_frame=None):
    """Load the canonical schema from XLSX/CSV, or from an existing DataFrame."""
    if source_frame is not None:
        return clean_pcm_properties(source_frame)
    if path is None:
        raise ValueError('Provide a workbook/CSV path or source_frame.')
    path = Path(path)
    frame = pd.read_csv(path) if path.suffix.lower() == '.csv' else pd.read_excel(path, sheet_name=sheet_name)
    return clean_pcm_properties(frame)

PCM_UNITS = {'Tm_low_C': 'degC', 'Tm_high_C': 'degC', 'Tf_low_C': 'degC', 'Tf_high_C': 'degC', 'L_Jkg': 'J/kg', 'k_solid_WmK': 'W/(m K)', 'k_liquid_WmK': 'W/(m K)', 'cp_solid_JkgK': 'J/(kg K)', 'cp_liquid_JkgK': 'J/(kg K)', 'rho_ref_kgm3': 'kg/m3', 'supercooling_C': 'K temperature difference'}

def inspect_pcm_workbook(path, sheet_name):
    """Validate literal canonical headers before pandas can mangle duplicates; reject formula caches."""
    if Path(path).suffix.lower() != '.xlsx':
        raise InputValidationError('Section 0 requires the reviewed code-ready .xlsx workbook.')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise InputValidationError(f'PCM workbook requires sheet {sheet_name!r}; found {workbook.sheetnames}.')
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows()
        first_row = next(rows, None)
        if first_row is None:
            raise InputValidationError('PCM worksheet is empty.')
        headers = [cell.value for cell in first_row]
        if any((not isinstance(header, str) or not header.strip() for header in headers)):
            raise InputValidationError('PCM headers must be nonblank literal strings in row 1.')
        if len(set(headers)) != len(headers):
            raise InputValidationError('PCM workbook contains duplicate headers in row 1.')
        missing = sorted(set(PCM_COLUMNS) - set(headers))
        if missing:
            raise InputValidationError(f'PCM workbook lacks canonical headers: {missing}. Keep unknown values blank in cells.')
        row_count = 0
        for row in rows:
            row_count += 1
            for cell in row:
                if cell.data_type == 'f':
                    raise InputValidationError(f'PCM formula at {sheet_name}!{cell.coordinate}: provide reviewed literal values; cached formulas may be missing or stale.')
        if row_count == 0:
            raise InputValidationError('PCM worksheet has no material rows.')
        return {'sheet_name': sheet_name, 'row_count': row_count, 'headers': headers, 'workbook_sheets': workbook.sheetnames}
    finally:
        workbook.close()

def audit_cleaned_pcm(cleaned):
    """Add stricter Section-0 checks to the inherited cleaner, with no imputation."""
    require_columns(cleaned, PCM_COLUMNS + ['model_readiness', 'input_issues'], context='cleaned PCM table')
    out = cleaned.copy()
    out['legacy_model_readiness'] = out['model_readiness']
    additional = [[] for _ in range(len(out))]
    for column in PCM_NUMERIC_COLUMNS:
        raw_column = 'raw_' + column
        boolean = out[raw_column].map(lambda v: isinstance(v, (bool, np.bool_))) if raw_column in out else pd.Series(False, index=out.index)
        finite = np.isfinite(out[column].to_numpy(dtype=float, na_value=np.nan))
        nonfinite = out[column].notna() & ~finite
        for i in out.index[boolean]:
            additional[i].append(f'{column}: boolean is not a physical value')
        for i in out.index[nonfinite]:
            additional[i].append(f'{column}: nonfinite normalized to missing; raw retained')
        if (boolean | nonfinite).any():
            out[column] = out[column].astype(float)
            out.loc[boolean | nonfinite, column] = np.nan
    for column in ('Tf_low_C', 'Tf_high_C'):
        for i in out.index[out[column].notna() & (out[column] <= -273.15)]:
            additional[i].append(f'{column}: require temperature above absolute zero')
    evidence_issues = []
    statuses = []
    for i, row in out.iterrows():
        evidence_missing = [column for column in ('evidence_tier', 'source', 'source_locator') if is_unknown(row[column])]
        evidence_issues.append('missing evidence metadata: ' + ', '.join(evidence_missing) if evidence_missing else '')
        legacy = row['legacy_model_readiness']
        if legacy == 'EXCLUDED':
            status = 'EXCLUDED'
        elif legacy == 'READY_EQUILIBRIUM_SCREEN' and (not additional[i]):
            status = 'RUNNABLE_SCHEMA_ONLY'
        else:
            status = 'NEEDS_DATA_OR_PHASE_MODEL'
        statuses.append(status)
    out['section0_readiness'] = statuses
    out.loc[(out['legacy_model_readiness'] == 'READY_EQUILIBRIUM_SCREEN') & (out['section0_readiness'] != 'RUNNABLE_SCHEMA_ONLY'), 'model_readiness'] = 'NEEDS_SOURCE_DATA'
    out['section0_additional_issues'] = ['; '.join(items) for items in additional]
    out['section0_issues'] = ['; '.join((s for s in (legacy, extra) if s)) for legacy, extra in zip(out['input_issues'], out['section0_additional_issues'])]
    out['evidence_metadata_issues'] = evidence_issues
    out['source_quality_status'] = 'NOT CHECKED'
    out.attrs.update(units=PCM_UNITS.copy(), quality_scope='SCHEMA_ONLY; PHYSICAL SOURCES NOT CHECKED')
    return out

def partition_pcm_tables(all_rows):
    """Return exhaustive, disjoint copies without ranking or discarding any material."""
    labels = {'runnable': 'RUNNABLE_SCHEMA_ONLY', 'needs_data': 'NEEDS_DATA_OR_PHASE_MODEL', 'excluded': 'EXCLUDED'}
    tables = {name: all_rows.loc[all_rows.section0_readiness == label].copy() for name, label in labels.items()}
    if sum((len(table) for table in tables.values())) != len(all_rows):
        raise InputValidationError('PCM readiness partition is incomplete.')
    return tables

def load_section0_pcm(config, registry):
    """Load only the pinned supplied workbook, then register its actual file and audit facts."""
    file_info = require_file(config['pcm_workbook_path'], label='Current PCM workbook', expected_sha256=config['pcm_expected_sha256'])
    workbook_info = inspect_pcm_workbook(config['pcm_workbook_path'], config['pcm_sheet_name'])
    all_rows = audit_cleaned_pcm(load_pcm_properties(config['pcm_workbook_path'], sheet_name=config['pcm_sheet_name']))
    if len(all_rows) != workbook_info['row_count']:
        raise InputValidationError('PCM loader row count differs from the literal workbook; trailing/blank rows need review.')
    if sha256_file(config['pcm_workbook_path']) != file_info['sha256']:
        raise InputValidationError('PCM workbook changed during loading; rerun from a stable file.')
    parts = partition_pcm_tables(all_rows)
    quality = 'PASS' if len(parts['runnable']) == len(all_rows) else 'PARTIAL'
    checks = [{'name': 'reviewed_file_identity', 'result': 'PASS', 'detail': 'SHA-256 matches configured current workbook.'}, {'name': 'canonical_schema_and_literal_values', 'result': 'PASS', 'detail': 'Unique literal headers; formula caches refused; all rows preserved.'}, {'name': 'material_input_completeness', 'result': quality, 'detail': f"{len(parts['runnable'])} schema-runnable; {len(parts['needs_data'])} need data/phase model; {len(parts['excluded'])} excluded."}]
    registry.register('pcm_workbook', source='User-supplied current code-ready PCM workbook; per-row citations retained', planned_source='Compiled PCM property workbook / pcm_properties', availability='PRESENT', data_nature='REAL', scope_kind='material', required_in_section0=True, units=PCM_UNITS, expected_units=PCM_UNITS, native_crs='NOT APPLICABLE', native_resolution='NOT APPLICABLE', file_info=file_info, quality_status=quality, source_quality_status='NOT CHECKED', checks=checks, notes='Only schema and artifact identity were checked. Citation validity, measurement conditions, units in original sources and physical suitability are not verified.')
    return (all_rows, parts, workbook_info)

FUTURE_INPUT_REQUIREMENTS = (('study_boundary', 'Reviewed Delhi administrative/study-domain vector boundary', 'Geometry with declared native CRS and source version', 'Required before the common grid extent is defined.'), ('era5_hourly', 'ECMWF/ERA5_LAND/HOURLY or a reviewed equivalent supplied file', {'T_air': 'degC after explicit conversion', 'u10': 'm/s', 'v10': 'm/s', 'solar_down': 'W/m2 interval mean', 'longwave_down': 'W/m2 interval mean'}, 'Full local-year coverage and source-native time/accumulation convention must be verified; not retrieved.'), ('landsat_observations', 'Landsat 8/9 Collection 2 Level 2; scenes and QA to be reviewed', 'LST degC, reflectance dimensionless after sourced scaling; native QA flags retained', 'Scene dates, masks, compositing method and native resolution are unselected; no composites created.'), ('ghsl_built_surface', 'JRC/GHSL/P2023A/GHS_BUILT_S or reviewed release/epoch', 'Native built surface m2; any derived fraction requires correct source-cell area', 'Release, observation epoch and native cell area require review; not a measured roof survey.'), ('roof_area_and_eligibility', 'Measured/sourced roof plan area and PCM installation eligibility', 'm2 and/or dimensionless fractions with documented denominator', 'Real georeferenced layers needed; GHSL must not silently substitute for surveyed roof eligibility.'), ('canopy_wind', 'Measured or independently justified canopy-height wind with height and processing metadata', {'Ux': 'm/s', 'Uy': 'm/s', 'height': 'm'}, 'Coverage and wind basis unverified.'), ('building_geometry', 'Sourced building/canopy geometry for the selected domain', 'm and georeferenced geometry', 'No universal building height is inserted.'), ('air_exchange_parameter_evidence', 'Literature/measurement/calibration records for future air-transport parameters', 'Parameter-specific SI units', 'Values, uncertainty and representativeness remain unset.'), ('roof_material_and_boundary_evidence', 'Sourced reference roof/assembly material properties and boundary conditions', 'Material-specific SI units', 'No thermal stack or bottom-boundary model is assumed here.'), ('cool_roof_optical_evidence', 'Sourced initial and aged coating optical measurements', 'Dimensionless reflectance/emissivity and dated state metadata', 'No optical or ageing values are inserted.'), ('pcm_retention_and_lifetime_evidence', 'Material/system-specific cycle retention, durability and service-life evidence', 'Fractions, cycles and years as explicitly reported', 'No degradation curve or lifetime is invented.'), ('assembly_mass_and_structural_evidence', 'Sourced installation mass and structural allowance records', 'kg/m2 with assembly scope', 'No structural feasibility decision is made.'), ('safety_and_compatibility_evidence', 'Source records for fire, toxicity, corrosion and leakage behaviour', 'Evidence/test classifications with conditions', 'No generic safety pass is assigned.'), ('cost_and_maintenance_evidence', 'Dated local quotations/records for supply, installation, maintenance, replacement and disposal', 'Currency with price year and denominator', 'No costs, totals or economics are calculated.'), ('economic_method_inputs', 'Explicit sourced financial assumptions and study horizon for any future economic section', 'Years, fractions and currency conventions', 'No discount rate, cost year or lifetime is assumed.'), ('independent_validation_observations', 'Field/CFD/benchmark observations with methods and uncertainty', 'Measurement-specific units', 'No empirical or CFD validation has been performed.'))

def register_future_inputs(registry, config):
    """Declare future dependencies as pending; planned metadata never become observed facts."""
    for dataset_id, planned_source, expected_units, notes in FUTURE_INPUT_REQUIREMENTS:
        spatial = dataset_id in {'study_boundary', 'landsat_observations', 'ghsl_built_surface', 'roof_area_and_eligibility', 'building_geometry'}
        registry.register(dataset_id, planned_source=planned_source, expected_units=expected_units, expected_city=config['city'], required_in_section0=False, expected_start=config['study_start_local'] if dataset_id == 'era5_hourly' else None, expected_end_exclusive=config['study_end_local_exclusive'] if dataset_id == 'era5_hourly' else None, target_crs=config['spatial']['crs'] if spatial else None, target_resolution=config['spatial']['resolution_m'] if spatial else None, availability='PENDING_FUTURE_STAGE', data_nature='UNKNOWN', quality_status='NOT CHECKED', source_quality_status='UNKNOWN', notes=notes)

def make_section0_readiness(registry, pcm_summary):
    """Show current material audit and every future input without requiring future files."""
    records = registry.to_frame()
    counts = dict(zip(pcm_summary['table'], pcm_summary['rows']))
    rows = []
    for record in records.to_dict('records'):
        current = record['required_in_section0']
        rows.append({'input': record['dataset_id'], 'required_now': 'YES' if current else 'NO', 'availability': record['availability'], 'quality_status': record['quality_status'], 'source_quality': record['source_quality_status'], 'section0_status': 'LOADED / AUDITED' if current and record['availability'] == 'PRESENT' else 'MISSING REQUIRED' if current else 'PENDING / NOT REQUIRED YET', 'detail': f"All {counts['PCM_ALL']}; runnable {counts['PCM_RUNNABLE']}; needs data {counts['PCM_NEEDS_DATA']}; excluded {counts['PCM_EXCLUDED']}. Schema only." if current else record['notes']})
    return pd.DataFrame(rows)

PATH_UTILITIES = {fn.__name__: fn for fn in (resolve_path, sha256_file, require_file, prepare_output_directories)}

TIME_UTILITIES = {fn.__name__: fn for fn in (to_utc_index, to_local_index, expected_hourly_intervals, annual_hourly_intervals, validate_hourly_coverage)}

SPATIAL_UTILITIES = {fn.__name__: fn for fn in (GridSpec, validate_grid_spec, define_common_grid, raster_grid_from_file, validate_grid_alignment, validate_raster_grid)}

VALIDATION_UTILITIES = {fn.__name__: fn for fn in (InputValidationError, is_unknown, require_fields, require_columns, require_finite_numeric, validate_config, require_registered_input)}

PCM_UTILITIES = {fn.__name__: fn for fn in (clean_pcm_properties, load_pcm_properties, inspect_pcm_workbook, audit_cleaned_pcm, partition_pcm_tables, load_section0_pcm)}

PRELIM_HANDOFF = None

PRELIM_READY = False

PRELIM_CONFIG = {'ee_project': os.environ.get('PCM_UHI_EE_PROJECT', ''), 'ee_auth_mode_colab': 'colab', 'boundary': {'collection': 'FAO/GAUL/2015/level1', 'version': '2015', 'filters': {'ADM0_NAME': CONFIG['country'], 'ADM1_NAME': CONFIG['city']}, 'required_feature_count': 1, 'geographic_crs': 'EPSG:4326', 'raster_membership': 'cell center; all_touched=False'}, 'landsat': {'collections': ['LANDSAT/LC08/C02/T1_L2', 'LANDSAT/LC09/C02/T1_L2'], 'windows': {'PREMONSOON': [f"{CONFIG['study_year']}-04-01", f"{CONFIG['study_year']}-06-01"], 'POSTMONSOON': [f"{CONFIG['study_year']}-10-01", f"{CONFIG['study_year']}-12-01"]}, 'date_timezone': 'UTC', 'processing_level': 'L2SP', 'cloud_cover_min_inclusive': 0, 'cloud_cover_max_exclusive': 30, 'qa_pixel_reject_bits': [0, 1, 2, 3, 4, 5], 'qa_radsat_require_zero': True, 'reflectance_scale': 2.75e-05, 'reflectance_offset': -0.2, 'st_scale_K_per_DN': 0.00341802, 'st_offset_K': 149.0, 'K_to_C': 273.15, 'index_method': 'explicit ratio on nonnegative reflectance; positive denominator', 'NDVI_bands': ['SR_B5', 'SR_B4'], 'NDBI_bands': ['SR_B6', 'SR_B5'], 'scene_to_grid': 'reduceResolution(mean); native support weighted by output-cell overlap', 'temporal_composite': 'median AFTER scene aggregation to common grid', 'minimum_scene_count': 1, 'minimum_cell_observations': 1, 'reduce_resolution_max_pixels': 4096, 'lst_validation_range_C': [-80.0, 90.0]}, 'ghsl': {'collection': 'JRC/GHSL/P2023A/GHS_BUILT_S', 'version': 'P2023A (R2023A)', 'epoch': 2020, 'band': 'built_surface', 'source_unit': 'm2 per 100 m cell', 'catalog_cell_area_m2': 10000.0, 'spatial_method': 'native built_surface/native projected cell area; overlap-weighted mean to common grid'}, 'hotspot': {'sigma_multiplier': 1.0, 'std_ddof': 0, 'built_quantile': 0.75, 'quantile_method': 'linear', 'connectivity': 4, 'minimum_component_cells': 5, 'minimum_joint_valid_cells': 100, 'threshold_population': 'LST: valid in both seasons; built: all valid boundary cells', 'selection_order': ['built_weighted_area_m2 descending', 'mean_pre_excess_C descending', 'n_cells descending', 'component_id ascending']}, 'weather': {'collection': 'ECMWF/ERA5_LAND/HOURLY', 'bands': ['temperature_2m', 'u_component_of_wind_10m', 'v_component_of_wind_10m', 'surface_solar_radiation_downwards_hourly', 'surface_thermal_radiation_downwards_hourly'], 'spatial_scope': 'SELECTED_HOTSPOT_NATIVE_ERA5_CELL', 'location_method': 'selected cell center nearest built-fraction-weighted hotspot centroid; row-major tie', 'extraction_method': 'Image.sample at hotspot point; exact native projection; no scale/factor/numPixels; dropNulls=False', 'catalog_nominal_resolution_m': 11132, 'max_pixels': 1000000, 'tile_scale': 2, 'source_radiation_interval': '[system:time_start - 1 hour, system:time_start)', 'source_radiation_units': 'J/m2', 'output_radiation_units': 'W/m2 interval mean', 'source_temperature_wind': 'instantaneous at system:time_start', 'target_temperature_wind': 'sample at interval midpoint; bracketed linear interpolation only', 'source_padding_hours': 1, 'shortwave_physical_qc': {'maximum_correctable_negative_magnitude_Jm2': 5.0, 'physical_minimum_Jm2': 0.0, 'operation': 'negative source increments in [-5,0) J/m2 set to 0 for model input only', 'larger_negative_policy': 'hard failure requiring provider source review', 'raw_policy': 'preserve every acquired value and acquisition artifact exactly', 'basis': 'explicit user-approved provider shortwave quantization/deaccumulation physical QC'}, 'air_validation_range_C': [-80.0, 60.0]}, 'pcm_screen': {'lst_percentile': 95.0, 'quantile_method': 'linear', 'round_decimals': 1, 'air_timezone': 'Asia/Kolkata', 'required_hours_per_local_day': 24, 'lower_statistic': 'mean of complete local daily Tmax from hourly hotspot T_air_C', 'air_missing_policy': 'fail on missing/duplicate/extra hours or days; no interpolation/substitution', 'rule': 'MELTING_INTERVAL_OVERLAP_INCLUSIVE_V1', 'pass_equation': 'Tm_high_C >= lower_C AND Tm_low_C <= upper_C; PCM_RUNNABLE only', 'boundary_contact_passes': True, 'table': 'Section-0 PCM_RUNNABLE only; all workbook rows retained for audit', 'review_status': 'REVIEW_REQUIRED_BEFORE_CATEGORY1'}, 'download': {'format': 'GEO_TIFF', 'tile_cells': 256, 'timeout_seconds': 180, 'output_dtype': 'float64', 'nodata_encoding': -9999.0, 'cache_policy': 'write acquired artifacts only; no automatic cache reads'}, 'figures': {'dpi': 180, 'map_size': [12, 6], 'relationship_size': [12, 5], 'hexbin_gridsize': 45}}

LANDSAT_PHYSICAL_CONTRACT = {'reflectance_scale': 2.75e-05, 'reflectance_offset': -0.2, 'st_scale_K_per_DN': 0.00341802, 'st_offset_K': 149.0, 'K_to_C': 273.15}

PRELIM_RUN_ID = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ') + '_' + uuid.uuid4().hex[:8]

PRELIM_RUN_DIR = CONFIG['output_dir'] / 'preliminary' / str(CONFIG['study_year']) / PRELIM_RUN_ID

PRELIM_SOURCE_METADATA = {}

PRELIM_ARTIFACTS = {}

PRELIM_RASTER_RECORDS = {}

PRELIM_CHECKS = []

PRELIM_EE_SESSION = None

def require_prelim(condition, message):
    if not bool(condition):
        raise InputValidationError(message)

def json_ready(value):
    if isinstance(value, (Path, pd.Timestamp, datetime)):
        return str(value)
    if isinstance(value, np.generic):
        return json_ready(value.item())
    if value is pd.NA or value is pd.NaT or (isinstance(value, float) and (not np.isfinite(value))):
        return None
    if isinstance(value, dict):
        return {str(k): json_ready(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_ready(v) for v in value]
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise TypeError(f'Unsupported metadata type: {type(value).__name__}')

def grid_metadata(grid):
    validate_grid_spec(grid, CONFIG)
    return {'crs': CRS.from_user_input(grid.crs).to_string(), 'shape': list(grid.shape), 'transform': list(tuple(grid.transform)[:6]), 'bounds_xy_m': list(rasterio.transform.array_bounds(*grid.shape, grid.transform)), 'resolution_m': [grid.transform.a, -grid.transform.e], 'axis_convention': 'north-up; row 0 north'}

def artifact_metadata(source, units, *, window=None, grid=None, status='VALIDATED_REAL_SOURCE'):
    return {'source': source, 'units': units, 'study_year': CONFIG['study_year'], 'data_window': window or [CONFIG['study_start_local'], CONFIG['study_end_local_exclusive']], 'timezone': CONFIG['timezone'], 'data_nature': 'REAL', 'status': status, 'run_id': PRELIM_RUN_ID, 'acquired_at_UTC': datetime.now(timezone.utc).isoformat(), 'grid': grid_metadata(grid) if grid is not None else None, 'config': json_ready(PRELIM_CONFIG)}

def track_artifact(path, metadata):
    info = require_file(path, label='Preliminary acquired/derived artifact')
    PRELIM_ARTIFACTS[str(path.relative_to(PRELIM_RUN_DIR))] = {**info, **json_ready(metadata)}
    return info

def save_json_artifact(name, value, metadata):
    path = PRELIM_RUN_DIR / name
    with path.open('x', encoding='utf-8') as handle:
        json.dump(json_ready(value), handle, indent=2, allow_nan=False)
    track_artifact(path, metadata)
    return path

def verify_cached_weather_table(path, original):
    restored = pd.read_csv(path)
    time_columns = [name for name in original.columns if name.endswith('_UTC') or name.endswith('_local')]
    for name in time_columns:
        require_prelim(to_utc_index(restored[name]).equals(to_utc_index(original[name])), f'Cached weather timestamp roundtrip failed: {name}')
    numeric_columns = list(original.select_dtypes(include=['number']).columns)
    require_finite_numeric(restored, numeric_columns, context='persisted annual weather roundtrip')
    for name in numeric_columns:
        require_prelim(np.allclose(restored[name], original[name], rtol=1e-12, atol=1e-12), f'Cached weather numeric roundtrip failed: {name}')
    for name in original.select_dtypes(include=['bool']).columns:
        require_prelim(restored[name].equals(original[name]), f'Cached weather flag changed: {name}')
    require_prelim(len(restored) == len(original), 'Cached weather row count changed.')

def save_table_artifact(name, frame, metadata):
    path = PRELIM_RUN_DIR / name
    with path.open('x', encoding='utf-8', newline='') as handle:
        frame.to_csv(handle, index=False)
    track_artifact(path, metadata)
    save_json_artifact(name + '.metadata.json', metadata, metadata)
    return path

def register_prelim_source(dataset_id, source, units, window, checks, *, native_crs, native_resolution, file_info=None, notes=''):
    require_prelim(PRELIM_EE_SESSION is not None, 'No authenticated Earth Engine session: source cannot be registered.')
    PROVENANCE_REGISTRY.register(dataset_id, source=source, availability='PRESENT', data_nature='REAL', observed_city=CONFIG['city'], expected_city=CONFIG['city'], data_start=window[0], data_end_exclusive=window[1], acquired_at_UTC=datetime.now(timezone.utc).isoformat(), units=units, native_crs=native_crs, native_resolution=native_resolution, target_crs=CONFIG['spatial']['crs'] if dataset_id in {'study_boundary', 'landsat_premonsoon', 'landsat_postmonsoon', 'ghsl_built_surface'} else None, target_resolution=CONFIG['spatial']['resolution_m'] if dataset_id in {'study_boundary', 'landsat_premonsoon', 'landsat_postmonsoon', 'ghsl_built_surface'} else None, quality_status='PASS', source_quality_status='PASS', checks=checks, file_info=file_info, notes=notes, replace=dataset_id in PROVENANCE_REGISTRY._records)
    return require_registered_input(dataset_id, PROVENANCE_REGISTRY, CONFIG)

def numeric_array_hash(array):
    values = np.ascontiguousarray(array)
    return hashlib.sha256(str(values.dtype).encode() + str(values.shape).encode() + values.tobytes()).hexdigest()

def download_common_band(image, band):
    require_prelim(PRELIM_EE_SESSION is not None and COMMON_GRID is not None, 'EE session and real grid required.')
    tile_size = PRELIM_CONFIG['download']['tile_cells']
    nodata = PRELIM_CONFIG['download']['nodata_encoding']
    result = np.full(COMMON_GRID.shape, np.nan, dtype=np.float64)
    covered = np.zeros(COMMON_GRID.shape, dtype=bool)
    prepared = image.select(band).unmask(nodata, sameFootprint=False).toDouble()
    for row in range(0, COMMON_GRID.shape[0], tile_size):
        for col in range(0, COMMON_GRID.shape[1], tile_size):
            height = min(tile_size, COMMON_GRID.shape[0] - row)
            width = min(tile_size, COMMON_GRID.shape[1] - col)
            transform = COMMON_GRID.transform * Affine.translation(col, row)
            expected = GridSpec((height, width), COMMON_GRID.crs, transform)
            try:
                url = prepared.getDownloadURL({'bands': [band], 'crs': COMMON_GRID.crs, 'crs_transform': list(tuple(transform)[:6]), 'dimensions': [width, height], 'format': PRELIM_CONFIG['download']['format']})
                response = requests.get(url, timeout=PRELIM_CONFIG['download']['timeout_seconds'])
                response.raise_for_status()
                with MemoryFile(response.content) as memfile:
                    with memfile.open() as dataset:
                        require_prelim(dataset.count == 1 and dataset.crs is not None, 'Invalid numeric raster response.')
                        actual = GridSpec((dataset.height, dataset.width), dataset.crs.to_wkt(), dataset.transform)
                        validate_grid_alignment(actual, expected, config=CONFIG)
                        values = dataset.read(1).astype(np.float64)
                        values[values == nodata] = np.nan
                result[row:row + height, col:col + width] = values
                covered[row:row + height, col:col + width] = True
            except Exception as error:
                raise RuntimeError(f'Real EE numeric download failed for {band}, tile row={row}, col={col}.') from error
    require_prelim(covered.all(), 'Incomplete tile transfer; raster rejected.')
    result[~DELHI_MASK] = np.nan
    validate_raster_grid(result, COMMON_GRID, COMMON_GRID, config=CONFIG)
    return result

def validate_field(array, name, *, lo=None, hi=None, integer=False, allow_empty=False):
    validate_raster_grid(array, COMMON_GRID, COMMON_GRID, config=CONFIG)
    require_prelim(not np.isinf(array).any(), f'{name}: infinity in raster.')
    values = array[DELHI_MASK & np.isfinite(array)]
    require_prelim(allow_empty or values.size > 0, f'{name}: no valid boundary pixels.')
    if lo is not None:
        require_prelim((values >= lo).all(), f'{name}: values below {lo}; no automatic clipping.')
    if hi is not None:
        require_prelim((values <= hi).all(), f'{name}: values above {hi}; no automatic clipping.')
    if integer:
        require_prelim(np.equal(values, np.round(values)).all(), f'{name}: observation counts must be integers.')
    return {'name': name, 'result': 'PASS', 'detail': f'{values.size} finite valid cells; range, missingness and full grid checked.'}

def save_real_raster(name, array, metadata):
    path = PRELIM_RUN_DIR / (name + '.tif')
    require_prelim(not path.exists(), f'Output already exists: {path}; rerun Section 1A for a new run.')
    nodata = PRELIM_CONFIG['download']['nodata_encoding']
    with rasterio.open(path, 'w', driver='GTiff', height=COMMON_GRID.shape[0], width=COMMON_GRID.shape[1], count=1, dtype='float64', crs=COMMON_GRID.crs, transform=COMMON_GRID.transform, nodata=nodata, compress='deflate') as dataset:
        dataset.write(np.where(np.isfinite(array), array, nodata).astype(np.float64), 1)
        dataset.set_band_description(1, name)
        dataset.update_tags(units=metadata['units'], source=metadata['source'], data_nature='REAL', run_id=PRELIM_RUN_ID, metadata_sidecar=name + '.metadata.json')
    actual_grid = raster_grid_from_file(path)
    validate_raster_grid(array, actual_grid, COMMON_GRID, config=CONFIG)
    with rasterio.open(path) as dataset:
        restored = dataset.read(1, masked=True).filled(np.nan)
    require_prelim(np.array_equal(restored, array, equal_nan=True), f'{name}: GeoTIFF numeric roundtrip failed.')
    array.setflags(write=False)
    info = track_artifact(path, metadata)
    full_metadata = {**metadata, 'numeric_sha256': numeric_array_hash(array), 'file_info': info}
    save_json_artifact(name + '.metadata.json', full_metadata, metadata)
    PRELIM_RASTER_RECORDS[name] = {**full_metadata, 'grid_spec': actual_grid}
    return info

def mask_and_scale_landsat(image):
    cfg = PRELIM_CONFIG['landsat']
    qa = image.select('QA_PIXEL')
    reject_mask = sum((1 << bit for bit in cfg['qa_pixel_reject_bits']))
    valid = qa.bitwiseAnd(reject_mask).eq(0)
    if cfg['qa_radsat_require_zero']:
        valid = valid.And(image.select('QA_RADSAT').eq(0))
    clean = image.updateMask(valid)
    optical = clean.select(['SR_B4', 'SR_B5', 'SR_B6']).multiply(cfg['reflectance_scale']).add(cfg['reflectance_offset'])
    lst = clean.select('ST_B10').multiply(cfg['st_scale_K_per_DN']).add(cfg['st_offset_K']).subtract(cfg['K_to_C']).rename('LST')

    def ratio(bands, name):
        a, b = (optical.select(bands[0]), optical.select(bands[1]))
        denominator = a.add(b)
        return a.subtract(b).divide(denominator).updateMask(a.gte(0).And(b.gte(0)).And(denominator.gt(0))).rename(name)
    fields = [lst, ratio(cfg['NDVI_bands'], 'NDVI'), ratio(cfg['NDBI_bands'], 'NDBI')]
    aggregated = []
    for field in fields:
        aggregated.append(field.reduceResolution(reducer=ee.Reducer.mean(), bestEffort=False, maxPixels=cfg['reduce_resolution_max_pixels']).reproject(COMMON_EE_PROJECTION))
    return ee.Image.cat(aggregated).copyProperties(image, ['system:time_start', 'system:index'])

def retrieve_landsat_season(season):
    cfg = PRELIM_CONFIG['landsat']
    window = cfg['windows'][season]
    collection = ee.ImageCollection(cfg['collections'][0]).merge(ee.ImageCollection(cfg['collections'][1]))
    collection = collection.filterBounds(AOI).filterDate(*window).filter(ee.Filter.eq('PROCESSING_LEVEL', cfg['processing_level'])).filter(ee.Filter.gte('CLOUD_COVER', cfg['cloud_cover_min_inclusive'])).filter(ee.Filter.lt('CLOUD_COVER', cfg['cloud_cover_max_exclusive'])).sort('system:time_start')
    n_scenes = int(collection.size().getInfo())
    require_prelim(n_scenes >= cfg['minimum_scene_count'], f'{season}: no usable real Landsat L2SP scenes.')
    scenes = collection.toList(n_scenes).getInfo()
    catalog_rows = []
    for scene in scenes:
        bands = {b['id']: b for b in scene['bands']}
        require_prelim(set(['SR_B4', 'SR_B5', 'SR_B6', 'ST_B10', 'QA_PIXEL', 'QA_RADSAT']) <= set(bands), f'{season}: expected Landsat science/QA band absent.')
        p = scene['properties']
        catalog_rows.append({'asset_id': scene['id'], 'system_time_start_ms': p['system:time_start'], 'acquisition_UTC': pd.Timestamp(p['system:time_start'], unit='ms', tz='UTC').isoformat(), 'CLOUD_COVER': p['CLOUD_COVER'], 'PROCESSING_LEVEL': p['PROCESSING_LEVEL'], 'LANDSAT_PRODUCT_ID': p.get('LANDSAT_PRODUCT_ID'), 'surface_temperature_algorithm': p.get('ALGORITHM_SOURCE_SURFACE_TEMPERATURE')})
    scene_path = save_table_artifact('landsat_' + season.lower() + '_scenes.csv', pd.DataFrame(catalog_rows), artifact_metadata('; '.join(cfg['collections']), 'scene catalog', window=window))
    prepared = collection.map(mask_and_scale_landsat)
    composite = prepared.select(['LST', 'NDVI', 'NDBI']).median().setDefaultProjection(COMMON_EE_PROJECTION).clip(AOI)
    counts = prepared.select(['LST', 'NDVI', 'NDBI']).count().setDefaultProjection(COMMON_EE_PROJECTION).clip(AOI)
    outputs, checks, availability = ({}, [], [])
    for band in ['LST', 'NDVI', 'NDBI']:
        values = download_common_band(composite, band)
        obs = download_common_band(counts, band)
        checks.append(validate_field(obs, f'{band}_OBS_COUNT_{season}', lo=0, hi=n_scenes, integer=True))
        supported = np.isfinite(obs) & (obs >= cfg['minimum_cell_observations'])
        require_prelim(not (np.isfinite(values) & ~supported).any(), f'{band}: composite has unsupported pixels.')
        name = f'LST_{season}_C' if band == 'LST' else f'{band}_{season}'
        limits = cfg['lst_validation_range_C'] if band == 'LST' else [-1.0, 1.0]
        checks.append(validate_field(values, name, lo=limits[0], hi=limits[1]))
        units = 'degC' if band == 'LST' else 'dimensionless'
        meta = artifact_metadata('; '.join(cfg['collections']), units, window=window, grid=COMMON_GRID)
        meta.update(season=season, source_scene_count=n_scenes, source_scene_catalog=str(scene_path), QA_method=cfg['qa_pixel_reject_bits'], aggregation=cfg['scene_to_grid'], composite=cfg['temporal_composite'])
        save_real_raster(name, values, meta)
        save_real_raster(f'{band}_OBS_COUNT_{season}', obs, {**meta, 'units': 'contributing scenes per common-grid cell'})
        outputs[name] = values
        outputs[f'{band}_OBS_COUNT_{season}'] = obs
        available = DELHI_MASK & np.isfinite(values)
        availability.append({'season': season, 'band': band, 'scene_count': n_scenes, 'valid_cells': int(available.sum()), 'boundary_cells': int(DELHI_MASK.sum()), 'valid_fraction': float(available.sum() / DELHI_MASK.sum()), 'min_count_on_valid': float(obs[available].min()), 'median_count_on_valid': float(np.median(obs[available])), 'max_count_on_valid': float(obs[available].max())})
    checks.append({'name': 'Landsat_C2_L2_physical_scaling', 'result': 'PASS', 'detail': 'Expected source science/QA bands present; configured factors match fixed official product contract.'})
    dataset_id = 'landsat_' + season.lower()
    register_prelim_source(dataset_id, '; '.join(cfg['collections']), {'LST': 'degC', 'NDVI': '1', 'NDBI': '1'}, window, checks, native_crs='Per-scene projections in preserved scene metadata', native_resolution='30 m product grid; thermal sensing coarser', file_info=require_file(scene_path, label='Landsat acquired scene catalog'), notes='Scene scaling and QA from official catalog; range checks are structural QC, not independent field validation.')
    PRELIM_SOURCE_METADATA[dataset_id] = {'window': window, 'scenes': scenes, 'scene_count': n_scenes, 'availability': availability, 'scaling': cfg}
    PRELIM_CHECKS.extend(checks)
    return (outputs, availability)

def select_persistent_hotspot(pre, post, built, boundary_mask, grid, cfg):
    """Select edge-connected regions using the original four-neighbour thresholds and deterministic ranking."""
    for field in [pre, post, built, boundary_mask]:
        validate_raster_grid(field, grid, grid, config=CONFIG)
    require_prelim(cfg['connectivity'] == 4, 'Only the explicitly specified 4-neighbour selector is implemented.')
    valid_lsts = boundary_mask & np.isfinite(pre) & np.isfinite(post)
    valid_built = boundary_mask & np.isfinite(built)
    require_prelim(valid_lsts.sum() >= cfg['minimum_joint_valid_cells'], 'Too few jointly valid seasonal LST cells.')
    require_prelim(valid_built.any(), 'No valid GHSL cells for the built quantile.')
    require_prelim(((built[valid_built] >= 0) & (built[valid_built] <= 1)).all(), 'Invalid built fraction.')
    pre_mean, post_mean = (float(np.mean(pre[valid_lsts])), float(np.mean(post[valid_lsts])))
    pre_std = float(np.std(pre[valid_lsts], ddof=cfg['std_ddof']))
    post_std = float(np.std(post[valid_lsts], ddof=cfg['std_ddof']))
    pre_cut = pre_mean + cfg['sigma_multiplier'] * pre_std
    post_cut = post_mean + cfg['sigma_multiplier'] * post_std
    built_cut = float(np.quantile(built[valid_built], cfg['built_quantile'], method=cfg['quantile_method']))
    valid = valid_lsts & valid_built
    candidate = valid & (pre > pre_cut) & (post > post_cut) & (built >= built_cut)
    visited = np.zeros(candidate.shape, dtype=bool)
    component_cells, rows = ({}, [])
    area = grid.transform.a * -grid.transform.e
    ny, nx = candidate.shape
    for r0, c0 in zip(*np.where(candidate)):
        if visited[r0, c0]:
            continue
        component_id = len(rows) + 1
        stack, cells_in_region = ([(int(r0), int(c0))], [])
        visited[r0, c0] = True
        while stack:
            r, c = stack.pop()
            cells_in_region.append((r, c))
            for rr, cc in [(r - 1, c), (r + 1, c), (r, c - 1), (r, c + 1)]:
                if 0 <= rr < ny and 0 <= cc < nx and candidate[rr, cc] and (not visited[rr, cc]):
                    visited[rr, cc] = True
                    stack.append((rr, cc))
        rr, cc = np.asarray(cells_in_region, dtype=int).T
        component_cells[component_id] = (rr, cc)
        rows.append({'component_id': component_id, 'n_cells': len(rr), 'physical_area_m2': len(rr) * area, 'built_weighted_area_m2': float(np.sum(built[rr, cc]) * area), 'mean_built_fraction': float(np.mean(built[rr, cc])), 'mean_pre_excess_C': float(np.mean(pre[rr, cc] - pre_cut)), 'mean_post_excess_C': float(np.mean(post[rr, cc] - post_cut)), 'survives_minimum': len(rr) >= cfg['minimum_component_cells']})
    columns = ['component_id', 'n_cells', 'physical_area_m2', 'built_weighted_area_m2', 'mean_built_fraction', 'mean_pre_excess_C', 'mean_post_excess_C', 'survives_minimum']
    audit = pd.DataFrame(rows, columns=columns)
    survivors = audit.loc[audit.survives_minimum].sort_values(['built_weighted_area_m2', 'mean_pre_excess_C', 'n_cells', 'component_id'], ascending=[False, False, False, True])
    summary = {'status': 'NO_QUALIFYING_REGION' if survivors.empty else 'SELECTED', 'pre_mean_C': pre_mean, 'pre_std_C': pre_std, 'pre_threshold_C': pre_cut, 'post_mean_C': post_mean, 'post_std_C': post_std, 'post_threshold_C': post_cut, 'built_threshold_fraction': built_cut, 'built_quantile': cfg['built_quantile'], 'std_ddof': cfg['std_ddof'], 'sigma_multiplier': cfg['sigma_multiplier'], 'joint_LST_valid_cells': int(valid_lsts.sum()), 'valid_built_population_cells': int(valid_built.sum()), 'joint_selection_valid_cells': int(valid.sum()), 'candidate_cells': int(candidate.sum()), 'candidate_regions': len(audit), 'surviving_regions': len(survivors), 'connectivity': cfg['connectivity'], 'minimum_component_cells': cfg['minimum_component_cells'], 'minimum_component_area_m2': cfg['minimum_component_cells'] * area, 'selection_order': cfg['selection_order'], 'grid': grid_metadata(grid)}
    if survivors.empty:
        return (None, summary, audit)
    selected = survivors.iloc[0].to_dict()
    rr, cc = component_cells[int(selected['component_id'])]
    mask = np.zeros(grid.shape, dtype=bool)
    mask[rr, cc] = True
    summary.update({key: json_ready(value) for key, value in selected.items()})
    x, y = rasterio.transform.xy(grid.transform, rr, cc, offset='center')
    summary['centroid_xy_m'] = [float(np.mean(x)), float(np.mean(y))]
    row_min, row_max, col_min, col_max = (int(rr.min()), int(rr.max()), int(cc.min()), int(cc.max()))
    summary['selected_cell_bounds_xy_m'] = [grid.transform.c + col_min * grid.transform.a, grid.transform.f + (row_max + 1) * grid.transform.e, grid.transform.c + (col_max + 1) * grid.transform.a, grid.transform.f + row_min * grid.transform.e]
    return (mask, summary, audit)

def normalize_native_projection(metadata):
    """Read numeric affine metadata; Projection.transform() is transform WKT."""
    require_prelim(isinstance(metadata, dict), 'Native projection metadata must be a dictionary.')
    definition = metadata.get('crs') or metadata.get('wkt')
    require_prelim(bool(definition), 'Native projection has neither CRS nor WKT.')
    crs = CRS.from_user_input(definition)
    values = metadata.get('transform')
    require_prelim(isinstance(values, (list, tuple)) and len(values) == 6, 'Native projection must contain six numeric affine coefficients from Describe/getInfo.')
    coefficients = np.asarray(values, dtype=float)
    require_prelim(np.isfinite(coefficients).all(), 'Nonfinite native affine coefficients.')
    affine = Affine(*coefficients)
    require_prelim(affine.determinant != 0, 'Singular native affine transform.')
    return {'crs': definition, 'transform': coefficients.tolist()}

def native_projections_equivalent(left, right):
    """Same coordinate reference system and affine, with the existing 1e-12 tolerance."""
    left, right = (normalize_native_projection(left), normalize_native_projection(right))
    return CRS.from_user_input(left['crs']).equals(CRS.from_user_input(right['crs'])) and np.allclose(left['transform'], right['transform'], rtol=0.0, atol=1e-12)

def derive_hotspot_weather_location(mask, built, grid):
    """
    Choose one reproducible 150 m hotspot cell to locate the coarse ERA5-Land
    weather sample.

    IMPORTANT:
    The 150 m grid is used ONLY to choose a representative hotspot location.
    The weather itself remains one native ERA5-Land grid cell.
    """
    for field in (mask, built):
        validate_raster_grid(field, grid, grid, config=CONFIG)
    require_prelim(mask.dtype == np.dtype(bool) and mask.any(), 'A selected boolean hotspot is required for weather.')
    weights = built[mask].astype(float)
    require_prelim(np.isfinite(weights).all() and ((weights >= 0) & (weights <= 1)).all() and (weights.sum() > 0), 'Selected hotspot lacks valid positive built support; no alternative weather location is selected.')
    rr, cc = np.where(mask)
    x, y = rasterio.transform.xy(grid.transform, rr, cc, offset='center')
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    centroid = np.array([np.average(x, weights=weights), np.average(y, weights=weights)])
    distances_sq = (x - centroid[0]) ** 2 + (y - centroid[1]) ** 2
    chosen = int(np.argmin(distances_sq))
    transformer = Transformer.from_crs(grid.crs, 'EPSG:4326', always_xy=True)
    lon, lat = transformer.transform(x[chosen], y[chosen])
    require_prelim(np.isfinite([lon, lat]).all() and -180 <= lon <= 180 and (-90 <= lat <= 90), 'Hotspot weather location could not be transformed.')
    return {'spatial_scope': PRELIM_CONFIG['weather']['spatial_scope'], 'location_method': PRELIM_CONFIG['weather']['location_method'], 'selected_grid_row': int(rr[chosen]), 'selected_grid_col': int(cc[chosen]), 'point_xy_m': [float(x[chosen]), float(y[chosen])], 'point_lon_lat': [float(lon), float(lat)], 'point_crs': 'EPSG:4326', 'weighted_centroid_xy_m': centroid.tolist(), 'selected_hotspot_cells': int(mask.sum()), 'location_grid': grid_metadata(grid), 'hotspot_mask_numeric_sha256': numeric_array_hash(mask), 'selected_built_numeric_sha256': numeric_array_hash(built[mask]), 'weather_resolution_note': '150 m grid locates hotspot only; weather remains one coarse native ERA5-Land cell.'}

def retrieve_ee_hourly_weather(location):
    cfg = PRELIM_CONFIG['weather']
    require_prelim(PRELIM_EE_SESSION is not None, 'Real Earth Engine session required for hotspot weather.')
    require_prelim(location == derive_hotspot_weather_location(HOTSPOT_MASK, BUILT_FRACTION, COMMON_GRID), 'Weather location does not match the reproducibly selected hotspot.')
    point = ee.Geometry.Point(location['point_lon_lat'])
    expected = ANNUAL_HOURLY_INTERVALS
    start = expected.interval_start_UTC.iloc[0].floor('h') - pd.Timedelta(hours=cfg['source_padding_hours'])
    stop = expected.interval_end_UTC.iloc[-1].ceil('h') + pd.Timedelta(hours=cfg['source_padding_hours'])
    rows, batches, cursor = ([], [], start)
    era5_native_projection, era5_native_provider_metadata, era5_nominal_scale_m = (None, None, None)
    reference_band = cfg['bands'][0]
    while cursor < stop:
        next_month = pd.Timestamp(year=cursor.year + (cursor.month == 12), month=cursor.month % 12 + 1, day=1, tz='UTC')
        end = min(next_month, stop)
        images = ee.ImageCollection(cfg['collection']).filterDate(cursor.isoformat(), end.isoformat()).select(cfg['bands']).sort('system:time_start')
        count, want = (int(images.size().getInfo()), int((end - cursor).total_seconds() / 3600))
        require_prelim(count == want and count > 0, f'ERA5 batch {cursor}–{end}: got {count}, expected {want}.')
        first_image = ee.Image(images.first())
        batch_projections = {band: first_image.select(band).projection().getInfo() for band in cfg['bands']}
        batch_reference = normalize_native_projection(batch_projections[reference_band])
        for band, metadata in batch_projections.items():
            require_prelim(native_projections_equivalent(metadata, batch_reference), f'ERA5 band {band!r} has a different native grid in batch {cursor}.')
        if era5_native_projection is None:
            era5_native_projection = batch_reference
            era5_native_provider_metadata = deepcopy(batch_projections[reference_band])
            era5_nominal_scale_m = float(first_image.select(reference_band).projection().nominalScale().getInfo())
            require_prelim(np.isfinite(era5_nominal_scale_m) and era5_nominal_scale_m > 0, 'Native ERA5 scale unavailable.')
        else:
            require_prelim(native_projections_equivalent(batch_reference, era5_native_projection), 'ERA5 native grid changed across batches.')
        source_projection = ee.Projection(era5_native_projection['crs'], era5_native_projection['transform'])

        def extract_image(image):
            image = ee.Image(image)
            coordinates = ee.Image.pixelCoordinates(source_projection).rename(['native_pixel_x', 'native_pixel_y'])
            indices = coordinates.floor().rename(['native_pixel_col', 'native_pixel_row'])
            samples = image.addBands(coordinates).addBands(indices).sample(region=point, projection=source_projection, dropNulls=False, geometries=True, tileScale=cfg['tile_scale'])
            sampled = ee.Feature(samples.first())
            center = sampled.geometry().coordinates()
            projection = image.select(reference_band).projection()
            band_projections = ee.Dictionary.fromLists(cfg['bands'], [ee.Algorithms.Describe(image.select(band).projection()) for band in cfg['bands']])
            description = ee.Dictionary(ee.Algorithms.Describe(projection))
            return sampled.set({'source_ms': image.get('system:time_start'), 'source_asset_id': image.id(), 'native_sample_count': samples.size(), 'native_pixel_center_lon': center.get(0), 'native_pixel_center_lat': center.get(1), 'native_grid_crs': projection.crs(), 'native_grid_transform': description.get('transform'), 'native_grid_transform_wkt': projection.transform(), 'native_band_projections': band_projections})
        features = ee.FeatureCollection(images.map(extract_image)).sort('source_ms').getInfo()['features']
        require_prelim(len(features) == want, 'ERA5 hotspot extraction lost rows.')
        rows.extend((feature['properties'] for feature in features))
        batches.append({'start_UTC': cursor.isoformat(), 'end_UTC_exclusive': end.isoformat(), 'count': count, 'band_native_projections': batch_projections})
        print(f'ERA5-Land hotspot {cursor:%Y-%m-%d} to {end:%Y-%m-%d}: {count} source hours')
        cursor = end
    raw = pd.DataFrame(rows)
    require_columns(raw, ['source_ms', 'source_asset_id', 'native_sample_count', 'native_pixel_center_lon', 'native_pixel_center_lat', 'native_pixel_x', 'native_pixel_y', 'native_pixel_col', 'native_pixel_row', 'native_grid_crs', 'native_grid_transform', 'native_grid_transform_wkt', 'native_band_projections', *cfg['bands']], context='real hotspot ERA5 source')
    raw['source_timestamp_UTC'] = pd.to_datetime(raw['source_ms'], unit='ms', utc=True)
    times = pd.DatetimeIndex(raw.source_timestamp_UTC)
    source_expectation = pd.date_range(start, stop, freq='h', inclusive='left')
    require_prelim(not times.has_duplicates and times.equals(source_expectation), 'ERA5 source hours missing, duplicate, extra or unordered; no repair performed.')
    require_prelim(raw.source_asset_id.notna().all() and raw.source_asset_id.is_unique, 'Missing/duplicate ERA5 asset identity.')
    require_finite_numeric(raw, [*cfg['bands'], 'native_sample_count', 'native_pixel_center_lon', 'native_pixel_center_lat', 'native_pixel_x', 'native_pixel_y', 'native_pixel_col', 'native_pixel_row'], context='real annual hotspot ERA5 raw inputs')
    require_prelim(raw.native_sample_count.eq(1).all(), 'Every source hour must sample exactly one native ERA5 cell.')
    for row in raw.itertuples(index=False):
        hour_reference = {'crs': row.native_grid_crs, 'transform': row.native_grid_transform}
        require_prelim(native_projections_equivalent(hour_reference, era5_native_projection), f'Source-hour ERA5 grid changed at {row.source_asset_id}.')
        require_prelim(set(row.native_band_projections) == set(cfg['bands']), 'Source-hour band projection records incomplete.')
        for band, metadata in row.native_band_projections.items():
            require_prelim(native_projections_equivalent(metadata, era5_native_projection), f'Source-hour ERA5 band {band!r} grid changed at {row.source_asset_id}.')
    pixel_indices = raw[['native_pixel_col', 'native_pixel_row']].to_numpy(dtype=float)
    pixel_coordinates = raw[['native_pixel_x', 'native_pixel_y']].to_numpy(dtype=float)
    require_prelim(np.allclose(pixel_coordinates, pixel_indices + 0.5, rtol=0.0, atol=1e-12), 'Native ERA5 sampled pixel coordinates do not match containing-cell centers.')
    require_prelim(np.equal(pixel_indices, np.floor(pixel_indices)).all() and np.equal(pixel_indices, pixel_indices[0]).all(), 'Native ERA5 pixel index changed across source hours.')
    centers = raw[['native_pixel_center_lon', 'native_pixel_center_lat']].to_numpy(dtype=float)
    require_prelim(np.allclose(centers, centers[0], rtol=0.0, atol=1e-12), 'Native ERA5 sampled cell changed across source hours.')
    era5_affine = Affine(*era5_native_projection['transform'])
    point_xy = Transformer.from_crs('EPSG:4326', era5_native_projection['crs'], always_xy=True).transform(*location['point_lon_lat'])
    containing_index = np.floor(~era5_affine * point_xy).astype(int)
    require_prelim(np.array_equal(pixel_indices[0], containing_index), 'Sampled ERA5 cell does not contain the hotspot point.')
    native_center_xy = era5_affine * tuple(pixel_indices[0] + 0.5)
    predicted_center = Transformer.from_crs(era5_native_projection['crs'], 'EPSG:4326', always_xy=True).transform(*native_center_xy)
    require_prelim(np.allclose(centers[0], predicted_center, rtol=0.0, atol=1e-12), 'Native ERA5 sampled center does not match its recorded affine/index.')
    raw['radiation_interval_start_UTC'] = times - pd.Timedelta(hours=1)
    raw['radiation_interval_end_UTC'] = times
    raw['source_timestamp_local'] = times.tz_convert(CONFIG['timezone'])
    support = {'spatial_scope': cfg['spatial_scope'], 'hotspot_location': location, 'extraction_method': cfg['extraction_method'], 'native_projection': era5_native_projection, 'native_provider_projection_metadata': era5_native_provider_metadata, 'native_crs_definition': era5_native_projection['crs'], 'native_transform': era5_native_projection['transform'], 'native_nominal_scale_m': era5_nominal_scale_m, 'catalog_nominal_resolution_m': cfg['catalog_nominal_resolution_m'], 'sampled_native_cell_center_lon_lat': centers[0].tolist(), 'native_pixel_col_row': pixel_indices[0].astype(int).tolist(), 'pixel_coordinate_convention': 'EE pixelCoordinates continuous center=(col+0.5,row+0.5); integer containing-cell indices=floor(coordinates).', 'native_cells_per_hour': 1, 'weather_resampled_to_common_grid': False, 'projection_validation': 'Every source hour and required band: semantic CRS + six numeric affine coefficients; fixed integer pixel and center containing hotspot point.', 'affine_absolute_tolerance': 1e-12, 'cell_center_absolute_tolerance_degrees': 1e-12, 'support_limitation': 'Single coarse regional reanalysis cell at hotspot location; no 150 m meteorological resolution.'}
    raw.attrs['spatial_support'] = support
    return (raw, {'collection': cfg['collection'], 'native_projection': era5_native_projection, 'native_provider_projection_metadata': era5_native_provider_metadata, 'native_crs_definition': era5_native_projection['crs'], 'native_transform': era5_native_projection['transform'], 'spatial_support': support, 'fetch_start_UTC': start.isoformat(), 'fetch_end_UTC_exclusive': stop.isoformat(), 'batches': batches, 'source_hour_count': len(raw), 'radiation_units': 'J/m2', 'source_interval': cfg['source_radiation_interval'], 'data_nature': 'REAL'})

def rebin_real_energy(times, energy, edges):
    t = to_utc_index(times).as_unit('ns').asi8
    e = to_utc_index(edges).as_unit('ns').asi8
    hour = 3600 * 10 ** 9
    require_prelim(len(t) == len(energy) and np.isfinite(energy).all(), 'Energy/time lengths or finite values invalid.')
    require_prelim(len(t) > 1 and np.all(np.diff(t) == hour), 'Source energy timestamps must be unique consecutive hours.')
    require_prelim(len(e) > 1 and np.all(np.diff(e) == hour), 'Target edges must be consecutive hours.')
    out = np.zeros(len(e) - 1, dtype=float)
    coverage = np.zeros(len(out), dtype=np.int64)
    i = j = 0
    while i < len(t) and j < len(out):
        left_source, right_source = (t[i] - hour, t[i])
        left_target, right_target = (e[j], e[j + 1])
        if right_source <= left_target:
            i += 1
            continue
        if left_source >= right_target:
            j += 1
            continue
        overlap = min(right_source, right_target) - max(left_source, left_target)
        if overlap > 0:
            out[j] += energy[i] * overlap / hour
            coverage[j] += overlap
        if right_source <= right_target:
            i += 1
        else:
            j += 1
    require_prelim(np.all(coverage == hour), 'Uncovered or multiply covered radiation interval; no substitute.')
    return out

def radiation_overlap_ledger(source_times, edges):
    t = to_utc_index(source_times).as_unit('ns').asi8
    e = to_utc_index(edges).as_unit('ns').asi8
    hour = 3600 * 10 ** 9
    records = []
    for j, (left, right) in enumerate(zip(e[:-1], e[1:])):
        first = int(np.searchsorted(t, left, side='right'))
        stop = int(np.searchsorted(t - hour, right, side='left'))
        for i in range(first, stop):
            overlap = max(0, min(int(t[i]), int(right)) - max(int(t[i] - hour), int(left)))
            if overlap:
                records.append({'target_row': j, 'source_row': i, 'source_timestamp_UTC': pd.Timestamp(t[i], tz='UTC'), 'overlap_seconds': overlap / 1000000000.0, 'source_energy_weight': overlap / hour})
    ledger = pd.DataFrame(records)
    require_prelim(not ledger.empty, 'No radiation overlaps exist.')
    coverage = ledger.groupby('target_row').overlap_seconds.sum().reindex(range(len(e) - 1))
    require_prelim(coverage.notna().all() and np.equal(coverage.to_numpy(), 3600).all(), 'Independent radiation overlap ledger lacks complete coverage.')
    return ledger

def apply_provider_shortwave_physical_qc(raw):
    cfg = PRELIM_CONFIG['weather']
    policy = cfg['shortwave_physical_qc']
    limit = policy['maximum_correctable_negative_magnitude_Jm2']
    require_prelim(limit == 5.0 and policy['physical_minimum_Jm2'] == 0.0, 'Shortwave QC differs from the explicitly approved 5 J/m2 policy.')
    require_finite_numeric(raw, [cfg['bands'][3]], context='raw provider shortwave physical QC')
    original = raw[cfg['bands'][3]].to_numpy(dtype=float, copy=True)
    require_prelim(not (original < -limit).any(), 'Provider shortwave increment below -5 J/m2; hard failure requiring source review.')
    corrected = original < 0.0
    model = original.copy()
    model[corrected] = 0.0
    magnitude = model - original
    require_prelim((model >= 0.0).all() and (magnitude >= 0.0).all() and np.array_equal(model[~corrected], original[~corrected]), 'Shortwave physical QC changed a nonnegative source value or failed its physical bound.')
    source_qc = raw[['source_timestamp_UTC', 'source_asset_id', 'radiation_interval_start_UTC', 'radiation_interval_end_UTC', 'native_pixel_center_lon', 'native_pixel_center_lat']].copy(deep=True)
    for name in ['source_timestamp_UTC', 'radiation_interval_start_UTC', 'radiation_interval_end_UTC']:
        source_qc[name] = to_utc_index(source_qc[name]).as_unit('ns')
    source_qc.insert(0, 'source_row', np.arange(len(raw), dtype=np.int64))
    source_qc['raw_shortwave_energy_Jm2'] = original
    source_qc['model_shortwave_energy_Jm2'] = model
    source_qc['shortwave_qc_correction_Jm2'] = magnitude
    source_qc['shortwave_qc_applied_flag'] = corrected
    source_qc['shortwave_qc_status'] = np.where(corrected, 'CORRECTED_SMALL_NEGATIVE_TO_ZERO', 'UNCHANGED_NONNEGATIVE')
    source_qc.attrs.update(data_nature='REAL', processing='PROVIDER_SHORTWAVE_PHYSICAL_QC', policy=deepcopy(policy), raw_source_file='era5_hourly_raw.csv', spatial_support=deepcopy(raw.attrs['spatial_support']))
    return source_qc

def prepare_annual_real_weather(raw):
    cfg = PRELIM_CONFIG['weather']
    weather = ANNUAL_HOURLY_INTERVALS.copy()
    weather.attrs = {}
    edges = pd.DatetimeIndex([weather.interval_start_UTC.iloc[0], *weather.interval_end_UTC])
    midpoints = pd.DatetimeIndex(weather.interval_start_UTC) + pd.Timedelta(minutes=30)
    source_times = pd.DatetimeIndex(raw.source_timestamp_UTC)
    require_prelim(not source_times.has_duplicates and source_times.is_monotonic_increasing and np.all(np.diff(source_times.as_unit('ns').asi8) == 3600 * 10 ** 9), 'Source hours invalid before weather conversion.')
    require_prelim(midpoints.min() >= source_times.min() and midpoints.max() <= source_times.max(), 'Instantaneous weather lacks bracketing source coverage; no extrapolation.')
    require_finite_numeric(raw, cfg['bands'], context='real annual weather conversion')
    for output, source in [('T_air_C', cfg['bands'][0]), ('u10_ms', cfg['bands'][1]), ('v10_ms', cfg['bands'][2])]:
        values = np.interp(midpoints.as_unit('ns').asi8, source_times.as_unit('ns').asi8, raw[source].to_numpy(dtype=float))
        weather[output] = values - 273.15 if output == 'T_air_C' else values
    weather['wind10_ms'] = np.hypot(weather.u10_ms, weather.v10_ms)
    source_qc = apply_provider_shortwave_physical_qc(raw)
    ledger = radiation_overlap_ledger(source_times, edges)
    source_rows = ledger.source_row.to_numpy(dtype=int)
    weights = ledger.source_energy_weight.to_numpy(dtype=float)
    for name in ['source_asset_id', 'native_pixel_center_lon', 'native_pixel_center_lat']:
        ledger[name] = raw[name].iloc[source_rows].to_numpy()
    for name in ['raw_shortwave_energy_Jm2', 'model_shortwave_energy_Jm2', 'shortwave_qc_correction_Jm2', 'shortwave_qc_applied_flag', 'shortwave_qc_status']:
        ledger[name] = source_qc[name].to_numpy()[source_rows]
    ledger['target_interval_start_UTC'] = weather.interval_start_UTC.iloc[ledger.target_row.to_numpy()].dt.as_unit('ns').to_numpy()
    ledger['target_interval_end_UTC'] = weather.interval_end_UTC.iloc[ledger.target_row.to_numpy()].dt.as_unit('ns').to_numpy()
    ledger['raw_shortwave_contribution_Jm2'] = ledger.raw_shortwave_energy_Jm2.to_numpy() * weights
    ledger['model_shortwave_contribution_Jm2'] = ledger.model_shortwave_energy_Jm2.to_numpy() * weights
    ledger['shortwave_correction_contribution_Jm2'] = ledger.shortwave_qc_correction_Jm2.to_numpy() * weights
    correction_ledger = ledger.loc[ledger.shortwave_qc_applied_flag].copy(deep=True).reset_index(drop=True)
    for output, source, energy_name in [('G_Wm2', cfg['bands'][3], 'shortwave_energy_Jm2'), ('LW_down_Wm2', cfg['bands'][4], 'longwave_energy_Jm2')]:
        source_energy = source_qc.model_shortwave_energy_Jm2.to_numpy(dtype=float) if output == 'G_Wm2' else raw[source].to_numpy(dtype=float)
        energy = rebin_real_energy(source_times, source_energy, edges)
        contributions = source_energy[ledger.source_row.to_numpy()] * ledger.source_energy_weight.to_numpy()
        reference = pd.Series(contributions).groupby(ledger.target_row.to_numpy()).sum().reindex(range(len(weather))).to_numpy()
        require_prelim(np.allclose(energy, reference, rtol=1e-12, atol=1e-08), 'Radiation rebin fails independent energy ledger.')
        weather[energy_name] = energy
        weather[output] = energy / CONFIG['hourly_interval_seconds']
    weather['instantaneous_sample_UTC'] = midpoints
    weather['instantaneous_sample_local'] = midpoints.tz_convert(CONFIG['timezone'])
    weather['shortwave_negative_flag'] = weather.G_Wm2 < 0
    weather['raw_shortwave_energy_Jm2'] = rebin_real_energy(source_times, source_qc.raw_shortwave_energy_Jm2.to_numpy(dtype=float), edges)
    weather['shortwave_qc_correction_Jm2'] = rebin_real_energy(source_times, source_qc.shortwave_qc_correction_Jm2.to_numpy(dtype=float), edges)
    weather['shortwave_qc_applied_flag'] = weather.shortwave_qc_correction_Jm2 > 0.0
    require_prelim((weather.G_Wm2 >= 0.0).all() and (weather.shortwave_energy_Jm2 >= 0.0).all(), 'Canonical downward shortwave must be nonnegative in every target interval.')
    require_prelim(np.allclose(weather.shortwave_energy_Jm2, weather.raw_shortwave_energy_Jm2 + weather.shortwave_qc_correction_Jm2, rtol=0.0, atol=1e-08), 'Adjusted energy does not reconcile with raw energy and correction.')
    reference_corrections = ledger.groupby('target_row').shortwave_correction_contribution_Jm2.sum().reindex(range(len(weather)), fill_value=0.0).to_numpy()
    require_prelim(np.allclose(weather.shortwave_qc_correction_Jm2, reference_corrections, rtol=0.0, atol=1e-12), 'Annual correction energy fails independent overlap ledger.')
    annual_correction = float(weather.shortwave_qc_correction_Jm2.sum())
    require_prelim(np.isclose(annual_correction, correction_ledger.shortwave_correction_contribution_Jm2.sum(), rtol=0.0, atol=1e-12), 'Total annual correction differs from correction-only ledger.')
    qc_summary = {'policy': deepcopy(cfg['shortwave_physical_qc']), 'source_hour_count': len(source_qc), 'corrected_padded_source_hours': int(source_qc.shortwave_qc_applied_flag.sum()), 'corrected_source_hours_overlapping_annual_targets': int(correction_ledger.source_row.nunique()), 'corrected_target_intervals': int(weather.shortwave_qc_applied_flag.sum()), 'correction_overlap_rows': len(correction_ledger), 'minimum_raw_shortwave_increment_Jm2': float(source_qc.raw_shortwave_energy_Jm2.min()), 'total_padded_source_correction_Jm2': float(source_qc.shortwave_qc_correction_Jm2.sum()), 'total_annual_correction_energy_Jm2': annual_correction, 'total_annual_model_shortwave_energy_Jm2': float(weather.shortwave_energy_Jm2.sum()), 'minimum_canonical_G_Wm2': float(weather.G_Wm2.min()), 'canonical_nonnegative_all_intervals': bool((weather.G_Wm2 >= 0.0).all()), 'raw_values_preserved': True, 'other_weather_fields_modified': False, 'source_qc_file': 'era5_shortwave_source_qc.csv', 'correction_ledger_file': 'era5_shortwave_correction_ledger.csv'}
    require_prelim((weather.LW_down_Wm2 > 0).all(), 'Nonpositive downward longwave radiation.')
    lo, hi = cfg['air_validation_range_C']
    require_prelim(weather.T_air_C.between(lo, hi).all(), 'Air temperature outside explicit QC bounds.')
    numeric = ['T_air_C', 'u10_ms', 'v10_ms', 'wind10_ms', 'G_Wm2', 'LW_down_Wm2', 'shortwave_energy_Jm2', 'longwave_energy_Jm2', 'raw_shortwave_energy_Jm2', 'shortwave_qc_correction_Jm2']
    require_finite_numeric(weather, numeric, context='canonical annual weather')
    coverage = validate_hourly_coverage(weather.timestamp_UTC, ANNUAL_HOURLY_INTERVALS, timestamp_label=CONFIG['timestamp_label'])
    require_prelim(len(weather) == EXPECTED_ANNUAL_HOURS, 'Unexpected annual weather row count.')
    units = {'T_air_C': 'degC', 'u10_ms': 'm/s', 'v10_ms': 'm/s', 'wind10_ms': 'm/s', 'G_Wm2': 'W/m2 interval mean', 'LW_down_Wm2': 'W/m2 interval mean', 'shortwave_energy_Jm2': 'J/m2 target interval', 'longwave_energy_Jm2': 'J/m2 target interval', 'raw_shortwave_energy_Jm2': 'J/m2 raw source overlap sum', 'shortwave_qc_correction_Jm2': 'J/m2 positive correction to target interval'}
    weather.attrs.update(data_nature='REAL', source=cfg['collection'], units=units, timestamp_label=CONFIG['timestamp_label'], timezone=CONFIG['timezone'], instantaneous_sample='midpoint', review_status='PRELIMINARY_ONLY')
    if 'spatial_support' in raw.attrs:
        weather.attrs['spatial_support'] = deepcopy(raw.attrs['spatial_support'])
    weather.attrs['shortwave_physical_qc'] = deepcopy(qc_summary)
    return (weather, ledger, coverage, source_qc, correction_ledger, qc_summary)

def derive_local_day_roof_screen_air(weather, raw):
    cfg = PRELIM_CONFIG['pcm_screen']
    tz = cfg['air_timezone']
    require_prelim(tz == CONFIG['timezone'] == 'Asia/Kolkata' and cfg['required_hours_per_local_day'] == 24, 'This screen requires complete 24-hour Asia/Kolkata local days.')
    validate_hourly_coverage(weather.timestamp_UTC, ANNUAL_HOURLY_INTERVALS, timestamp_label=CONFIG['timestamp_label'])
    require_finite_numeric(weather, ['T_air_C'], context='hourly hotspot screen air')
    require_columns(weather, ['instantaneous_sample_UTC', 'instantaneous_sample_local'], context='hourly air sample times')
    samples = to_utc_index(weather.instantaneous_sample_UTC)
    expected_samples = pd.DatetimeIndex(ANNUAL_HOURLY_INTERVALS.interval_start_UTC) + pd.Timedelta(minutes=30)
    require_prelim(samples.equals(expected_samples) and to_utc_index(weather.instantaneous_sample_local).equals(samples), 'Screen hourly sample instants differ from the canonical midpoint calendar.')
    require_prelim('spatial_support' in weather.attrs and 'spatial_support' in raw.attrs and (weather.attrs['spatial_support'] == raw.attrs['spatial_support']) and (weather.attrs['spatial_support']['spatial_scope'] == PRELIM_CONFIG['weather']['spatial_scope']), 'Screen air must use the canonical selected-hotspot native-cell weather.')
    require_finite_numeric(raw, [PRELIM_CONFIG['weather']['bands'][0]], context='acquired hotspot raw air')
    source_times = to_utc_index(raw.source_timestamp_UTC)
    require_prelim(not source_times.has_duplicates and source_times.is_monotonic_increasing and np.all(np.diff(source_times.as_unit('ns').asi8) == 3600 * 10 ** 9), 'Screen raw air has missing/duplicate/unordered hours; no interpolation permitted.')
    source_rows = source_times.get_indexer(samples)
    require_prelim((source_rows >= 0).all(), 'Screen samples are not acquired native hourly observations; no interpolation permitted.')
    raw_C = raw[PRELIM_CONFIG['weather']['bands'][0]].to_numpy(dtype=float)[source_rows] - 273.15
    require_prelim(np.array_equal(weather.T_air_C.to_numpy(dtype=float), raw_C), 'Screen air is not identical to acquired hourly temperature converted to degC.')
    start, stop = PRELIM_CONFIG['landsat']['windows']['PREMONSOON']
    local_start, local_stop = (pd.Timestamp(start).tz_localize(tz), pd.Timestamp(stop).tz_localize(tz))
    require_prelim(local_start == local_start.normalize() and local_stop == local_stop.normalize() and (local_start < local_stop), 'Pre-monsoon screen dates must be ordered local midnight boundaries.')
    local_samples = samples.tz_convert(tz)
    included = (local_samples >= local_start) & (local_samples < local_stop)
    observed = local_samples[included]
    expected_hours = pd.date_range(local_start + pd.Timedelta(minutes=30), local_stop, freq='h', inclusive='left')
    require_prelim(observed.equals(expected_hours), 'Incomplete local pre-monsoon hourly coverage; no repairs permitted.')
    seasonal = pd.DataFrame({'local_date': observed.strftime('%Y-%m-%d'), 'T_air_C': weather.T_air_C.to_numpy(dtype=float)[included]})
    daily = seasonal.groupby('local_date', sort=False).agg(contributing_hourly_observations=('T_air_C', 'size'), local_daily_Tmax_C=('T_air_C', 'max')).reset_index()
    expected_dates = pd.date_range(local_start, local_stop, freq='D', inclusive='left').strftime('%Y-%m-%d').tolist()
    require_prelim(daily.local_date.tolist() == expected_dates, 'Local pre-monsoon days missing, extra or unordered.')
    complete = daily.contributing_hourly_observations.eq(cfg['required_hours_per_local_day'])
    daily['completeness_status'] = np.where(complete, 'COMPLETE', 'INCOMPLETE')
    require_prelim(complete.all(), 'Incomplete local daily Tmax support; no missing-hour/day substitution permitted.')
    require_finite_numeric(daily, ['local_daily_Tmax_C'], context='local daily air maxima')
    unrounded = float(daily.local_daily_Tmax_C.mean())
    lo, hi = PRELIM_CONFIG['weather']['air_validation_range_C']
    require_prelim(np.isfinite(unrounded) and lo <= unrounded <= hi, 'Local daily-max mean outside air QC bounds.')
    metadata = {'source': PRELIM_CONFIG['weather']['collection'], 'derived_from': 'ANNUAL_WEATHER.T_air_C', 'raw_source_file': 'era5_hourly_raw.csv', 'annual_source_file': 'annual_weather.csv', 'temperature_band': PRELIM_CONFIG['weather']['bands'][0], 'timezone': tz, 'window': [start, stop], 'window_basis': 'local calendar dates; [local midnight start, local midnight stop)', 'local_start_inclusive': local_start.isoformat(), 'local_end_exclusive': local_stop.isoformat(), 'UTC_start_inclusive': local_start.tz_convert('UTC').isoformat(), 'UTC_end_exclusive': local_stop.tz_convert('UTC').isoformat(), 'hourly_sample_semantics': 'instantaneous canonical midpoint; exact acquired UTC source hour; no interpolation', 'first_sample_local': observed[0].isoformat(), 'last_sample_local': observed[-1].isoformat(), 'expected_day_count': len(expected_dates), 'observed_day_count': len(daily), 'required_hours_per_day': cfg['required_hours_per_local_day'], 'expected_hour_count': len(expected_hours), 'observed_hour_count': len(seasonal), 'complete_local_days': bool(complete.all()), 'missing_policy': cfg['air_missing_policy'], 'unrounded_lower_C': unrounded, 'statistic': cfg['lower_statistic'], 'spatial_support': deepcopy(weather.attrs['spatial_support']), 'data_nature': 'REAL'}
    return (unrounded, daily, metadata)

def screen_workbook_roof(all_rows, runnable_rows, lower, upper):
    """Binary inclusive overlap screen; retain all audited workbook decisions."""
    require_prelim(np.isfinite([lower, upper]).all() and lower <= upper, 'Derived screen bounds nonfinite or reversed; review input, no replacement window.')
    require_columns(all_rows, ['pcm_id', 'Tm_low_C', 'Tm_high_C', 'section0_readiness'], context='Section-0 PCM audit')
    require_finite_numeric(runnable_rows, ['Tm_low_C', 'Tm_high_C'], context='Section-0 runnable melting ranges')
    require_prelim(all_rows.pcm_id.is_unique and runnable_rows.pcm_id.is_unique, 'PCM IDs not unique.')
    runnable_ids = set(runnable_rows.pcm_id)
    require_prelim(runnable_ids == set(all_rows.loc[all_rows.section0_readiness == 'RUNNABLE_SCHEMA_ONLY', 'pcm_id']), 'PCM_RUNNABLE disagrees with the unchanged Section-0 partition.')
    screen = all_rows[['pcm_id', 'material_name', 'Tm_low_C', 'Tm_high_C', 'section0_readiness', 'source_quality_status', 'section0_issues']].copy()
    screen['screen_lower_C'] = lower
    screen['screen_upper_C'] = upper
    flags, status, reasons = ([], [], [])
    for row in screen.itertuples():
        if row.pcm_id not in runnable_ids:
            flags.append(pd.NA)
            status.append('NOT_EVALUATED')
            reasons.append('Outside Section-0 model-input-ready table: ' + row.section0_readiness + '; ' + str(row.section0_issues))
            continue
        reasons_row = []
        if row.Tm_high_C < lower:
            reasons_row.append('melting interval wholly below plausible thermal window')
        if row.Tm_low_C > upper:
            reasons_row.append('melting interval wholly above plausible thermal window')
        passed = not reasons_row
        flags.append(passed)
        status.append('PASS' if passed else 'FAIL')
        reasons.append('melting interval overlaps inclusive plausible thermal window; thermal relevance only' if passed else '; '.join(reasons_row))
    screen['pass'] = pd.array(flags, dtype='boolean')
    screen['screen_status'] = status
    screen['reason'] = reasons
    screen['overlap_low_C'] = np.maximum(screen.Tm_low_C, lower)
    screen['overlap_high_C'] = np.minimum(screen.Tm_high_C, upper)
    screen['overlap_width_C'] = np.maximum(0.0, screen.overlap_high_C - screen.overlap_low_C)
    screen.loc[screen.screen_status == 'NOT_EVALUATED', ['overlap_low_C', 'overlap_high_C', 'overlap_width_C']] = np.nan
    screen['screen_method'] = CUMULATIVE_SCREEN_METHOD
    screen['review_status'] = PRELIM_CONFIG['pcm_screen']['review_status']
    ids = screen.loc[screen.screen_status == 'PASS', 'pcm_id'].tolist()
    candidates = runnable_rows.loc[runnable_rows.pcm_id.isin(ids)].copy()
    candidates['screen_lower_C'], candidates['screen_upper_C'] = (lower, upper)
    candidates['review_status'] = PRELIM_CONFIG['pcm_screen']['review_status']
    decisions = screen.set_index('pcm_id')
    for name in ['overlap_low_C', 'overlap_high_C', 'overlap_width_C', 'pass', 'screen_status', 'reason', 'screen_method']:
        candidates[name] = candidates.pcm_id.map(decisions[name])
    require_prelim(len(screen) == len(all_rows) and len(ids) == len(candidates), 'PCM screen audit loses workbook rows.')
    return (screen, candidates, ids)

def run_overlap_screen_tests():
    cases = [('wholly_below', 30.0, 35.0, False), ('wholly_above', 48.0, 52.0, False), ('partial_from_below', 35.0, 40.0, True), ('partial_from_above', 44.0, 48.0, True), ('completely_inside', 40.0, 44.0, True), ('screen_inside_PCM', 30.0, 52.0, True), ('lower_touch', 37.8, 37.8, True), ('upper_touch', 46.1, 46.1, True)]
    table = pd.DataFrame([{'pcm_id': name, 'material_name': 'MECHANICAL_TEST_ONLY', 'Tm_low_C': low, 'Tm_high_C': high, 'section0_readiness': 'RUNNABLE_SCHEMA_ONLY', 'source_quality_status': 'NOT CHECKED', 'section0_issues': ''} for name, low, high, passed in cases])
    screen, candidates, ids = screen_workbook_roof(table, table.copy(), 37.8, 46.1)
    assert screen['pass'].tolist() == [passed for name, low, high, passed in cases]
    assert ids == candidates.pcm_id.tolist() == [name for name, low, high, passed in cases if passed]
    assert screen.loc[screen.pcm_id.isin(['lower_touch', 'upper_touch']), 'overlap_width_C'].eq(0).all()
    return {'status': 'PASS', 'count': len(cases), 'scientific_run': False, 'data_nature': 'MECHANICAL_TEST_ONLY', 'cases': [{'name': name, 'expected_pass': passed, 'result': 'PASS'} for name, low, high, passed in cases]}

def save_prelim_figure(fig, filename, metadata):
    path = PRELIM_RUN_DIR / filename
    require_prelim(not path.exists(), f'Figure file already exists: {path}')
    fig.savefig(path, dpi=PRELIM_CONFIG['figures']['dpi'], bbox_inches='tight')
    track_artifact(path, metadata)
    save_json_artifact(filename + '.metadata.json', metadata, metadata)
    FIGURE_PATHS.append(str(path))
    plt.show()
    plt.close(fig)

def dress_projected_map(ax):
    polygons = [DELHI_PROJECTED_SHAPE] if DELHI_PROJECTED_SHAPE.geom_type == 'Polygon' else list(DELHI_PROJECTED_SHAPE.geoms)
    for polygon in polygons:
        for ring in [polygon.exterior, *polygon.interiors]:
            x, y = ring.xy
            ax.plot(np.asarray(x) / 1000, np.asarray(y) / 1000, color='#222222', linewidth=0.7)
    bounds = grid_metadata(COMMON_GRID)['bounds_xy_m']
    x0, y0, x1, y1 = np.asarray(bounds) / 1000
    ax.set_xlim(x0, x1)
    ax.set_ylim(y0, y1)
    ax.set_aspect('equal')
    ax.set_xlabel('UTM easting (km; EPSG:32643)')
    ax.set_ylabel('UTM northing (km)')
    ax.annotate('N', xy=(0.93, 0.92), xytext=(0.93, 0.8), xycoords='axes fraction', ha='center', arrowprops={'arrowstyle': '-|>', 'color': '#222222'})
    bar_km = 10
    sx, sy = (x0 + 0.07 * (x1 - x0), y0 + 0.06 * (y1 - y0))
    ax.plot([sx, sx + bar_km], [sy, sy], color='#222222', linewidth=2)
    ax.text(sx + bar_km / 2, sy + 0.015 * (y1 - y0), '10 km', ha='center', fontsize=9)

def draw_prelim_map(ax, array, title, cmap, limits):
    left, bottom, right, top = np.asarray(grid_metadata(COMMON_GRID)['bounds_xy_m']) / 1000
    image = ax.imshow(np.ma.masked_invalid(array), extent=[left, right, bottom, top], origin='upper', cmap=cmap, vmin=limits[0], vmax=limits[1], interpolation='nearest')
    dress_projected_map(ax)
    ax.set_title(title)
    return image

C1_IST = CONFIG['timezone']

C1_MODEL_VERSION = 'category1-cumulative-exact-L0-coexistence-2.0.0'

CATEGORY1_HANDOFF = None

CATEGORY1_READY = False

C1_SCIENTIFIC_STATUS = 'NOT_RUN'

CATEGORY1_CONFIG = {'year': 2025, 'A_m2': 1.0, 'd_m': 0.02, 'alpha_s': 0.6, 'epsilon_LW': 0.9, 'bottom_boundary': 'adiabatic', 'surface_boundary': 'environment', 'wind_reference_height_m': 10.0, 'N_cells': 10, 'dt_s': 60.0, 'T_ref_C': 0.0, 'initial_temperature_C': 26.0, 'initial_liquid_fraction_at_transition': 0.0, 'T_base_C': 26.0, 'f_sat': 0.99, 'f_solid': 0.01, 'r_tol_per_h': 0.005, 'reset_limit': 0.1, 'morning_checkpoint': '06:00', 'weather_time_label': 'midpoint', 'report_start_IST': None, 'report_end_IST': None, 'solver_max_iterations': 40, 'solver_atol_Wm2': 1e-07, 'solver_rtol': 1e-09, 'minimum_dt_s': 0.9375, 'energy_atol_Jm2': 0.05, 'energy_rtol': 1e-07, 'prescribed_top_flux_Wm2': 0.0, 'save_cell_states': True, 'output_dir': 'runtime/category1'}

CATEGORY1_VALIDATION = {'surface_atol_Wm2': 1e-07, 'step_energy_atol_Jm2': 0.05, 'annual_energy_atol_Jm2': 0.05, 'annual_energy_rtol': 1e-07, 'latent_accounting_atol_Jm2': 0.05, 'latent_accounting_rtol': 1e-07, 'continuity_atol_Jkg': 1e-08, 'automatic_scientific_cache_reuse': False}

C1_SHARED_CONFIG = deepcopy(CONFIG)

C1_PROJECT_ROOT = CONFIG['project_root']

CATEGORY1_IO = {'input_route': 'memory', 'preliminary_artifact_dir': None, 'workbook_path': str(CONFIG['pcm_workbook_path']), 'workbook_sheet': CONFIG['pcm_sheet_name'], 'expected_workbook_sha256': CONFIG['pcm_expected_sha256'], 'output_dir': str(CONFIG['output_dir'] / 'category1' / CUMULATIVE_RUN_ID), 'persist_real_outputs': True}

C1_SIGMA_Wm2K4 = 5.670374419e-08

C1_PCM_COLUMNS = list(PCM_COLUMNS)

C1_PCM_REQUIRED_NUMBERS = list(PCM_REQUIRED_NUMBERS)

C1_STATE_NAMES = ('ACTIVE', 'SATURATED', 'UNACTIVATED', 'OTHER NO UPTAKE', 'REGENERATING', 'IDLE')

C1_STATE_DURATION_COLUMNS = ('active_h', 'saturated_h', 'unactivated_h', 'other_no_uptake_h', 'regenerating_h', 'idle_h')

def c1_make_run_config(overrides=None):
    """Validate modelling assumptions without putting them into material rows."""
    cfg = dict(CATEGORY1_CONFIG)
    unknown = set(overrides or {}) - set(cfg)
    if unknown:
        raise ValueError(f'Unknown configuration keys: {sorted(unknown)}')
    cfg.update(overrides or {})
    for name in ('A_m2', 'd_m', 'dt_s', 'minimum_dt_s'):
        if not np.isfinite(cfg[name]) or cfg[name] <= 0:
            raise ValueError(f'{name} must be finite and positive.')
    n = cfg['N_cells']
    if int(n) != n or n < 1:
        raise ValueError('N_cells must be a positive integer; use 10 for the benchmark.')
    cfg['N_cells'] = int(n)
    if cfg['dt_s'] > 3600 or cfg['minimum_dt_s'] > cfg['dt_s']:
        raise ValueError('Require minimum_dt_s <= dt_s <= 3600.')
    for name in ('alpha_s', 'epsilon_LW', 'f_sat', 'f_solid', 'reset_limit', 'initial_liquid_fraction_at_transition'):
        if not 0 <= cfg[name] <= 1:
            raise ValueError(f'{name} must be in [0, 1].')
    if cfg['f_solid'] >= cfg['f_sat'] or cfg['r_tol_per_h'] < 0:
        raise ValueError('Invalid phase classification thresholds.')
    if cfg['bottom_boundary'] != 'adiabatic':
        raise NotImplementedError('The Category 1 benchmark has an adiabatic bottom.')
    if cfg['surface_boundary'] not in ('environment', 'zero_flux', 'prescribed_flux'):
        raise ValueError('Unknown surface boundary.')
    if cfg['wind_reference_height_m'] != 10:
        raise ValueError('Input winds are at 10 m; a height conversion needs a new model.')
    if cfg['weather_time_label'] not in ('start', 'midpoint', 'end'):
        raise ValueError('weather_time_label must be start, midpoint, or end.')
    if cfg['solver_max_iterations'] < 1:
        raise ValueError('solver_max_iterations must be positive.')
    for name in ('solver_atol_Wm2', 'solver_rtol', 'energy_atol_Jm2', 'energy_rtol'):
        if cfg[name] < 0 or not np.isfinite(cfg[name]):
            raise ValueError(f'Invalid {name}.')
    hh, mm = map(int, cfg['morning_checkpoint'].split(':'))
    if not (0 <= hh < 24 and mm == 0):
        raise ValueError('Use an exact hourly checkpoint, e.g. 06:00.')
    if not np.isfinite(cfg['initial_temperature_C']) or cfg['initial_temperature_C'] <= -273.15:
        raise ValueError('Invalid initial temperature.')
    return cfg

C1_APPROVED_PRELIM_SOURCE_IDENTITIES = [{'source_content_sha256': '261c4883b2b47ed321dc625b7ce523cbbfe21be5e17be75e22ea7e0aaf354b38', 'executable_source_sha256': '731a0d11adddf8f5d4924a18370e5350ed96ed5dd66c0884a071674d880f2f13', 'fingerprint_policy': 'Recorded computational source fingerprint; exact source retained privately.'}]

@njit
def c1_h_from_temperature(T, p, ref, plateau_fraction=0.0):
    """Specific enthalpy; sensible cp is integrated, latent heat is added once."""
    lo, hi, latent, cs, cl = (p[0], p[1], p[2], p[5], p[6])
    hlow = cs * (lo - ref)
    width = hi - lo
    if width == 0.0:
        if T < lo:
            return cs * (T - ref)
        if T == lo:
            return hlow + latent * plateau_fraction
        return hlow + latent + cl * (T - hi)
    if T <= lo:
        return cs * (T - ref)
    if T < hi:
        z = T - lo
        return hlow + cs * z + (cl - cs) * z * z / (2.0 * width) + latent * z / width
    hhigh = hlow + 0.5 * (cs + cl) * width + latent
    return hhigh + cl * (T - hi)

@njit
def c1_state_from_h_scalar(h, p, ref):
    """Invert h to temperature, liquid fraction, and derivative dT/dh."""
    lo, hi, latent, cs, cl = (p[0], p[1], p[2], p[5], p[6])
    width = hi - lo
    hlow = cs * (lo - ref)
    hhigh = hlow + 0.5 * (cs + cl) * width + latent
    if h < hlow:
        return (lo + (h - hlow) / cs, 0.0, 1.0 / cs)
    if h > hhigh:
        return (hi + (h - hhigh) / cl, 1.0, 1.0 / cl)
    if width == 0.0:
        if latent == 0.0:
            return (lo, 0.0, 1.0 / cs)
        return (lo, min(1.0, max(0.0, (h - hlow) / latent)), 0.0)
    y = h - hlow
    b = cs + latent / width
    a = (cl - cs) / (2.0 * width)
    z = 2.0 * y / (b + np.sqrt(max(0.0, b * b + 4.0 * a * y)))
    f = min(1.0, max(0.0, z / width))
    cp_app = cs + (cl - cs) * f + latent / width
    return (lo + z, f, 1.0 / cp_app)

@njit
def c1_cell_state(h, p, ref):
    n = len(h)
    T = np.empty(n)
    f = np.empty(n)
    dTdh = np.empty(n)
    k = np.empty(n)
    for i in range(n):
        T[i], f[i], dTdh[i] = c1_state_from_h_scalar(h[i], p, ref)
        k[i] = p[3] * (1.0 - f[i]) + p[4] * f[i]
    return (T, f, dTdh, k)

def c1_material_array(material, allow_sensible_only=False):
    row = dict(material)
    issues = []
    for c in C1_PCM_REQUIRED_NUMBERS:
        try:
            valid = c in row and np.isfinite(float(row[c]))
        except (TypeError, ValueError):
            valid = False
        if not valid:
            issues.append(c)
    if issues:
        raise ValueError(f'Missing/nonfinite material inputs: {issues}')
    for name in C1_PCM_REQUIRED_NUMBERS:
        row[name] = float(row[name])
    if not -273.15 < row['Tm_low_C'] <= row['Tm_high_C']:
        raise ValueError('Invalid melting bounds.')
    for c in C1_PCM_REQUIRED_NUMBERS[2:]:
        if row[c] <= 0 and (not (c == 'L_Jkg' and row[c] == 0 and allow_sensible_only)):
            raise ValueError(f'{c} must be positive.')
    if row.get('phase_model') not in ('equilibrium', 'equilibrium_screen'):
        raise NotImplementedError('This version implements reversible equilibrium screening only.')
    return np.array([row[c] for c in ['Tm_low_C', 'Tm_high_C', 'L_Jkg', 'k_solid_WmK', 'k_liquid_WmK', 'cp_solid_JkgK', 'cp_liquid_JkgK', 'rho_ref_kgm3']], dtype=float)

def c1_enthalpy_from_state(temperature_C, material, run_config=None, liquid_fraction=None):
    """Specific h in J/kg; explicit f is needed at an isothermal transition."""
    cfg = c1_make_run_config(run_config)
    p = c1_material_array(material, allow_sensible_only=True)
    temperatures = np.atleast_1d(np.asarray(temperature_C, dtype=float))
    if not np.isfinite(temperatures).all() or (temperatures <= -273.15).any():
        raise ValueError('Temperature must be finite and above absolute zero.')
    fractions = np.full(len(temperatures), cfg['initial_liquid_fraction_at_transition'])
    if liquid_fraction is not None:
        fractions = np.broadcast_to(np.asarray(liquid_fraction, dtype=float), temperatures.shape)
        if ((fractions < 0) | (fractions > 1)).any():
            raise ValueError('liquid_fraction must lie in [0,1].')
        if p[0] != p[1]:
            equilibrium_f = np.clip((temperatures - p[0]) / (p[1] - p[0]), 0.0, 1.0)
            if not np.allclose(fractions, equilibrium_f, atol=1e-10):
                raise ValueError('Explicit f disagrees with the selected equilibrium temperature curve.')
    return np.array([c1_h_from_temperature(t, p, cfg['T_ref_C'], f) for t, f in zip(temperatures, fractions)])

def c1_state_from_enthalpy(h_Jkg, material, run_config=None):
    cfg = c1_make_run_config(run_config)
    p = c1_material_array(material, allow_sensible_only=True)
    h = np.atleast_1d(np.asarray(h_Jkg, dtype=float))
    if not np.isfinite(h).all():
        raise ValueError('Cell enthalpies must be finite.')
    T, f, _, _ = c1_cell_state(h, p, cfg['T_ref_C'])
    if (T <= -273.15).any():
        raise ValueError('Enthalpy implies a temperature at/below absolute zero.')
    return {'T_C': T, 'h_Jkg': h.copy(), 'f': f, 'history': {'phase_model': 'equilibrium_screen', 'last_direction': np.zeros(len(h), dtype=np.int8)}}

@njit
def c1_convective_coefficient(wind10_ms):
    """Provisional 10 m screening relation, W/m²/K. Change this one function to test alternatives."""
    return 5.7 + 3.8 * wind10_ms

def c1_wind_speed(u10_ms, v10_ms):
    return np.hypot(u10_ms, v10_ms)

def c1_surface_energy_balance(T_surface_C, T_air_C, wind10_ms, G_Wm2, LW_down_Wm2, run_config=None):
    cfg = c1_make_run_config(run_config)
    hc = c1_convective_coefficient(wind10_ms)
    solar = cfg['alpha_s'] * G_Wm2
    convection = hc * (T_air_C - T_surface_C)
    longwave = cfg['epsilon_LW'] * (LW_down_Wm2 - C1_SIGMA_Wm2K4 * (T_surface_C + 273.15) ** 4)
    return {'q_solar_Wm2': solar, 'q_convection_Wm2': convection, 'q_longwave_Wm2': longwave, 'q_top_Wm2': solar + convection + longwave, 'hc_Wm2K': hc}

@njit
def c1_surface(Tcell, kcell, dx, air, wind, solar, lw, alpha, eps, mode, prescribed):
    """Coupled surface boundary. mode=0 environment, 1 zero flux, 2 prescribed."""
    hc = c1_convective_coefficient(wind)
    gc = 2.0 * kcell / dx
    if mode != 0:
        q = 0.0 if mode == 1 else prescribed
        return (Tcell + q / gc, q, 0.0, hc, 0.0)
    low = -273.149999
    high = max(100.0, max(Tcell, air)) + 100.0
    for _ in range(20):
        fhigh = alpha * solar + hc * (air - high) + eps * (lw - C1_SIGMA_Wm2K4 * (high + 273.15) ** 4) - gc * (high - Tcell)
        if fhigh <= 0.0:
            break
        high += 100.0
    ts = min(high, max(low, Tcell))
    for _ in range(50):
        rad = eps * (lw - C1_SIGMA_Wm2K4 * (ts + 273.15) ** 4)
        residual = alpha * solar + hc * (air - ts) + rad - gc * (ts - Tcell)
        if abs(residual) < 1e-09:
            break
        if residual > 0:
            low = ts
        else:
            high = ts
        deriv = -hc - 4.0 * eps * C1_SIGMA_Wm2K4 * (ts + 273.15) ** 3 - gc
        candidate = ts - residual / deriv
        ts = candidate if low < candidate < high else 0.5 * (low + high)
    q = gc * (ts - Tcell)
    he = hc + 4.0 * eps * C1_SIGMA_Wm2K4 * (ts + 273.15) ** 3
    effective_g = gc * he / (gc + he)
    mismatch = alpha * solar + hc * (air - ts) + eps * (lw - C1_SIGMA_Wm2K4 * (ts + 273.15) ** 4) - q
    return (ts, q, effective_g, hc, mismatch)

@njit
def c1_residual(h, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed):
    T, f, dTdh, k = c1_cell_state(h, p, ref)
    ts, q, ge, hc, mismatch = c1_surface(T[0], k[0], dx, air, wind, solar, lw, alpha, eps, mode, prescribed)
    faces = np.empty(len(h) - 1)
    r = rho * dx * (h - old_h) / dt
    r[0] -= q
    scale = max(1.0, abs(q))
    for i in range(len(h) - 1):
        faces[i] = 2.0 * k[i] * k[i + 1] / (k[i] + k[i + 1]) / dx
        flux = faces[i] * (T[i] - T[i + 1])
        r[i] += flux
        r[i + 1] -= flux
        scale = max(scale, abs(flux))
    return (r, dTdh, faces, ge, ts, q, hc, mismatch, scale)

@njit
def c1_tridiagonal(lower, diagonal, upper, rhs):
    n = len(diagonal)
    c = np.zeros(n)
    d = np.empty(n)
    result = np.empty(n)
    if n > 1:
        c[0] = upper[0] / diagonal[0]
    d[0] = rhs[0] / diagonal[0]
    for i in range(1, n):
        denominator = diagonal[i] - lower[i - 1] * c[i - 1]
        if i < n - 1:
            c[i] = upper[i] / denominator
        d[i] = (rhs[i] - lower[i - 1] * d[i - 1]) / denominator
    result[-1] = d[-1]
    for i in range(n - 2, -1, -1):
        result[i] = d[i] - c[i] * result[i + 1]
    return result

@njit
def c1_implicit_step(old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed, max_iterations, atol, rtol):
    """Backward Euler with the original line search and a phase-jump corrector.

    The zero-latent, zero-width law has a discontinuous conductivity when its
    equilibrium phase coordinate jumps. A monotone line search cannot traverse
    that jump. Only after the original search stagnates, try the full Newton
    crossing followed by one correction on its new branch. Accept the corrected
    trial only if the ORIGINAL residual decreases. Every accepted timestep must
    still meet the ORIGINAL convergence tolerance. No phase smoothing is used.
    The corrector consumes the same configured nonlinear iteration budget.
    """
    h = old_h.copy()
    n = len(h)
    iterations = 0
    while iterations < max_iterations:
        r, dTdh, faces, ge, ts, q, hc, mismatch, scale = c1_residual(h, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)
        norm = np.max(np.abs(r))
        if np.isfinite(norm) and norm <= atol + rtol * scale:
            return (h, ts, q, hc, True, iterations, norm, mismatch)
        diagonal = np.full(n, rho * dx / dt)
        lower = np.empty(n - 1)
        upper = np.empty(n - 1)
        diagonal[0] += ge * dTdh[0]
        for i in range(n - 1):
            diagonal[i] += faces[i] * dTdh[i]
            diagonal[i + 1] += faces[i] * dTdh[i + 1]
            upper[i] = -faces[i] * dTdh[i + 1]
            lower[i] = -faces[i] * dTdh[i]
        change = c1_tridiagonal(lower, diagonal, upper, -r)
        iterations += 1
        damping = 1.0
        improved = False
        for _ in range(18):
            trial = h + damping * change
            trial_r = c1_residual(trial, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)[0]
            trial_norm = np.max(np.abs(trial_r))
            if np.isfinite(trial_norm) and trial_norm < norm:
                h = trial
                improved = True
                break
            damping *= 0.5
        if not improved and iterations < max_iterations and (p[2] == 0.0) and (p[0] == p[1]) and (p[3] != p[4]):
            candidate = h + change
            T, f, _, _ = c1_cell_state(h, p, ref)
            trial_T, trial_f, _, _ = c1_cell_state(candidate, p, ref)
            if np.all(np.isfinite(candidate)) and np.all(trial_T > -273.15) and np.any(trial_f != f):
                cr, cdT, cfaces, cge, _, _, _, _, _ = c1_residual(candidate, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)
                diagonal = np.full(n, rho * dx / dt)
                diagonal[0] += cge * cdT[0]
                for i in range(n - 1):
                    diagonal[i] += cfaces[i] * cdT[i]
                    diagonal[i + 1] += cfaces[i] * cdT[i + 1]
                    upper[i] = -cfaces[i] * cdT[i + 1]
                    lower[i] = -cfaces[i] * cdT[i]
                corrected = candidate + c1_tridiagonal(lower, diagonal, upper, -cr)
                iterations += 1
                corrected_T = c1_cell_state(corrected, p, ref)[0]
                corrected_r = c1_residual(corrected, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)[0]
                corrected_norm = np.max(np.abs(corrected_r))
                if np.all(np.isfinite(corrected)) and np.all(corrected_T > -273.15) and np.isfinite(corrected_norm) and (corrected_norm < norm):
                    h = corrected
                    improved = True
        if not improved:
            break
    r, _, _, _, ts, q, hc, mismatch, scale = c1_residual(h, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)
    norm = np.max(np.abs(r))
    ok = np.isfinite(norm) and norm <= atol + rtol * scale
    return (h, ts, q, hc, ok, iterations, norm, mismatch)

@njit
def c1_classify(air, base, f, rate, fsat, fsolid, rtol):
    if air > base:
        if rate > rtol:
            return 0
        if f >= fsat:
            return 1
        if f <= fsolid:
            return 2
        return 3
    if rate < -rtol:
        return 4
    return 5

def c1_solve_pcm_timestep(state, material, weather_step, dt_s=None, run_config=None):
    """One reviewable update: boundary input minus output equals stored-energy change."""
    cfg = c1_make_run_config(run_config)
    p = c1_material_array(material, allow_sensible_only=True)
    old = np.asarray(state['h_Jkg'], dtype=float)
    n = len(old)
    dx = cfg['d_m'] / n
    dt = float(dt_s or cfg['dt_s'])
    if n != cfg['N_cells']:
        raise ValueError('State cell count does not match N_cells.')
    air = float(weather_step['T_air_C'])
    if 'wind10_ms' in weather_step:
        wind = float(weather_step['wind10_ms'])
    elif 'u10_ms' in weather_step and 'v10_ms' in weather_step:
        wind = float(np.hypot(weather_step['u10_ms'], weather_step['v10_ms']))
    else:
        raise ValueError('Supply wind10_ms or both u10_ms and v10_ms; missing wind is not zero.')
    args = [air, wind, float(weather_step['G_Wm2']), float(weather_step['LW_down_Wm2'])]
    if not np.isfinite(old).all() or not np.isfinite(args).all() or dt <= 0:
        raise ValueError('Nonfinite state/weather or invalid timestep.')
    if air <= -273.15 or wind < 0 or args[2] < 0 or (args[3] < 0):
        raise ValueError('Temperature, wind or radiation outside physical bounds.')
    if (c1_cell_state(old, p, cfg['T_ref_C'])[0] <= -273.15).any():
        raise ValueError('Initial enthalpy implies a temperature at/below absolute zero.')
    mode = {'environment': 0, 'zero_flux': 1, 'prescribed_flux': 2}[cfg['surface_boundary']]
    old_f = c1_cell_state(old, p, cfg['T_ref_C'])[1]
    special = p[0] == p[1] and p[2] == 0.0
    if special and 'f' in state:
        supplied = np.asarray(state['f'], dtype=float)
        boundary = old == p[5] * (p[0] - cfg['T_ref_C'])
        c1_require(supplied.shape == old.shape and np.isfinite(supplied).all() and ((supplied >= 0) & (supplied <= 1)).all(), 'Invalid coexistence state fractions.')
        c1_require(np.array_equal(supplied[~boundary], old_f[~boundary]), 'Fraction disagrees with the pure-phase law.')
        old_f = supplied.copy()
    if special:
        h, new_f, ts, q, hc, ok, iterations, residual, mismatch, used_graph, null_count = c1_zero_latent_coexistence_step(old, old_f, p, cfg['T_ref_C'], p[7], dx, dt, *args, cfg['alpha_s'], cfg['epsilon_LW'], mode, cfg['prescribed_top_flux_Wm2'], cfg['solver_max_iterations'], cfg['solver_atol_Wm2'], cfg['solver_rtol'])
    else:
        h, ts, q, hc, ok, iterations, residual, mismatch = c1_implicit_step(old, p, cfg['T_ref_C'], p[7], dx, dt, *args, cfg['alpha_s'], cfg['epsilon_LW'], mode, cfg['prescribed_top_flux_Wm2'], cfg['solver_max_iterations'], cfg['solver_atol_Wm2'], cfg['solver_rtol'])
        used_graph = 0
        null_count = 0
    if not ok:
        raise RuntimeError(f'Implicit timestep did not converge: max residual {residual:g} W/m².')
    new = c1_state_from_enthalpy(h, material, cfg)
    if special:
        new['f'] = new_f.copy()
    new['history'] = dict(state.get('history', {}))
    new['history'].update(phase_model='equilibrium_screen', last_direction=np.sign(new['f'] - old_f).astype(np.int8))
    stored = p[7] * dx * np.sum(h - old)
    return (new, {'T_surface_C': ts, 'q_top_Wm2': q, 'hc_Wm2K': hc, 'stored_energy_change_Jm2': stored, 'boundary_energy_Jm2': q * dt, 'energy_residual_Jm2': stored - q * dt, 'iterations': iterations, 'surface_balance_residual_Wm2': mismatch, 'coexistence_used': used_graph, 'passive_fraction_nullity': null_count})

C1_HOURLY_NUMERIC_COLUMNS = ('T_surface_C', 'T_surface_end_C', 'T_PCM_C', 'T_PCM_end_C', 'h_mean_Jkg', 'H_area_start_Jm2', 'H_area_Jm2', 'q_top_mean_Wm2', 'q_solar_mean_Wm2', 'q_convection_mean_Wm2', 'q_longwave_mean_Wm2', 'hc_Wm2K', 'wind_mean_ms', 'f_start', 'f_end', 'f_min', 'f_max', 'f_mean', 'latent_uptake_Jm2', 'latent_release_Jm2', 'cooling_relevant_latent_uptake_Jm2', 'cell_latent_uptake_Jm2', 'cell_latent_release_Jm2', 'active_h', 'saturated_h', 'unactivated_h', 'other_no_uptake_h', 'regenerating_h', 'idle_h', 'cooling_relevant_h', 'CDH_C_h', 'energy_residual_Jm2', 'energy_relative_residual', 'solver_retry_count', 'solver_max_iterations_used', 'minimum_accepted_dt_s', 'surface_balance_max_residual_Wm2', 'q_top_positive_energy_Jm2', 'q_top_negative_energy_Jm2')

@njit
def c1_instantaneous_weather(values, hour_index, elapsed_s):
    """Interpolate instantaneous values at interval midpoints, hold coverage edges."""
    x = hour_index + elapsed_s / 3600.0 - 0.5
    if x <= 0:
        return values[0]
    if x >= len(values) - 1:
        return values[-1]
    left = int(x)
    fraction = x - left
    return values[left] * (1.0 - fraction) + values[left + 1] * fraction

@njit
def c1_coexistence_state(z, p, ref):
    n = len(z)
    T = np.empty(n)
    h = np.empty(n)
    f = np.empty(n)
    dT = np.empty(n)
    dh = np.empty(n)
    dk = np.empty(n)
    k = np.empty(n)
    hl = p[5] * (p[0] - ref)
    for j in range(n):
        if z[j] < 0.0:
            T[j] = p[0] + z[j]
            h[j] = hl + p[5] * z[j]
            f[j] = 0.0
            dT[j] = 1.0
            dh[j] = p[5]
            dk[j] = 0.0
        elif z[j] > 1.0:
            T[j] = p[0] + z[j] - 1.0
            h[j] = hl + p[2] + p[6] * (z[j] - 1.0)
            f[j] = 1.0
            dT[j] = 1.0
            dh[j] = p[6]
            dk[j] = 0.0
        else:
            T[j] = p[0]
            h[j] = hl + p[2] * z[j]
            f[j] = z[j]
            dT[j] = 0.0
            dh[j] = p[2]
            dk[j] = p[4] - p[3]
        k[j] = p[3] * (1.0 - f[j]) + p[4] * f[j]
    return (T, h, f, dT, dh, dk, k)

@njit
def c1_coexistence_residual(z, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed):
    T, h, f, dT, dh, dk, k = c1_coexistence_state(z, p, ref)
    ts, q, ge, hc, mismatch = c1_surface(T[0], k[0], dx, air, wind, solar, lw, alpha, eps, mode, prescribed)
    r = rho * dx * (h - old_h) / dt
    J = np.zeros((len(z), len(z)))
    for j in range(len(z)):
        J[j, j] = rho * dx / dt * dh[j]
    r[0] -= q
    if mode == 0:
        he = hc + 4 * eps * C1_SIGMA_Wm2K4 * (ts + 273.15) ** 3
        gc = 2 * k[0] / dx
        J[0, 0] += ge * dT[0] - he / (he + gc) * 2 / dx * dk[0] * (ts - T[0])
    scale = max(1.0, abs(q))
    for j in range(len(z) - 1):
        total = k[j] + k[j + 1]
        g = 2 * k[j] * k[j + 1] / total / dx
        difference = T[j] - T[j + 1]
        flux = g * difference
        r[j] += flux
        r[j + 1] -= flux
        left = g * dT[j] + 2 * k[j + 1] ** 2 / (total ** 2 * dx) * dk[j] * difference
        right = -g * dT[j + 1] + 2 * k[j] ** 2 / (total ** 2 * dx) * dk[j + 1] * difference
        J[j, j] += left
        J[j, j + 1] += right
        J[j + 1, j] -= left
        J[j + 1, j + 1] -= right
        scale = max(scale, abs(flux))
    return (r, J, T, h, f, ts, q, hc, mismatch, scale)

@njit
def c1_coexistence_nullity(J):
    singular_values = np.linalg.svd(J)[1]
    rank = 0
    if len(singular_values) and singular_values[0] > 0.0:
        rank = np.sum(singular_values > singular_values[0] * 1e-14)
    zero_columns = 0
    for j in range(J.shape[1]):
        if np.max(np.abs(J[:, j])) == 0.0:
            zero_columns += 1
    return (J.shape[1] - rank, zero_columns)

@njit
def c1_zero_latent_coexistence_step(old_h, old_f, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed, max_iterations, atol, rtol):
    assert p[0] == p[1] and p[2] == 0.0
    ordinary = c1_implicit_step(old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed, max_iterations, atol, rtol)
    hl = p[5] * (p[0] - ref)
    if ordinary[4]:
        T, f, _, _ = c1_cell_state(ordinary[0], p, ref)
        if not np.any(ordinary[0] == hl):
            return (ordinary[0], f, ordinary[1], ordinary[2], ordinary[3], True, ordinary[5], ordinary[6], ordinary[7], 0, 0)
        for j in range(len(f)):
            if ordinary[0][j] == hl and old_h[j] == hl:
                f[j] = old_f[j]
        z = np.empty(len(old_h))
        for j in range(len(z)):
            z[j] = (ordinary[0][j] - hl) / p[5] if ordinary[0][j] < hl else 1 + (ordinary[0][j] - hl) / p[6] if ordinary[0][j] > hl else f[j]
        rr = c1_coexistence_residual(z, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)
        norm = np.max(np.abs(rr[0]))
        if norm <= atol + rtol * rr[9]:
            nullity, zero_columns = c1_coexistence_nullity(rr[1])
            if nullity == zero_columns:
                return (rr[3], rr[4], rr[5], rr[6], rr[7], True, ordinary[5], norm, rr[8], 0, nullity)
    z = np.empty(len(old_h))
    for j in range(len(z)):
        z[j] = (old_h[j] - hl) / p[5] if old_h[j] < hl else 1 + (old_h[j] - hl) / p[6] if old_h[j] > hl else old_f[j]
    iterations = 0
    null_fraction_count = 0
    for iterations in range(max_iterations):
        r, J, T, h, f, ts, q, hc, mismatch, scale = c1_coexistence_residual(z, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)
        norm = np.max(np.abs(r))
        if np.isfinite(norm) and norm <= atol + rtol * scale:
            nullity, zero_columns = c1_coexistence_nullity(J)
            return (h, f, ts, q, hc, nullity == zero_columns, iterations, norm, mismatch, 1, nullity)
        delta = np.linalg.lstsq(J, -r, rcond=1e-14)[0]
        null_fraction_count = 0
        for j in range(len(z)):
            if np.max(np.abs(J[:, j])) == 0.0:
                delta[j] = 0.0
                null_fraction_count += 1
        improved = False
        damping = 1.0
        for _ in range(24):
            trial = z + damping * delta
            value = c1_coexistence_residual(trial, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)
            trial_norm = np.max(np.abs(value[0]))
            if np.all(value[2] > -273.15) and np.isfinite(trial_norm) and (trial_norm < norm):
                z = trial
                improved = True
                break
            damping *= 0.5
        if not improved:
            break
    r, J, T, h, f, ts, q, hc, mismatch, scale = c1_coexistence_residual(z, old_h, p, ref, rho, dx, dt, air, wind, solar, lw, alpha, eps, mode, prescribed)
    norm = np.max(np.abs(r))
    ok = np.isfinite(norm) and norm <= atol + rtol * scale
    nullity, zero_columns = c1_coexistence_nullity(J)
    return (h, f, ts, q, hc, ok and nullity == zero_columns, iterations + 1, norm, mismatch, 1, nullity)

@njit
def c1_integrate_hours(air_values, u_values, v_values, G, LW, initial_h, p, ref, thickness, nominal_dt, min_dt, alpha, eps, mode, prescribed, max_iterations, atol, rtol, base, fsat, fsolid, rate_tolerance, initial_fraction):
    hours = len(G)
    n = len(initial_h)
    dx = thickness / n
    rho = p[7]
    capacity = rho * thickness * p[2]
    h = initial_h.copy()
    master = np.zeros((hours, 45))
    snapshots = np.empty((hours + 1, n))
    last_direction = np.zeros(n, dtype=np.int8)
    snapshots[0, :] = h
    fraction_snapshots = np.empty((hours + 1, n))
    fraction_snapshots[0, :] = initial_fraction
    numerical_diagnostics = np.zeros((hours, 2))
    f = initial_fraction.copy()
    coexistence_case = p[0] == p[1] and p[2] == 0.0
    for hour in range(hours):
        T, derived_f, _, _ = c1_cell_state(h, p, ref)
        if not coexistence_case:
            f = derived_f
        fbar = np.mean(f)
        master[hour, 5] = rho * dx * np.sum(h)
        master[hour, 13] = fbar
        master[hour, 15] = fbar
        master[hour, 16] = fbar
        master[hour, 35] = nominal_dt
        master[hour, 43] = np.nan
        master[hour, 44] = np.nan
        elapsed = 0.0
        while elapsed < 3600.0 - 1e-09:
            dt = min(nominal_dt, 3600.0 - elapsed)
            while True:
                air = c1_instantaneous_weather(air_values, hour, elapsed + 0.5 * dt)
                u = c1_instantaneous_weather(u_values, hour, elapsed + 0.5 * dt)
                v = c1_instantaneous_weather(v_values, hour, elapsed + 0.5 * dt)
                wind = np.sqrt(u * u + v * v)
                if coexistence_case:
                    new_h, step_fraction, ts, q, hc, ok, iterations, residual, mismatch, used_graph, null_count = c1_zero_latent_coexistence_step(h, f, p, ref, rho, dx, dt, air, wind, G[hour], LW[hour], alpha, eps, mode, prescribed, max_iterations, atol, rtol)
                    if ok:
                        numerical_diagnostics[hour, 0] += used_graph
                        numerical_diagnostics[hour, 1] += null_count
                else:
                    new_h, ts, q, hc, ok, iterations, residual, mismatch = c1_implicit_step(h, p, ref, rho, dx, dt, air, wind, G[hour], LW[hour], alpha, eps, mode, prescribed, max_iterations, atol, rtol)
                if ok:
                    break
                master[hour, 33] += 1
                dt *= 0.5
                if dt < min_dt - 1e-10:
                    return (master, snapshots, h, last_direction, hour, elapsed, residual, fraction_snapshots, numerical_diagnostics)
            new_T, new_f, _, _ = c1_cell_state(new_h, p, ref)
            if coexistence_case:
                new_f = step_fraction
            new_fbar = np.mean(new_f)
            dfbar = new_fbar - fbar
            uptake = capacity * max(dfbar, 0.0)
            release = capacity * max(-dfbar, 0.0)
            master[hour, 18] += uptake
            master[hour, 19] += release
            if air > base:
                master[hour, 20] += uptake
                master[hour, 29] += dt / 3600.0
            for i in range(n):
                local_delta = rho * dx * p[2] * (new_f[i] - f[i])
                last_direction[i] = 1 if new_f[i] > f[i] else -1 if new_f[i] < f[i] else 0
                master[hour, 21] += max(local_delta, 0.0)
                master[hour, 22] += max(-local_delta, 0.0)
            master[hour, 39] = max(master[hour, 39], residual)
            step_energy_error = rho * dx * np.sum(new_h - h) - q * dt
            master[hour, 40] = max(master[hour, 40], abs(step_energy_error))
            if new_fbar >= fsat:
                master[hour, 41] += dt / 3600.0
                if fbar < fsat:
                    master[hour, 42] += 1
                if np.isnan(master[hour, 43]):
                    master[hour, 43] = elapsed + dt
                master[hour, 44] = elapsed + dt
            rate = dfbar / (dt / 3600.0)
            state = c1_classify(air, base, new_fbar, rate, fsat, fsolid, rate_tolerance)
            master[hour, 23 + state] += dt / 3600.0
            master[hour, 30] += max(air - base, 0.0) * dt / 3600.0
            weight = dt / 3600.0
            master[hour, 0] += ts * weight
            master[hour, 1] = ts
            master[hour, 2] += 0.5 * (np.mean(T) + np.mean(new_T)) * weight
            master[hour, 3] = np.mean(new_T)
            master[hour, 7] += q * weight
            if mode == 0:
                master[hour, 8] += alpha * G[hour] * weight
                master[hour, 9] += hc * (air - ts) * weight
                master[hour, 10] += eps * (LW[hour] - C1_SIGMA_Wm2K4 * (ts + 273.15) ** 4) * weight
            master[hour, 11] += hc * weight
            master[hour, 12] += wind * weight
            master[hour, 15] = min(master[hour, 15], new_fbar)
            master[hour, 16] = max(master[hour, 16], new_fbar)
            master[hour, 17] += 0.5 * (fbar + new_fbar) * weight
            master[hour, 34] = max(master[hour, 34], iterations)
            master[hour, 35] = min(master[hour, 35], dt)
            master[hour, 36] = max(master[hour, 36], abs(mismatch))
            master[hour, 37] += max(q, 0.0) * dt
            master[hour, 38] += max(-q, 0.0) * dt
            h = new_h
            T = new_T
            f = new_f
            fbar = new_fbar
            elapsed += dt
        master[hour, 4] = np.mean(h)
        master[hour, 6] = rho * dx * np.sum(h)
        master[hour, 14] = fbar
        error = master[hour, 6] - master[hour, 5] - 3600.0 * master[hour, 7]
        master[hour, 31] = error
        master[hour, 32] = abs(error) / max(1.0, master[hour, 37] + master[hour, 38])
        snapshots[hour + 1, :] = h
        fraction_snapshots[hour + 1, :] = f
        if mode != 0:
            for c in range(8, 11):
                master[hour, c] = np.nan
    return (master, snapshots, h, last_direction, -1, 0.0, 0.0, fraction_snapshots, numerical_diagnostics)

class Category1InputError(ValueError):
    """Required real Preliminary input or identity is unavailable/invalid."""

def c1_require(condition, message):
    if not bool(condition):
        raise Category1InputError(message)

def c1_sha256_file(path):
    path = Path(path)
    c1_require(path.is_file(), f'Required input file missing: {path}')
    return sha256_file(path)

def c1_json_value(value):
    if isinstance(value, dict):
        return {str(k): c1_json_value(v) for k, v in value.items()}
    if isinstance(value, (set, frozenset)):
        return [c1_json_value(v) for v in sorted(value, key=str)]
    if isinstance(value, (list, tuple, np.ndarray)):
        return [c1_json_value(v) for v in value]
    if isinstance(value, (np.integer, np.floating, np.bool_)):
        value = value.item()
    if value is pd.NA or value is pd.NaT or value is None:
        return None
    if isinstance(value, float) and (not np.isfinite(value)):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)

def c1_digest(value):
    return hashlib.sha256(json.dumps(c1_json_value(value), sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()

def c1_table_digest(frame):
    return c1_digest({'columns': list(frame.columns), 'records': frame.to_dict('records'), 'attrs': frame.attrs})

def c1_explicit_utc(values, label):
    try:
        c1_require(all((pd.notna(v) and pd.Timestamp(v).tzinfo is not None for v in values)), f'{label}: timestamps require explicit UTC/local timezone offsets.')
        result = pd.DatetimeIndex(pd.to_datetime(values, utc=True, errors='raise')).as_unit('ns')
    except (TypeError, ValueError, OverflowError) as error:
        raise Category1InputError(f'{label}: invalid/unlabelled timestamps.') from error
    c1_require(not result.hasnans, f'{label}: missing timestamp.')
    return result

def c1_require_columns(frame, columns, label):
    c1_require(isinstance(frame, pd.DataFrame), f'{label}: require a DataFrame.')
    c1_require(not frame.empty, f'{label} is empty. Preliminary requires review.')
    c1_require(not frame.columns.duplicated().any(), f'{label}: duplicate headers.')
    c1_require(set(columns) <= set(frame.columns), f'{label}: missing columns {sorted(set(columns) - set(frame.columns))}.')

def c1_numeric(frame, columns, label):
    for name in columns:
        c1_require(not frame[name].map(lambda v: isinstance(v, (bool, np.bool_))).any(), f'{label}: boolean physical value in {name}.')
        try:
            x = pd.to_numeric(frame[name], errors='raise').to_numpy(dtype=float)
        except (TypeError, ValueError) as error:
            raise Category1InputError(f'{label}: nonnumeric {name}.') from error
        c1_require(np.isfinite(x).all(), f'{label}: missing/nonfinite {name}; no filling is allowed.')

def c1_require_real(metadata, label):
    c1_require(isinstance(metadata, dict) and metadata.get('data_nature') == 'REAL', f'{label}: explicit REAL provenance required; mechanical/demo inputs are refused.')

    def walk(value):
        if isinstance(value, dict):
            for k, v in value.items():
                if k in {'data_nature', 'scientific_reality_status', 'data_mode'}:
                    c1_require(v == 'REAL', f'{label}: non-real {k} in lineage.')
                if k in {'synthetic', 'demo', 'test_only', 'fixture', 'is_synthetic', 'is_demo'}:
                    c1_require(v is False or v is None, f'{label}: test/demo lineage marker {k}.')
                if k in {'source', 'origin', 'lineage', 'status'} and isinstance(v, str):
                    c1_require(not re.search('(?i)(?:^|[ _:/-])(synthetic|demo|fixture|mechanical_test)(?:$|[ _:/-])', v), f'{label}: forbidden scientific lineage {v!r}.')
                walk(v)
        elif isinstance(value, (list, tuple)):
            for v in value:
                walk(v)
    walk(metadata)

C1_WEATHER_COLUMNS = ['interval_start_UTC', 'interval_end_UTC', 'interval_start_local', 'interval_end_local', 'timestamp_UTC', 'timestamp_local', 'instantaneous_sample_UTC', 'instantaneous_sample_local', 'T_air_C', 'u10_ms', 'v10_ms', 'wind10_ms', 'G_Wm2', 'LW_down_Wm2', 'shortwave_energy_Jm2', 'longwave_energy_Jm2', 'shortwave_negative_flag']

C1_WEATHER_UNITS = {'T_air_C': 'degC', 'u10_ms': 'm/s', 'v10_ms': 'm/s', 'wind10_ms': 'm/s', 'G_Wm2': 'W/m2 interval mean', 'LW_down_Wm2': 'W/m2 interval mean', 'shortwave_energy_Jm2': 'J/m2 target interval', 'longwave_energy_Jm2': 'J/m2 target interval'}

C1_MATERIAL_UNITS = {'Tm_low_C': 'degC', 'Tm_high_C': 'degC', 'L_Jkg': 'J/kg', 'k_solid_WmK': 'W/m/K', 'k_liquid_WmK': 'W/m/K', 'cp_solid_JkgK': 'J/kg/K', 'cp_liquid_JkgK': 'J/kg/K', 'rho_ref_kgm3': 'kg/m3'}

def c1_weather_shape(frame, metadata):
    """Structural checks only. This function never attests scientific reality."""
    c1_require_columns(frame, C1_WEATHER_COLUMNS, 'ANNUAL_WEATHER')
    c1_require(metadata.get('study_year') == 2025, 'Weather metadata must declare study_year=2025.')
    semantics = metadata.get('time_semantics', metadata)
    c1_require(semantics.get('timezone') == C1_IST, 'Weather timezone must be Asia/Kolkata.')
    label = semantics.get('timestamp_label')
    c1_require(label in {'start', 'midpoint', 'end'}, 'Declare timestamp_label=start/midpoint/end.')
    c1_require(semantics.get('instantaneous_sample') == 'midpoint', 'Air and wind require midpoint sample semantics.')
    units = metadata.get('units', semantics.get('units'))
    c1_require(isinstance(units, dict), 'Weather requires explicit per-column units.')
    for name, unit in C1_WEATHER_UNITS.items():
        c1_require(units.get(name) == unit, f'Weather unit mismatch for {name}: expected {unit}.')
    expected = pd.date_range(pd.Timestamp('2025-01-01', tz=C1_IST), pd.Timestamp('2026-01-01', tz=C1_IST), freq='h', inclusive='left').tz_convert('UTC')
    c1_require(len(frame) == 8760, f'Expected exactly 8760 target intervals, received {len(frame)}.')
    starts = c1_explicit_utc(frame.interval_start_UTC, 'interval_start_UTC')
    c1_require(not starts.has_duplicates and starts.is_monotonic_increasing and starts.equals(expected), 'Weather hours missing/duplicated/out of order/wrong study calendar; no sorting or filling allowed.')
    ends = expected + pd.Timedelta(hours=1)
    mids = expected + pd.Timedelta(minutes=30)
    labels = expected + pd.Timedelta(minutes={'start': 0, 'midpoint': 30, 'end': 60}[label])
    for name, reference in [('interval_end_UTC', ends), ('interval_start_local', expected), ('interval_end_local', ends), ('timestamp_UTC', labels), ('timestamp_local', labels), ('instantaneous_sample_UTC', mids), ('instantaneous_sample_local', mids)]:
        c1_require(c1_explicit_utc(frame[name], name).equals(reference), f'Weather interval/sample mismatch: {name}.')
        if name.endswith('_local'):
            c1_require(all((pd.Timestamp(v).utcoffset().total_seconds() == 19800 for v in frame[name])), f'{name}: local timestamps must carry the Asia/Kolkata +05:30 offset.')
    c1_numeric(frame, C1_WEATHER_UNITS, 'ANNUAL_WEATHER')
    c1_require((frame.T_air_C > -273.15).all(), 'Invalid air temperature.')
    c1_require((frame.G_Wm2 >= 0).all(), 'Negative shortwave input: upstream review required; no clipping in Category 1.')
    c1_require((frame.LW_down_Wm2 > 0).all(), 'Nonpositive downward longwave input.')
    c1_require(np.allclose(frame.wind10_ms, np.hypot(frame.u10_ms, frame.v10_ms), rtol=1e-08, atol=1e-08), 'wind10_ms disagrees with 10 m wind components.')
    for flux, energy in [('G_Wm2', 'shortwave_energy_Jm2'), ('LW_down_Wm2', 'longwave_energy_Jm2')]:
        c1_require(np.allclose(frame[flux].to_numpy(float) * 3600, frame[energy].to_numpy(float), rtol=1e-12, atol=1e-08), f'{flux}: hourly radiation does not conserve the provided interval energy.')
    c1_require(frame.shortwave_negative_flag.map(lambda v: isinstance(v, (bool, np.bool_))).all(), 'shortwave_negative_flag must be explicit boolean values.')
    c1_require(np.array_equal(frame.shortwave_negative_flag.to_numpy(bool), frame.G_Wm2.to_numpy() < 0), 'shortwave_negative_flag disagrees with radiation values.')
    return {'status': 'PASS', 'hours': 8760, 'study_year': 2025, 'start_UTC': expected[0].isoformat(), 'end_UTC_exclusive': ends[-1].isoformat(), 'timestamp_label': label, 'solver_label': 'instantaneous_sample_UTC at midpoint'}

def c1_validate_weather(frame, metadata):
    c1_require_real(metadata, 'ANNUAL_WEATHER metadata')
    if frame.attrs:
        c1_require_real(frame.attrs, 'ANNUAL_WEATHER attributes')
    c1_require(metadata.get('source') == 'ECMWF/ERA5_LAND/HOURLY', 'Only Preliminary ERA5-Land weather is accepted.')
    support = metadata.get('spatial_support', metadata.get('time_semantics', {}).get('spatial_support'))
    c1_require(isinstance(support, dict), 'Hotspot-representative weather support metadata required.')
    c1_require(support.get('spatial_scope') == 'SELECTED_HOTSPOT_NATIVE_ERA5_CELL' and support.get('weather_resampled_to_common_grid') is False and (support.get('native_cells_per_hour') == 1), "Weather must describe the selected hotspot's native ERA5-Land cell.")
    location = support.get('hotspot_location', {})
    xy = location.get('point_lon_lat')
    c1_require(isinstance(xy, (list, tuple)) and len(xy) == 2 and np.isfinite(np.asarray(xy, float)).all(), 'Hotspot extraction coordinate missing.')
    c1_require(location.get('selected_hotspot_cells', 0) > 0 and bool(location.get('hotspot_mask_numeric_sha256')), 'Weather support lacks selected-hotspot identity.')
    return c1_weather_shape(frame, metadata)

def c1_inspect_workbook(io_config):
    path = Path(io_config['workbook_path']).expanduser()
    if not path.is_absolute():
        path = C1_PROJECT_ROOT / path
    digest = c1_sha256_file(path)
    c1_require(digest == io_config['expected_workbook_sha256'], 'PCM workbook does not match Section-0 pinned SHA-256.')
    literal = inspect_pcm_workbook(path, io_config['workbook_sheet'])
    c1_require(literal['headers'] == C1_PCM_COLUMNS, 'Workbook literal canonical schema differs from the inspected schema.')
    rows = pd.read_excel(path, sheet_name=io_config['workbook_sheet'])
    c1_require(rows.pcm_id.notna().all() and rows.pcm_id.is_unique, 'Workbook IDs missing or duplicate.')
    runnable = 0
    for _, row in rows.iterrows():
        try:
            c1_numeric(pd.DataFrame([row]), C1_PCM_REQUIRED_NUMBERS, 'workbook inspection')
            c1_material_array(row)
            runnable += 1
        except (Category1InputError, ValueError, NotImplementedError):
            pass
    c1_require(c1_sha256_file(path) == digest, 'Workbook changed during inspection.')
    return {'rows': rows, 'receipt': {'path': str(path.resolve()), 'sha256': digest, 'row_count': len(rows), 'complete_equilibrium_model_rows': runnable, 'shortlist_created': False, 'property_values_filled': False, 'physical_source_quality': 'NOT CHECKED'}}

def c1_candidate_shape(frame, ids, metadata, screen, workbook):
    c1_require_columns(frame, C1_PCM_COLUMNS + ['screen_lower_C', 'screen_upper_C', 'review_status'], 'PCM_PRELIM_CANDIDATES')
    c1_require(isinstance(ids, (list, tuple)) and len(ids) > 0, 'Missing/empty PCM_PRELIM_CANDIDATE_IDS; Preliminary requires review.')
    c1_require(all((isinstance(v, str) and v.strip() == v and v for v in ids)), 'Candidate IDs must be literal nonempty strings.')
    c1_require(frame.pcm_id.notna().all() and frame.pcm_id.is_unique and (list(frame.pcm_id) == list(ids)), 'Candidate IDs must match the ordered Preliminary shortlist exactly.')
    pinned = workbook['receipt']['sha256']
    c1_require(metadata.get('workbook_sha256') == pinned and screen.get('workbook_sha256') == pinned, 'Candidate provenance must match the pinned workbook.')
    c1_require(screen.get('passing_rows') == len(ids), 'Screen passing-row count disagrees with candidate IDs.')
    lower, upper = (screen.get('lower_C'), screen.get('upper_C'))
    c1_require(lower is not None and upper is not None and np.isfinite([lower, upper]).all() and (lower <= upper), 'Missing/reversed Preliminary transition-temperature screen bounds.')
    c1_require(screen.get('screen_method') == CUMULATIVE_SCREEN_METHOD and screen.get('pass_equation') == CUMULATIVE_SCREEN_EQUATION and (screen.get('boundary_contact_passes') is True), 'Category 1 requires the new inclusive-overlap screen lineage.')
    c1_numeric(frame, C1_PCM_REQUIRED_NUMBERS + ['screen_lower_C', 'screen_upper_C'], 'PCM candidates')
    known = workbook['rows'].set_index('pcm_id')
    for _, row in frame.iterrows():
        c1_require(row.pcm_id in known.index, f'Candidate {row.pcm_id} absent from the pinned workbook.')
        c1_require(row.screen_lower_C == lower and row.screen_upper_C == upper and (row.Tm_low_C <= row.Tm_high_C) and (row.Tm_high_C >= lower) and (row.Tm_low_C <= upper), f'{row.pcm_id}: Preliminary screen mismatch.')
        c1_require(row.review_status == screen.get('review_status'), f'{row.pcm_id}: Preliminary review lineage mismatch.')
        for name in C1_PCM_COLUMNS[1:]:
            original, supplied = (known.loc[row.pcm_id, name], row[name])
            if pd.isna(original) and pd.isna(supplied):
                continue
            if isinstance(original, (int, float, np.number)) and (not isinstance(original, (bool, np.bool_))):
                try:
                    same = np.isclose(float(original), float(supplied), rtol=1e-12, atol=1e-12)
                except (TypeError, ValueError):
                    same = False
            else:
                same = c1_json_value(original) == c1_json_value(supplied)
            c1_require(same, f'{row.pcm_id}: {name} changed/filled relative to the literal workbook.')
        c1_material_array(row)
        if 'section0_readiness' in frame:
            c1_require(row.section0_readiness == 'RUNNABLE_SCHEMA_ONLY', f'{row.pcm_id}: Section-0 row not runnable.')
    return {'status': 'PASS', 'candidate_ids': list(ids), 'workbook_sha256': pinned, 'physical_source_quality': 'NOT CHECKED', 'preliminary_review_status': screen.get('review_status')}

def c1_validate_candidates(frame, ids, metadata, screen, workbook):
    c1_require_real(metadata, 'PCM candidates')
    if frame.attrs.get('data_nature') is not None:
        c1_require_real(frame.attrs, 'PCM candidate attributes')
    return c1_candidate_shape(frame, ids, metadata, screen, workbook)

def c1_read_preliminary_directory(directory):
    folder = Path(directory).expanduser().resolve()
    manifest = json.loads((folder / 'preliminary_manifest.json').read_text())
    c1_require(manifest.get('status') in {'PRELIMINARY_COMPLETE_REVIEW_REQUIRED', 'PRELIMINARY_COMPLETE'}, 'Preliminary manifest is incomplete or invalid.')
    c1_require(manifest.get('global_config', {}).get('data_mode') == 'REAL', 'Preliminary manifest is not real-data mode.')
    c1_require(manifest.get('screen_method') == CUMULATIVE_SCREEN_METHOD, 'Refuse previous strict-screen lineage.')
    c1_require(_source_identity_matches(manifest.get('cumulative_source_identity'), CUMULATIVE_SOURCE_IDENTITY) or any(_source_identity_matches(manifest.get('cumulative_source_identity'), approved) for approved in C1_APPROVED_PRELIM_SOURCE_IDENTITIES), 'Preliminary code identity differs; rerun scientific pipeline.')
    records = manifest.get('artifacts', {})
    names = ['annual_weather.csv', 'annual_weather.csv.metadata.json', 'pcm_prelim_candidates.csv', 'pcm_prelim_candidates.csv.metadata.json', 'pcm_prelim_candidate_ids.json', 'pcm_screen_derivation.json']
    files = {}
    for name in names:
        c1_require(name in records, f'Manifest missing artifact identity: {name}.')
        record = records[name]
        c1_require_real(record, name)
        c1_require(record.get('run_id') == manifest.get('run_id') and bool(record.get('sha256')), f'{name}: run/hash mismatch.')
        c1_require(c1_sha256_file(folder / name) == record['sha256'], f'{name}: stale/changed artifact SHA-256.')
        files[name] = record['sha256']
    weather_meta = json.loads((folder / 'annual_weather.csv.metadata.json').read_text())
    candidates_meta = json.loads((folder / 'pcm_prelim_candidates.csv.metadata.json').read_text())
    screen = json.loads((folder / 'pcm_screen_derivation.json').read_text())
    for meta in (weather_meta, candidates_meta):
        c1_require(meta.get('run_id') == manifest['run_id'], 'Preliminary sidecar from a different run.')
    candidates_meta['workbook_sha256'] = screen.get('workbook_sha256')
    weather = pd.read_csv(folder / 'annual_weather.csv', float_precision='round_trip')
    schema = candidates_meta.get('candidate_table_schema')
    c1_require(isinstance(schema, dict), 'Candidate CSV requires explicit empty-text/missing-value schema.')
    candidates = pd.read_csv(folder / 'pcm_prelim_candidates.csv', float_precision='round_trip', keep_default_na=False, dtype={'pcm_id': str})
    c1_require(list(candidates.columns) == schema.get('columns'), 'Candidate CSV schema differs from hashed sidecar.')
    for name in candidates.columns:
        null_rows = schema['null_rows'][name]
        c1_require(all((isinstance(row, int) and 0 <= row < len(candidates) for row in null_rows)), 'Invalid declared candidate missing-value positions.')
        if null_rows:
            candidates.loc[null_rows, name] = np.nan
    for name in schema['numeric_columns']:
        candidates[name] = pd.to_numeric(candidates[name], errors='raise')
    for name in schema['boolean_columns']:
        c1_require(candidates[name].isin([True, False, 'True', 'False']).all(), 'Invalid declared candidate boolean values.')
        candidates[name] = candidates[name].map(lambda value: value is True or value == 'True').astype(bool)
    ids = json.loads((folder / 'pcm_prelim_candidate_ids.json').read_text())
    return (weather, candidates, ids, {'weather': weather_meta, 'candidates': candidates_meta, 'screen_derivation': screen, 'preliminary_run_id': manifest['run_id'], 'artifact_directory': str(folder), 'cumulative_run_id': manifest.get('cumulative_run_id'), 'cumulative_source_identity': manifest.get('cumulative_source_identity'), 'consumed_file_sha256': files, 'manifest_sha256': c1_sha256_file(folder / 'preliminary_manifest.json')})

def c1_same_table_values(memory, persisted, label):
    c1_require(list(memory.columns) == list(persisted.columns) and len(memory) == len(persisted), f'{label}: memory/file schema mismatch.')
    for name in memory.columns:
        if name.endswith(('_UTC', '_local')):
            same = c1_explicit_utc(memory[name], name).equals(c1_explicit_utc(persisted[name], name))
        elif pd.api.types.is_numeric_dtype(memory[name]) and (not pd.api.types.is_bool_dtype(memory[name])):
            same = np.allclose(memory[name].to_numpy(float), persisted[name].to_numpy(float), rtol=1e-12, atol=1e-12, equal_nan=True)
        else:
            same = [c1_json_value(v) for v in memory[name]] == [c1_json_value(v) for v in persisted[name]]
        c1_require(same, f'{label}: in-memory values differ from verified persisted column {name}.')

@dataclass(frozen=True)
class C1ValidatedInputs:
    weather: object
    candidates: object
    ids: tuple
    metadata: dict
    config: dict
    validation: dict
    identity: str
    workbook_receipt: dict
    _seal: object

_C1_REAL_INPUT_SEAL = object()

def c1_validate_input_bundle(weather, candidates, ids, metadata, io_config):
    c1_require(isinstance(metadata, dict), 'Explicit Preliminary metadata is required.')
    c1_require(C1_SHARED_CONFIG.get('data_mode', 'REAL') == 'REAL', 'Section-0 mode must be REAL.')
    for key, value in {'study_year': 2025, 'timezone': C1_IST, 'city': 'Delhi', 'hourly_interval_seconds': 3600}.items():
        c1_require(C1_SHARED_CONFIG.get(key, value) == value, f'Section-0 {key} incompatible with this Category-1 contract.')
    cfg = c1_make_run_config(CATEGORY1_CONFIG)
    c1_require(cfg['year'] == 2025 and cfg['d_m'] == 0.02 and (cfg['N_cells'] >= 2), 'Scientific model requires 2025, fixed 20 mm depth, and a depth-resolved grid.')
    c1_require(cfg['surface_boundary'] == 'environment' and cfg['bottom_boundary'] == 'adiabatic' and cfg['save_cell_states'], 'Scientific boundary/state configuration is invalid.')
    c1_require(cfg['weather_time_label'] == 'midpoint' and cfg['report_start_IST'] is None and (cfg['report_end_IST'] is None), 'Scientific solver must consume midpoint samples and report the full local calendar year.')
    workbook = c1_inspect_workbook(io_config)
    wcheck = c1_validate_weather(weather, metadata.get('weather', {}))
    pcheck = c1_validate_candidates(candidates, ids, metadata.get('candidates', {}), metadata.get('screen_derivation', {}), workbook)
    all_rows = audit_cleaned_pcm(load_pcm_properties(io_config['workbook_path'], sheet_name=io_config['workbook_sheet']))
    runnable = partition_pcm_tables(all_rows)['runnable']
    screen = metadata['screen_derivation']
    expected_screen, expected_candidates, expected_ids = screen_workbook_roof(all_rows, runnable, screen['lower_C'], screen['upper_C'])
    c1_require(expected_ids == list(ids), 'Supplied candidate population differs from workbook-derived inclusive overlap selection.')
    c1_same_table_values(expected_candidates, candidates, 'workbook-derived candidate population')
    c1_require((_source_identity_matches(metadata.get('cumulative_source_identity'), CUMULATIVE_SOURCE_IDENTITY) or any(_source_identity_matches(metadata.get('cumulative_source_identity'), approved) for approved in C1_APPROVED_PRELIM_SOURCE_IDENTITIES)) and bool(metadata.get('preliminary_run_id')), 'Exact cumulative code and Preliminary run identity required.')
    identity = c1_digest({'weather': c1_table_digest(weather), 'candidates': c1_table_digest(candidates), 'ids': list(ids), 'metadata': metadata, 'config': cfg, 'acceptance_limits': CATEGORY1_VALIDATION, 'workbook': workbook['receipt']['sha256']})
    return C1ValidatedInputs(weather.copy(deep=True), candidates.copy(deep=True), tuple(ids), deepcopy(metadata), cfg, {'weather': wcheck, 'candidates': pcheck}, identity, workbook['receipt'], _C1_REAL_INPUT_SEAL)

def c1_load_inputs(namespace, io_config):
    """Explicit memory-first cumulative route; never repair or silently switch inputs."""
    route = io_config.get('input_route', 'memory')
    c1_require(route in {'memory', 'persisted'}, 'Choose memory or explicit persisted recovery.')
    if route == 'memory':
        objects = namespace.get('PRELIM_HANDOFF')
        c1_require(namespace.get('PRELIM_READY') is True and isinstance(objects, dict), 'Completed in-memory PRELIM_HANDOFF required.')
        names = ['ANNUAL_WEATHER', 'PCM_PRELIM_CANDIDATES', 'PCM_PRELIM_CANDIDATE_IDS']
        c1_require(all((objects.get(name) is not None for name in names)), 'Incomplete in-memory Preliminary handoff; no fallback.')
        weather, candidates, ids = [objects[name] for name in names]
        directory = objects.get('PRELIM_ARTIFACT_DIR')
        c1_require(directory is not None, 'Preliminary published artifact identity missing.')
        pw, pp, pi, metadata = c1_read_preliminary_directory(directory)
        c1_require(objects.get('PRELIM_RUN_ID') == metadata['preliminary_run_id'], 'In-memory/persisted Preliminary run identity mismatch.')
        c1_require(objects.get('CUMULATIVE_RUN_ID') == metadata['cumulative_run_id'] == CUMULATIVE_RUN_ID, 'In-memory cumulative run identity mismatch.')
        c1_require(objects.get('PCM_SCREEN_METHOD') == CUMULATIVE_SCREEN_METHOD, 'In-memory overlap lineage missing.')
        c1_same_table_values(weather, pw, 'ANNUAL_WEATHER')
        c1_same_table_values(candidates, pp, 'PCM_PRELIM_CANDIDATES')
        c1_require(list(ids) == pi, 'Memory/persisted candidate ID mismatch.')
        c1_require(c1_json_value(objects['PCM_SCREEN_DERIVATION']) == metadata['screen_derivation'], 'Memory/persisted screen derivation mismatch.')
    else:
        directory = io_config['preliminary_artifact_dir']
        c1_require(directory is not None, 'Explicit recovery requires one exact verified Preliminary directory.')
        weather, candidates, ids, metadata = c1_read_preliminary_directory(directory)
    metadata['input_route'] = route
    time_utils = namespace.get('TIME_UTILITIES', {})
    validators = namespace.get('VALIDATION_UTILITIES', {})
    c1_require(callable(time_utils.get('annual_hourly_intervals')) and callable(time_utils.get('validate_hourly_coverage')), 'Execute approved Section-0 time utilities first.')
    label = metadata['weather'].get('time_semantics', metadata['weather']).get('timestamp_label')
    expectation = time_utils['annual_hourly_intervals'](2025, timezone_name=C1_IST, timestamp_label=label)
    time_utils['validate_hourly_coverage'](weather.timestamp_UTC, expectation, timestamp_label=label)
    c1_require(callable(validators.get('require_finite_numeric')), 'Execute approved Section-0 validators first.')
    validators['require_finite_numeric'](weather, list(C1_WEATHER_UNITS), context='Category-1 annual weather')
    validators['require_finite_numeric'](candidates, C1_PCM_REQUIRED_NUMBERS, context='Category-1 Preliminary candidates')
    return c1_validate_input_bundle(weather, candidates, ids, metadata, io_config)

def c1_prepare_model_weather(frame, cfg=None):
    """Solver-only adapter: explicitly midpoint-labelled, continuous, no reordering/filling."""
    cfg = c1_make_run_config(cfg)
    required = ['datetime_UTC', 'T_air_C', 'u10_ms', 'v10_ms', 'G_Wm2', 'LW_down_Wm2']
    c1_require_columns(frame, required, 'solver weather')
    x = frame.copy().reset_index(drop=True)
    c1_require(cfg['weather_time_label'] == 'midpoint', 'Solver requires explicit midpoint weather labels.')
    labels = c1_explicit_utc(x.datetime_UTC, 'solver midpoint datetime_UTC')
    c1_require(not labels.has_duplicates and labels.is_monotonic_increasing, 'Solver weather must already be ordered and unique.')
    c1_require(len(labels) == 1 or np.all(np.diff(labels.as_unit('ns').asi8) == 3600 * 10 ** 9), 'Solver weather has missing hours.')
    c1_numeric(x, required[1:], 'solver weather')
    c1_require((x.T_air_C > -273.15).all() and (x.G_Wm2 >= 0).all() and (x.LW_down_Wm2 >= 0).all(), 'Solver weather violates physical bounds.')
    x['datetime_UTC'] = labels
    x['wind10_ms'] = np.hypot(x.u10_ms, x.v10_ms)
    if 'wind10_ms' in frame:
        c1_require(np.allclose(frame.wind10_ms, x.wind10_ms, rtol=1e-08, atol=1e-08), 'Solver wind speed/components disagree.')
    x['interval_start_UTC'] = labels - pd.Timedelta(minutes=30)
    x['interval_end_UTC'] = labels + pd.Timedelta(minutes=30)
    x['interval_start_IST'] = x.interval_start_UTC.dt.tz_convert(C1_IST)
    x['interval_end_IST'] = x.interval_end_UTC.dt.tz_convert(C1_IST)
    c1_require(all((t.minute == t.second == t.microsecond == t.nanosecond == 0 for t in x.interval_start_IST)), 'Solver intervals must align with exact local hours.')
    x['datetime_IST'] = labels.tz_convert(C1_IST)
    x['source_quality_flag'] = x.get('quality_flag', pd.Series('OK', index=x.index)).fillna('UNSPECIFIED')
    return x

def c1_solver_weather_from_annual(weather):
    x = weather[['T_air_C', 'u10_ms', 'v10_ms', 'wind10_ms', 'G_Wm2', 'LW_down_Wm2']].copy()
    x.insert(0, 'datetime_UTC', c1_explicit_utc(weather.instantaneous_sample_UTC, 'instantaneous_sample_UTC'))
    x.attrs = deepcopy(weather.attrs)
    return x

def c1_report_bounds(cfg):
    start = pd.Timestamp(cfg['report_start_IST'] or f"{cfg['year']}-01-01")
    end = pd.Timestamp(cfg['report_end_IST'] or f"{cfg['year'] + 1}-01-01")
    start = start.tz_localize(C1_IST) if start.tz is None else start.tz_convert(C1_IST)
    end = end.tz_localize(C1_IST) if end.tz is None else end.tz_convert(C1_IST)
    if end <= start:
        raise ValueError('Reporting end must be after start.')
    if start.minute != 0 or end.minute != 0:
        raise ValueError('Reporting bounds must be aligned with whole IST hours.')
    return (start, end)

def c1_hash_payload(payload):
    text = json.dumps(payload, sort_keys=True, default=str, allow_nan=True)
    return hashlib.sha256(text.encode()).hexdigest()

def c1_simulate_pcm(material, weather_hourly, run_config=None, initial_state=None):
    """Simulate any consecutive hourly window; all table columns use accepted substeps.

    PCM state carries through every hour. Use earlier weather for warmup if wanted;
    reporting boundaries filter summaries and never reset physical state.
    """
    cfg = c1_make_run_config(run_config)
    material = dict(material)
    p = c1_material_array(material, allow_sensible_only=True)
    w = c1_prepare_model_weather(weather_hourly, cfg)
    n = cfg['N_cells']
    if initial_state is None:
        initial_h = np.full(n, c1_h_from_temperature(cfg['initial_temperature_C'], p, cfg['T_ref_C'], cfg['initial_liquid_fraction_at_transition']))
    else:
        initial_h = np.asarray(initial_state['h_Jkg'], dtype=float).copy()
        if len(initial_h) != n or not np.isfinite(initial_h).all():
            raise ValueError('initial_state must have one finite enthalpy for each cell.')
        if (c1_cell_state(initial_h, p, cfg['T_ref_C'])[0] <= -273.15).any():
            raise ValueError('Initial enthalpy implies an invalid absolute temperature.')
    if len(w) > 1000 and (not C1_NUMBA_AVAILABLE):
        warnings.warn('Numba is unavailable: the full-year loop will be slower. In Colab run %pip install numba, then rerun model definitions.')
    mode = {'environment': 0, 'zero_flux': 1, 'prescribed_flux': 2}[cfg['surface_boundary']]
    initial_fraction = c1_cell_state(initial_h, p, cfg['T_ref_C'])[1]
    if initial_state is None and p[0] == p[1] and (p[2] == 0.0):
        at_boundary = initial_h == p[5] * (p[0] - cfg['T_ref_C'])
        initial_fraction[at_boundary] = cfg['initial_liquid_fraction_at_transition']
    if initial_state is not None and p[0] == p[1] and (p[2] == 0.0):
        supplied = np.asarray(initial_state.get('f', initial_fraction), dtype=float)
        c1_require(supplied.shape == initial_fraction.shape and np.isfinite(supplied).all() and ((supplied >= 0) & (supplied <= 1)).all(), 'Invalid initial coexistence fractions.')
        at_boundary = initial_h == p[5] * (p[0] - cfg['T_ref_C'])
        c1_require(np.array_equal(supplied[~at_boundary], initial_fraction[~at_boundary]), 'Non-boundary initial phase disagreement.')
        initial_fraction = supplied.copy()
    numeric, h_snapshots, final_h, last_direction, failed_hour, failed_elapsed, residual, fraction_snapshots, numerical_diagnostics = c1_integrate_hours(w.T_air_C.to_numpy(float), w.u10_ms.to_numpy(float), w.v10_ms.to_numpy(float), w.G_Wm2.to_numpy(float), w.LW_down_Wm2.to_numpy(float), initial_h, p, cfg['T_ref_C'], cfg['d_m'], cfg['dt_s'], cfg['minimum_dt_s'], cfg['alpha_s'], cfg['epsilon_LW'], mode, cfg['prescribed_top_flux_Wm2'], cfg['solver_max_iterations'], cfg['solver_atol_Wm2'], cfg['solver_rtol'], cfg['T_base_C'], cfg['f_sat'], cfg['f_solid'], cfg['r_tol_per_h'], initial_fraction)
    if failed_hour >= 0:
        timestamp = w.interval_start_IST.iloc[failed_hour] + pd.Timedelta(seconds=failed_elapsed)
        raise RuntimeError(f"PCM {material.get('pcm_id')} failed to converge at {timestamp}; residual={residual:g} W/m². No incomplete result was saved.")
    hourly = pd.DataFrame(numeric, columns=C1_HOURLY_NUMERIC_COLUMNS)
    weather_hash = hashlib.sha256(pd.util.hash_pandas_object(w[['interval_start_UTC', 'T_air_C', 'u10_ms', 'v10_ms', 'G_Wm2', 'LW_down_Wm2']], index=False).values.tobytes()).hexdigest()
    run_hash = c1_hash_payload({'config': cfg, 'material': {k: material.get(k) for k in C1_PCM_COLUMNS}, 'weather_hash': weather_hash})[:12]
    run_id = f"{material.get('pcm_id', 'diagnostic')}_{run_hash}"
    hourly.insert(0, 'datetime_end_IST', w.interval_end_IST.to_numpy())
    hourly.insert(0, 'datetime_IST', w.interval_start_IST.to_numpy())
    hourly.insert(0, 'datetime_end_UTC', w.interval_end_UTC.to_numpy())
    hourly.insert(0, 'datetime_UTC', w.interval_start_UTC.to_numpy())
    hourly.insert(0, 'pcm_id', material.get('pcm_id', 'diagnostic'))
    hourly.insert(0, 'run_id', run_id)
    for name in ('T_air_C', 'u10_ms', 'v10_ms', 'wind10_ms', 'G_Wm2', 'LW_down_Wm2'):
        hourly[name] = w[name].to_numpy()
    hourly['wind_ms'] = hourly.wind10_ms
    hourly['H_total_J'] = hourly.H_area_Jm2 * cfg['A_m2']
    hourly['q_bottom_mean_Wm2'] = 0.0
    hourly['latent_net_Jm2'] = hourly.latent_uptake_Jm2 - hourly.latent_release_Jm2
    hourly['fraction_change_rate_mean_per_h'] = hourly.f_end - hourly.f_start
    hourly['latent_capacity_Jm2'] = p[7] * cfg['d_m'] * p[2]
    hourly['phase_model'] = 'equilibrium_screen'
    hourly['evidence_tier'] = material.get('evidence_tier', np.nan)
    hourly['regeneration_interpretation'] = 'REVERSIBLE_SCREEN_NOT_VERIFIED_FREEZING'
    duration = hourly[list(C1_STATE_DURATION_COLUMNS)].to_numpy()
    hourly['dominant_state'] = np.asarray(C1_STATE_NAMES)[duration.argmax(axis=1)]
    hourly['dominant_state_tie'] = np.isclose(duration, duration.max(axis=1)[:, None], atol=1e-12).sum(axis=1) > 1
    flow_scale = hourly.q_top_positive_energy_Jm2 + hourly.q_top_negative_energy_Jm2
    hourly['energy_check_pass'] = hourly.energy_residual_Jm2.abs() <= cfg['energy_atol_Jm2'] + cfg['energy_rtol'] * flow_scale
    report_start, report_end = c1_report_bounds(cfg)
    hourly['in_reporting_period'] = (hourly.datetime_IST >= report_start) & (hourly.datetime_end_IST <= report_end)
    flags = []
    for i, row in hourly.iterrows():
        flag = ['EQUILIBRIUM_SCREEN']
        if row.solver_retry_count:
            flag.append('TIMESTEP_REDUCED')
        if not row.energy_check_pass:
            flag.append('ENERGY_BALANCE_FAILED')
        if row.dominant_state_tie:
            flag.append('DOMINANT_STATE_TIE')
        if i in (0, len(hourly) - 1):
            flag.append('INSTANTANEOUS_WEATHER_EDGE_HELD')
        if i == 0 and initial_state is None:
            flag.append('INITIAL_STATE_ASSUMPTION')
        if p[0] == p[1]:
            flag.append('ISOTHERMAL_TRANSITION')
        if cfg['surface_boundary'] != 'environment':
            flag.append('DIAGNOSTIC_BOUNDARY')
        if pd.isna(material.get('source', np.nan)):
            flag.append('SOURCE_NOT_PROVIDED')
        if pd.isna(material.get('evidence_tier', np.nan)):
            flag.append('EVIDENCE_TIER_NOT_SET')
        if not row.in_reporting_period:
            flag.append('COVERAGE_ONLY')
        if w.source_quality_flag.iloc[i] != 'OK':
            flag.append('SOURCE_' + str(w.source_quality_flag.iloc[i]))
        flags.append(';'.join(flag))
    hourly['quality_flags'] = flags
    boundaries = pd.DatetimeIndex([w.interval_start_UTC.iloc[0], *w.interval_end_UTC])
    snapshots = {'datetime_UTC': boundaries, 'h_Jkg': h_snapshots}
    if cfg['save_cell_states']:
        temp, fractions, _, _ = c1_cell_state(h_snapshots.ravel(), p, cfg['T_ref_C'])
        snapshots['T_C'] = temp.reshape(h_snapshots.shape)
        snapshots['f'] = fraction_snapshots
    final = c1_state_from_enthalpy(final_h, material, cfg)
    final['f'] = fraction_snapshots[-1].copy()
    final['history'] = dict(initial_state.get('history', {})) if initial_state is not None else {}
    final['history'].update(phase_model='equilibrium_screen', last_direction=last_direction)
    final['history']['note'] = 'History retained for extension; equilibrium curve has no hysteresis.'
    result = {'hourly': hourly, 'cell_snapshots': snapshots, 'final_state': final, 'run_id': run_id, 'material': material, 'run_config': cfg, 'weather_hash': weather_hash}
    result['numerical_extension'] = {'method': 'EXACT_ZERO_WIDTH_L0_COEXISTENCE_V1', 'applied_to_case': bool(p[0] == p[1] and p[2] == 0.0), 'coexistence_solver_steps': int(numerical_diagnostics[:, 0].sum()), 'null_fraction_step_count': int(numerical_diagnostics[:, 1].sum()), 'passive_fraction_rule': 'retain previous fraction in an exact null Jacobian column; consistent with zero net heat as L tends to zero'}
    result['daily'] = c1_make_daily_summary(hourly, cfg)
    result['monthly'] = c1_make_monthly_summary(hourly, cfg)
    result['annual'] = c1_make_annual_summary(hourly, cfg)
    return result

def c1_simulate_pcm_year(material, weather_hourly, run_config=None, initial_state=None):
    """Require complete reporting-year coverage, with optional warmup and next-day tail."""
    cfg = c1_make_run_config(run_config)
    w = c1_prepare_model_weather(weather_hourly, cfg)
    start, end = c1_report_bounds(cfg)
    if w.interval_start_IST.iloc[0] > start or w.interval_end_IST.iloc[-1] < end:
        raise ValueError(f'Full-year simulation requires coverage from {start} through {end}.')
    result = c1_simulate_pcm(material, weather_hourly, cfg, initial_state)
    expected_hours = (end - start).total_seconds() / 3600.0
    if result['hourly'].in_reporting_period.sum() != expected_hours:
        raise ValueError('Reporting period has uncovered or misaligned hourly intervals.')
    return result

def c1_make_hourly_summary(simulation_result, include_coverage=False):
    """Hourly quantities were accumulated at every accepted sub-hour timestep."""
    hourly = simulation_result['hourly']
    return (hourly if include_coverage else hourly[hourly.in_reporting_period]).copy()

def c1_make_daily_summary(hourly, run_config=None):
    """One row per reporting day, at the exact configured IST morning checkpoint.

    A calendar year of weather alone leaves the last 06:00-to-06:00 cycle partial.
    Its next-morning/reset fields remain unavailable until the extra six hours exist.
    """
    cfg = c1_make_run_config(run_config)
    start, end = c1_report_bounds(cfg)
    hh, mm = map(int, cfg['morning_checkpoint'].split(':'))
    checkpoint_offset = pd.Timedelta(hours=hh, minutes=mm)
    days = pd.date_range(start.normalize(), end.normalize(), freq='D')
    rows = []
    for (run_id, pcm_id), group in hourly.groupby(['run_id', 'pcm_id'], sort=False):
        group = group.sort_values('datetime_IST')
        for day in days:
            cycle_start = day + checkpoint_offset
            cycle_end = cycle_start + pd.Timedelta(days=1)
            if not start <= cycle_start < end:
                continue
            g = group[(group.datetime_IST >= cycle_start) & (group.datetime_end_IST <= cycle_end)]
            has_start = len(g) > 0 and g.datetime_IST.iloc[0] == cycle_start
            has_end = len(g) > 0 and g.datetime_end_IST.iloc[-1] == cycle_end
            complete = has_start and has_end and (len(g) == 24)
            morning = float(g.f_start.iloc[0]) if has_start else np.nan
            next_morning = float(g.f_end.iloc[-1]) if has_end else np.nan
            incomplete = next_morning > cfg['reset_limit'] if np.isfinite(next_morning) else pd.NA
            row = {'run_id': run_id, 'pcm_id': pcm_id, 'cycle_start_IST': cycle_start, 'cycle_end_IST': cycle_end, 'observed_hours': len(g), 'is_complete_cycle': complete, 'f_morning': morning, 'f_next_morning': next_morning, 'f_cycle_min': g.f_min.min(), 'f_cycle_max': g.f_max.max(), 'next_headroom_fraction': 1.0 - next_morning, 'incomplete_reset': incomplete, 'reset_status': 'UNKNOWN_PARTIAL_CYCLE' if not np.isfinite(next_morning) else 'INCOMPLETE_RESET_SCREEN' if incomplete else 'RESET_SCREEN', 'regeneration_interpretation': 'REVERSIBLE_SCREEN_NOT_VERIFIED_FREEZING', 'energy_residual_Jm2': g.energy_residual_Jm2.sum(min_count=1), 'max_abs_hourly_energy_residual_Jm2': g.energy_residual_Jm2.abs().max(), 'latent_capacity_Jm2': group.latent_capacity_Jm2.iloc[0], 'quality_flags': 'COMPLETE_CYCLE' if complete else 'PARTIAL_CYCLE'}
            for name in ('latent_uptake_Jm2', 'latent_release_Jm2', 'cooling_relevant_latent_uptake_Jm2', 'cell_latent_uptake_Jm2', 'cell_latent_release_Jm2', 'cooling_relevant_h', 'CDH_C_h', *C1_STATE_DURATION_COLUMNS):
                row[name] = g[name].sum(min_count=1)
            capacity = row['latent_capacity_Jm2']
            row['equivalent_full_uptake_cycles'] = row['latent_uptake_Jm2'] / capacity if capacity > 0 else np.nan
            rows.append(row)
    return pd.DataFrame(rows)

def c1_period_record(g, cfg):
    """Aggregate energy/duration totals; percentages use summed sub-hour denominators."""
    g = g.sort_values('datetime_UTC')
    first, last = (g.iloc[0], g.iloc[-1])
    row = {'run_id': first.run_id, 'pcm_id': first.pcm_id, 'period_start_IST': first.datetime_IST, 'period_end_IST': last.datetime_end_IST, 'observed_hours': len(g), 'T_PCM_mean_C': g.T_PCM_C.mean(), 'f_start': first.f_start, 'f_end': last.f_end, 'f_min': g.f_min.min(), 'f_max': g.f_max.max(), 'f_mean': g.f_mean.mean(), 'latent_capacity_Jm2': first.latent_capacity_Jm2, 'stored_energy_change_Jm2': last.H_area_Jm2 - first.H_area_start_Jm2, 'boundary_energy_Jm2': g.q_top_mean_Wm2.sum() * 3600.0, 'max_abs_hourly_energy_residual_Jm2': g.energy_residual_Jm2.abs().max(), 'energy_check_pass': bool(g.energy_check_pass.all()), 'solver_retry_count': int(g.solver_retry_count.sum()), 'phase_model': first.phase_model, 'evidence_tier': first.evidence_tier, 'regeneration_interpretation': first.regeneration_interpretation}
    for name in ('latent_uptake_Jm2', 'latent_release_Jm2', 'cooling_relevant_latent_uptake_Jm2', 'cell_latent_uptake_Jm2', 'cell_latent_release_Jm2', 'cooling_relevant_h', 'CDH_C_h', 'q_top_positive_energy_Jm2', 'q_top_negative_energy_Jm2', *C1_STATE_DURATION_COLUMNS):
        row[name] = g[name].sum()
    row['non_cooling_relevant_h'] = len(g) - row['cooling_relevant_h']
    for name in C1_STATE_DURATION_COLUMNS[:4]:
        label = name[:-2] + '_pct_of_cooling_time'
        row[label] = 100.0 * row[name] / row['cooling_relevant_h'] if row['cooling_relevant_h'] > 0 else np.nan
    for name in C1_STATE_DURATION_COLUMNS[4:]:
        label = name[:-2] + '_pct_of_non_cooling_time'
        row[label] = 100.0 * row[name] / row['non_cooling_relevant_h'] if row['non_cooling_relevant_h'] > 0 else np.nan
    row['latent_net_Jm2'] = row['latent_uptake_Jm2'] - row['latent_release_Jm2']
    capacity = row['latent_capacity_Jm2']
    row['equivalent_full_uptake_cycles'] = row['latent_uptake_Jm2'] / capacity if capacity > 0 else np.nan
    row['equivalent_full_release_cycles'] = row['latent_release_Jm2'] / capacity if capacity > 0 else np.nan
    row['energy_residual_Jm2'] = row['stored_energy_change_Jm2'] - row['boundary_energy_Jm2']
    row['latent_fraction_accounting_residual_Jm2'] = row['latent_net_Jm2'] - capacity * (row['f_end'] - row['f_start'])
    return row

def c1_make_monthly_summary(hourly, run_config=None):
    cfg = c1_make_run_config(run_config)
    x = hourly[hourly.in_reporting_period].copy()
    if x.empty:
        return pd.DataFrame()
    x['year'] = x.datetime_IST.dt.year
    x['month'] = x.datetime_IST.dt.month
    rows = []
    for (run_id, pcm_id, year, month), group in x.groupby(['run_id', 'pcm_id', 'year', 'month'], sort=False):
        row = c1_period_record(group, cfg)
        row.update(year=year, month=month)
        rows.append(row)
    return pd.DataFrame(rows)

def c1_make_annual_summary(hourly, run_config=None):
    cfg = c1_make_run_config(run_config)
    x = hourly[hourly.in_reporting_period]
    rows = []
    for _, group in x.groupby(['run_id', 'pcm_id'], sort=False):
        row = c1_period_record(group, cfg)
        daily = c1_make_daily_summary(group if len(group) == len(hourly) else hourly[(hourly.run_id == row['run_id']) & (hourly.pcm_id == row['pcm_id'])], cfg)
        known = daily.f_next_morning.notna() if not daily.empty else pd.Series(dtype=bool)
        row['complete_operational_cycles'] = int(daily.is_complete_cycle.sum()) if not daily.empty else 0
        row['partial_operational_cycles'] = int((~daily.is_complete_cycle).sum()) if not daily.empty else 0
        row['known_next_morning_checkpoints'] = int(known.sum())
        row['incomplete_reset_cycles'] = int(daily.loc[known, 'incomplete_reset'].sum()) if known.any() else 0
        row['incomplete_reset_pct_of_known_cycles'] = 100.0 * row['incomplete_reset_cycles'] / row['known_next_morning_checkpoints'] if row['known_next_morning_checkpoints'] else np.nan
        rows.append(row)
    return pd.DataFrame(rows)

def c1_check_energy_balance(hourly, run_config=None):
    """Independent reconciliation from hourly stored energy and boundary integral."""
    cfg = c1_make_run_config(run_config)
    delta = hourly.H_area_Jm2 - hourly.H_area_start_Jm2
    boundary = hourly.q_top_mean_Wm2 * 3600.0
    error = delta - boundary
    scale = hourly.q_top_positive_energy_Jm2 + hourly.q_top_negative_energy_Jm2
    return pd.DataFrame({'datetime_UTC': hourly.datetime_UTC, 'stored_change_Jm2': delta, 'boundary_input_Jm2': boundary, 'residual_Jm2': error, 'pass': error.abs() <= cfg['energy_atol_Jm2'] + cfg['energy_rtol'] * scale})

def c1_json_ready(value):
    if isinstance(value, dict):
        return {str(k): c1_json_ready(v) for k, v in value.items()}
    if isinstance(value, (tuple, list, np.ndarray)):
        return [c1_json_ready(v) for v in value]
    if isinstance(value, (np.integer, np.floating, np.bool_)):
        value = value.item()
    if value is None or (isinstance(value, float) and (not np.isfinite(value))):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)

def c1_live_code_identity():
    """Fingerprint executed function bodies too, so an edited definition invalidates identity."""
    from types import CodeType

    def code_value(code):
        return {'bytecode': code.co_code.hex(), 'constants': [code_value(v) if isinstance(v, CodeType) else c1_json_value(v) for v in code.co_consts], 'names': code.co_names, 'variables': code.co_varnames, 'argcount': code.co_argcount}
    bodies = {}
    for name, value in sorted(globals().items()):
        if name.startswith('c1_'):
            function = getattr(value, 'py_func', value)
            if hasattr(function, '__code__'):
                bodies[name] = code_value(function.__code__)
    return c1_digest({'computational_source_sha256': C1_RECORDED_SOURCE_SHA256, 'executed_functions': bodies, 'model_version': C1_MODEL_VERSION, 'sigma': C1_SIGMA_Wm2K4, 'weather_units': C1_WEATHER_UNITS, 'material_units': C1_MATERIAL_UNITS})

def c1_reuse_scientific_cache(*args, **kwargs):
    raise Category1InputError('Automatic scientific-cache reuse is disabled. Recompute from verified Preliminary inputs; filenames never establish identity.')

def c1_matched_control(pcm_result, solver_weather):
    c1_require(pcm_result['run_config']['surface_boundary'] == 'environment', 'Matched scientific control requires the environmental boundary.')
    control_material = deepcopy(pcm_result['material'])
    control_material['L_Jkg'] = 0.0
    initial_T = pcm_result['cell_snapshots']['T_C'][0].copy()
    control_h = c1_enthalpy_from_state(initial_T, control_material, pcm_result['run_config'])
    return c1_simulate_pcm(control_material, solver_weather, pcm_result['run_config'], initial_state={'h_Jkg': control_h})

def c1_validate_run(result, *, require_annual=True):
    cfg, material = (result['run_config'], result['material'])
    hourly, states = (result['hourly'], result['cell_snapshots'])
    c1_require(not hourly.empty, 'Empty result cannot be validated.')
    if require_annual:
        expected = pd.date_range(pd.Timestamp('2025-01-01', tz=C1_IST), pd.Timestamp('2026-01-01', tz=C1_IST), freq='h', inclusive='left').tz_convert('UTC')
        c1_require(len(hourly) == 8760 and c1_explicit_utc(hourly.datetime_UTC, 'result time').equals(expected), 'Scientific result does not have all 8760 ordered target intervals.')
    start = c1_explicit_utc(hourly.datetime_UTC, 'result start')
    end = c1_explicit_utc(hourly.datetime_end_UTC, 'result end')
    c1_require(not start.has_duplicates and start.is_monotonic_increasing and (end == start + pd.Timedelta(hours=1)).all(), 'Result intervals are invalid.')
    c1_require(len(start) == 1 or (start[1:] == end[:-1]).all(), 'Result intervals have gaps.')
    boundaries = pd.DatetimeIndex([start[0], *end])
    c1_require(pd.DatetimeIndex(states['datetime_UTC']).equals(boundaries), 'Cell states and hourly intervals disagree.')
    for name in ('T_C', 'h_Jkg', 'f'):
        c1_require(np.asarray(states[name]).shape == (len(hourly) + 1, cfg['N_cells']) and np.isfinite(states[name]).all(), f'Missing/nonfinite or mismatched depth states: {name}.')
    h, fraction = (states['h_Jkg'], states['f'])
    c1_require(((fraction >= 0) & (fraction <= 1)).all(), 'Cell liquid fraction outside [0,1].')
    inverse = c1_state_from_enthalpy(h.ravel(), material, cfg)
    expected_f = inverse['f'].reshape(h.shape)
    phase_matches = np.isclose(expected_f, fraction, rtol=0, atol=1e-12)
    coexistence = np.zeros(h.shape, dtype=bool)
    if float(material['L_Jkg']) == 0.0 and material['Tm_low_C'] == material['Tm_high_C']:
        h_transition = float(material['cp_solid_JkgK']) * (float(material['Tm_low_C']) - cfg['T_ref_C'])
        coexistence = (h == h_transition) & (states['T_C'] == float(material['Tm_low_C']))
    c1_require(np.allclose(inverse['T_C'].reshape(h.shape), states['T_C'], rtol=0, atol=1e-09) and np.all(phase_matches | coexistence), 'Saved state fails the approved equilibrium graph.')
    rho_dx = float(material['rho_ref_kgm3']) * cfg['d_m'] / cfg['N_cells']
    stored = rho_dx * h.sum(axis=1)
    c1_require(np.allclose(stored[:-1], hourly.H_area_start_Jm2, rtol=0, atol=1e-06) and np.allclose(stored[1:], hourly.H_area_Jm2, rtol=0, atol=1e-06), 'Hourly energy disagrees with depth enthalpies.')
    c1_require(np.allclose(hourly.H_area_start_Jm2.to_numpy()[1:], hourly.H_area_Jm2.to_numpy()[:-1], rtol=0, atol=1e-06), 'PCM state energy was reset or discontinuous between hours.')
    c1_require(np.allclose(fraction[:-1].mean(axis=1), hourly.f_start, rtol=0, atol=1e-12) and np.allclose(fraction[1:].mean(axis=1), hourly.f_end, rtol=0, atol=1e-12), 'Hourly f disagrees with continuous state.')
    c1_require(np.allclose(result['final_state']['h_Jkg'], h[-1], rtol=0, atol=CATEGORY1_VALIDATION['continuity_atol_Jkg']), 'Final state differs from final hour boundary.')
    numeric = list(C1_HOURLY_NUMERIC_COLUMNS)
    numeric = [n for n in numeric if n not in {'first_saturated_substep_end_s', 'last_saturated_substep_end_s'}]
    if cfg['surface_boundary'] != 'environment':
        numeric = [n for n in numeric if n not in {'q_solar_mean_Wm2', 'q_convection_mean_Wm2', 'q_longwave_mean_Wm2'}]
    c1_numeric(hourly, numeric, 'hourly result')
    boundary = 3600.0 * hourly.q_top_mean_Wm2.to_numpy(float)
    residual = np.diff(stored) - boundary
    absolute_flow = (hourly.q_top_positive_energy_Jm2 + hourly.q_top_negative_energy_Jm2).to_numpy(float)
    c1_require((absolute_flow >= 0).all(), 'Invalid absolute boundary energy scale.')
    tolerance = cfg['energy_atol_Jm2'] + cfg['energy_rtol'] * absolute_flow
    c1_require((np.abs(residual) <= tolerance).all(), 'HARD FAIL: per-hour PCM energy conservation.')
    c1_require(np.allclose(residual, hourly.energy_residual_Jm2, rtol=0, atol=1e-06), 'Saved energy residual differs from independent reconstruction.')
    annual_error = float(stored[-1] - stored[0] - boundary.sum())
    annual_scale = float(absolute_flow.sum())
    annual_tolerance = CATEGORY1_VALIDATION['annual_energy_atol_Jm2'] + CATEGORY1_VALIDATION['annual_energy_rtol'] * annual_scale
    c1_require(abs(annual_error) <= annual_tolerance, 'HARD FAIL: annual PCM energy conservation.')
    c1_require(hourly.max_step_energy_residual_Jm2.max() <= CATEGORY1_VALIDATION['step_energy_atol_Jm2'], 'HARD FAIL: accepted-step energy conservation.')
    c1_require(hourly.surface_balance_max_residual_Wm2.max() <= CATEGORY1_VALIDATION['surface_atol_Wm2'], 'HARD FAIL: surface-energy-balance residual.')
    c1_require(np.allclose(hourly[list(C1_STATE_DURATION_COLUMNS)].sum(axis=1), 1.0, rtol=0, atol=1e-10), 'Accepted substep state durations do not sum to each hour.')
    capacity = float(material['rho_ref_kgm3']) * cfg['d_m'] * float(material['L_Jkg'])
    latent_net = float(hourly.latent_uptake_Jm2.sum() - hourly.latent_release_Jm2.sum())
    fraction_delta_energy = capacity * float(fraction[-1].mean() - fraction[0].mean())
    latent_error = latent_net - fraction_delta_energy
    latent_scale = float(hourly.latent_uptake_Jm2.sum() + hourly.latent_release_Jm2.sum())
    c1_require(abs(latent_error) <= CATEGORY1_VALIDATION['latent_accounting_atol_Jm2'] + CATEGORY1_VALIDATION['latent_accounting_rtol'] * latent_scale, 'HARD FAIL: latent/fraction accounting.')
    if float(material['L_Jkg']) == 0.0:
        c1_require(latent_scale == 0.0 and capacity == 0.0, 'L=0 control produced latent energy.')
    if cfg['surface_boundary'] == 'environment':
        closure = hourly.q_solar_mean_Wm2 + hourly.q_convection_mean_Wm2 + hourly.q_longwave_mean_Wm2 - hourly.q_top_mean_Wm2
        c1_require(closure.abs().max() <= CATEGORY1_VALIDATION['surface_atol_Wm2'], 'Hourly surface flux components do not close.')
    return {'status': 'PASS', 'solver_status': 'ALL_ACCEPTED_STEPS_CONVERGED', 'unresolved_solver_failures': 0, 'weather_hash': result['weather_hash'], 'run_id': result['run_id'], 'annual_stored_change_Jm2': float(stored[-1] - stored[0]), 'annual_boundary_input_Jm2': float(boundary.sum()), 'annual_energy_residual_Jm2': annual_error, 'annual_energy_relative_residual': abs(annual_error) / max(1.0, annual_scale), 'annual_absolute_boundary_energy_Jm2': annual_scale, 'annual_energy_tolerance_Jm2': annual_tolerance, 'max_abs_hourly_energy_residual_Jm2': float(np.max(np.abs(residual))), 'max_step_energy_residual_Jm2': float(hourly.max_step_energy_residual_Jm2.max()), 'max_surface_residual_Wm2': float(hourly.surface_balance_max_residual_Wm2.max()), 'max_cell_residual_Wm2': float(hourly.solver_max_cell_residual_Wm2.max()), 'solver_retry_count': int(hourly.solver_retry_count.sum()), 'max_nonlinear_iterations_used': int(hourly.solver_max_iterations_used.max()), 'minimum_accepted_dt_s': float(hourly.minimum_accepted_dt_s.min()), 'latent_accounting_residual_Jm2': latent_error, 'state_continuity_status': 'PASS'}

def c1_validate_pair(pcm, control):
    c1_require(pcm['run_config'] == control['run_config'], 'PCM/L0 configuration differs.')
    pmat, bmat = (pcm['material'], control['material'])
    c1_require(set(pmat) == set(bmat), 'PCM/L0 material schemas differ.')
    c1_require(float(pmat['L_Jkg']) > 0 and float(bmat['L_Jkg']) == 0, 'Control must differ only by setting L=0.')
    for name in pmat:
        if name != 'L_Jkg':
            c1_require(c1_json_value(pmat[name]) == c1_json_value(bmat[name]), f'PCM/L0 differs in {name}.')
    c1_require(pcm['weather_hash'] == control['weather_hash'], 'PCM/L0 weather identity differs.')
    p, b = (pcm['hourly'], control['hourly'])
    c1_require(len(p) > 0 and len(p) == len(b), 'PCM/L0 lengths differ or are empty.')
    for name in ['datetime_UTC', 'datetime_end_UTC']:
        c1_require(c1_explicit_utc(p[name], name).equals(c1_explicit_utc(b[name], name)), 'Delta H interval alignment mismatch.')
    c1_require(p.pcm_id.eq(pmat['pcm_id']).all() and b.pcm_id.eq(pmat['pcm_id']).all(), 'Delta H PCM identity mismatch.')
    c1_require(p.run_id.eq(pcm['run_id']).all() and b.run_id.eq(control['run_id']).all(), 'Stale hourly run identity.')
    for name in ['T_air_C', 'u10_ms', 'v10_ms', 'wind10_ms', 'G_Wm2', 'LW_down_Wm2']:
        c1_require(np.array_equal(p[name].to_numpy(), b[name].to_numpy()), f'PCM/L0 forcing differs in {name}.')
    c1_require(np.allclose(pcm['cell_snapshots']['T_C'][0], control['cell_snapshots']['T_C'][0], rtol=0, atol=1e-10), 'PCM/L0 initial temperature profiles differ.')
    return {'status': 'PASS', 'difference': 'L_Jkg only', 'initial_equivalence': 'identical cell temperatures; own enthalpy laws'}

def c1_delta_h(pcm, control):
    """Form interval-average roof-to-air H_PCM minus matched H_L0, preserving the convection sign."""
    c1_validate_pair(pcm, control)
    p, b = (pcm['hourly'], control['hourly'])
    c1_numeric(p, ['q_convection_mean_Wm2'], 'PCM convection')
    c1_numeric(b, ['q_convection_mean_Wm2'], 'L0 convection')
    out = p[['datetime_UTC', 'datetime_end_UTC', 'datetime_IST', 'datetime_end_IST', 'pcm_id']].copy()
    out['H_PCM_Wm2'] = -p.q_convection_mean_Wm2.to_numpy(float)
    out['H_L0_Wm2'] = -b.q_convection_mean_Wm2.to_numpy(float)
    out['delta_H_Wm2'] = out.H_PCM_Wm2 - out.H_L0_Wm2
    out['pcm_run_id'] = pcm['run_id']
    out['L0_run_id'] = control['run_id']
    out['weather_hash'] = pcm['weather_hash']
    out['baseline_kind'] = 'no_latent'
    out.attrs = {'units': {'H_PCM_Wm2': 'W/m2', 'H_L0_Wm2': 'W/m2', 'delta_H_Wm2': 'W/m2'}, 'sign_convention': 'positive roof into atmosphere; delta_H = H_PCM - H_L0', 'data_nature': p.attrs.get('data_nature', 'UNATTESTED_LOW_LEVEL_RESULT')}
    return out

def c1_calendar_daily(hourly):
    records = []
    for day, group in hourly.groupby(hourly.datetime_IST.dt.date, sort=False):
        records.append({'date_IST': str(day), 'observed_hours': len(group), 'f_min': float(group.f_min.min()), 'f_max': float(group.f_max.max()), 'f_mean': float(group.f_mean.mean()), 'activation_h': float(group.active_h.sum()), 'regenerating_h': float(group.regenerating_h.sum()), 'raw_saturated_h': float(group.raw_saturated_h.sum()), 'raw_saturation_onsets': int(group.raw_saturation_onsets.sum()), 'latent_uptake_Jm2': float(group.latent_uptake_Jm2.sum()), 'latent_release_Jm2': float(group.latent_release_Jm2.sum())})
    return pd.DataFrame(records)

def c1_saturation_intervals(hourly):
    fields = ['datetime_UTC', 'datetime_end_UTC', 'datetime_IST', 'datetime_end_IST', 'pcm_id', 'f_start', 'f_end', 'raw_saturated_h', 'raw_saturation_onsets', 'first_saturated_substep_end_s', 'last_saturated_substep_end_s']
    out = hourly.loc[hourly.raw_saturated_h > 0, fields].copy()
    out['first_saturated_evaluation_UTC'] = out.datetime_UTC + pd.to_timedelta(out.first_saturated_substep_end_s, unit='s')
    out['last_saturated_evaluation_UTC'] = out.datetime_UTC + pd.to_timedelta(out.last_saturated_substep_end_s, unit='s')
    out['time_resolution_note'] = 'accepted-step endpoint f>=f_sat; duration assigned to that step; not a continuous crossing-time root'
    return out

def c1_annual_candidate_summary(pcm, control, delta, validation):
    row = pcm['annual'].iloc[0].to_dict()
    material, daily = (pcm['material'], pcm['daily'])
    calendar = c1_calendar_daily(pcm['hourly'])
    known = daily.f_next_morning.notna()
    resets = int((daily.loc[known, 'f_next_morning'] <= pcm['run_config']['reset_limit']).sum())
    row.update({name: material[name] for name in ['material_name', 'Tm_low_C', 'Tm_high_C', 'L_Jkg']})
    row.update(activation_days=int((calendar.activation_h > 0).sum()), regenerating_days=int((calendar.regenerating_h > 0).sum()), raw_saturation_days=int((calendar.raw_saturated_h > 0).sum()), raw_saturation_h=float(calendar.raw_saturated_h.sum()), raw_saturation_onsets=int(calendar.raw_saturation_onsets.sum()), resetting_known_cycles=resets, reset_fraction_of_known_cycles=resets / int(known.sum()) if known.any() else np.nan, H_PCM_annual_mean_Wm2=float(delta.H_PCM_Wm2.mean()), H_L0_annual_mean_Wm2=float(delta.H_L0_Wm2.mean()), delta_H_annual_mean_Wm2=float(delta.delta_H_Wm2.mean()), delta_H_min_Wm2=float(delta.delta_H_Wm2.min()), delta_H_max_Wm2=float(delta.delta_H_Wm2.max()), delta_H_annual_integral_Jm2=float(delta.delta_H_Wm2.sum() * 3600), H_PCM_annual_integral_Jm2=float(delta.H_PCM_Wm2.sum() * 3600), H_L0_annual_integral_Jm2=float(delta.H_L0_Wm2.sum() * 3600), solver_status=validation['pcm']['solver_status'], control_solver_status=validation['L0']['solver_status'], energy_conservation_status='PASS', matched_control_status='PASS', physical_source_quality='NOT CHECKED')
    for name in ['annual_energy_residual_Jm2', 'annual_energy_relative_residual', 'annual_energy_tolerance_Jm2']:
        row[name] = validation['pcm'][name]
        row['L0_' + name] = validation['L0'][name]
    return row

def c1_result_seal(handoff):
    content = {}
    for pcm_id, result in handoff['by_pcm'].items():
        content[pcm_id] = {name: c1_table_digest(result[name]) for name in ['hourly', 'matched_control', 'delta_H', 'daily', 'monthly', 'calendar_daily', 'saturation_events']}
        for name in ['summary', 'matched_control_summary', 'material', 'control_material', 'initial_condition', 'numerical_extension']:
            content[pcm_id][name] = c1_digest(result[name])
        content[pcm_id]['validation'] = c1_digest(result['validation'])
        content[pcm_id]['identity'] = c1_digest(result['identity'])
        for case in ['pcm', 'L0']:
            states = result['cell_snapshots'][case]
            content[pcm_id][case + '_state_times'] = c1_digest(list(states['datetime_UTC']))
            content[pcm_id][case + '_states'] = {name: hashlib.sha256(np.asarray(states[name]).tobytes()).hexdigest() for name in ['T_C', 'h_Jkg', 'f']}
    return c1_digest({'by_pcm': content, 'annual_summary': c1_table_digest(handoff['annual_summary']), 'config': handoff['config'], 'acceptance_limits': handoff['acceptance_limits'], 'units': handoff['units'], 'provenance': handoff['provenance'], 'validation': handoff['validation'], 'candidate_ids': handoff['candidate_ids']})

def c1_output_units():
    units = {name: 'J/m2' if name.endswith('_Jm2') else 'W/m2' if name.endswith('_Wm2') else 'J/kg' if name.endswith('_Jkg') else 'degC' if name.endswith('_C') else 'h' if name.endswith('_h') else 's' if name.endswith('_s') else '1' for name in C1_HOURLY_NUMERIC_COLUMNS}
    units.update(C1_WEATHER_UNITS)
    units.update({'CDH_C_h': 'degC*h', 'hc_Wm2K': 'W/m2/K', 'wind_mean_ms': 'm/s', 'wind_ms': 'm/s', 'H_total_J': 'J', 'H_PCM_Wm2': 'W/m2', 'H_L0_Wm2': 'W/m2', 'delta_H_Wm2': 'W/m2', 'q_bottom_mean_Wm2': 'W/m2', 'fraction_change_rate_mean_per_h': '1/h', 'h_Jkg': 'J/kg', 'T_C': 'degC', 'f': '1', 'x_center_m': 'm', 'datetime_UTC_ns': 'ns since Unix epoch UTC', 'timestamps': 'timezone-aware ISO 8601', 'sign_convention': 'q_* positive into roof; H_* positive roof into atmosphere; delta_H=H_PCM-H_L0'})
    return units

def c1_table_units(frame):
    base = c1_output_units()
    result = {}
    for name in frame.columns:
        canonical_name = name[4:] if name.startswith('raw_') else name
        if canonical_name in C1_MATERIAL_UNITS:
            unit = C1_MATERIAL_UNITS[canonical_name]
        elif name in base:
            unit = base[name]
        elif name.endswith(('_UTC', '_IST', '_local')) or name in {'date_IST'}:
            unit = 'timezone-aware ISO 8601' if name != 'date_IST' else 'local calendar date Asia/Kolkata'
        elif name in C1_MATERIAL_UNITS:
            unit = C1_MATERIAL_UNITS[name]
        elif name.endswith('_Jm2'):
            unit = 'J/m2'
        elif name.endswith('_Wm2'):
            unit = 'W/m2'
        elif name.endswith('_Jkg'):
            unit = 'J/kg'
        elif name.endswith('_JkgK'):
            unit = 'J/kg/K'
        elif name.endswith('_WmK'):
            unit = 'W/m/K'
        elif name.endswith('_kgm3'):
            unit = 'kg/m3'
        elif name.endswith('_C'):
            unit = 'degC'
        elif '_pct_' in name or name.endswith('_pct'):
            unit = '%'
        elif name.endswith('_h') or name == 'observed_hours':
            unit = 'h'
        elif name.endswith('_s'):
            unit = 's'
        elif pd.api.types.is_numeric_dtype(frame[name]) or pd.api.types.is_bool_dtype(frame[name]):
            unit = '1'
        else:
            unit = 'text / source metadata'
        result[name] = unit
    return result

def c1_run_category1(inputs):
    c1_require(isinstance(inputs, C1ValidatedInputs) and inputs._seal is _C1_REAL_INPUT_SEAL, 'Scientific run requires the validated REAL input contract; mechanical fixtures are refused.')
    revalidated = c1_validate_input_bundle(inputs.weather, inputs.candidates, list(inputs.ids), inputs.metadata, CATEGORY1_IO)
    c1_require(revalidated.identity == inputs.identity and revalidated.config == inputs.config, 'Inputs/config changed after validation.')
    solver_weather = c1_solver_weather_from_annual(inputs.weather)
    live_code = c1_live_code_identity()
    environment = {'python': sys.version.split()[0], 'numpy': np.__version__, 'pandas': pd.__version__, 'numba_used': C1_NUMBA_AVAILABLE}
    if C1_NUMBA_AVAILABLE:
        import numba
        environment['numba'] = numba.__version__
    by_pcm = {}
    for _, material in inputs.candidates.iterrows():
        pcm_id = material.pcm_id
        print(f'Category 1: {pcm_id} — real annual PCM + matched L=0 control', flush=True)
        pcm = c1_simulate_pcm_year(material, solver_weather, inputs.config)
        control = c1_matched_control(pcm, solver_weather)
        for result in (pcm, control):
            table = result['hourly']
            capacity = float(result['material']['rho_ref_kgm3']) * inputs.config['d_m'] * float(result['material']['L_Jkg'])
            table['latent_stored_start_Jm2'] = capacity * table.f_start
            table['latent_stored_Jm2'] = capacity * table.f_end
            table['sensible_stored_start_Jm2'] = table.H_area_start_Jm2 - table.latent_stored_start_Jm2
            table['sensible_stored_Jm2'] = table.H_area_Jm2 - table.latent_stored_Jm2
            table.attrs = {'data_nature': 'REAL', 'units': c1_output_units(), 'input_identity': inputs.identity}
        validation = {'pcm': c1_validate_run(pcm), 'L0': c1_validate_run(control), 'pair': c1_validate_pair(pcm, control)}
        delta = c1_delta_h(pcm, control)
        row_hash = c1_digest({name: material[name] for name in C1_PCM_COLUMNS})
        identity = {'pcm_id': pcm_id, 'pcm_row_sha256': row_hash, 'workbook_sha256': inputs.workbook_receipt['sha256'], 'weather_content_sha256': c1_table_digest(inputs.weather), 'weather_solver_hash': pcm['weather_hash'], 'weather_artifact_sha256': inputs.metadata.get('consumed_file_sha256', {}).get('annual_weather.csv'), 'study_year': 2025, 'config_sha256': c1_digest(inputs.config), 'code_sha256': live_code, 'computational_source_sha256': C1_RECORDED_SOURCE_SHA256, 'model_version': C1_MODEL_VERSION, 'environment': environment, 'input_identity': inputs.identity, 'pcm_run_id': pcm['run_id'], 'L0_run_id': control['run_id']}
        pair_id = c1_digest(identity)
        identity['pair_id'] = pair_id
        for table in [pcm['hourly'], control['hourly'], delta]:
            table['pair_id'] = pair_id
            table['code_sha256'] = live_code
            table['study_year'] = 2025
            table['scientific_reality_status'] = 'REAL'
        pcm['hourly']['H_PCM_Wm2'] = delta.H_PCM_Wm2.to_numpy()
        control['hourly']['H_L0_Wm2'] = delta.H_L0_Wm2.to_numpy()
        summary = c1_annual_candidate_summary(pcm, control, delta, validation)
        summary.update(pair_id=pair_id, code_sha256=live_code, scientific_reality_status='REAL')
        by_pcm[pcm_id] = {'hourly': pcm['hourly'], 'matched_control': control['hourly'], 'delta_H': delta, 'H_PCM': delta.set_index('datetime_UTC').H_PCM_Wm2, 'H_L0': delta.set_index('datetime_UTC').H_L0_Wm2, 'DELTA_H': delta.set_index('datetime_UTC').delta_H_Wm2, 'daily': pcm['daily'], 'monthly': pcm['monthly'], 'calendar_daily': c1_calendar_daily(pcm['hourly']), 'saturation_events': c1_saturation_intervals(pcm['hourly']), 'summary': summary, 'matched_control_summary': control['annual'].iloc[0].to_dict(), 'validation': validation, 'identity': identity, 'material': c1_json_value(dict(material)), 'control_material': c1_json_value(control['material']), 'cell_snapshots': {'pcm': pcm['cell_snapshots'], 'L0': control['cell_snapshots']}, 'numerical_extension': {'pcm': pcm['numerical_extension'], 'L0': control['numerical_extension']}, 'initial_condition': {'pcm_T_C': pcm['cell_snapshots']['T_C'][0].tolist(), 'pcm_f': pcm['cell_snapshots']['f'][0].tolist(), 'pcm_h_Jkg': pcm['cell_snapshots']['h_Jkg'][0].tolist(), 'L0_T_C': control['cell_snapshots']['T_C'][0].tolist(), 'L0_f': control['cell_snapshots']['f'][0].tolist(), 'L0_h_Jkg': control['cell_snapshots']['h_Jkg'][0].tolist(), 'spinup': 'NONE'}}
    c1_require(list(by_pcm) == list(inputs.ids) and bool(by_pcm), 'Not all Preliminary candidates completed.')
    handoff = {'status': 'COMPLETE_REAL_ANNUAL', 'scientific_reality_status': 'REAL', 'candidate_ids': list(inputs.ids), 'by_pcm': by_pcm, 'annual_summary': pd.DataFrame([r['summary'] for r in by_pcm.values()]), 'config': deepcopy(inputs.config), 'acceptance_limits': deepcopy(CATEGORY1_VALIDATION), 'units': c1_output_units(), 'provenance': {'inputs': deepcopy(inputs.metadata), 'workbook': deepcopy(inputs.workbook_receipt), 'input_identity': inputs.identity, 'code_sha256': live_code, 'environment': environment}, 'validation': {'status': 'PASS', 'inputs': deepcopy(inputs.validation), 'all_candidates_completed': True}, 'artifact_directory': None, 'manifest': None, '_real_input_seal': _C1_REAL_INPUT_SEAL}
    handoff['result_use'] = 'DEVELOPMENT'
    handoff['candidate_population_final'] = False
    handoff['cumulative_run_id'] = CUMULATIVE_RUN_ID
    handoff['preliminary_input_identity'] = inputs.identity
    handoff['weather_identity'] = {'content_sha256': c1_table_digest(inputs.weather), 'artifact_sha256': inputs.metadata['consumed_file_sha256']['annual_weather.csv']}
    handoff['workbook_identity'] = deepcopy(inputs.workbook_receipt)
    handoff['provenance']['cumulative_source_identity'] = deepcopy(CUMULATIVE_SOURCE_IDENTITY)
    handoff['provenance']['cumulative_run_id'] = CUMULATIVE_RUN_ID
    handoff['provenance']['result_use'] = 'DEVELOPMENT'
    handoff['result_content_sha256'] = c1_result_seal(handoff)
    return handoff

def c1_persist_category1(handoff, inputs, output_directory):
    """Write only fully verified real annual outputs, without loading any result cache."""
    c1_require(isinstance(inputs, C1ValidatedInputs) and inputs._seal is _C1_REAL_INPUT_SEAL, 'Persistence refuses mechanical/unvalidated input bundles.')
    c1_require(isinstance(handoff, dict) and handoff.get('_real_input_seal') is _C1_REAL_INPUT_SEAL and (handoff.get('status') == 'COMPLETE_REAL_ANNUAL') and (handoff.get('scientific_reality_status') == 'REAL'), 'Persistence requires a completed REAL annual handoff; fixtures cannot be saved as scientific outputs.')
    current = c1_validate_input_bundle(inputs.weather, inputs.candidates, list(inputs.ids), inputs.metadata, CATEGORY1_IO)
    c1_require(current.identity == inputs.identity == handoff['provenance']['input_identity'], 'Scientific inputs changed before persistence.')
    c1_require(c1_live_code_identity() == handoff['provenance']['code_sha256'], 'Executed code changed before persistence; recompute.')
    c1_require(c1_result_seal(handoff) == handoff['result_content_sha256'], 'Result content changed after hard validation.')
    c1_require(handoff['candidate_ids'] == list(inputs.ids) and list(handoff['by_pcm']) == list(inputs.ids), 'Scientific candidate population mismatch.')
    for pcm_id, record in handoff['by_pcm'].items():
        for key in ['hourly', 'matched_control', 'delta_H']:
            frame = record[key]
            c1_require_real(frame.attrs, f'{pcm_id}/{key} attributes')
            c1_require(frame.scientific_reality_status.eq('REAL').all() and len(frame) == 8760 and frame.pcm_id.eq(pcm_id).all(), 'Scientific table lineage/annual identity mismatch.')
        delta = record['delta_H']
        c1_require(np.array_equal(delta.H_PCM_Wm2.to_numpy(), -record['hourly'].q_convection_mean_Wm2.to_numpy()) and np.array_equal(delta.H_L0_Wm2.to_numpy(), -record['matched_control'].q_convection_mean_Wm2.to_numpy()) and np.array_equal(delta.delta_H_Wm2.to_numpy(), (delta.H_PCM_Wm2 - delta.H_L0_Wm2).to_numpy()), 'Delta H sign/value mismatch before persistence.')
        for name, column in [('H_PCM', 'H_PCM_Wm2'), ('H_L0', 'H_L0_Wm2'), ('DELTA_H', 'delta_H_Wm2')]:
            c1_require(record[name].equals(delta.set_index('datetime_UTC')[column]), 'Handoff flux view changed before persistence.')
    parent = Path(output_directory).expanduser().resolve()
    parent.mkdir(parents=True, exist_ok=True)
    token = uuid.uuid4().hex
    staging = parent / ('.staging_' + token)
    final = parent / ('category1_2025_' + token)
    staging.mkdir(exist_ok=False)
    records = {}
    common = {'result_use': 'DEVELOPMENT', 'candidate_population_final': False, 'cumulative_run_id': CUMULATIVE_RUN_ID, 'cumulative_source_identity': CUMULATIVE_SOURCE_IDENTITY, 'screen_method': CUMULATIVE_SCREEN_METHOD, 'study_year': 2025, 'scientific_reality_status': 'REAL', 'data_nature': 'REAL', 'config': handoff['config'], 'acceptance_limits': handoff['acceptance_limits'], 'model_version': C1_MODEL_VERSION, 'code_sha256': handoff['provenance']['code_sha256'], 'workbook_sha256': inputs.workbook_receipt['sha256'], 'weather_content_sha256': c1_table_digest(inputs.weather), 'weather_artifact_sha256': inputs.metadata.get('consumed_file_sha256', {}).get('annual_weather.csv'), 'input_identity': inputs.identity, 'units': handoff['units'], 'provenance': handoff['provenance']}

    def register(path, kind, pcm_id=None, identity=None):
        records[str(path.relative_to(staging))] = {**common, 'kind': kind, 'pcm_id': pcm_id, 'pair_identity': identity, 'sha256': c1_sha256_file(path)}

    def save_csv(path, frame, kind, pcm_id=None, identity=None):
        frame.to_csv(path, index=False)
        restored = pd.read_csv(path, float_precision='round_trip')
        c1_require(len(restored) == len(frame) and list(restored.columns) == list(frame.columns), f'CSV roundtrip schema failed: {path.name}.')
        for name in frame.select_dtypes(include='number').columns:
            c1_require(np.allclose(frame[name].to_numpy(float), restored[name].to_numpy(float), rtol=1e-12, atol=1e-12, equal_nan=True), f'CSV numeric roundtrip failed: {path.name}/{name}.')
        for name in [n for n in frame.columns if n.endswith(('_UTC', '_IST')) and (not n.startswith('date_'))]:
            c1_require(c1_explicit_utc(frame[name], name).equals(c1_explicit_utc(restored[name], name)), f'CSV timestamp roundtrip failed: {name}.')
        register(path, kind, pcm_id, identity)
        records[str(path.relative_to(staging))]['column_units'] = c1_table_units(frame)

    def save_json(path, value, kind, pcm_id=None, identity=None):
        path.write_text(json.dumps(c1_json_value(value), indent=2, ensure_ascii=False, allow_nan=False), encoding='utf-8')
        register(path, kind, pcm_id, identity)
    try:
        save_csv(staging / 'annual_summary.csv', handoff['annual_summary'], 'annual candidate comparison')
        save_csv(staging / 'annual_weather_used.csv', inputs.weather, 'exact Preliminary annual forcing')
        save_csv(staging / 'pcm_prelim_candidates_used.csv', inputs.candidates, 'exact Preliminary shortlist')
        save_json(staging / 'category1_config.json', {'model': handoff['config'], 'acceptance': handoff['acceptance_limits']}, 'configuration')
        save_json(staging / 'input_provenance.json', common, 'input/model provenance')
        folders = {}
        for i, (pcm_id, result) in enumerate(handoff['by_pcm'].items(), 1):
            sub = staging / f'pcm_{i:03d}'
            sub.mkdir()
            folders[pcm_id] = sub.name
            identity = result['identity']
            for key, filename in [('hourly', 'hourly_pcm.csv'), ('matched_control', 'hourly_L0.csv'), ('delta_H', 'delta_H.csv'), ('daily', 'daily_operation.csv'), ('monthly', 'monthly_summary.csv'), ('calendar_daily', 'calendar_daily.csv'), ('saturation_events', 'saturation_intervals.csv')]:
                save_csv(sub / filename, result[key], key, pcm_id, identity)
            save_json(sub / 'annual_summary.json', {'pcm': result['summary'], 'L0': result['matched_control_summary']}, 'per-case annual metrics', pcm_id, identity)
            save_json(sub / 'validation.json', result['validation'], 'hard numerical validation', pcm_id, identity)
            save_json(sub / 'metadata.json', {**common, 'identity': identity, 'material': result['material'], 'control_material': result['control_material'], 'initial_condition': result['initial_condition'], 'numerical_extension': result['numerical_extension']}, 'PCM/run metadata', pcm_id, identity)
            for case, filename in [('pcm', 'cell_states_pcm.npz'), ('L0', 'cell_states_L0.npz')]:
                states = result['cell_snapshots'][case]
                centers = (np.arange(handoff['config']['N_cells']) + 0.5) * handoff['config']['d_m'] / handoff['config']['N_cells']
                state_path = sub / filename
                np.savez_compressed(state_path, datetime_UTC_ns=pd.DatetimeIndex(states['datetime_UTC']).as_unit('ns').asi8, T_C=states['T_C'], h_Jkg=states['h_Jkg'], f=states['f'], x_center_m=centers)
                with np.load(state_path, allow_pickle=False) as restored:
                    for name in ['T_C', 'h_Jkg', 'f']:
                        c1_require(np.array_equal(restored[name], states[name]), 'NPZ state roundtrip mismatch.')
                register(state_path, f'{case} boundary depth states', pcm_id, identity)
        manifest = {**common, 'status': 'COMPLETE_REAL_ANNUAL', 'created_at_UTC': datetime.now(timezone.utc).isoformat(), 'candidate_ids': handoff['candidate_ids'], 'by_pcm_directories': folders, 'validation': handoff['validation'], 'result_content_sha256': handoff['result_content_sha256'], 'cache_policy': 'NO_AUTOMATIC_REUSE; RECOMPUTE', 'artifacts': records, 'manifest_note': 'Written last; does not hash itself. All q into roof; H and delta_H roof into atmosphere.'}
        (staging / 'category1_manifest.json').write_text(json.dumps(c1_json_value(manifest), indent=2, allow_nan=False), encoding='utf-8')
        for relative, record in records.items():
            c1_require(c1_sha256_file(staging / relative) == record['sha256'], 'Output file changed while writing manifest.')
        staging.rename(final)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    handoff['artifact_directory'] = str(final)
    handoff['manifest'] = manifest
    for pcm_id, folder in manifest['by_pcm_directories'].items():
        handoff['by_pcm'][pcm_id]['persisted_artifacts'] = {name: str(final / relative) for relative in manifest['artifacts'] if relative.startswith(folder + '/') for name in [Path(relative).name]}
    return final

def c1_plot_style():
    return {'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False, 'axes.grid': True, 'grid.alpha': 0.18, 'figure.dpi': 110, 'savefig.dpi': 170}

def c1_plot_liquid_fraction(by_pcm, cfg):
    with plt.rc_context(c1_plot_style()):
        fig, axes = plt.subplots(2, 1, figsize=(11, 7), constrained_layout=True)
        palette = ['#28638a', '#ac7626', '#765290', '#577044', '#ba6251']
        for i, (pcm_id, record) in enumerate(by_pcm.items()):
            daily = record['calendar_daily']
            times = pd.to_datetime(daily.date_IST)
            axes[0].plot(times, daily.f_mean, label=pcm_id, color=palette[i % len(palette)], linestyle=['-', '--', ':', '-.'][i // len(palette) % 4], linewidth=1.2)
            monthly = record['monthly']
            axes[1].plot(monthly.month, monthly.f_mean, marker='o', label=pcm_id, color=palette[i % len(palette)], linestyle=['-', '--', ':', '-.'][i // len(palette) % 4])
        axes[0].set(title='Daily mean depth-averaged liquid fraction · Delhi 2025', ylabel='Liquid fraction (1)', ylim=(-0.02, 1.02))
        axes[1].set(title='Monthly mean liquid fraction · same continuous annual run', xlabel='Local calendar month', ylabel='Liquid fraction (1)', xticks=np.arange(1, 13), ylim=(-0.02, 1.02))
        axes[0].legend(ncols=min(4, len(by_pcm)), fontsize=9)
        return fig

def c1_plot_morning_reset(by_pcm, cfg):
    count = len(by_pcm)
    with plt.rc_context(c1_plot_style()):
        fig, axes = plt.subplots(count, 1, figsize=(11, max(3, 2.1 * count)), squeeze=False, constrained_layout=True)
        for ax, (pcm_id, record) in zip(axes.flat, by_pcm.items()):
            daily = record['daily']
            times = pd.DatetimeIndex(daily.cycle_end_IST).tz_localize(None)
            ax.plot(times, daily.f_next_morning, color='#28638a', linewidth=1.1, label='Next-morning f')
            ax.axhline(cfg['reset_limit'], color='#ac7626', linestyle='--', label=f"Incomplete reset above {cfg['reset_limit']:.2f}")
            known = int(daily.f_next_morning.notna().sum())
            ax.set(title=f'{pcm_id} · 06:00 IST next-morning state · {known} known checkpoints', ylabel='Liquid fraction (1)', ylim=(-0.02, 1.02))
            ax.legend(loc='upper right', fontsize=9)
        axes[-1, 0].set_xlabel('Next-morning local date · unavailable final checkpoint remains a gap')
        return fig

def c1_plot_mechanism(record, start_IST, end_IST):
    start, end = (pd.Timestamp(start_IST), pd.Timestamp(end_IST))
    start = start.tz_localize(C1_IST) if start.tzinfo is None else start.tz_convert(C1_IST)
    end = end.tz_localize(C1_IST) if end.tzinfo is None else end.tz_convert(C1_IST)
    c1_require(end > start, 'Plot window must have positive length.')
    p = record['hourly']
    mask = (p.datetime_IST >= start) & (p.datetime_end_IST <= end)
    c1_require(mask.any(), 'Mechanism plot has no observed rows in the declared window.')
    x = p.loc[mask]
    delta = record['delta_H'].loc[mask]
    t = pd.DatetimeIndex(x.datetime_IST).tz_localize(None) + pd.Timedelta(minutes=30)
    t_end = pd.DatetimeIndex(x.datetime_end_IST).tz_localize(None)
    states = record['cell_snapshots']['pcm']
    state_rows = np.flatnonzero(mask.to_numpy()) + 1
    with plt.rc_context(c1_plot_style()):
        fig, axes = plt.subplots(4, 1, figsize=(11, 10), sharex=True, constrained_layout=True)
        axes[0].plot(t, x.T_air_C, color='#777777', linestyle='--', label='Air midpoint sample')
        axes[0].plot(t, x.T_surface_C, color='#ac7626', label='Surface hourly mean')
        axes[0].plot(t_end, states['T_C'][state_rows, 0], color='#28638a', linestyle=':', label='Top cell at hour end')
        axes[0].plot(t_end, states['T_C'][state_rows, -1], color='#765290', linestyle='-.', label='Bottom cell at hour end')
        axes[0].set_ylabel('Temperature (°C)')
        axes[1].plot(t, x.f_mean, color='#28638a', label='Layer mean, hourly')
        axes[1].plot(t_end, states['f'][state_rows, 0], color='#ac7626', linestyle=':', label='Top cell at hour end')
        axes[1].plot(t_end, states['f'][state_rows, -1], color='#765290', linestyle='--', label='Bottom cell at hour end')
        axes[1].set(ylabel='Liquid fraction (1)', ylim=(-0.02, 1.02))
        axes[2].plot(t, delta.H_PCM_Wm2, color='#28638a', label='H_PCM')
        axes[2].plot(t, delta.H_L0_Wm2, color='#777777', linestyle='--', label='H_L0')
        axes[2].axhline(0, color='#333333', linewidth=0.6)
        axes[2].set_ylabel('Roof→air H (W/m²)')
        axes[3].plot(t, delta.delta_H_Wm2, color='#ac7626', label='ΔH = H_PCM − H_L0')
        axes[3].axhline(0, color='#333333', linewidth=0.6)
        axes[3].set(ylabel='ΔH (W/m²)', xlabel='Local time (Asia/Kolkata) · flux means at interval midpoint; cell states at hour end')
        for ax in axes:
            ax.legend(loc='upper left', ncols=2, fontsize=9)
        fig.suptitle(f'{x.pcm_id.iloc[0]} · physical mechanism · {start:%d %b}–{end:%d %b %Y}')
        return fig

def c1_plot_comparison(summary):
    with plt.rc_context(c1_plot_style()):
        fig, axes = plt.subplots(1, 3, figsize=(12, max(3.5, 0.35 * len(summary))), constrained_layout=True)
        labels = summary.pcm_id.astype(str)
        axes[0].barh(labels, summary.equivalent_full_uptake_cycles, color='#28638a')
        axes[0].set(title='Actual latent cycling', xlabel='Equivalent uptake cycles (1)', xlim=(0, None))
        axes[1].barh(labels, summary.raw_saturation_h, color='#ac7626')
        axes[1].set(title='Near-full-liquid duration', xlabel='Hours with f ≥ 0.99 (h)', xlim=(0, None))
        axes[2].barh(labels, summary.reset_fraction_of_known_cycles, color='#577044')
        axes[2].set(title='Next-morning reset', xlabel='Fraction of known cycles (1)', xlim=(0, 1))
        fig.suptitle('Candidate comparison · annual Delhi 2025 · separate metrics, no weighted score')
        return fig

C3_DELIVERY_CODE_SHA256 = '4197a0620730109cd0f176b968ceedb4d124e990e9f0eb2c78add168156ec68c'

C3_VERSION = 'C3-FV-MUSCL-SSPRK2-STRANG-1.0'

C3_LIMITATION = 'Effective vertically mixed urban air-layer PCM-minus-control anomaly; one-way exploratory model; not pedestrian temperature, absolute air temperature, resolved canyons, measured mitigation or human heat exposure.'

C3_BASE = dict(H_c_m=40.0, tau_s=1800.0, K_h_m2s=20.0, r_U=0.6, eta=0.7, mixing='constant', boundary='open')

C3_RANGES = dict(H_c_m=(20.0, 80.0), tau_s=(900.0, 3600.0), K_h_m2s=(5.0, 50.0), r_U=(0.3, 1.0))

C3_METRICS = ('mean_K', 'G_cool_K_m2_s', 'G_warm_K_m2_s', 'max_cooling_K', 'max_warming_K')

class Category3Error(RuntimeError):
    pass

def c3_require(condition, message):
    if not bool(condition):
        raise Category3Error(message)

def c3_json(value):
    if isinstance(value, dict):
        return {str(k): c3_json(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [c3_json(v) for v in value]
    if isinstance(value, (Path, pd.Timestamp)):
        return str(value)
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return c3_json(value.item())
    if isinstance(value, float) and (not np.isfinite(value)):
        return None
    return value

def c3_digest(value):
    return hashlib.sha256(json.dumps(c3_json(value), sort_keys=True, separators=(',', ':'), allow_nan=False).encode()).hexdigest()

def c3_array_hash(a):
    a = np.ascontiguousarray(a)
    return hashlib.sha256(str(a.dtype).encode() + str(a.shape).encode() + a.tobytes()).hexdigest()

def c3_file_hash(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def c3_write_json(p, v):
    p = Path(p)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(p.suffix + '.tmp')
    tmp.write_text(json.dumps(c3_json(v), indent=2, allow_nan=False))
    tmp.replace(p)

def c3_live_identity(namespace):

    def code(c):
        return {'bytes': c.co_code.hex(), 'constants': [code(v) if inspect.iscode(v) else repr(v) for v in c.co_consts], 'names': c.co_names, 'vars': c.co_varnames}
    bodies = {}
    for n, f in namespace.items():
        if n.startswith('c3_') and callable(f):
            p = getattr(f, 'py_func', f)
            if hasattr(p, '__code__'):
                bodies[n] = code(p.__code__)
    return c3_digest({'functions': bodies, 'version': C3_VERSION, 'base': C3_BASE, 'ranges': C3_RANGES, 'source': namespace.get('C3_DELIVERY_CODE_SHA256'), 'environment': {'python': platform.python_version(), 'numpy': np.__version__, 'pandas': pd.__version__, 'numba': __import__('numba').__version__, 'scipy': __import__('scipy').__version__, 'shapely': __import__('shapely').__version__, 'pyproj': __import__('pyproj').__version__}})

_C3_REAL_INPUT_SEAL = object()

@dataclass
class C3Inputs:
    weather: object
    by_pcm: dict
    ids: tuple
    hotspot: object
    cells: list
    grid: object
    location: dict
    identity: dict
    upstream: object
    nature: str = 'REAL'
    _seal: object = None

def c3_validate_handoff(ns):
    h = ns.get('CATEGORY1_HANDOFF')
    p = ns.get('PRELIM_HANDOFF')
    w = ns.get('C1_VALIDATED_INPUTS')
    c3_require(ns.get('PRELIM_READY') is True and ns.get('CATEGORY1_READY') is True and isinstance(h, dict) and isinstance(p, dict), 'Category 3 requires completed real PRELIM_READY and CATEGORY1_READY handoffs.')
    c3_require(h.get('status') == 'COMPLETE_REAL_ANNUAL' and h.get('scientific_reality_status') == 'REAL' and (h.get('_real_input_seal') is ns.get('_C1_REAL_INPUT_SEAL')) and (ns.get('_C1_REAL_INPUT_SEAL') is not None), 'Missing real Category-1 validation seal; synthetic/partial inputs are refused.')
    c3_require(w is not None and w._seal is ns['_C1_REAL_INPUT_SEAL'], 'Validated Category-1 input bundle required.')
    current = ns['c1_validate_input_bundle'](w.weather, w.candidates, list(w.ids), w.metadata, ns['CATEGORY1_IO'])
    c3_require(current.identity == w.identity == h['provenance']['input_identity'], 'Upstream inputs changed.')
    c3_require(ns['c1_live_code_identity']() == h['provenance']['code_sha256'], 'Category-1 executed code changed.')
    c3_require(ns['c1_result_seal'](h) == h['result_content_sha256'], 'Category-1 result seal changed.')
    ids = tuple(w.ids)
    c3_require(ids and tuple(h['by_pcm']) == ids and (tuple(h['candidate_ids']) == ids), 'Exhaustive candidate population mismatch.')
    for pid in ids:
        r = h['by_pcm'][pid]
        d = r['delta_H']
        c3_require(len(d) == 8760 and d.attrs.get('data_nature') == 'REAL' and d.scientific_reality_status.eq('REAL').all(), 'Only complete real annual Category-1 forcing is accepted.')
        c3_require(np.array_equal(d.H_PCM_Wm2, -r['hourly'].q_convection_mean_Wm2) and np.array_equal(d.H_L0_Wm2, -r['matched_control'].q_convection_mean_Wm2) and np.array_equal(d.delta_H_Wm2, d.H_PCM_Wm2 - d.H_L0_Wm2), 'Roof-to-air sign/value mismatch.')
        c3_require(pd.DatetimeIndex(pd.to_datetime(d.datetime_UTC, utc=True)).equals(pd.DatetimeIndex(pd.to_datetime(w.weather.interval_start_UTC, utc=True))), 'Forcing/weather hour mismatch.')
        c3_require(pd.DatetimeIndex(pd.to_datetime(d.datetime_end_UTC, utc=True)).equals(pd.DatetimeIndex(pd.to_datetime(w.weather.interval_end_UTC, utc=True))), 'Forcing/weather interval-end mismatch.')
    mask = np.asarray(p['HOTSPOT_MASK'], bool)
    g = p['COMMON_GRID']
    tr = g.transform
    c3_require(str(g.crs) == 'EPSG:32643' and abs(tr.a - 150) < 1e-08 and (abs(tr.e + 150) < 1e-08) and (tr.b == 0) and (tr.d == 0), 'Required aligned 150 m EPSG:32643 grid is absent.')
    c3_require(mask.shape == g.shape and mask.sum() == 44, 'Approved 44-cell hotspot must remain unchanged.')
    location = p['HOTSPOT_WEATHER_LOCATION']
    c3_require(ns['numeric_array_hash'](mask) == location['hotspot_mask_numeric_sha256'], 'Hotspot numeric hash differs from approved weather-location record.')
    cells = []
    for row, col in np.argwhere(mask):
        x, y = tr * (int(col), int(row))
        cells.append(box(x, y - 150, x + 150, y))
    hotspot = unary_union(cells)
    c3_require(abs(hotspot.area - 44 * 22500) < 1e-06, 'Hotspot union geometry/area mismatch.')
    return C3Inputs(w.weather.copy(deep=True), h['by_pcm'], ids, hotspot, cells, g, deepcopy(location), {'category1_result_seal': h['result_content_sha256'], 'category1_code': h['provenance']['code_sha256'], 'input_identity': w.identity, 'workbook': h['workbook_identity'], 'weather': h['weather_identity'], 'preliminary': w.metadata, 'hotspot_wkb_sha256': hashlib.sha256(hotspot.wkb).hexdigest(), 'grid': ns['grid_metadata'](g), 'location': deepcopy(location), 'ordered_cell_geometry_sha256': c3_digest([c.wkb.hex() for c in cells]), 'forcing_sha256': {pid: ns['c1_table_digest'](h['by_pcm'][pid]['delta_H']) for pid in ids}}, w, _seal=_C3_REAL_INPUT_SEAL)

def c3_check_consumed_inputs(inputs, ns, roofs=None, pid=None):
    c3_require(isinstance(inputs, C3Inputs) and inputs._seal is _C3_REAL_INPUT_SEAL and (inputs.nature == 'REAL'), 'Category 3 requires its validated real input seal; manual/synthetic bundles are refused.')
    c3_require(ns['c1_table_digest'](inputs.weather) == inputs.identity['weather']['content_sha256'], 'Consumed weather changed after validation; refuse stale scientific identity.')
    c3_require(c3_json(inputs.location) == c3_json(inputs.identity['location']) and c3_digest([c.wkb.hex() for c in inputs.cells]) == inputs.identity['ordered_cell_geometry_sha256'], 'Consumed location/cell geometry changed after validation.')
    if pid is not None:
        c3_require(ns['c1_table_digest'](inputs.by_pcm[pid]['delta_H']) == inputs.identity['forcing_sha256'][pid], 'Category-1 forcing changed after validation; no cached/scientific substitution.')
    if roofs is not None:
        meta = roofs['metadata']
        c3_require(hashlib.sha256(roofs['geometry'].wkb).hexdigest() == meta['roof_geometry_sha256'] and c3_array_hash(roofs['areas']) == meta['cell_roof_area_sha256'] and (c3_digest(meta) == roofs['identity']), 'Validated roof geometry/areas/provenance changed.')

def c3_reporting_root(ns):
    inputs = ns.get('C3_INPUTS')
    roofs = ns.get('C3_ROOFS')
    cfg = ns['C3_EXECUTION']
    c3_require(inputs is not None and roofs is not None, 'Validated inputs/buildings required for scoped report persistence.')
    controls = {k: v for k, v in cfg.items() if k not in {'batch_size', 'upstream_resume_dir', 'chunk_hours', 'resume_verified_chunks'}}
    key = c3_digest({'inputs': inputs.identity, 'roofs': roofs['identity'], 'source': ns['CUMULATIVE_SOURCE_IDENTITY'], 'code': c3_live_identity(ns), 'execution_controls': controls})
    folder = Path(ns['CONFIG']['output_dir']) / 'category3' / 'reports' / ('study_' + key)
    folder.mkdir(parents=True, exist_ok=True)
    return folder

C3_BUILDING_ASSET = 'GOOGLE/Research/open-buildings/v3/polygons'

C3_BUILDING_CATALOG = 'https://developers.google.com/earth-engine/datasets/catalog/GOOGLE_Research_open-buildings_v3_polygons'

def c3_acquire_roofs(inputs, ns, cfg):
    c3_check_consumed_inputs(inputs, ns)
    to_ll = Transformer.from_crs('EPSG:32643', 'EPSG:4326', always_xy=True).transform
    to_xy = Transformer.from_crs('EPSG:4326', 'EPSG:32643', always_xy=True).transform
    threshold = float(cfg['building_confidence_min'])
    c3_require(0.65 <= threshold <= 1, 'Open Buildings confidence threshold must be in [0.65,1].')
    key = c3_digest({'asset': C3_BUILDING_ASSET, 'geometry': inputs.identity['hotspot_wkb_sha256'], 'threshold': threshold, 'method': 'union_projected_polygons_v1'})
    folder = Path(ns['CONFIG']['cache_dir']) / 'category3_roofs' / key
    folder.mkdir(parents=True, exist_ok=True)
    gp = folder / 'footprints.geojson'
    mp = folder / 'acquisition.json'
    if gp.exists() and mp.exists():
        meta = json.loads(mp.read_text())
        c3_require(meta['key'] == key and c3_file_hash(gp) == meta['sha256'], 'Building cache identity failed.')
        features = json.loads(gp.read_text())['features']
    else:
        if ns.get('PRELIM_EE_SESSION') is None:
            import ee
            ee.Authenticate(auth_mode='colab' if ns.get('IN_COLAB') else 'localhost')
            ee.Initialize(project=ns['PRELIM_CONFIG']['ee_project'])
            ns['ee'] = ee
            ns['PRELIM_EE_SESSION'] = {'project': ns['PRELIM_CONFIG']['ee_project'], 'purpose': 'Open Buildings only after explicit upstream restore'}
        ee = ns['ee']
        aoi = ee.Geometry(mapping(geometry_transform(to_ll, inputs.hotspot)))
        try:
            fc = ee.FeatureCollection(C3_BUILDING_ASSET).filterBounds(aoi).filter(ee.Filter.gte('confidence', threshold)).sort('system:index')
            count = int(fc.size().getInfo())
            c3_require(0 < count <= cfg['building_max_features'], 'Empty or excessive Open Buildings extraction; review source/limit.')
            features = []
            for offset in range(0, count, cfg['building_page_size']):
                page = fc.toList(min(cfg['building_page_size'], count - offset), offset).getInfo()
                c3_require(len(page) == min(cfg['building_page_size'], count - offset), 'Incomplete building acquisition page.')
                features.extend(page)
            c3_require(len(features) == count, 'Building acquisition is incomplete.')
        except Exception as e:
            raise Category3Error('Google Open Buildings unavailable or extraction failed. No roof source substitution is permitted.') from e
        c3_write_json(gp, {'type': 'FeatureCollection', 'features': features})
        meta = {'key': key, 'sha256': c3_file_hash(gp), 'data_nature': 'REAL', 'asset': C3_BUILDING_ASSET, 'catalog': C3_BUILDING_CATALOG, 'inference_epoch': 'May 2023', 'study_year': 2025, 'epoch_mismatch': 'Static inferred 2023 footprints used for 2025; not surveyed 2025 roofs.', 'confidence_min': threshold, 'features': count, 'scope': 'treatment hotspot only', 'acquired_UTC': pd.Timestamp.now(tz='UTC').isoformat(), 'license': 'CC-BY-4.0; attribution Google Research Open Buildings', 'footprint_proxy': 'roof-plan proxy, with no retrofit/structural certification'}
        c3_write_json(mp, meta)
    geometries = []
    seen = set()
    raw_area = 0.0
    for f in features:
        pr = f['properties']
        geom = shape(f['geometry'])
        c3_require({'confidence', 'area_in_meters', 'full_plus_code', 'longitude_latitude'} <= set(pr), 'Official building schema missing.')
        c3_require(threshold <= float(pr['confidence']) <= 1 and float(pr['area_in_meters']) > 0, 'Invalid building confidence/area.')
        c3_require(geom.geom_type in {'Polygon', 'MultiPolygon'} and geom.is_valid and (not geom.is_empty), 'Invalid footprint polygon; stop for review.')
        ident = f.get('id', c3_digest(f))
        c3_require(ident not in seen, 'Duplicate feature ID from pagination; stop for review.')
        seen.add(ident)
        projected = geometry_transform(to_xy, geom)
        clipped = projected.intersection(inputs.hotspot)
        if not clipped.is_empty:
            geometries.append(clipped)
            raw_area += clipped.area
    c3_require(geometries, 'No valid buildings intersect treatment cells.')
    roof = unary_union(geometries)
    areas = np.array([roof.intersection(c).area for c in inputs.cells], float)
    fractions = areas / 22500.0
    c3_require(np.isfinite(fractions).all() and ((fractions >= 0) & (fractions <= 1 + 1e-12)).all(), 'Roof fractions outside [0,1].')
    c3_require(abs(areas.sum() - roof.area) <= max(1e-05, roof.area * 1e-10), 'Cell intersections do not reconcile to union roof area.')
    meta.update(union_area_m2=roof.area, overlap_removed_m2=raw_area - roof.area, roof_geometry_sha256=hashlib.sha256(roof.wkb).hexdigest(), cell_roof_area_sha256=c3_array_hash(areas), intersection='valid projected polygon union, exact cell intersections; no raster approximation')
    c3_write_json(folder / 'roof_validation.json', meta)
    return {'geometry': roof, 'areas': areas, 'fractions': fractions, 'metadata': meta, 'identity': c3_digest(meta)}

def c3_layout(inputs, roofs, coverage, arrangement, eta):
    c3_require(0 <= coverage <= 1 and 0 < eta <= 1, 'Invalid coverage/usable fraction.')
    eligible = np.asarray(roofs['areas'], float) * eta
    c3_require(np.isfinite(eligible).all() and (eligible >= 0).all() and (eligible <= 22500).all(), 'Invalid eligible roof areas.')
    centers = np.array([[g.centroid.x, g.centroid.y] for g in inputs.cells])
    ids = np.flatnonzero(eligible > 0).tolist()
    target = eligible.sum() * coverage
    allocation = np.zeros(len(eligible))
    if coverage == 1:
        allocation = eligible.copy()
        arrangement = 'all_roofs'
        order = ids
    elif arrangement == 'dispersed':
        allocation = coverage * eligible
        order = ids
    elif arrangement == 'roof_area_priority':
        order = sorted(ids, key=lambda i: (-eligible[i], i))
    elif arrangement == 'clustered':
        centroid = np.average(centers, axis=0, weights=eligible)
        seed = min(ids, key=lambda i: (np.sum((centers[i] - centroid) ** 2), i))
        order = [seed]
        remaining = set(ids) - {seed}
        while remaining:
            adjacent = [i for i in remaining if any((np.isclose(np.abs(centers[i] - centers[j]).sum(), 150) for j in order))]
            if adjacent:
                nxt = min(adjacent, key=lambda i: (np.sum((centers[i] - centers[seed]) ** 2), i))
            else:
                nxt = min(remaining, key=lambda i: (min((np.sum((centers[i] - centers[j]) ** 2) for j in order)), np.sum((centers[i] - centers[seed]) ** 2), i))
            order.append(nxt)
            remaining.remove(nxt)
    else:
        raise Category3Error('Unknown deployment arrangement.')
    if coverage < 1 and arrangement != 'dispersed':
        remaining_area = target
        for i in order:
            allocation[i] = min(eligible[i], max(0.0, remaining_area))
            remaining_area -= allocation[i]
    tolerance = max(1e-06, 1e-10 * max(target, 1))
    c3_require(abs(allocation.sum() - target) <= tolerance and (allocation <= eligible + tolerance).all() and (allocation >= 0).all(), 'Equal-treated-roof-area constraint failed.')
    fraction = np.divide(allocation, eligible, out=np.zeros_like(allocation), where=eligible > 0)
    return {'coverage': coverage, 'arrangement': arrangement, 'eta': eta, 'areas_m2': allocation, 'c': fraction, 'target_m2': target, 'actual_m2': float(allocation.sum()), 'eligible_total_m2': float(eligible.sum()), 'order': order, 'tolerance_m2': tolerance, 'identity': c3_digest({'areas': allocation, 'eta': eta, 'coverage': coverage, 'arrangement': arrangement})}

def c3_domain(inputs, roofs, layout, buffer_m, dx=150.0):
    cache_key = c3_digest({'roofs': roofs['identity'], 'layout': layout['identity'], 'buffer': buffer_m, 'dx': dx})
    if not hasattr(inputs, '_domain_cache'):
        inputs._domain_cache = {}
    if cache_key in inputs._domain_cache:
        return inputs._domain_cache[cache_key]
    c3_require(dx in (150.0, 75.0), 'Only approved alignment or nested 75 m verification grid permitted.')
    xmin, ymin, xmax, ymax = inputs.hotspot.bounds
    left = math.floor((xmin - buffer_m) / dx) * dx
    bottom = math.floor((ymin - buffer_m) / dx) * dx
    right = math.ceil((xmax + buffer_m) / dx) * dx
    top = math.ceil((ymax + buffer_m) / dx) * dx
    nx = int(round((right - left) / dx))
    ny = int(round((top - bottom) / dx))
    source = np.zeros((ny, nx))
    weights = np.zeros((3, ny, nx))
    radius = np.zeros((ny, nx))
    fixed = inputs.hotspot.buffer(2000.0)
    union_bounds = fixed.bounds
    centroid = inputs.hotspot.centroid
    for i in range(ny):
        y = bottom + i * dx
        for j in range(nx):
            x = left + j * dx
            cell = box(x, y, x + dx, y + dx)
            radius[i, j] = cell.centroid.distance(centroid)
            if x + dx >= union_bounds[0] and x <= union_bounds[2] and (y + dx >= union_bounds[1]) and (y <= union_bounds[3]):
                weights[0, i, j] = inputs.hotspot.intersection(cell).area
                weights[1, i, j] = fixed.intersection(cell).area
            if weights[0, i, j] > 0:
                for k, parent in enumerate(inputs.cells):
                    overlap = cell.intersection(parent)
                    if not overlap.is_empty and overlap.area > 0:
                        source[i, j] += roofs['geometry'].intersection(overlap).area * layout['eta'] * layout['c'][k] / (dx * dx)
                weights[2, i, j] = source[i, j] * dx * dx
    c3_require(abs(source.sum() * dx * dx - layout['actual_m2']) <= layout['tolerance_m2'], 'Atmospheric source treated-area mismatch.')
    c3_require(np.allclose(weights[:2].sum(axis=(1, 2)), [inputs.hotspot.area, fixed.area], rtol=1e-10, atol=1e-05), 'Fixed receptor areas are truncated or inconsistent.')
    result = dict(source=source, weights=weights, radius=radius, dx=dx, shape=(ny, nx), bounds=(left, bottom, right, top), buffer_m=buffer_m, identity=c3_digest({'bounds': (left, bottom, right, top), 'dx': dx, 'source': c3_array_hash(source), 'weights': c3_array_hash(weights)}))
    if len(inputs._domain_cache) >= 8:
        inputs._domain_cache.pop(next(iter(inputs._domain_cache)))
    inputs._domain_cache[cache_key] = result
    return result

def c3_solar_elevation(times, lon, lat):
    t = pd.DatetimeIndex(pd.to_datetime(times, utc=True))
    hour = t.hour + t.minute / 60 + t.second / 3600 + t.microsecond / 3600000000.0
    gamma = 2 * np.pi / (365 + np.asarray(t.is_leap_year, int)) * (t.dayofyear - 1 + (hour - 12) / 24)
    e = 229.18 * (7.5e-05 + 0.001868 * np.cos(gamma) - 0.032077 * np.sin(gamma) - 0.014615 * np.cos(2 * gamma) - 0.040849 * np.sin(2 * gamma))
    d = 0.006918 - 0.399912 * np.cos(gamma) + 0.070257 * np.sin(gamma) - 0.006758 * np.cos(2 * gamma) + 0.000907 * np.sin(2 * gamma) - 0.002697 * np.cos(3 * gamma) + 0.00148 * np.sin(3 * gamma)
    ha = np.deg2rad((hour * 60 + e + 4 * lon) / 4 - 180)
    la = np.deg2rad(lat)
    return np.rad2deg(np.arcsin(np.clip(np.sin(la) * np.sin(d) + np.cos(la) * np.cos(d) * np.cos(ha), -1, 1)))

def c3_time_segments(weather, location):
    lon, lat = location['point_lon_lat']
    starts = pd.DatetimeIndex(pd.to_datetime(weather.interval_start_UTC, utc=True))
    segments = np.zeros((len(starts), 4))
    day_seconds = np.zeros(len(starts))
    for k, t in enumerate(starts):
        a, b = (float(c3_solar_elevation([t], lon, lat)[0]), float(c3_solar_elevation([t + pd.Timedelta(hours=1)], lon, lat)[0]))
        if a * b < 0:
            cross = brentq(lambda s: float(c3_solar_elevation([t + pd.Timedelta(seconds=s)], lon, lat)[0]), 0, 3600, xtol=0.01)
            segments[k] = [cross, float(a > 0), 3600 - cross, float(b > 0)]
            day_seconds[k] = cross * (a > 0) + (3600 - cross) * (b > 0)
        else:
            segments[k] = [3600, float(a + b > 0), 0, 0]
            day_seconds[k] = 3600 * (a + b > 0)
    return (segments, day_seconds)

def c3_projected_wind(weather, location):
    lon, lat = location['point_lon_lat']
    proj = Transformer.from_crs('EPSG:4326', 'EPSG:32643', always_xy=True)
    geod = Geod(ellps='WGS84')
    x, y = proj.transform(lon, lat)
    ex, ey = proj.transform(*geod.fwd(lon, lat, 90, 10)[:2])
    nx, ny = proj.transform(*geod.fwd(lon, lat, 0, 10)[:2])
    east = np.array([ex - x, ey - y])
    east /= np.linalg.norm(east)
    north = np.array([nx - x, ny - y])
    north /= np.linalg.norm(north)
    c3_require(abs(np.dot(east, north)) < 1e-05, 'Projection basis is not orthogonal.')
    return weather.u10_ms.to_numpy()[:, None] * east + weather.v10_ms.to_numpy()[:, None] * north

@njit
def c3_minmod(a, b):
    if a * b <= 0:
        return 0.0
    return math.copysign(min(abs(a), abs(b)), a)

@njit
def c3_transport(T, u, v, K, dx, boundary):
    """Conservative MUSCL/minmod advection and horizontal diffusion on the original atmospheric grid."""
    ny, nx = T.shape
    sx = np.zeros_like(T)
    sy = np.zeros_like(T)
    for i in range(ny):
        for j in range(1, nx - 1):
            sx[i, j] = c3_minmod(T[i, j] - T[i, j - 1], T[i, j + 1] - T[i, j])
    for i in range(1, ny - 1):
        for j in range(nx):
            sy[i, j] = c3_minmod(T[i, j] - T[i - 1, j], T[i + 1, j] - T[i, j])
    fx = np.zeros((ny, nx + 1))
    fy = np.zeros((ny + 1, nx))
    for i in range(ny):
        for j in range(1, nx):
            face = T[i, j - 1] + 0.5 * sx[i, j - 1] if u >= 0 else T[i, j] - 0.5 * sx[i, j]
            fx[i, j] = u * face - K * (T[i, j] - T[i, j - 1]) / dx
        fx[i, 0] = u * T[i, 0] if u < 0 else 0.0
        fx[i, nx] = u * T[i, nx - 1] if u > 0 else 0.0
        if boundary == 1:
            fx[i, 0] -= 2 * K * T[i, 0] / dx
            fx[i, nx] += 2 * K * T[i, nx - 1] / dx
    for j in range(nx):
        for i in range(1, ny):
            face = T[i - 1, j] + 0.5 * sy[i - 1, j] if v >= 0 else T[i, j] - 0.5 * sy[i, j]
            fy[i, j] = v * face - K * (T[i, j] - T[i - 1, j]) / dx
        fy[0, j] = v * T[0, j] if v < 0 else 0.0
        fy[ny, j] = v * T[ny - 1, j] if v > 0 else 0.0
        if boundary == 1:
            fy[0, j] -= 2 * K * T[0, j] / dx
            fy[ny, j] += 2 * K * T[ny - 1, j] / dx
    rhs = np.empty_like(T)
    for i in range(ny):
        for j in range(nx):
            rhs[i, j] = -(fx[i, j + 1] - fx[i, j] + fy[i + 1, j] - fy[i, j]) / dx
    outward = (fx[:, nx].sum() - fx[:, 0].sum() + fy[ny, :].sum() - fy[0, :].sum()) * dx
    return (rhs, outward)

@njit
def c3_step(T, source, u, v, K, dx, tau, dt, boundary):
    """Advance the anomaly with SSPRK2 transport and exact Strang source/exchange splitting."""
    a = math.exp(-dt / (2 * tau)) if tau < 1e+99 else 1.0
    b = tau * -math.expm1(-dt / (2 * tau)) if tau < 1e+99 else dt / 2
    first = T * a + source * b
    source_integral = source.sum() * dx * dx * dt
    exchange_first = (T.sum() - first.sum()) * dx * dx + source_integral / 2
    rhs0, out0 = c3_transport(first, u, v, K, dx, boundary)
    stage = first + dt * rhs0
    rhs1, out1 = c3_transport(stage, u, v, K, dx, boundary)
    transported = 0.5 * first + 0.5 * (stage + dt * rhs1)
    final = transported * a + source * b
    exchange_second = (transported.sum() - final.sum()) * dx * dx + source_integral / 2
    lateral = dt * 0.5 * (out0 + out1)
    exchange = exchange_first + exchange_second
    residual = (final.sum() - T.sum()) * dx * dx - source_integral + exchange + lateral
    return (final, source_integral, exchange, lateral, residual)

@njit
def c3_stable_dt(u, v, K, dx, boundary, cfl, cap):
    rate = 2 * (abs(u) + abs(v)) / dx + (6 if boundary == 1 else 4) * K / (dx * dx)
    return min(cap, cfl / rate) if rate > 0 else cap

C3_RECEPTORS = ('hotspot', 'buffer2km', 'treated_roof')

C3_RCOLS = ('mean_K', 'unweighted_mean_K', 'G_cool_K_m2_s', 'G_warm_K_m2_s', 'max_cooling_K', 'max_warming_K', 'cooling_area_mean_m2', 'warming_area_mean_m2', 'extent_cooling_max_m', 'extent_warming_max_m', 'day_integral_K_m2_s', 'night_integral_K_m2_s', 'day_cool_K_m2_s', 'night_warm_K_m2_s')

C3_BCOLS = ('source_J', 'storage_change_J', 'exchange_J', 'lateral_out_J', 'residual_J', 'max_step_residual_J', 'budget_tolerance_J', 'substeps', 'dt_min_s', 'dt_max_s', 'state_start_J', 'state_end_J')

@njit
def c3_chunk(T, field_integral, peak_cool, peak_warm, weights, radius, F, dh, wind, segments, params, dx, cfl, cap, threshold):
    """Integrate hourly receptor means, gross responses, and instantaneous cell extrema across adaptive substeps."""
    H, tau0, K, rU, mixing, boundary = params
    C = 1.2 * 1005 * H
    hours = len(dh)
    records = np.zeros((hours, 3 * 14 + 12 + 24))
    best_c = peak_cool.min()
    best_w = peak_warm.max()
    for h in range(hours):
        start_energy = T.sum() * dx * dx * C
        budget = np.zeros(7)
        ns = 0
        dtmin = 1e+99
        dtmax = 0.0
        rc = np.zeros((3, 14))
        rc[:, 4] = 0.0
        rc[:, 5] = 0.0
        threshold_stats = np.zeros((2, 3, 4))
        S = dh[h] * F / C
        u = wind[h, 0] * rU
        v = wind[h, 1] * rU
        for seg in range(2):
            duration = segments[h, seg * 2]
            day = segments[h, seg * 2 + 1] > 0
            tau = (1200.0 if day else 3600.0) if mixing == 1 else tau0
            stable = c3_stable_dt(u, v, K, dx, int(boundary), cfl, cap)
            if duration <= 0:
                continue
            steps = int(math.ceil(duration / stable))
            dt = duration / steps
            for step in range(steps):
                new, src, ex, lat, res = c3_step(T, S, u, v, K, dx, tau, dt, int(boundary))
                ns += 1
                dtmin = min(dtmin, dt)
                dtmax = max(dtmax, dt)
                budget[0] += src * C
                budget[2] += ex * C
                budget[3] += lat * C
                budget[4] += res * C
                budget[5] = max(budget[5], abs(res * C))
                budget[6] += abs(src * C) + abs(ex * C) + abs(lat * C) + abs((new.sum() - T.sum()) * dx * dx * C)
                for i in range(T.shape[0]):
                    for j in range(T.shape[1]):
                        field_integral[i, j] += 0.5 * (T[i, j] + new[i, j]) * dt
                        for r in range(3):
                            w = weights[r, i, j]
                            if w <= 0:
                                continue
                            z = 0.5 * (T[i, j] + new[i, j])
                            cool = 0.5 * (max(-T[i, j], 0) + max(-new[i, j], 0))
                            warm = 0.5 * (max(T[i, j], 0) + max(new[i, j], 0))
                            rc[r, 0] += z * w * dt
                            rc[r, 1] += z * dt
                            rc[r, 2] += cool * w * dt
                            rc[r, 3] += warm * w * dt
                            rc[r, 4] = max(rc[r, 4], -T[i, j], -new[i, j])
                            rc[r, 5] = max(rc[r, 5], T[i, j], new[i, j])
                            rc[r, 6] += 0.5 * (int(T[i, j] < -threshold) + int(new[i, j] < -threshold)) * w * dt
                            rc[r, 7] += 0.5 * (int(T[i, j] > threshold) + int(new[i, j] > threshold)) * w * dt
                            for ei in range(2):
                                factor = 0.5 / 0.7 if ei == 0 else 0.9 / 0.7
                                th = threshold / factor
                                threshold_stats[ei, r, 0] += 0.5 * (int(T[i, j] < -th) + int(new[i, j] < -th)) * w * dt / 3600
                                threshold_stats[ei, r, 1] += 0.5 * (int(T[i, j] > th) + int(new[i, j] > th)) * w * dt / 3600
                                if min(T[i, j], new[i, j]) < -th:
                                    threshold_stats[ei, r, 2] = max(threshold_stats[ei, r, 2], radius[i, j])
                                if max(T[i, j], new[i, j]) > th:
                                    threshold_stats[ei, r, 3] = max(threshold_stats[ei, r, 3], radius[i, j])
                            if min(T[i, j], new[i, j]) < -threshold:
                                rc[r, 8] = max(rc[r, 8], radius[i, j])
                            if max(T[i, j], new[i, j]) > threshold:
                                rc[r, 9] = max(rc[r, 9], radius[i, j])
                            if day:
                                rc[r, 10] += z * w * dt
                                rc[r, 12] += cool * w * dt
                            else:
                                rc[r, 11] += z * w * dt
                                rc[r, 13] += warm * w * dt
                T = new
                if T.min() < best_c:
                    best_c = T.min()
                    peak_cool = T.copy()
                if T.max() > best_w:
                    best_w = T.max()
                    peak_warm = T.copy()
        end_energy = T.sum() * dx * dx * C
        budget[1] = end_energy - start_energy
        residual = budget[1] - budget[0] + budget[2] + budget[3]
        tolerance = 1e-05 + 1e-10 * budget[6]
        for r in range(3):
            area = weights[r].sum()
            count = (weights[r] > 0).sum()
            if area > 0:
                rc[r, 0] /= area * 3600
            if count > 0:
                rc[r, 1] /= count * 3600
            rc[r, 6] /= 3600
            rc[r, 7] /= 3600
            for k in range(14):
                records[h, r * 14 + k] = rc[r, k]
        tail = np.array([budget[0], budget[1], budget[2], budget[3], residual, budget[5], tolerance, ns, dtmin, dtmax, start_energy, end_energy])
        records[h, 42:54] = tail
        records[h, 54:] = threshold_stats.ravel()
    return (T, field_integral, peak_cool, peak_warm, records)

def c3_parameter_vector(p):
    c3_require(all((np.isfinite(p[k]) and p[k] > 0 for k in ['H_c_m', 'tau_s', 'r_U', 'eta'])) and p['K_h_m2s'] >= 0, 'Invalid atmospheric parameters.')
    c3_require(p['mixing'] in {'constant', 'daynight'} and p['boundary'] in {'open', 'zero'}, 'Invalid boundary/mixing scenario.')
    c3_require(all((C3_RANGES[k][0] <= p[k] <= C3_RANGES[k][1] for k in C3_RANGES)) and 0.5 <= p['eta'] <= 0.9, 'Atmospheric/roof assumptions must stay within approved exploratory ranges.')
    return np.array([p['H_c_m'], p['tau_s'], p['K_h_m2s'], p['r_U'], float(p['mixing'] == 'daynight'), float(p['boundary'] == 'zero')])

def c3_hourly_table(a, weather, day_seconds, cutoff):
    columns = [r + '_' + c for r in C3_RECEPTORS for c in C3_RCOLS] + list(C3_BCOLS) + [f'eta{eta}_{r}_{c}' for eta in [0.5, 0.9] for r in C3_RECEPTORS for c in ['cooling_area_mean_m2', 'warming_area_mean_m2', 'extent_cooling_max_m', 'extent_warming_max_m']]
    t = pd.DataFrame(a, columns=columns)
    t['datetime_UTC'] = pd.to_datetime(weather.interval_start_UTC, utc=True).to_numpy()
    local = pd.DatetimeIndex(pd.to_datetime(weather.interval_start_UTC, utc=True)).tz_convert('Asia/Kolkata') + pd.Timedelta(minutes=30)
    t['local_month'] = local.month
    t['local_hour'] = local.hour
    t['day_seconds'] = day_seconds
    t['night_seconds'] = 3600 - day_seconds
    t['T_air_C'] = weather.T_air_C.to_numpy(float)
    t['extreme_heat'] = t.T_air_C >= cutoff
    return (t, cutoff)

def c3_summarize(t, domain):
    """Aggregate saved hourly receptor statistics without changing spatial or temporal weighting."""
    out = {}
    duration = len(t) * 3600
    for r, k in zip(C3_RECEPTORS, range(3)):
        area = float(domain['weights'][k].sum())
        prefix = r + '_'
        v = {'mean_K': float(t[prefix + 'mean_K'].mean()), 'unweighted_mean_K': float(t[prefix + 'unweighted_mean_K'].mean()), 'G_cool_K_m2_s': float(t[prefix + 'G_cool_K_m2_s'].sum()), 'G_warm_K_m2_s': float(t[prefix + 'G_warm_K_m2_s'].sum()), 'max_cooling_K': float(t[prefix + 'max_cooling_K'].max()), 'max_warming_K': float(t[prefix + 'max_warming_K'].max()), 'cooling_hours_spatial_mean': int((t[prefix + 'mean_K'] < 0).sum()), 'warming_hours_spatial_mean': int((t[prefix + 'mean_K'] > 0).sum()), 'any_cell_cooling_hours': int((t[prefix + 'max_cooling_K'] > 0).sum()), 'any_cell_warming_hours': int((t[prefix + 'max_warming_K'] > 0).sum()), 'cooling_area_mean_m2': float(t[prefix + 'cooling_area_mean_m2'].mean()), 'warming_area_mean_m2': float(t[prefix + 'warming_area_mean_m2'].mean()), 'extent_cooling_max_m': float(t[prefix + 'extent_cooling_max_m'].max()), 'extent_warming_max_m': float(t[prefix + 'extent_warming_max_m'].max()), 'evaluation_area_m2': area, 'duration_s': duration, 'day_mean_K': float(t[prefix + 'day_integral_K_m2_s'].sum() / max(area * t.day_seconds.sum(), 1e-100)), 'night_mean_K': float(t[prefix + 'night_integral_K_m2_s'].sum() / max(area * t.night_seconds.sum(), 1e-100)), 'day_G_cool_K_m2_s': float(t[prefix + 'day_cool_K_m2_s'].sum()), 'night_G_warm_K_m2_s': float(t[prefix + 'night_warm_K_m2_s'].sum()), 'extreme_heat_mean_K': float(t.loc[t.extreme_heat, prefix + 'mean_K'].mean()), 'peak_weather_hour_mean_K': float(t.loc[t.T_air_C.idxmax(), prefix + 'mean_K'])}
        v['G_net_K_m2_s'] = v['G_cool_K_m2_s'] - v['G_warm_K_m2_s']
        c3_require(abs(v['G_net_K_m2_s'] + v['mean_K'] * area * duration) <= 1e-05 + 1e-10 * (v['G_cool_K_m2_s'] + v['G_warm_K_m2_s']), 'Gross cooling/warming integrals disagree with signed area-time mean.')
        previous = t[prefix + 'mean_K'].lt(0).shift(1).rolling(6, min_periods=1).max().fillna(0).astype(bool)
        delayed = (t[prefix + 'mean_K'] > 0) & previous
        v['delayed_warming_hours_6h'] = int(delayed.sum())
        v['delayed_G_warm_K_m2_s'] = float(t.loc[delayed, prefix + 'G_warm_K_m2_s'].sum())
        out[r] = v
    c3_require(np.isfinite(t.select_dtypes('number').to_numpy()).all(), 'Nonfinite scientific output.')
    c3_require((t.residual_J.abs() <= t.budget_tolerance_J).all() and (t.max_step_residual_J <= t.budget_tolerance_J).all(), 'Discrete atmospheric energy closure failed; completion withheld.')
    if len(t) > 1:
        c3_require(np.allclose(t.state_start_J.iloc[1:], t.state_end_J.iloc[:-1], rtol=0, atol=1e-05), 'Atmospheric temporal continuity failed.')
    return out

def c3_run(inputs, roofs, layout, params, buffer_m, ns, cfg, *, dx=150.0, dt_factor=1.0, override_dh=None, label='oneway', hours=None):
    c3_require(inputs.nature == 'REAL', 'Scientific run refuses synthetic provenance.')
    c3_require(np.isfinite(dt_factor) and 0 < dt_factor <= 1 and np.isfinite(cfg['cfl']) and (0 < cfg['cfl'] <= 0.45) and np.isfinite(cfg['dt_cap_s']) and (0 < cfg['dt_cap_s'] <= 3600) and (cfg['chunk_hours'] >= 1) and (cfg['effect_threshold_K'] > 0), 'Invalid numerical timestep/chunk/threshold configuration; unsafe execution refused.')
    c3_require(c3_upstream_seal_valid(inputs, ns), 'Scientific input seal is missing.')
    c3_check_consumed_inputs(inputs, ns, roofs, params['pcm_id'])
    c3_require(params['eta'] == layout['eta'], 'Source-map eta and atmospheric scenario eta differ.')
    eligible = roofs['areas'] * layout['eta']
    expected_c = np.divide(layout['areas_m2'], eligible, out=np.zeros_like(eligible), where=eligible > 0)
    c3_require(np.allclose(expected_c, layout['c'], rtol=0, atol=1e-12) and c3_digest({'areas': layout['areas_m2'], 'eta': layout['eta'], 'coverage': layout['coverage'], 'arrangement': layout['arrangement']}) == layout['identity'], 'Deployment changed after area/identity validation.')
    pid = params['pcm_id']
    d = inputs.by_pcm[pid]['delta_H']
    n = 8760 if hours is None else int(hours)
    c3_require(n in (48, 8760), 'Scientific run must be the annual calendar or declared 48-hour pilot.')
    offset = int(cfg.get('pilot_start_hour', 0)) if n == 48 else 0
    c3_require(0 <= offset <= 8760 - n, 'Pilot window outside canonical year.')
    domain = c3_domain(inputs, roofs, layout, buffer_m, dx)
    dh = d.delta_H_Wm2.to_numpy(float)[offset:offset + n] if override_dh is None else np.asarray(override_dh, float)[offset:offset + n]
    c3_require(dh.shape == (n,) and np.isfinite(dh).all(), 'Invalid real/derived forcing series.')
    weather = inputs.weather.iloc[offset:offset + n]
    wind = c3_projected_wind(weather, inputs.location)
    segment_key = c3_digest({'weather': inputs.identity['weather'], 'location': inputs.location, 'hours': n, 'start_hour': offset})
    if not hasattr(inputs, '_solar_cache'):
        inputs._solar_cache = {}
    if segment_key not in inputs._solar_cache:
        inputs._solar_cache[segment_key] = c3_time_segments(weather, inputs.location)
    segments, days = inputs._solar_cache[segment_key]
    identity = {'input': inputs.identity, 'roofs': roofs['identity'], 'layout': layout['identity'], 'parameters': params, 'domain': domain['identity'], 'dt_factor': dt_factor, 'cfl': cfg['cfl'], 'dt_cap_s': cfg['dt_cap_s'], 'effect_threshold_K': cfg['effect_threshold_K'], 'hours': n, 'start_hour': offset, 'forcing_sha256': c3_array_hash(dh), 'forcing_nature': 'DERIVED_FEEDBACK_SCENARIO' if override_dh is not None else 'REAL_CATEGORY1_HOURLY_MEAN', 'label': label, 'code': c3_live_identity(ns), 'computational_source': ns['CUMULATIVE_SOURCE_IDENTITY'], 'rho_air_kgm3': 1.2, 'cp_air_JkgK': 1005.0, 'initial': 'ZERO_ANOMALY_NO_SPINUP'}
    eid = c3_digest(identity)
    folder = Path(ns['CONFIG']['output_dir']) / 'category3' / eid
    folder.mkdir(parents=True, exist_ok=True)
    cached = folder / 'experiment_manifest.json'
    if cached.exists() and cfg['resume_verified_chunks']:
        m = json.loads(cached.read_text())
        if m['identity'] == c3_json(identity) and m.get('validation', {}).get('energy') == 'PASS' and (m.get('validation', {}).get('continuity') == 'PASS'):
            compact = {'experiment_id': eid, 'identity': identity, 'params': params, 'layout': layout, 'domain': domain, 'summary': m['summary'], 'folder': str(folder), 'hours': n, 'validation': m['validation'], 'extreme_heat_threshold_C': m['extreme_heat_threshold_C'], 'limitation': C3_LIMITATION, 'elapsed_wall_s': m.get('timing', {}).get('actual_wall_s', 0.0), 'cache_reused': True}
            print('Verified accepted result cache:', eid[:12])
            return c3_load_result(c3_compact_result(compact))
    T = np.zeros(domain['shape'])
    integ = T.copy()
    pc = T.copy()
    pw = T.copy()
    records = []
    chain = c3_array_hash(T)
    start = time.perf_counter()
    c3_write_json(folder / 'run_identity.json', {'status': 'PARTIAL', 'identity': identity, 'experiment_id': eid, 'limitation': C3_LIMITATION})
    for a in range(0, n, cfg['chunk_hours']):
        b = min(a + cfg['chunk_hours'], n)
        npz = folder / f'chunk_{a:05d}_{b:05d}.npz'
        meta = npz.with_suffix('.json')
        if npz.exists() and meta.exists() and cfg['resume_verified_chunks']:
            m = json.loads(meta.read_text())
            c3_require(m['experiment_id'] == eid and m['incoming_state_hash'] == chain and (m['sha256'] == c3_file_hash(npz)), 'Checkpoint mismatch; refuse stale/modified scientific cache.')
            with np.load(npz, allow_pickle=False) as z:
                T = z['T']
                integ = z['integral']
                pc = z['peak_cool']
                pw = z['peak_warm']
                arr = z['records']
        else:
            T, integ, pc, pw, arr = c3_chunk(T, integ, pc, pw, domain['weights'], domain['radius'], domain['source'], dh[a:b], wind[a:b], segments[a:b], c3_parameter_vector(params), dx, cfg['cfl'] * dt_factor, cfg['dt_cap_s'] * dt_factor, cfg['effect_threshold_K'])
            c3_require(np.isfinite(arr).all() and np.isfinite(T).all(), 'Nonfinite atmospheric chunk.')
            c3_require((abs(arr[:, 46]) <= arr[:, 48]).all(), 'Chunk discrete energy closure failed.')
            temp = npz.with_suffix('.tmp.npz')
            np.savez_compressed(temp, T=T, integral=integ, peak_cool=pc, peak_warm=pw, records=arr)
            temp.replace(npz)
            c3_write_json(meta, {'status': 'PARTIAL_CHECKPOINT', 'experiment_id': eid, 'incoming_state_hash': chain, 'sha256': c3_file_hash(npz), 'hour_start': a, 'hour_end': b})
        chain = c3_array_hash(T)
        records.append(arr)
        print(f'{pid} {label}: {b}/{n} h · {eid[:10]}', flush=True)
    table, cutoff = c3_hourly_table(np.vstack(records), weather, days, float(inputs.weather.T_air_C.quantile(0.95)))
    summary = c3_summarize(table, domain)
    result = {'experiment_id': eid, 'identity': identity, 'params': params, 'layout': layout, 'domain': domain, 'hourly': table, 'summary': summary, 'mean_field_K': integ / (n * 3600), 'peak_cool_field_K': pc, 'peak_warm_field_K': pw, 'final_state_K': T, 'extreme_heat_threshold_C': cutoff, 'folder': str(folder), 'hours': n, 'validation': {'energy': 'PASS', 'continuity': 'PASS'}, 'elapsed_wall_s': time.perf_counter() - start, 'limitation': C3_LIMITATION}
    c3_persist_result(result)
    return result

def c3_persist_result(r):
    folder = Path(r['folder'])
    t = r['hourly']
    t.to_csv(folder / 'hourly_statistics.csv', index=False)
    np.savez_compressed(folder / 'spatial_diagnostics.npz', mean_K=r['mean_field_K'], peak_cool_K=r['peak_cool_field_K'], peak_warm_K=r['peak_warm_field_K'], final_K=r['final_state_K'], source=r['domain']['source'], weights=r['domain']['weights'])
    c3_write_json(folder / 'summary.json', r['summary'])
    for column in ['local_month', 'local_hour', 'extreme_heat']:
        grouped = []
        for key, g in t.groupby(column):
            for rec in C3_RECEPTORS:
                grouped.append({'group': str(key), 'receptor': rec, 'hours': len(g), 'mean_K': g[rec + '_mean_K'].mean(), 'G_cool_K_m2_s': g[rec + '_G_cool_K_m2_s'].sum(), 'G_warm_K_m2_s': g[rec + '_G_warm_K_m2_s'].sum()})
        pd.DataFrame(grouped).to_csv(folder / (column + '_summary.csv'), index=False)
    seasons = {1: 'winter', 2: 'winter', 3: 'premonsoon', 4: 'premonsoon', 5: 'premonsoon', 6: 'monsoon', 7: 'monsoon', 8: 'monsoon', 9: 'monsoon', 10: 'postmonsoon', 11: 'postmonsoon', 12: 'winter'}
    t.assign(season=t.local_month.map(seasons)).groupby('season')[[x + '_mean_K' for x in C3_RECEPTORS]].mean().to_csv(folder / 'seasonal_summary.csv')
    c3_write_json(folder / 'experiment_manifest.json', {'status': 'PILOT_COMPLETE_NOT_ANNUAL' if r['hours'] == 48 else 'ANNUAL_COMPUTED_NUMERICAL_REFINEMENTS_PENDING', 'experiment_id': r['experiment_id'], 'identity': r['identity'], 'validation': r['validation'], 'summary': r['summary'], 'extreme_heat_criterion': 'Tair >= 95th percentile of full annual canonical real hourly weather; no atmospheric outcome selection', 'extreme_heat_threshold_C': r['extreme_heat_threshold_C'], 'thermal_extent_threshold_K': r['identity']['effect_threshold_K'], 'units': {'temperature': 'K anomaly', 'G': 'K m2 s', 'energy': 'J', 'extent': 'm; radial distance from hotspot geometric centroid'}, 'timing': {'actual_wall_s': r['elapsed_wall_s'], 'cache_reused': r.get('cache_reused', False), 'transport_substeps': int(t.substeps.sum()), 'grid_cells': int(np.prod(r['domain']['shape'])), 'estimate_same_grid_annual_s_from_pilot': r['elapsed_wall_s'] * 8760 / 48 if r['hours'] == 48 else None, 'estimate_caveat': 'Pilot projection includes compilation/cache overhead and different annual wind/step counts; not an annual benchmark.'}, 'limitations': r['limitation'], 'result_use': 'DEVELOPMENT', 'candidate_population_final': False, 'artifacts': {p.name: c3_file_hash(p) for p in folder.iterdir() if p.is_file() and p.name != 'experiment_manifest.json' and ('.tmp' not in p.name)}})

def c3_compare_metrics(a, b, relative_limit, kind):
    rows = []
    for rec in C3_RECEPTORS:
        for metric in C3_METRICS:
            x = float(a['summary'][rec][metric])
            y = float(b['summary'][rec][metric])
            diff = abs(y - x)
            area = a['summary'][rec]['evaluation_area_m2']
            duration = a['summary'][rec]['duration_s']
            absolute = 0.0001 * (area * duration if metric.startswith('G_') else 1.0)
            relative = diff / max(abs(y), 1e-100)
            rows.append({'test': kind, 'receptor': rec, 'metric': metric, 'coarse': x, 'refined': y, 'absolute_change': diff, 'relative_change': relative, 'relative_limit': relative_limit, 'absolute_limit': absolute, 'pass': bool(diff <= absolute or relative < relative_limit)})
    return rows

def c3_validate_annual(r, inputs, roofs, ns, cfg):
    if r['validation'].get('numerical_status') == 'PASS':
        return r
    c3_require(r['hours'] == 8760, 'Only annual runs qualify for annual numerical acceptance.')
    p = r['params']
    l = r['layout']
    buffer = r['domain']['buffer_m']
    rows = []
    dt = c3_run(inputs, roofs, l, p, buffer, ns, cfg, dt_factor=0.5, label='timestep_refinement')
    rows += c3_compare_metrics(r, dt, 0.02, 'timestep')
    grid = c3_run(inputs, roofs, l, p, buffer, ns, cfg, dx=75.0, label='grid_refinement')
    rows += c3_compare_metrics(r, grid, 0.05, 'grid')
    larger = c3_run(inputs, roofs, l, p, buffer + 3000, ns, cfg, label='domain_refinement')
    rows += c3_compare_metrics(r, larger, 0.05, 'domain')
    if p['boundary'] == 'open':
        zero = c3_run(inputs, roofs, l, {**p, 'boundary': 'zero'}, buffer + 3000, ns, cfg, label='boundary_sensitivity')
        rows += c3_compare_metrics(larger, zero, 0.05, 'boundary')
    table = pd.DataFrame(rows)
    table.to_csv(Path(r['folder']) / 'numerical_convergence.csv', index=False)
    r['validation'].update(numerical_status='PASS' if table['pass'].all() else 'FAIL', convergence_rows=rows, refined_experiments={'timestep': dt['experiment_id'], 'grid': grid['experiment_id'], 'domain': larger['experiment_id']})
    c3_persist_result(r)
    c3_require(table['pass'].all(), 'Annual numerical acceptance failed. Inspect numerical_convergence.csv; enlarge domain/refine resolution or timestep explicitly and rerun. No ready handoff.')
    return r

def c3_select_domain(inputs, roofs, ns, cfg):
    rows = []
    accepted = []
    for pid in inputs.ids:
        for coverage, arr in [(1.0, 'all_roofs'), (0.5, 'clustered')]:
            if cfg.get('domain_selection_case') and [pid, coverage] != cfg['domain_selection_case']:
                continue
            layout = c3_layout(inputs, roofs, coverage, 'dispersed' if coverage == 1 else arr, 0.7)
            p = {**C3_BASE, 'pcm_id': pid, 'H_c_m': 20.0, 'tau_s': 3600.0, 'K_h_m2s': 5.0, 'r_U': 0.3, 'name': 'domain_selection'}
            previous = None
            found = None
            for buffer in cfg['domain_candidates_m']:
                result = c3_run(inputs, roofs, layout, p, buffer, ns, cfg, label='domain_selection')
                if previous is not None:
                    pair = c3_compare_metrics(previous, result, 0.05, 'domain_selection')
                    rows.extend([{**v, 'pcm_id': pid, 'coverage': coverage, 'arrangement': arr, 'smaller_buffer_m': previous['domain']['buffer_m'], 'larger_buffer_m': buffer} for v in pair])
                    if all((v['pass'] for v in pair)) and buffer >= 10000:
                        found = buffer
                        break
                previous = result
            c3_require(found is not None, 'No domain converged through configured larger buffers; extend domain_candidates_m and rerun. No ready handoff.')
            accepted.append(found)
    folder = c3_reporting_root(ns)
    folder.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_csv(folder / ('domain_selection_' + c3_digest(cfg.get('domain_selection_case', 'all'))[:12] + '.csv'), index=False)
    return max(accepted)

def c3_run_baselines(inputs, roofs, ns, cfg, buffer):
    records = []
    n = 48 if cfg['mode'] == 'REAL_48H_PILOT' else 8760
    coverage_levels = cfg['coverage_fractions'] if cfg['mode'] == 'REAL_FULL_EXPERIMENT' else [1.0]
    for pid in inputs.ids:
        for coverage in coverage_levels:
            for arr in ['dispersed'] if coverage == 1 else cfg['arrangements']:
                layout = c3_layout(inputs, roofs, coverage, arr, 0.7)
                p = {**C3_BASE, 'pcm_id': pid, 'name': 'baseline'}
                r = c3_run(inputs, roofs, layout, p, buffer, ns, cfg, hours=n)
                if n == 8760:
                    c3_validate_annual(r, inputs, roofs, ns, cfg)
                records.append(c3_compact_result(r))
    return records

def c3_save_figure(fig, path):
    import matplotlib.pyplot as plt
    fig.savefig(path, dpi=220, bbox_inches='tight')
    plt.close(fig)

def c3_plot_forcing(inputs, folder):
    import matplotlib.pyplot as plt
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    rows = []
    common_limit = 1e-09
    for pid in inputs.ids:
        delta = inputs.by_pcm[pid]['delta_H']
        local = pd.DatetimeIndex(pd.to_datetime(delta.datetime_UTC, utc=True)).tz_convert('Asia/Kolkata') + pd.Timedelta(minutes=30)
        composite = pd.DataFrame({'month': local.month, 'hour': local.hour, 'dh': delta.delta_H_Wm2.to_numpy()}).groupby(['month', 'hour']).dh.mean()
        common_limit = max(common_limit, float(abs(composite).max()))
    for i, pid in enumerate(inputs.ids):
        d = inputs.by_pcm[pid]['delta_H']
        local = pd.DatetimeIndex(pd.to_datetime(d.datetime_UTC, utc=True)).tz_convert('Asia/Kolkata') + pd.Timedelta(minutes=30)
        t = pd.DataFrame({'month': local.month, 'hour': local.hour, 'delta_H': d.delta_H_Wm2.to_numpy()})
        mat = t.pivot_table(index='month', columns='hour', values='delta_H', aggfunc='mean')
        limit = common_limit
        fig, ax = plt.subplots(figsize=(10, 4), layout='constrained')
        im = ax.imshow(mat, cmap='RdBu_r', vmin=-limit, vmax=limit, aspect='auto', origin='lower', extent=(-0.5, 23.5, 0.5, 12.5))
        ax.set(xlabel='Local interval midpoint hour · IST', ylabel='Month', title=f'{pid}: signed real Category-1 roof forcing\nΔH = HPCM − HL0; negative reduces heat release')
        ax.set_yticks(range(1, 13))
        fig.colorbar(im, ax=ax, label='Hourly mean ΔH (W m⁻²)')
        c3_save_figure(fig, folder / f'forcing_{i:03d}.png')
        t.groupby(['month', 'hour']).delta_H.mean().to_csv(folder / f'forcing_month_hour_{i:03d}.csv')
        rows.append({'pcm_id': pid, 'signed_annual_mean_Wm2': d.delta_H_Wm2.mean(), 'annual_integral_Jm2': d.delta_H_Wm2.sum() * 3600, 'cooling_forcing_hours': int(d.delta_H_Wm2.lt(0).sum()), 'warming_forcing_hours': int(d.delta_H_Wm2.gt(0).sum())})
    pd.DataFrame(rows).to_csv(folder / 'category1_forcing_signed_summary.csv', index=False)

def c3_plot_roofs_layouts(inputs, roofs, folder):
    import matplotlib.pyplot as plt
    from matplotlib.patches import Polygon
    from matplotlib.collections import PatchCollection
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)

    def plot_values(ax, values, title):
        patches = [Polygon(np.asarray(g.exterior.coords)) for g in inputs.cells]
        collection = PatchCollection(patches, cmap='cividis', edgecolor='white', linewidth=0.4)
        collection.set_array(np.asarray(values))
        collection.set_clim(0, 1)
        ax.add_collection(collection)
        ax.autoscale()
        ax.set_aspect('equal')
        ax.set(title=title, xlabel='UTM easting (m)', ylabel='UTM northing (m)')
        ax.ticklabel_format(useOffset=False, style='plain')
        return collection
    fig, ax = plt.subplots(figsize=(7, 6), layout='constrained')
    im = plot_values(ax, roofs['fractions'], 'Real inferred 2023 building footprint / 150 m cell\nRoof-plan proxy in unchanged treatment hotspot')
    fig.colorbar(im, ax=ax, label='Footprint fraction')
    c3_save_figure(fig, folder / 'roof_fraction.png')
    for coverage in [0.1, 0.25, 0.5, 0.75, 1.0]:
        arrangements = ['dispersed'] if coverage == 1 else ['dispersed', 'clustered', 'roof_area_priority']
        fig, axes = plt.subplots(1, len(arrangements), figsize=(6 * len(arrangements), 5), squeeze=False, layout='constrained')
        for ax, arr in zip(axes[0], arrangements):
            l = c3_layout(inputs, roofs, coverage, arr, 0.7)
            im = plot_values(ax, l['c'], f"{l['arrangement']} · {coverage:.0%}\n{l['actual_m2']:,.1f} m² treated · η=0.7")
        fig.colorbar(im, ax=axes[0].tolist(), label='Fraction of eligible roof area treated')
        c3_save_figure(fig, folder / f'layouts_{coverage:.2f}.png')

def c3_plot_result(r, color_limit=None):
    import matplotlib.pyplot as plt
    from matplotlib.colors import TwoSlopeNorm
    folder = Path(r['folder'])
    bounds = r['domain']['bounds']
    fields = [r['mean_field_K'], r['peak_cool_field_K'], r['peak_warm_field_K']]
    limit = max(max((abs(a).max() for a in fields)), 1e-09) if color_limit is None else max(color_limit, 1e-09)
    fig, axes = plt.subplots(1, 3, figsize=(17, 5), layout='constrained')
    for ax, a, title in zip(axes, fields, ['Annual signed mean' if r['hours'] == 8760 else '48-hour pilot signed mean', 'Snapshot at strongest cooling (zero if none)', 'Snapshot at strongest warming (zero if none)']):
        im = ax.imshow(a, origin='lower', extent=(bounds[0], bounds[2], bounds[1], bounds[3]), cmap='RdBu_r', norm=TwoSlopeNorm(vmin=-limit, vcenter=0, vmax=limit))
        ax.set(title=title, xlabel='UTM easting (m)', ylabel='UTM northing (m)')
        ax.ticklabel_format(useOffset=False, style='plain')
    fig.suptitle(f"{r['params']['pcm_id']} · {r['layout']['arrangement']} · {r['layout']['coverage']:.0%}\nModelled effective mixed-air-layer anomaly · prescribed sensitivity · DEVELOPMENT", fontsize=12)
    fig.colorbar(im, ax=axes, label='PCM − control air-layer ΔT (K); blue cooling, red warming')
    c3_save_figure(fig, folder / 'atmospheric_maps.png')
    t = r['hourly']
    monthly = t.groupby('local_month')[[f'hotspot_G_{name}_K_m2_s' for name in ['cool', 'warm']]].sum() / 3600000000.0
    fig, axes = plt.subplots(1, 3, figsize=(16, 4), layout='constrained')
    monthly.plot.bar(ax=axes[0], color=['#2166ac', '#b35806'])
    axes[0].set(title='Gross monthly cooling and warming', xlabel='Month', ylabel='K km² h')
    daynight = [r['summary']['hotspot'][k] for k in ['day_mean_K', 'night_mean_K']]
    axes[1].bar(['Day', 'Night'], daynight, color=['#2166ac', '#b35806'])
    axes[1].axhline(0, color='black', lw=0.8)
    axes[1].set(title='Solar elevation >0 vs ≤0', ylabel='Signed mean ΔT (K)')
    t.groupby('local_hour').hotspot_mean_K.mean().plot(ax=axes[2], color='#404040')
    axes[2].axhline(0, color='black', lw=0.8)
    axes[2].set(title='Local-hour composite; every hour retained', xlabel='IST interval midpoint hour', ylabel='Signed mean ΔT (K)')
    c3_save_figure(fig, folder / 'temporal_metrics.png')
    fig, ax = plt.subplots(figsize=(10, 4), layout='constrained')
    ax.plot(t.datetime_UTC, t.residual_J.abs(), lw=0.7, label='Absolute discrete residual')
    ax.plot(t.datetime_UTC, t.budget_tolerance_J, lw=0.7, label='Acceptance tolerance')
    ax.set(yscale='symlog', ylabel='J per hour', xlabel='UTC', title='Discrete source/storage/exchange/lateral closure')
    ax.legend()
    c3_save_figure(fig, folder / 'energy_closure.png')
    path = folder / 'numerical_convergence.csv'
    if path.exists():
        c = pd.read_csv(path)
        v = c[c.metric == 'mean_K']
        fig, ax = plt.subplots(figsize=(9, 4), layout='constrained')
        for test, g in v.groupby('test'):
            ax.plot(g.receptor, g.absolute_change, marker='o', label=test)
        ax.axhline(0.0001, color='#444444', ls='--', label='Near-zero absolute criterion')
        ax.set(ylabel='Absolute signed-mean change (K)', title='Annual fixed-receptor numerical refinement')
        ax.legend()
        c3_save_figure(fig, folder / 'numerical_convergence.png')

def c3_comparison_table(results):
    rows = []
    for r in results:
        for rec, metrics in r['summary'].items():
            rows.append({'experiment_id': r['experiment_id'], 'pcm_id': r['params']['pcm_id'], 'scenario': r['params']['name'], 'coverage': r['layout']['coverage'], 'arrangement': r['layout']['arrangement'], 'eta': r['layout']['eta'], 'actual_treated_area_m2': r['layout']['actual_m2'], 'receptor': rec, **metrics, 'numerical_status': r['validation'].get('numerical_status', 'NOT_ASSESSED'), 'result_use': 'DEVELOPMENT'})
    return pd.DataFrame(rows)

def c3_plot_comparisons(results, folder):
    import matplotlib.pyplot as plt
    folder = Path(folder)
    table = c3_comparison_table(results)
    table.to_csv(folder / 'experiment_comparisons.csv', index=False)
    for pid in table.pcm_id.unique():
        base = table[(table.pcm_id == pid) & (table.receptor == 'hotspot') & (table.scenario == 'baseline')]
        if not base.empty:
            fig, ax = plt.subplots(figsize=(9, 5), layout='constrained')
            for name, g in base.groupby('arrangement'):
                ax.plot(g.coverage, g.mean_K, marker='o', label=name)
            ax.axhline(0, color='black', lw=0.8)
            ax.legend()
            ax.set(title=f'{pid} · equal actual treated roof area at each coverage', xlabel='Fraction of eligible roof area treated', ylabel='Annual signed hotspot ΔT (K)')
            c3_save_figure(fig, folder / f'coverage_{c3_digest(pid)[:10]}.png')
        sens = table[(table.pcm_id == pid) & (table.receptor == 'hotspot') & (table.coverage == 1)]
        if len(sens) > 1:
            fig, ax = plt.subplots(figsize=(9, max(5, len(sens) * 0.22)), layout='constrained')
            ax.barh(sens.scenario, sens.mean_K, color='#2166ac')
            ax.axvline(0, color='black', lw=0.8)
            ax.set(title=f'{pid} · prescribed atmospheric/roof/mixing sensitivities\nDeterministic parameter exploration, not a confidence interval', xlabel='Annual signed ΔT (K)')
            c3_save_figure(fig, folder / f'sensitivity_{c3_digest(pid)[:10]}.png')
    mixing = table[(table.receptor == 'hotspot') & table.scenario.isin(['baseline', 'daynight_20_60min'])]
    if mixing.scenario.eq('daynight_20_60min').any():
        fig, ax = plt.subplots(figsize=(10, 5), layout='constrained')
        for name, g in mixing.groupby('scenario'):
            grouped = g.groupby(['pcm_id', 'coverage', 'arrangement']).mean_K.mean()
            ax.plot([str(k) for k in grouped.index], grouped.to_numpy(), marker='o', label=name)
        ax.axhline(0, color='black', lw=0.8)
        ax.tick_params(axis='x', rotation=45)
        ax.set(title='Constant 30 min vs illustrative solar day/night 20/60 min mixing', ylabel='Annual signed hotspot ΔT (K)')
        ax.legend()
        c3_save_figure(fig, folder / 'daynight_mixing_comparison.png')
    fb = folder / 'feedback_adequacy.csv'
    if fb.exists():
        a = pd.read_csv(fb)
        a = a[(a.receptor == 'hotspot') & (a.metric == 'mean_K')]
        fig, ax = plt.subplots(figsize=(8, 5), layout='constrained')
        for beta, g in a.groupby('beta'):
            ax.scatter(g.one_way, g.corrected, label=f'β={beta}', marker='o' if beta == 0.5 else 'x')
        limit = max(abs(a[['one_way', 'corrected']].to_numpy()).max(), 1e-06)
        ax.plot([-limit, limit], [-limit, limit], color='#333333', ls='--')
        ax.set(xlabel='One-way annual ΔT (K)', ylabel='Offline feedback annual ΔT (K)', title='Exploratory roof feedback · unchanged original L=0')
        ax.legend()
        c3_save_figure(fig, folder / 'feedback_comparison.png')

def c3_save_upstream_checkpoint(ns):
    import gzip
    inputs = c3_validate_handoff(ns)
    h = ns['CATEGORY1_HANDOFF']
    p = ns['PRELIM_HANDOFF']
    w = ns['C1_VALIDATED_INPUTS']
    identity = {'source': ns['CUMULATIVE_SOURCE_IDENTITY'], 'category1_code': ns['c1_live_code_identity'](), 'input_identity': w.identity, 'result_seal': h['result_content_sha256']}
    folder = Path(ns['CONFIG']['cache_dir']) / 'category3_upstream' / c3_digest(identity)
    folder.mkdir(parents=True, exist_ok=True)
    counter = [0]

    def pack(value):
        if isinstance(value, pd.DataFrame):
            return {'__kind__': 'frame', 'columns': list(value.columns), 'records': ns['c1_json_value'](value.to_dict('records')), 'attrs': ns['c1_json_value'](value.attrs), 'dtypes': {n: str(d) for n, d in value.dtypes.items()}, 'content_digest': ns['c1_table_digest'](value)}
        if isinstance(value, (pd.Index, pd.Series)):
            return {'__kind__': 'index', 'values': ns['c1_json_value'](list(value))}
        if isinstance(value, np.ndarray):
            counter[0] += 1
            name = f'array_{counter[0]:04d}.npz'
            np.savez_compressed(folder / name, value=value)
            return {'__kind__': 'array', 'file': name, 'sha256': c3_file_hash(folder / name)}
        if isinstance(value, dict):
            return {str(k): pack(v) for k, v in value.items() if k not in {'_real_input_seal', 'H_PCM', 'H_L0', 'DELTA_H'}}
        if isinstance(value, (list, tuple)):
            return [pack(v) for v in value]
        return ns['c1_json_value'](value)
    payload = {'identity': identity, 'cumulative_run_id': ns['CUMULATIVE_RUN_ID'], 'handoff': pack(h), 'validated_inputs': pack({'weather': w.weather, 'candidates': w.candidates, 'ids': w.ids, 'metadata': w.metadata}), 'preliminary': pack({k: p[k] for k in ['HOTSPOT_MASK', 'HOTSPOT_WEATHER_LOCATION', 'PRELIM_ARTIFACT_DIR', 'PRELIM_RUN_ID', 'PCM_SCREEN_METHOD', 'PCM_SCREEN_DERIVATION', 'PCM_PRELIM_CANDIDATE_IDS', 'CUMULATIVE_SOURCE_IDENTITY']}), 'grid': ns['grid_metadata'](p['COMMON_GRID'])}
    temp = folder / 'upstream.tmp.json.gz'
    with gzip.open(temp, 'wt', encoding='utf-8') as f:
        json.dump(payload, f, allow_nan=False)
    temp.replace(folder / 'upstream.json.gz')
    c3_write_json(folder / 'checkpoint_manifest.json', {'status': 'VALIDATED_REAL_UPSTREAM_CHECKPOINT', 'identity': identity, 'artifacts': {f.name: c3_file_hash(f) for f in folder.iterdir() if f.is_file() and f.name != 'checkpoint_manifest.json'}, 'policy': 'No pickle or executable deserialization; exact hashes, source/code/input/result seals rechecked on explicit restore.'})
    print('Upstream restart checkpoint:', folder)
    return str(folder)

def c3_restore_upstream_checkpoint(ns, path):
    import gzip
    folder = Path(path)
    m = json.loads((folder / 'checkpoint_manifest.json').read_text())
    ident = m['identity']
    c3_require(m['status'] == 'VALIDATED_REAL_UPSTREAM_CHECKPOINT', 'Incomplete upstream checkpoint.')
    c3_require(ident['source'] == ns['CUMULATIVE_SOURCE_IDENTITY'] and ident['category1_code'] == ns['c1_live_code_identity'](), 'Computational/Category-1 code changed. Checkpoint cannot be reused; rerun upstream.')
    for name, digest in m['artifacts'].items():
        c3_require(Path(name).name == name and c3_file_hash(folder / name) == digest, 'Upstream checkpoint hash/path failed.')
    with gzip.open(folder / 'upstream.json.gz', 'rt', encoding='utf-8') as f:
        payload = json.load(f)

    def unpack(v):
        if isinstance(v, list):
            return [unpack(a) for a in v]
        if not isinstance(v, dict):
            return v
        kind = v.get('__kind__')
        if kind == 'array':
            c3_require(Path(v['file']).name == v['file'] and c3_file_hash(folder / v['file']) == v['sha256'], 'Array identity/path failed.')
            with np.load(folder / v['file'], allow_pickle=False) as z:
                return z['value'].copy()
        if kind == 'index':
            return pd.DatetimeIndex(pd.to_datetime(v['values'], utc=True))
        if kind == 'frame':
            frame = pd.DataFrame(v['records'], columns=v['columns'])
            for name, dtype in v['dtypes'].items():
                if dtype.startswith('datetime64'):
                    series = pd.to_datetime(frame[name], utc=True)
                    if ', ' in dtype:
                        series = series.dt.tz_convert(dtype.split(', ', 1)[1].rstrip(']'))
                    frame[name] = series
                elif dtype != 'object':
                    frame[name] = frame[name].astype(dtype)
            frame.attrs = v['attrs']
            c3_require(ns['c1_table_digest'](frame) == v['content_digest'], 'Upstream frame serialization changed.')
            return frame
        return {k: unpack(a) for k, a in v.items()}
    raw = unpack(payload['validated_inputs'])
    h = unpack(payload['handoff'])
    p = unpack(payload['preliminary'])
    g = payload['grid']
    ns['CUMULATIVE_RUN_ID'] = payload['cumulative_run_id']
    verified_weather, verified_candidates, verified_ids, verified_meta = ns['c1_read_preliminary_directory'](raw['metadata']['artifact_directory'])
    ns['c1_same_table_values'](raw['weather'], verified_weather, 'restored weather')
    ns['c1_same_table_values'](raw['candidates'], verified_candidates, 'restored candidate population')
    c3_require(raw['ids'] == verified_ids and raw['metadata']['manifest_sha256'] == verified_meta['manifest_sha256'], 'Original Preliminary artifacts changed.')
    w = ns['c1_validate_input_bundle'](raw['weather'], raw['candidates'], raw['ids'], raw['metadata'], ns['CATEGORY1_IO'])
    c3_require(w.identity == ident['input_identity'], 'Upstream input identity changed; rerun.')
    h['_real_input_seal'] = ns['_C1_REAL_INPUT_SEAL']
    for r in h['by_pcm'].values():
        d = r['delta_H']
        r['H_PCM'] = d.set_index('datetime_UTC').H_PCM_Wm2
        r['H_L0'] = d.set_index('datetime_UTC').H_L0_Wm2
        r['DELTA_H'] = d.set_index('datetime_UTC').delta_H_Wm2
    c3_require(ns['c1_result_seal'](h) == ident['result_seal'] == h['result_content_sha256'], 'Restored Category-1 result seal mismatch.')
    p['COMMON_GRID'] = ns['GridSpec'](tuple(g['shape']), ns['CRS'].from_user_input(g['crs']), ns['Affine'](*g['transform']))
    p['CUMULATIVE_RUN_ID'] = ns['CUMULATIVE_RUN_ID']
    p['ANNUAL_WEATHER'] = w.weather
    p['PCM_PRELIM_CANDIDATES'] = w.candidates
    ns.update(PRELIM_HANDOFF=p, PRELIM_READY=True, CATEGORY1_HANDOFF=h, CATEGORY1_READY=True, C1_VALIDATED_INPUTS=w)
    c3_validate_handoff(ns)
    print('Exact verified upstream restored. Satellite/weather retrieval and Category-1 annual recomputation skipped.')

def c3_compact_result(r):
    drop = {'hourly', 'mean_field_K', 'peak_cool_field_K', 'peak_warm_field_K', 'final_state_K'}
    out = {k: v for k, v in r.items() if k not in drop}
    out['domain'] = {k: v for k, v in r['domain'].items() if k not in {'source', 'weights', 'radius'}}
    return out

def c3_load_result(r):
    if 'hourly' in r:
        return r
    folder = Path(r['folder'])
    m = json.loads((folder / 'experiment_manifest.json').read_text())
    c3_require(m['experiment_id'] == r['experiment_id'], 'Persisted result identity mismatch.')
    for name in ['hourly_statistics.csv', 'spatial_diagnostics.npz', 'summary.json']:
        c3_require(name in m['artifacts'] and c3_file_hash(folder / name) == m['artifacts'][name], 'Persisted result artifact mismatch.')
    if 'numerical_convergence.csv' in m['artifacts']:
        c3_require(c3_file_hash(folder / 'numerical_convergence.csv') == m['artifacts']['numerical_convergence.csv'], 'Convergence artifact changed.')
    out = {**r, 'hourly': pd.read_csv(folder / 'hourly_statistics.csv', float_precision='round_trip')}
    out['hourly']['datetime_UTC'] = pd.to_datetime(out['hourly'].datetime_UTC, utc=True)
    with np.load(folder / 'spatial_diagnostics.npz', allow_pickle=False) as z:
        out.update(mean_field_K=z['mean_K'], peak_cool_field_K=z['peak_cool_K'], peak_warm_field_K=z['peak_warm_K'], final_state_K=z['final_K'])
        out['domain'] = {**r['domain'], 'source': z['source'], 'weights': z['weights']}
    return out

def c3_refresh_manifest(r):
    path = Path(r['folder']) / 'experiment_manifest.json'
    m = json.loads(path.read_text())
    m['validation'] = r['validation']
    m['status'] = 'ANNUAL_NUMERICALLY_ACCEPTED' if r['validation'].get('numerical_status') == 'PASS' else m['status']
    m['artifacts'] = {p.name: c3_file_hash(p) for p in path.parent.iterdir() if p.is_file() and p.name != path.name and ('.tmp' not in p.name)}
    c3_write_json(path, m)

_C3_HISTORICAL_SEAL = object()

_c3_fresh_validate_handoff = c3_validate_handoff

def c3_verify_members(folder, manifest, ns, checked):
    for name, rec in manifest['artifacts'].items():
        path = (folder / name).resolve()
        c3_require(path.is_relative_to(folder.resolve()) and path.is_file(), f'Missing cached artifact: {path}')
        c3_require(c3_file_hash(path) == rec['sha256'], f'Cached artifact SHA-256 mismatch: {path}')
        checked[str(path)] = rec['sha256']
    return checked

def c3_verify_historical_files(inputs):
    for name, expected in inputs.upstream.file_hashes.items():
        c3_require(Path(name).is_file() and c3_file_hash(name) == expected, f'Historical upstream artifact missing/changed: {name}')

def c3_upstream_seal_valid(inputs, ns):
    if inputs.upstream is None:
        return False
    if inputs.upstream._seal is _C3_HISTORICAL_SEAL:
        return inputs._seal is _C3_REAL_INPUT_SEAL and inputs.identity.get('upstream_execution_route') == 'VERIFIED_HISTORICAL_ARTIFACTS'
    return ns.get('_C1_REAL_INPUT_SEAL') is not None and inputs.upstream._seal is ns['_C1_REAL_INPUT_SEAL']




def _private_metadata_value(record, role):
    """Read an explicitly configured private schema field without public labels."""
    fields = json.loads(os.environ.get('PCM_UHI_PRIVATE_METADATA_FIELDS', '{}'))
    if not isinstance(fields, dict):
        raise ValueError('PCM_UHI_PRIVATE_METADATA_FIELDS must be a JSON field mapping.')
    field = fields.get(role)
    return record.get(field) if field else None

def _source_fingerprint(identity):
    """Read scientific source fingerprints from current or private input metadata."""
    if not isinstance(identity, dict):
        return None
    content = identity.get('source_content_sha256', _private_metadata_value(identity, 'source_content'))
    executable = identity.get('executable_source_sha256')
    if not content or not executable:
        return None
    return (content, executable)


def _source_identity_matches(actual, expected):
    """Compare exact scientific fingerprints independently of source-file labels."""
    fingerprint = _source_fingerprint(actual)
    return fingerprint is not None and fingerprint == _source_fingerprint(expected)

def _private_dependency_path(environment):
    """Require a real user-supplied private dependency; invent no fallback path."""
    value = os.environ.get(environment, '').strip()
    if not value:
        raise ValueError(f'Set {environment} to the authorized private dependency path.')
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f'{environment} must identify an existing private dependency.')
    return path

def c3_historical_source_check(path, expected):
    source_record = json.loads(Path(path).read_text())
    cells = [c for c in source_record['cells'] if 'CUMULATIVE_FINGERPRINT_CELL' not in ''.join(c['source'])]
    normalized = [{'cell_type': c['cell_type'], 'source': ''.join(c['source'])} for c in cells]
    normalized_hash = hashlib.sha256(json.dumps(normalized, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    executable_hash = hashlib.sha256('\n\n'.join((''.join(c['source']) for c in cells if c['cell_type'] == 'code')).encode()).hexdigest()
    c3_require((normalized_hash, executable_hash) == _source_fingerprint(expected), 'Recorded computational source identity mismatch.')

def c3_load_historical(ns, root):
    started = time.perf_counter()
    root = Path(root).expanduser().resolve()
    preliminary = root / 'preliminary'
    bundle_path = root / 'upstream_bundle_manifest.json'
    c3_require(bundle_path.is_file(), f'Missing portable input manifest: {bundle_path}')
    bundle = json.loads(bundle_path.read_text())
    category1 = root / bundle['category1_directory']
    checked = {str(bundle_path): c3_file_hash(bundle_path)}
    for relative, sha in bundle['files'].items():
        path = (root / relative).resolve()
        c3_require(path.is_relative_to(root) and path.is_file(), f'Missing packaged historical file: {path}')
        c3_require(c3_file_hash(path) == sha, f'Packaged historical file SHA-256 mismatch: {path}')
        checked[str(path)] = sha
    pm = json.loads((preliminary / 'preliminary_manifest.json').read_text())
    m = json.loads((category1 / 'category1_manifest.json').read_text())
    completed = json.loads((_private_dependency_path('PCM_UHI_PRIVATE_COMPLETION_RECEIPT')).read_text())
    c3_verify_members(preliminary, pm, ns, checked)
    c3_verify_members(category1, m, ns, checked)
    c3_require(completed['status'] == 'COMPLETE_REAL_THROUGH_CATEGORY1' and completed['prelim_ready'] is True and (completed['category1_ready'] is True) and (completed['annual_hours_per_case'] == 8760) and (completed['annual_case_count'] == 2 * len(m['candidate_ids'])) and (completed['synthetic_fallback_used'] is False), 'Historical scientific completion is absent or incomplete.')
    c3_require(c3_file_hash(category1 / 'category1_manifest.json') == completed['category1_manifest_sha256'] and c3_file_hash(preliminary / 'preliminary_manifest.json') == completed['original_preliminary_manifest_sha256'], 'Historical completion references different manifests.')
    c3_require(m['status'] == 'COMPLETE_REAL_ANNUAL' and m['data_nature'] == 'REAL' and (m['scientific_reality_status'] == 'REAL') and (m['validation']['status'] == 'PASS') and (m['validation']['all_candidates_completed'] is True), 'Cached Category-1 science is not complete REAL annual science.')
    c3_require(m['model_version'] == ns['C1_MODEL_VERSION'] == 'category1-cumulative-exact-L0-coexistence-2.0.0' and completed['constitutive_extension'].startswith('EXACT_ZERO_WIDTH_L0_COEXISTENCE_V1'), 'Approved exact-zero-latent coexistence model identity missing.')
    for key in ['code_sha256', 'input_identity', 'model_version']:
        c3_require(m[key] == completed[key], f'Historical completion mismatch: {key}')
    c3_require(m['cumulative_source_identity'] == completed['source_identity'] and m['result_content_sha256'] == completed['scientific_result_content_sha256'] and (pm['cumulative_source_identity'] == completed['original_preliminary_source_identity']) and (pm['run_id'] == completed['original_preliminary_run_id']), 'Historical source/result lineage mismatch.')
    c3_historical_source_check(_private_dependency_path('PCM_UHI_PRIVATE_ROOF_SOURCE'), m['cumulative_source_identity'])
    c3_historical_source_check(_private_dependency_path('PCM_UHI_PRIVATE_PRELIM_SOURCE'), pm['cumulative_source_identity'])
    c3_require(pm['status'] in {'PRELIMINARY_COMPLETE_REVIEW_REQUIRED', 'PRELIMINARY_COMPLETE'} and pm['global_config']['data_mode'] == 'REAL' and all((v['result'] == 'PASS' for v in pm['validation'])), 'Saved Preliminary scientific validation failed.')
    weather, candidates, ids, metadata = ns['c1_read_preliminary_directory'](preliminary)
    c3_require(ids == m['candidate_ids'] == completed['candidate_ids'] and tuple(m['by_pcm_directories']) == tuple(ids), 'Candidate population / saved pair membership differs.')
    workbook_path = Path(ns['PCM_WORKBOOK_PATH']).resolve()
    c3_require(c3_file_hash(workbook_path) == m['workbook_sha256'] == ns['CONFIG']['pcm_expected_sha256'], 'Workbook changed; existing upstream results cannot be reused for a new workbook.')
    io = {**ns['CATEGORY1_IO'], 'workbook_path': str(workbook_path)}
    workbook = ns['c1_inspect_workbook'](io)
    ns['c1_validate_weather'](weather, metadata['weather'])
    ns['c1_validate_candidates'](candidates, ids, metadata['candidates'], metadata['screen_derivation'], workbook)
    rows = ns['audit_cleaned_pcm'](ns['load_pcm_properties'](workbook_path, sheet_name=io['workbook_sheet']))
    runnable = ns['partition_pcm_tables'](rows)['runnable']
    screen = metadata['screen_derivation']
    _, derived, dynamic_ids = ns['screen_workbook_roof'](rows, runnable, screen['lower_C'], screen['upper_C'])
    c3_require(dynamic_ids == ids, 'Saved population differs from current workbook inclusive-overlap screen.')
    ns['c1_same_table_values'](derived, candidates, 'historical candidates / dynamic workbook screen')
    candidates = derived
    weather.attrs = deepcopy(metadata['weather']['time_semantics'])
    for name in ['interval_start_UTC', 'interval_end_UTC', 'timestamp_UTC', 'instantaneous_sample_UTC']:
        weather[name] = pd.to_datetime(weather[name], utc=True)
    for name in ['interval_start_local', 'interval_end_local', 'timestamp_local', 'instantaneous_sample_local']:
        weather[name] = pd.to_datetime(weather[name], utc=True).dt.tz_convert('Asia/Kolkata')
    c3_require(ns['c1_table_digest'](weather) == m['weather_content_sha256'], 'Saved weather content identity mismatch.')
    used = pd.read_csv(category1 / 'annual_weather_used.csv', float_precision='round_trip')
    ns['c1_same_table_values'](weather, used, 'Preliminary vs Category-1 consumed weather')
    c3_require(completed['weather_identity'] == {'content_sha256': m['weather_content_sha256'], 'artifact_sha256': m['weather_artifact_sha256']} and m['weather_artifact_sha256'] == metadata['consumed_file_sha256']['annual_weather.csv'] and (completed['consumed_input_sha256'] == metadata['consumed_file_sha256']), 'Consumed historical weather/input hashes differ.')
    stored_meta = m['provenance']['inputs']
    for key in ['weather', 'candidates', 'screen_derivation', 'preliminary_run_id', 'consumed_file_sha256', 'manifest_sha256', 'cumulative_source_identity']:
        c3_require(ns['c1_json_value'](stored_meta[key]) == ns['c1_json_value'](metadata[key]), f'Historical input metadata differs: {key}')
    reconstructed_input = ns['c1_digest']({'weather': ns['c1_table_digest'](weather), 'candidates': ns['c1_table_digest'](candidates), 'ids': ids, 'metadata': stored_meta, 'config': m['config'], 'acceptance_limits': m['acceptance_limits'], 'workbook': m['workbook_sha256']})
    c3_require(reconstructed_input == m['input_identity'], 'Historical input-content seal mismatch.')
    c3_require(m['acceptance_limits'] == ns['CATEGORY1_VALIDATION'], 'Category-1 validation limits changed.')
    c3_require({k: v for k, v in m['config'].items() if k != 'output_dir'} == {k: v for k, v in ns['c1_make_run_config'](ns['CATEGORY1_CONFIG']).items() if k != 'output_dir'}, 'Current Category-1 physical configuration differs from cached science.')
    expected = ns['c1_explicit_utc'](weather.interval_start_UTC, 'weather start')
    end = expected + pd.Timedelta(hours=1)
    solver_weather = ns['c1_solver_weather_from_annual'](weather)
    prepared = ns['c1_prepare_model_weather'](solver_weather, m['config'])
    solver_hash = hashlib.sha256(pd.util.hash_pandas_object(prepared[['interval_start_UTC', 'T_air_C', 'u10_ms', 'v10_ms', 'G_Wm2', 'LW_down_Wm2']], index=False).values.tobytes()).hexdigest()
    by_pcm = {}
    independent = []
    for pid in ids:
        folder = category1 / m['by_pcm_directories'][pid]
        meta = json.loads((folder / 'metadata.json').read_text())
        identity = meta['identity']
        validation = json.loads((folder / 'validation.json').read_text())
        pair = {k: v for k, v in identity.items() if k != 'pair_id'}
        c3_require(ns['c1_digest'](pair) == identity['pair_id'] and identity['pcm_id'] == pid and (identity['input_identity'] == m['input_identity']) and (identity['code_sha256'] == m['code_sha256']) and (identity['model_version'] == m['model_version']) and (identity['weather_solver_hash'] == solver_hash) and (identity['workbook_sha256'] == m['workbook_sha256']) and (identity['config_sha256'] == ns['c1_digest'](m['config'])) and (identity.get('computational_source_sha256', _private_metadata_value(identity, 'pair_source')) == m['cumulative_source_identity']['executable_source_sha256']), f'{pid}: material/control pair source identity mismatch.')
        candidate = candidates.loc[candidates.pcm_id == pid].iloc[0]
        c3_require(identity['pcm_row_sha256'] == ns['c1_digest']({k: candidate[k] for k in ns['C1_PCM_COLUMNS']}), f'{pid}: material row hash differs.')
        for relative, record in m['artifacts'].items():
            if relative.startswith(m['by_pcm_directories'][pid] + '/'):
                c3_require(record['pcm_id'] == pid and record['pair_identity'] == identity, f'{relative}: wrong pair membership.')
        for case in ['pcm', 'L0']:
            c3_require(validation[case]['status'] == 'PASS' and validation[case]['unresolved_solver_failures'] == 0 and (validation[case]['state_continuity_status'] == 'PASS') and (meta['numerical_extension'][case]['method'] == 'EXACT_ZERO_WIDTH_L0_COEXISTENCE_V1'), f'{pid}/{case}: failed saved validation.')
        material = meta['material']
        control = meta['control_material']
        c3_require(material['pcm_id'] == pid and control['pcm_id'] == pid and (material['L_Jkg'] > 0) and (control['L_Jkg'] == 0) and ({k: v for k, v in material.items() if k != 'L_Jkg'} == {k: v for k, v in control.items() if k != 'L_Jkg'}), f'{pid}: control is not matched L=0.')
        r = {}
        cases = {}
        for key, filename, case, mat in [('hourly', 'hourly_pcm.csv', 'pcm', material), ('matched_control', 'hourly_L0.csv', 'L0', control)]:
            frame = pd.read_csv(folder / filename, float_precision='round_trip')
            for col in ['datetime_UTC', 'datetime_end_UTC', 'datetime_IST', 'datetime_end_IST']:
                frame[col] = pd.to_datetime(frame[col], utc=True)
                if col.endswith('IST'):
                    frame[col] = frame[col].dt.tz_convert('Asia/Kolkata')
            c3_require(len(frame) == 8760 and ns['c1_explicit_utc'](frame.datetime_UTC, 'result start').equals(expected) and ns['c1_explicit_utc'](frame.datetime_end_UTC, 'result end').equals(end), f'{pid}/{case}: wrong 8760-hour chronology.')
            c3_require(frame.pcm_id.eq(pid).all() and frame.run_id.eq(identity['pcm_run_id' if case == 'pcm' else 'L0_run_id']).all() and frame.pair_id.eq(identity['pair_id']).all() and frame.scientific_reality_status.eq('REAL').all(), f'{pid}/{case}: stale or non-real result.')
            for col in ['T_air_C', 'u10_ms', 'v10_ms', 'wind10_ms', 'G_Wm2', 'LW_down_Wm2']:
                same = np.allclose(frame[col], prepared[col], rtol=1e-12, atol=1e-12) if col == 'wind10_ms' else np.array_equal(frame[col].to_numpy(), weather[col].to_numpy())
                c3_require(same, f'{pid}/{case}: weather values differ in {col}.')
            frame.attrs = {'data_nature': 'REAL', 'units': m['units'], 'input_identity': m['input_identity']}
            with np.load(folder / ('cell_states_' + case + '.npz'), allow_pickle=False) as z:
                states = {k: z[k].copy() for k in ['T_C', 'h_Jkg', 'f', 'x_center_m']}
                states['datetime_UTC'] = pd.to_datetime(z['datetime_UTC_ns'], utc=True)
            result = {'hourly': frame, 'cell_snapshots': states, 'material': mat, 'run_config': m['config'], 'weather_hash': solver_hash, 'run_id': identity['pcm_run_id' if case == 'pcm' else 'L0_run_id'], 'final_state': {k: states[k][-1] for k in ['T_C', 'h_Jkg', 'f']}}
            v = ns['c1_validate_run'](result)
            independent.append({'pcm_id': pid, 'case': case, 'validation': v})
            r[key] = frame
            cases[case] = result
        ns['c1_validate_pair'](cases['pcm'], cases['L0'])
        d = pd.read_csv(folder / 'delta_H.csv', float_precision='round_trip')
        for col in ['datetime_UTC', 'datetime_end_UTC', 'datetime_IST', 'datetime_end_IST']:
            d[col] = pd.to_datetime(d[col], utc=True)
            if col.endswith('IST'):
                d[col] = d[col].dt.tz_convert('Asia/Kolkata')
        units = m['artifacts'][m['by_pcm_directories'][pid] + '/delta_H.csv']['column_units']
        c3_require(all((units[k] == 'W/m2' for k in ['H_PCM_Wm2', 'H_L0_Wm2', 'delta_H_Wm2'])) and m['units']['sign_convention'] == 'q_* positive into roof; H_* positive roof into atmosphere; delta_H=H_PCM-H_L0', f'{pid}: flux units/sign convention wrong.')
        c3_require(len(d) == 8760 and ns['c1_explicit_utc'](d.datetime_UTC, 'delta start').equals(expected) and ns['c1_explicit_utc'](d.datetime_end_UTC, 'delta end').equals(end) and d.pcm_id.eq(pid).all() and d.pair_id.eq(identity['pair_id']).all() and d.pcm_run_id.eq(identity['pcm_run_id']).all() and d.L0_run_id.eq(identity['L0_run_id']).all() and d.weather_hash.eq(solver_hash).all() and d.scientific_reality_status.eq('REAL').all() and np.isfinite(d.delta_H_Wm2).all(), f'{pid}: ΔH pair/weather/time identity mismatch.')
        c3_require(np.array_equal(d.H_PCM_Wm2, -r['hourly'].q_convection_mean_Wm2) and np.array_equal(d.H_L0_Wm2, -r['matched_control'].q_convection_mean_Wm2) and np.array_equal(d.delta_H_Wm2, d.H_PCM_Wm2 - d.H_L0_Wm2), f'{pid}: roof-to-air ΔH sign/values mismatch.')
        d.attrs = {'units': {k: 'W/m2' for k in ['H_PCM_Wm2', 'H_L0_Wm2', 'delta_H_Wm2']}, 'sign_convention': 'positive roof into atmosphere; delta_H = H_PCM - H_L0', 'data_nature': 'REAL'}
        annual = json.loads((folder / 'annual_summary.json').read_text())
        r.update(delta_H=d, material=material, control_material=control, identity=identity, validation=validation, summary=annual['pcm'], matched_control_summary=annual['L0'], initial_condition=meta['initial_condition'], numerical_extension=meta['numerical_extension'], cell_snapshots={k: v['cell_snapshots'] for k, v in cases.items()})
        by_pcm[pid] = r
    grid_data = json.loads((preliminary / 'common_grid.json').read_text())
    grid = ns['GridSpec'](tuple(grid_data['shape']), grid_data['crs'], ns['Affine'](*grid_data['transform']))
    ns['validate_grid_spec'](grid, ns['CONFIG'])
    with ns['rasterio'].open(preliminary / 'HOTSPOT_MASK.tif') as ds:
        c3_require(ds.shape == grid.shape and ds.transform == grid.transform and (ds.crs == ns['CRS'].from_user_input(grid.crs)), 'Hotspot raster CRS/grid mismatch.')
        raw = ds.read(1)
        mask = np.isfinite(raw) & (raw == 1)
    location = json.loads((preliminary / 'hotspot_weather_location.json').read_text())
    c3_require(mask.sum() == 44 and ns['numeric_array_hash'](mask) == location['hotspot_mask_numeric_sha256'] and (grid_data == location['location_grid'] == metadata['weather']['spatial_support']['hotspot_location']['location_grid']) and (location == metadata['weather']['spatial_support']['hotspot_location']), 'Hotspot mask/grid/weather-location identity mismatch.')
    cells = []
    for row, col in np.argwhere(mask):
        x, y = grid.transform * (int(col), int(row))
        cells.append(box(x, y - 150, x + 150, y))
    hotspot = unary_union(cells)
    c3_require(abs(hotspot.area - 990000) < 1e-06, 'Hotspot cell union area mismatch.')
    upstream = SimpleNamespace(_seal=_C3_HISTORICAL_SEAL, file_hashes=checked, config=m['config'], metadata=stored_meta)
    identity = {'upstream_execution_route': 'VERIFIED_HISTORICAL_ARTIFACTS', 'category1_result_seal': m['result_content_sha256'], 'category1_artifact_manifest_sha256': c3_file_hash(category1 / 'category1_manifest.json'), 'category1_source_identity': m['cumulative_source_identity'], 'preliminary_source_identity': pm['cumulative_source_identity'], 'category1_code': m['code_sha256'], 'input_identity': m['input_identity'], 'workbook': completed['workbook_identity'], 'weather': completed['weather_identity'], 'preliminary': stored_meta, 'grid': grid_data, 'location': deepcopy(location), 'hotspot_wkb_sha256': hashlib.sha256(hotspot.wkb).hexdigest(), 'ordered_cell_geometry_sha256': c3_digest([c.wkb.hex() for c in cells]), 'forcing_sha256': {pid: ns['c1_table_digest'](by_pcm[pid]['delta_H']) for pid in ids}, 'historical_file_hashes': {str(Path(k).relative_to(root)): v for k, v in checked.items()}}
    inputs = C3Inputs(weather, by_pcm, tuple(ids), hotspot, cells, grid, deepcopy(location), identity, upstream, _seal=_C3_REAL_INPUT_SEAL)
    ns['CATEGORY1_HISTORICAL_HANDOFF'] = {'status': 'VERIFIED_HISTORICAL_REAL_ANNUAL', 'config': m['config'], 'candidate_ids': ids, 'manifest': m, 'artifact_directory': str(category1), 'fresh_execution_in_this_process': False}
    ns['C3_HISTORICAL_INPUTS'] = inputs
    receipt = {'status': 'PASS_VERIFIED_HISTORICAL_UPSTREAM', 'fresh_preliminary_or_category1_execution': False, 'category3_scientific_execution': False, 'category3_ready': False, 'candidate_ids': ids, 'hours_per_case': 8760, 'verified_files': identity['historical_file_hashes'], 'independent_saved_state_checks': independent, 'historical_source': m['cumulative_source_identity'], 'computational_source_identity': ns['CUMULATIVE_SOURCE_IDENTITY'], 'elapsed_s': time.perf_counter() - started}
    ns['C3_HISTORICAL_VALIDATION'] = receipt
    c3_write_json(Path(ns['CONFIG']['output_dir']) / 'category3' / 'historical_upstream_validation.json', receipt)
    print(f"Cached upstream PASS: {len(ids)} materials, {2 * len(ids)} existing annual cases, {len(checked)} files; {receipt['elapsed_s']:.1f} s. No upstream physics recomputed.")
    return inputs

def c3_validate_handoff(ns):
    if ns.get('CACHE_FIRST', True):
        inputs = ns.get('C3_HISTORICAL_INPUTS')
        c3_require(inputs is not None and c3_upstream_seal_valid(inputs, ns), 'Run Stage A cached input verification first.')
        c3_verify_historical_files(inputs)
        c3_check_consumed_inputs(inputs, ns)
        return inputs
    return _c3_fresh_validate_handoff(ns)

def c3_choose_pilot(inputs, ns):
    weather = inputs.weather
    start = ns['c1_explicit_utc'](weather.interval_start_UTC, 'pilot weather')
    local = start.tz_convert('Asia/Kolkata')
    eligible = np.flatnonzero((local.hour == 0) & (local.minute == 0))
    candidates = [(float(weather.T_air_C.iloc[i:i + 48].mean()), int(i)) for i in eligible if i + 48 <= len(weather)]
    mean, offset = max(candidates, key=lambda v: (v[0], -v[1]))
    return {'selection_rule': 'Hottest mean-Tair 48h window beginning at IST midnight; earliest start breaks ties. Weather only; no PCM cooling selection.', 'start_hour': offset, 'hours': 48, 'start_UTC': str(start[offset]), 'end_UTC_exclusive': str(start[offset + 47] + pd.Timedelta(hours=1)), 'start_IST': str(local[offset]), 'mean_T_air_C': mean, 'weather_content_sha256': inputs.identity['weather']['content_sha256']}

def c3_stage_scope(inputs, roofs, ns, cfg):
    return c3_digest({'inputs': inputs.identity, 'roofs': roofs['identity'], 'source': ns['CUMULATIVE_SOURCE_IDENTITY'], 'code': c3_live_identity(ns), 'base': C3_BASE, 'buffer_m': cfg['baseline_buffer_m'], 'cfl': cfg['cfl'], 'dt_cap_s': cfg['dt_cap_s'], 'threshold': cfg['effect_threshold_K'], 'pilot': c3_choose_pilot(inputs, ns)})

def c3_stage_a_geometry(inputs, roofs, ns, cfg):
    c3_check_consumed_inputs(inputs, ns, roofs)
    layout = c3_layout(inputs, roofs, 1.0, 'dispersed', C3_BASE['eta'])
    domain = c3_domain(inputs, roofs, layout, cfg['baseline_buffer_m'], 150.0)
    c3_require(abs(domain['source'].sum() * 150 ** 2 - layout['actual_m2']) <= max(1e-06, layout['actual_m2'] * 1e-10), 'Real source-area conversion mismatch.')
    wind = c3_projected_wind(inputs.weather, inputs.location)
    steps = [c3_stable_dt(C3_BASE['r_U'] * u, C3_BASE['r_U'] * v, C3_BASE['K_h_m2s'], 150.0, 0, cfg['cfl'], cfg['dt_cap_s']) for u, v in wind]
    c3_require(np.isfinite(steps).all() and min(steps) > 0, 'Actual weather stability bound invalid.')
    receipt = {'status': 'PASS_STAGE_A', 'category3_scientific_execution': False, 'category3_ready': False, 'math_tests': ns['C3_MECHANICAL_RECEIPT']['status'], 'roof_geometry': roofs['identity'], 'hotspot_area_m2': inputs.hotspot.area, 'footprint_union_m2': roofs['geometry'].area, 'baseline_treated_roof_m2': layout['actual_m2'], 'source_integral_m2': float(domain['source'].sum() * 150 ** 2), 'actual_weather_dt_min_s': float(min(steps)), 'pilot_selection': c3_choose_pilot(inputs, ns)}
    c3_require(ns['C3_MECHANICAL_RECEIPT']['status'] == 'PASS', 'Atmospheric mathematical tests failed.')
    folder = Path(ns['CONFIG']['output_dir']) / 'category3' / 'stage_receipts' / c3_stage_scope(inputs, roofs, ns, cfg)
    c3_write_json(folder / 'stage_A.json', receipt)
    return (folder, receipt)

C4_INHERITED_C3_SOURCE_IDENTITY = {'source_content_sha256': '641c531fa52dd68a46e1b72cdf5445ec2d281c8f5eee9c14281292e07140f94b', 'executable_source_sha256': '0704954d7da567a165afca2c6c3e7b8595071795bac1041e470dfa2a0d5d36e9', 'fingerprint_policy': 'Recorded computational source fingerprint; exact source retained privately.'}

C4_CONFIG = {'output_dir': str(Path(CONFIG['project_root']) / 'PCM_UHI_FINAL_Category4_outputs'), 'assembly_ids': {}, 'building_ids': {}, 'assembly_inventory_complete': {}, 'objectives': {}}

C4_EVIDENCE_CSV = None

C4_COMPONENTS_CSV = None

C4_OBJECTIVES_JSON = None

C4_CATEGORY3_RECEIPT_PATH = None

C4_CATEGORY3_RECEIPT_SHA256 = None

C4_CATEGORY3_ARTIFACT_ROOT = None

C4_CRITERIA = ('fire_safety', 'leakage_encapsulation', 'corrosion_compatibility', 'structural_load', 'conductivity_heat_transfer', 'thermal_cycling_stability', 'nighttime_regeneration', 'durability_weather_exposure')

C4_COMPONENTS = ('encapsulation', 'enhancement', 'fins', 'supports', 'mounting', 'protective_layers')

C4_CLASSES = {'SUPPORTED', 'FAILED', 'INSUFFICIENT_EVIDENCE'}

C4_EVIDENCE_COLUMNS = ['pcm_id', 'criterion', 'classification', 'source', 'locator', 'source_path', 'source_sha256', 'material_name', 'assembly_id', 'conditions', 'explanation', 'missing_evidence']

C4_COMPARATOR = 'PCM minus matched L=0 mathematical control; no commercial cool-roof comparison'

C4_OBJECTIVE_UNITS = {'active_pcm_kg_m2': 'kg/m2', 'total_assembly_kg_m2': 'kg/m2', 'active_h': 'h', 'activation_days': 'day', 'saturated_h': 'h', 'raw_saturation_h': 'h', 'equivalent_full_uptake_cycles': 'cycle', 'equivalent_full_release_cycles': 'cycle', 'incomplete_reset_cycles': 'cycle', 'incomplete_reset_pct_of_known_cycles': '%', 'delta_H_annual_mean_Wm2': 'W/m2', 'delta_H_annual_integral_Jm2': 'J/m2', 'annual_signed_temperature_effect_K': 'K', 'gross_cooling_K_m2_s': 'K m2 s', 'gross_warming_K_m2_s': 'K m2 s', 'extreme_hot_period_mean_K': 'K', 'night_signed_mean_K': 'K', 'night_gross_warming_K_m2_s': 'K m2 s', 'atmospheric_ensemble_min_K': 'K', 'atmospheric_ensemble_max_K': 'K', 'daynight_mixing_mean_min_K': 'K', 'daynight_mixing_mean_max_K': 'K'}

def c4_missing(value):
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().upper() in {'', 'TODO', 'UNKNOWN', 'NOT AVAILABLE', 'NAN'}
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False

def c4_require_inputs(mapping, names, context):
    missing = [k for k in names if c4_missing(mapping.get(k))]
    if missing:
        raise ValueError(f'{context}: missing supplied inputs {missing}')

def c4_finite(value, name, *, positive=False, nonnegative=False):
    if c4_missing(value) or isinstance(value, (bool, np.bool_)):
        raise ValueError(f'{name} must be a supplied finite number')
    try:
        x = float(value)
    except (TypeError, ValueError):
        raise ValueError(f'{name} must be numeric') from None
    if not np.isfinite(x) or (positive and x <= 0) or (nonnegative and x < 0):
        raise ValueError(f'Invalid {name}: {value!r}')
    return x

def c4_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def c4_source_check(record, material, assembly_id):
    """An archived, hashed source and exact applicability are required for a claim.

    File identity does not prove the interpretation; explanation/locator must be
    reviewed against the source. Generic material-family claims are not accepted.
    """
    needed = ('source', 'locator', 'source_path', 'source_sha256', 'material_name', 'assembly_id', 'conditions', 'explanation')
    if any((c4_missing(record.get(k)) for k in needed)):
        return (False, 'Missing source, locator, applicability or conditions')
    if record['material_name'] != material['material_name'] or not assembly_id or record['assembly_id'] != assembly_id:
        return (False, 'Evidence is not bound to this exact formulation and declared assembly')
    path = Path(str(record['source_path'])).expanduser()
    if not path.is_file() or c4_hash(path) != record['source_sha256']:
        return (False, 'Source file unavailable or SHA-256 mismatch')
    return (True, '')

def c4_system_mass(material, thickness_m, treated_area_m2=None, components=None, *, inventory_complete=False, assembly_id=None):
    rho = c4_finite(material.get('rho_ref_kgm3'), 'density', positive=True)
    d = c4_finite(thickness_m, 'thickness', positive=True)
    area = np.nan if c4_missing(treated_area_m2) else c4_finite(treated_area_m2, 'treated area', nonnegative=True)
    active = rho * d
    records = list(components or [])
    names = [r.get('component') for r in records]
    if len(set(names)) != len(names):
        raise ValueError('Duplicate assembly component')
    known = []
    rejected = []
    for r in records:
        valid, reason = c4_source_check(r, material, assembly_id)
        if valid and (not c4_missing(r.get('kg_m2'))):
            known.append(c4_finite(r['kg_m2'], r['component'], nonnegative=True))
        else:
            rejected.append(str(r.get('component')) + ': ' + (reason or 'mass missing'))
    complete = inventory_complete is True and set(C4_COMPONENTS).issubset(names) and (not rejected) and bool(records)
    extra = sum(known) if known else np.nan
    total = active + extra if complete else np.nan
    return {'density_kg_m3': rho, 'thickness_m': d, 'active_pcm_kg_m2': active, 'treated_area_m2': area, 'active_pcm_mass_kg': active * area, 'known_additional_component_kg_m2': extra, 'assembly_inventory_complete': complete, 'total_assembly_kg_m2': total, 'total_assembly_mass_kg': total * area, 'assembly_missing': '; '.join(rejected or ([] if complete else ['Complete sourced assembly inventory and declared design required']))}

def c4_pcm_diagnostics(result):
    a = result['summary']
    names = ('active_h', 'activation_days', 'saturated_h', 'raw_saturation_h', 'equivalent_full_uptake_cycles', 'equivalent_full_release_cycles', 'incomplete_reset_cycles', 'known_next_morning_checkpoints', 'incomplete_reset_pct_of_known_cycles', 'regeneration_interpretation', 'delta_H_annual_mean_Wm2', 'delta_H_annual_integral_Jm2')
    return {k: a.get(k, np.nan) for k in names}

def c4_feasibility_table(inputs, config, evidence=None, components=None, areas=None):
    evidence = pd.DataFrame(columns=C4_EVIDENCE_COLUMNS) if evidence is None else evidence.copy()
    if not evidence.empty:
        c4_require_inputs({k: k for k in evidence.columns}, C4_EVIDENCE_COLUMNS, 'Evidence schema')
        if evidence.duplicated(['pcm_id', 'criterion']).any():
            raise ValueError('Duplicate feasibility evidence')
        if not evidence.pcm_id.isin(inputs.ids).all() or not evidence.criterion.isin(C4_CRITERIA).all():
            raise ValueError('Evidence contains unknown candidate or criterion')
        if not evidence.classification.isin(C4_CLASSES).all():
            raise ValueError('Unknown evidence classification')
    rows = []
    masses = {}
    for pid in inputs.ids:
        r = inputs.by_pcm[pid]
        mat = r['material']
        diag = c4_pcm_diagnostics(r)
        assembly = config.get('assembly_ids', {}).get(pid)
        mass = c4_system_mass(mat, inputs.upstream.config['d_m'], (areas or {}).get(pid), (components or {}).get(pid), inventory_complete=config.get('assembly_inventory_complete', {}).get(pid) is True, assembly_id=assembly)
        masses[pid] = mass
        for criterion in C4_CRITERIA:
            match = evidence.loc[(evidence.pcm_id == pid) & (evidence.criterion == criterion)]
            rec = match.iloc[0].to_dict() if len(match) else {}
            status = 'INSUFFICIENT_EVIDENCE'
            source = ''
            locator = ''
            scope = 'FIELD_ASSEMBLY'
            conditions = ''
            reason = ''
            explanation = 'No candidate-specific installation evidence supplied.'
            missing = 'Candidate-specific test evidence for the declared rooftop assembly and exposure conditions'
            if rec:
                valid, reason = c4_source_check(rec, mat, assembly)
                if valid:
                    status = rec['classification']
                    source = rec['source']
                    locator = rec['locator']
                    conditions = rec['conditions']
                    explanation = rec['explanation']
                    missing = rec['missing_evidence']
                    if status == 'INSUFFICIENT_EVIDENCE' and c4_missing(missing):
                        missing = 'Further validation required'
                else:
                    missing = reason
            model_source = f"verified Category-1 pair {r['identity']['pair_id']}; input {inputs.identity['input_identity']}"
            if criterion == 'nighttime_regeneration' and (not rec):
                n = diag['known_next_morning_checkpoints']
                incomplete = diag['incomplete_reset_cycles']
                if not c4_missing(n) and n > 0 and (not c4_missing(incomplete)):
                    status = 'SUPPORTED' if incomplete == 0 else 'FAILED'
                    scope = 'NUMERICAL_MODEL_ONLY'
                    source = model_source
                    locator = 'annual_summary.pcm: incomplete_reset_cycles; known_next_morning_checkpoints'
                    conditions = f"{inputs.upstream.config['year']}; checkpoint {inputs.upstream.config['morning_checkpoint']}; reset_limit={inputs.upstream.config['reset_limit']}; {diag['regeneration_interpretation']}"
                    explanation = f'{int(incomplete)} incomplete resets in {int(n)} known model checkpoints. This evaluates the inherited numerical reset criterion.'
                    missing = 'Real formulation-specific freezing, supercooling and field regeneration validation; numerical reset is not durability evidence'
            if criterion == 'thermal_cycling_stability' and (not rec):
                source = model_source
                locator = 'annual_summary.pcm: equivalent_full_uptake_cycles; equivalent_full_release_cycles'
                conditions = str(diag['regeneration_interpretation'])
                explanation = f"Model equivalent uptake/release cycles: {diag['equivalent_full_uptake_cycles']}, {diag['equivalent_full_release_cycles']}. No material ageing experiment was simulated."
                missing = 'Measured latent/property retention and assembly integrity over relevant thermal cycles'
            if criterion == 'conductivity_heat_transfer' and (not rec):
                source = mat.get('source', '')
                locator = mat.get('source_locator', '')
                conditions = mat.get('notes', '')
                explanation = f"Inherited k_solid={mat.get('k_solid_WmK')} and k_liquid={mat.get('k_liquid_WmK')} W/(m K). Layer conduction is modelled; interface and assembly heat-transfer performance are unvalidated."
            if criterion == 'structural_load':
                explanation += f" Active loading is {mass['active_pcm_kg_m2']:.6g} kg/m². Total assembly loading is {('known' if mass['assembly_inventory_complete'] else 'unknown')}."
                building = config.get('building_ids', {}).get(pid)
                if not mass['assembly_inventory_complete'] or not building or rec.get('building_id') != building:
                    status = 'INSUFFICIENT_EVIDENCE'
                    missing = 'Complete sourced assembly loading and building-specific structural capacity assessment'
            rows.append(dict(pcm_id=pid, material_name=mat['material_name'], criterion=criterion, classification=status, evidence_scope=scope, source=source, locator=locator, source_path=rec.get('source_path', ''), source_sha256=rec.get('source_sha256', ''), assembly_id=assembly, conditions=conditions, explanation=explanation, missing_evidence=missing, source_rejection=reason, **mass))
    return (pd.DataFrame(rows), masses)

def c4_read_completed_category3(ns, path, expected_sha256, artifact_root=None):
    """Read a completed snapshot only; no cache refresh, checkpoint or solver calls.

    expected SHA comes from the original execution receipt or a trusted copied-file
    receipt. artifact_root permits rebasing experiment folders after Drive export.
    """
    c4_require_inputs({'path': path, 'sha': expected_sha256}, ('path', 'sha'), 'Category-3 receipt')
    path = Path(path).expanduser().resolve()
    if c4_hash(path) != expected_sha256:
        raise ValueError('Category-3 completion receipt hash mismatch')
    m = json.loads(path.read_text())
    inputs = ns['c3_validate_handoff'](ns)
    full = m.get('status') == 'COMPLETE_REAL_THROUGH_CATEGORY3_DEVELOPMENT' and m.get('CATEGORY3_READY') is True
    baseline = m.get('status') == 'COMPLETE_REAL_BASELINE_ANNUAL'
    if not (full or baseline):
        raise ValueError('Category-3 result is absent, partial, pilot-only or failed')
    upstream = m.get('upstream_identity', _private_metadata_value(m, 'upstream_identity'))
    if ns['c3_json'](upstream) != ns['c3_json'](inputs.identity):
        raise ValueError('Category-3 upstream identity differs')
    source = m.get('computational_source_identity', _private_metadata_value(m, 'source_identity'))
    allowed = [ns['CUMULATIVE_SOURCE_IDENTITY'], ns['C4_INHERITED_C3_SOURCE_IDENTITY']]
    if not any(_source_identity_matches(source, approved) for approved in allowed):
        raise ValueError('Category-3 computational source identity is not an approved unchanged implementation')
    if m.get('result_use') != 'DEVELOPMENT' or m.get('candidate_population_final') is not False:
        raise ValueError('Category-3 development lineage unexpectedly changed')
    cfg = ns['C3_EXECUTION']
    live = ns['c3_live_identity'](ns)
    if full:
        if m.get('code_identity') != live:
            raise ValueError('Category-3 executed scientific code/environment differs')
        keys = ['cfl', 'dt_cap_s', 'effect_threshold_K', 'coverage_fractions', 'arrangements', 'sensitivity_patterns', 'design_seed', 'feedback_iterations', 'domain_candidates_m']
        if any((m['execution_config'].get(k) != cfg.get(k) for k in keys)):
            raise ValueError('Category-3 experiment configuration differs')
        fb = m.get('validation', {}).get('feedback', {})
        if not (fb.get('thresholds_pass') is True and fb.get('observed_corrections_decrease_and_small') is True):
            raise ValueError('Category-3 feedback hard gates incomplete')
    elif m.get('baseline_parameters') != ns['C3_BASE'] or m.get('buffer_m') != cfg['baseline_buffer_m']:
        raise ValueError('Category-3 baseline settings differ')
    results = []
    checked = {str(path): expected_sha256}
    layouts = {}
    domains = {}
    if full:
        found = {}
        for lp in (path.parent / 'stage_ledgers').glob('*.json'):
            ledger = json.loads(lp.read_text())
            li = ledger.get('identity', {})
            if li.get('upstream') != ns['c3_json'](inputs.identity) or li.get('source') != source or li.get('code') != live:
                continue
            if ns['c3_digest'](li) != lp.stem:
                raise ValueError('Category-3 stage ledger identity seal differs')
            stage = ledger['stage']
            if stage in found:
                raise ValueError('Ambiguous completed Category-3 stage ledgers')
            if ledger.get('status') != 'COMPLETE' or not ledger.get('next_index') == ledger.get('task_count') == len(ledger.get('completed', [])) == len(li.get('tasks', [])):
                raise ValueError('Category-3 stage task ledger incomplete')
            if [e['task'] for e in ledger['completed']] != li['tasks']:
                raise ValueError('Category-3 task sequence differs')
            found[stage] = ledger
            checked[str(lp)] = c4_hash(lp)
            for entry in ledger['completed']:
                for result in entry.get('results', []):
                    layouts[result['experiment_id']] = result['layout']
                    domains[result['experiment_id']] = result['domain']
        if set(found) != {'domain', 'baseline', 'sensitivity', 'feedback'}:
            raise ValueError('Mandatory Category-3 stage ledgers missing')
        names = set(ns['c3_experiment_matrix'](cfg['design_seed']).name) - {'baseline'}
        for pid in inputs.ids:
            b = [t for t in found['baseline']['identity']['tasks'] if t['pcm_id'] == pid]
            expected = {(c, a) for c in cfg['coverage_fractions'] for a in (['dispersed'] if c == 1 else cfg['arrangements'])}
            if {(t['coverage'], t['arrangement']) for t in b} != expected or len(b) != len(expected):
                raise ValueError('Deployment sensitivity incomplete')
            s = [t for t in found['sensitivity']['identity']['tasks'] if t['pcm_id'] == pid]
            for coverage, arr in cfg['sensitivity_patterns']:
                g = [t for t in s if t['coverage'] == coverage and t['arrangement'] == arr]
                if {t['parameters']['name'] for t in g if 'parameters' in t} != names or {t['eta'] for t in g if 'eta' in t} != {0.5, 0.9}:
                    raise ValueError('Atmospheric/day-night/roof availability experiments incomplete')
            fb = [t for t in found['feedback']['identity']['tasks'] if t['pcm_id'] == pid]
            expected_fb = {(c, a, p, beta) for c, a in [(1.0, 'dispersed'), (0.5, 'dispersed'), (0.5, 'clustered')] for p in ['feedback_baseline', 'feedback_large_response'] for beta in [0.5, 1.0]}
            if {(t['coverage'], t['arrangement'], t['atmosphere'], t['beta']) for t in fb} != expected_fb:
                raise ValueError('Feedback groups incomplete')
        if set(layouts) != {r['experiment_id'] for r in m['experiments']}:
            raise ValueError('Category-3 completed experiment set differs from task ledgers')
    for ref in m.get('experiments', []):
        old = Path(ref['manifest'])
        p = Path(artifact_root) / ref['experiment_id'] / old.name if artifact_root else old
        p = p.expanduser().resolve()
        sha = ref.get('manifest_sha256', ref.get('sha256'))
        if c4_hash(p) != sha:
            raise ValueError('Category-3 experiment manifest hash mismatch')
        em = json.loads(p.read_text())
        identity = em['identity']
        val = em['validation']
        if em['experiment_id'] != ref['experiment_id'] or ns['c3_digest'](identity) != ref['experiment_id']:
            raise ValueError('Category-3 experiment identity seal differs')
        if identity['input'] != ns['c3_json'](inputs.identity) or identity['code'] != live or identity.get('computational_source', _private_metadata_value(identity, 'experiment_source')) != source:
            raise ValueError('Category-3 result source/input/code identity differs')
        if identity['hours'] != 8760 or identity.get('start_hour', 0) != 0 or em['status'] != 'ANNUAL_NUMERICALLY_ACCEPTED':
            raise ValueError('Category-3 annual numerical completion missing')
        if any((val.get(k) != 'PASS' for k in ['energy', 'continuity', 'numerical_status'])):
            raise ValueError('Failed Category-3 hard gate')
        if any((identity.get(k) != cfg[k] for k in ['cfl', 'dt_cap_s', 'effect_threshold_K'])):
            raise ValueError('Category-3 numerical config differs')
        pid = identity['parameters']['pcm_id']
        if pid not in inputs.ids:
            raise ValueError('Unknown Category-3 candidate')
        for relative, sha2 in em['artifacts'].items():
            a = (p.parent / relative).resolve()
            if not a.is_relative_to(p.parent) or not a.is_file() or c4_hash(a) != sha2:
                raise ValueError('Category-3 artifact missing, modified or outside result folder')
            checked[str(a)] = sha2
        checked[str(p)] = sha
        t = pd.read_csv(p.parent / 'hourly_statistics.csv', float_precision='round_trip')
        if len(t) != 8760 or not pd.DatetimeIndex(pd.to_datetime(t.datetime_UTC, utc=True)).equals(pd.DatetimeIndex(pd.to_datetime(inputs.weather.interval_start_UTC, utc=True))):
            raise ValueError('Category-3 annual weather chronology differs')
        convergence = pd.read_csv(p.parent / 'numerical_convergence.csv')
        if 'pass' not in convergence or convergence.empty or (not convergence['pass'].map(lambda x: x is True or str(x).lower() == 'true').all()):
            raise ValueError('Category-3 numerical convergence receipt failed')
        with np.load(p.parent / 'spatial_diagnostics.npz', allow_pickle=False) as z:
            summary = ns['c3_summarize'](t, {'weights': z['weights']})
            area = float(z['source'].sum() * 150 ** 2)
            if full:
                dom = domains[em['experiment_id']]
                lay = layouts[em['experiment_id']]
                if dom['dx'] != 150 or dom['identity'] != identity['domain'] or ns['c3_digest']({'bounds': dom['bounds'], 'dx': dom['dx'], 'source': ns['c3_array_hash'](z['source']), 'weights': ns['c3_array_hash'](z['weights'])}) != identity['domain']:
                    raise ValueError('Category-3 saved domain identity mismatch')
                if lay['identity'] != identity['layout'] or ns['c3_digest']({'areas': lay['areas_m2'], 'eta': lay['eta'], 'coverage': lay['coverage'], 'arrangement': lay['arrangement']}) != identity['layout'] or (not np.isclose(area, lay['actual_m2'], rtol=1e-10, atol=1e-05)):
                    raise ValueError('Category-3 saved deployment identity mismatch')
        for rec in summary:
            for key, value in summary[rec].items():
                if not np.isclose(value, em['summary'][rec][key], rtol=1e-10, atol=1e-08):
                    raise ValueError('Category-3 summary differs from hourly data')
        if identity.get('forcing_nature') in {'REAL_CATEGORY1_HOURLY_MEAN', 'DERIVED_FEEDBACK_SCENARIO'}:
            dh = inputs.by_pcm[pid]['delta_H'].delta_H_Wm2.to_numpy(float) if identity['forcing_nature'] == 'REAL_CATEGORY1_HOURLY_MEAN' else pd.read_csv(p.parent / 'derived_delta_H.csv', float_precision='round_trip').delta_H_Wm2.to_numpy(float)
            if identity['forcing_sha256'] != ns['c3_array_hash'](dh):
                raise ValueError('Category-3 forcing identity differs')
            nonzero = np.abs(dh) > 0
            if not nonzero.any():
                raise ValueError('Cannot recover actual area from zero forcing; explicit layout evidence required')
            implied_area = float(np.median(t.source_J.to_numpy()[nonzero] / (dh[nonzero] * 3600)))
            if area <= 0 or not np.isclose(area, implied_area, rtol=1e-10, atol=1e-05) or (not np.allclose(t.source_J, dh * area * 3600, rtol=1e-10, atol=1e-05)):
                raise ValueError('Category-3 treated-area/source mismatch')
        results.append({'experiment_id': em['experiment_id'], 'params': identity['parameters'], 'identity': identity, 'summary': summary, 'treated_area_m2': area, 'manifest': str(p), 'layout': layouts.get(em['experiment_id'], {'coverage': m.get('coverage'), 'arrangement': 'all_roofs', 'eta': identity['parameters']['eta']})})
    if {r['params']['pcm_id'] for r in results} != set(inputs.ids):
        raise ValueError('Incomplete Category-3 candidate population')
    if any((c4_hash(p) != sha for p, sha in checked.items())):
        raise ValueError('Category-3 files changed during read; use a completed stable snapshot')
    return {'full_complete': full, 'status': m['status'], 'results': results, 'manifest': m, 'receipt_path': str(path), 'receipt_sha256': expected_sha256, 'checked_files': checked}

def c4_decide(metrics, feasible, c3_full_complete, objective):
    if not objective:
        return ('INDETERMINATE', 'Evaluation objectives and acceptable uncertainty have not been declared')
    c4_require_inputs(objective, ('objective_id', 'description', 'rationale', 'checks', 'max_sensitivity_spread_K'), 'Deployment objective')
    if not c3_full_complete:
        return ('INDETERMINATE', 'Mandatory Category-3 deployment/atmospheric/day-night/feedback experiments incomplete')
    if not isinstance(objective['checks'], list) or not objective['checks']:
        raise ValueError('At least one declared objective check required')
    spread = metrics.get('atmospheric_sensitivity_spread_K')
    limit = c4_finite(objective['max_sensitivity_spread_K'], 'declared uncertainty tolerance', nonnegative=True)
    if c4_missing(spread) or spread > limit:
        return ('INDETERMINATE', 'Atmospheric uncertainty is missing or exceeds the declared acceptable spread')
    ops = {'<': operator.lt, '<=': operator.le, '>': operator.gt, '>=': operator.ge, '==': operator.eq}
    findings = []
    for check in objective['checks']:
        c4_require_inputs(check, ('metric', 'operator', 'value', 'unit', 'rationale'), 'Objective check')
        key = check['metric']
        if key not in C4_OBJECTIVE_UNITS or check['unit'] != C4_OBJECTIVE_UNITS[key]:
            raise ValueError('Unknown objective metric or incompatible unit')
        if key not in metrics or c4_missing(metrics[key]):
            return ('INDETERMINATE', f'Required objective metric missing: {key}')
        if check['operator'] not in ops:
            raise ValueError('Unsupported objective operator')
        ok = ops[check['operator']](c4_finite(metrics[key], key), c4_finite(check['value'], 'objective value'))
        findings.append((key, ok))
    failed = [k for k, ok in findings if not ok]
    if failed:
        return ('NOT_PROMISING_UNDER_TESTED_ASSUMPTIONS', f"Declared objective {objective['objective_id']} failed: {', '.join(failed)}; no categorical Delhi conclusion")
    if feasible.classification.eq('FAILED').any():
        return ('NOT_PROMISING_UNDER_TESTED_ASSUMPTIONS', 'A sourced, applicable feasibility constraint failed: ' + ', '.join(feasible.loc[feasible.classification.eq('FAILED'), 'criterion']))
    if not feasible.classification.eq('SUPPORTED').all() or feasible.evidence_scope.eq('NUMERICAL_MODEL_ONLY').any() or feasible.missing_evidence.map(lambda x: not c4_missing(x)).any():
        return ('INDETERMINATE', 'Installation safety, durability or field regeneration evidence remains incomplete')
    return ('CONDITIONALLY_PROMISING', f"Declared objective {objective['objective_id']} met within tested assumptions, stated uncertainty and applicable feasibility evidence")

def c4_run(ns, config, evidence=None, components=None, completed_c3=None):
    ns.update(CATEGORY4_READY=False, CATEGORY4_HANDOFF=None, CATEGORY4_PARTIAL_STATUS='VALIDATING')
    inputs = ns['c3_validate_handoff'](ns)
    for pid in inputs.ids:
        ns['c3_check_consumed_inputs'](inputs, ns, pid=pid)
    results = (completed_c3 or {}).get('results', [])
    areas = {}
    base = {}
    if ns.get('C3_ROOFS') is not None:
        ns['c3_check_consumed_inputs'](inputs, ns, ns['C3_ROOFS'])
        layout = ns['c3_layout'](inputs, ns['C3_ROOFS'], 1.0, 'dispersed', ns['C3_BASE']['eta'])
        areas = {pid: layout['actual_m2'] for pid in inputs.ids}
    for pid in inputs.ids:
        rows = [r for r in results if r['params']['pcm_id'] == pid and r['params'].get('name') == 'baseline' and (r['identity']['parameters'].get('eta') == ns['C3_BASE']['eta']) and (r['layout']['coverage'] == 1.0)]
        rows = [r for r in rows if r['identity']['label'] in {'baseline', 'oneway'}]
        if rows:
            best = max(rows, key=lambda r: r['treated_area_m2'] or 0)
            base[pid] = best
            areas[pid] = best['treated_area_m2']
    feasibility, masses = c4_feasibility_table(inputs, config, evidence, components, areas)
    assessments = []
    missing = []
    for pid in inputs.ids:
        mat = inputs.by_pcm[pid]['material']
        metrics = {**masses[pid], **c4_pcm_diagnostics(inputs.by_pcm[pid])}
        atm_keys = {'annual_signed_temperature_effect_K': 'mean_K', 'gross_cooling_K_m2_s': 'G_cool_K_m2_s', 'gross_warming_K_m2_s': 'G_warm_K_m2_s', 'extreme_hot_period_mean_K': 'extreme_heat_mean_K', 'night_signed_mean_K': 'night_mean_K', 'night_gross_warming_K_m2_s': 'night_G_warm_K_m2_s'}
        summary = base.get(pid, {}).get('summary', {}).get('hotspot', {})
        metrics.update({k: summary.get(v, np.nan) for k, v in atm_keys.items()})
        pr = [r for r in results if r['params']['pcm_id'] == pid]
        sensitivity = [r['summary']['hotspot']['mean_K'] for r in pr if r['layout']['coverage'] == 1.0 and r['params'].get('eta') == ns['C3_BASE']['eta'] and (r['identity']['forcing_nature'] == 'REAL_CATEGORY1_HOURLY_MEAN') and (r['params'].get('name') != 'baseline')]
        if sensitivity and pid in base:
            sensitivity.append(base[pid]['summary']['hotspot']['mean_K'])
        metrics['atmospheric_sensitivity_spread_K'] = max(sensitivity) - min(sensitivity) if sensitivity else np.nan
        metrics['atmospheric_ensemble_min_K'] = min(sensitivity) if sensitivity else np.nan
        metrics['atmospheric_ensemble_max_K'] = max(sensitivity) if sensitivity else np.nan
        mixing = [r['summary']['hotspot']['mean_K'] for r in pr if r['layout']['coverage'] == 1.0 and r['params'].get('mixing') == 'daynight']
        metrics['daynight_mixing_mean_min_K'] = min(mixing) if mixing else np.nan
        metrics['daynight_mixing_mean_max_K'] = max(mixing) if mixing else np.nan
        deployment = [{'experiment_id': r['experiment_id'], 'layout_identity': r['identity']['layout'], 'coverage': r['layout']['coverage'], 'arrangement': r['layout']['arrangement'], 'treated_area_m2': r['treated_area_m2'], 'annual_mean_K': r['summary']['hotspot']['mean_K']} for r in pr if r['params'].get('name') == 'baseline']
        f = feasibility.loc[feasibility.pcm_id == pid]
        full = bool(completed_c3 and completed_c3['full_complete'])
        objective = config.get('objectives', {}).get(pid)
        decision, reason = c4_decide(metrics, f, full, objective)
        positive = metrics['delta_H_annual_mean_Wm2'] > 0 if not c4_missing(metrics['delta_H_annual_mean_Wm2']) else None
        finding = 'Positive annual signed roof sensible-heat difference under tested roof assumptions' if positive is True else 'Nonpositive annual signed roof sensible-heat difference under tested roof assumptions' if positive is False else 'Roof annual difference unavailable'
        row = dict(pcm_id=pid, material_name=mat['material_name'], **metrics, decision=decision, decision_reason=reason, model_specific_finding=finding, category3_annual_available=pid in base, category3_full_complete=full, feedback_adequacy='PASS_EXISTING_C3_GATES' if full else 'INCOMPLETE', deployment_sensitivity=json.dumps(deployment) if deployment else None, objectives=json.dumps(objective) if objective else None, comparator=C4_COMPARATOR, result_use='DEVELOPMENT')
        assessments.append(row)
        for _, x in f.iterrows():
            if x.classification == 'INSUFFICIENT_EVIDENCE' or not c4_missing(x.missing_evidence):
                missing.append(dict(pcm_id=pid, stage='4A', criterion=x.criterion, missing_evidence=x.missing_evidence, source=x.source))
        if not full:
            missing.append(dict(pcm_id=pid, stage='4B', criterion='completed_atmospheric_assessment', missing_evidence='Verified full Category-3 completion, deployment/atmospheric ensemble, day-night sensitivity and feedback adequacy', source=''))
        if not objective:
            missing.append(dict(pcm_id=pid, stage='4B', criterion='declared_objectives', missing_evidence='Declared evaluation objectives, units, rationale and acceptable atmospheric sensitivity spread', source=''))
    assessment = pd.DataFrame(assessments)
    missing_table = pd.DataFrame(missing, columns=['pcm_id', 'stage', 'criterion', 'missing_evidence', 'source'])
    output = Path(config['output_dir']).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    tables = {'category4_pcm_feasibility.csv': feasibility, 'category4_candidate_assessment.csv': assessment, 'category4_missing_evidence.csv': missing_table}
    paper = assessment[['pcm_id', 'material_name', 'active_pcm_kg_m2', 'total_assembly_kg_m2', 'active_h', 'equivalent_full_uptake_cycles', 'saturated_h', 'incomplete_reset_cycles', 'delta_H_annual_mean_Wm2', 'annual_signed_temperature_effect_K', 'gross_cooling_K_m2_s', 'gross_warming_K_m2_s', 'decision']]
    tables['category4_paper_results.csv'] = paper
    for name, frame in tables.items():
        frame.to_csv(output / name, index=False, na_rep='')
    assumptions = {k: inputs.upstream.config[k] for k in ['year', 'd_m', 'alpha_s', 'epsilon_LW', 'bottom_boundary'] if k in inputs.upstream.config}
    assumptions['phase_models'] = {pid: inputs.by_pcm[pid]['material']['phase_model'] for pid in inputs.ids}
    text = f"Engineering feasibility uses verified dynamically screened Category-1 candidates and mass/area = density × thickness. Roof assumptions: {json.dumps(assumptions)}. Comparator: {C4_COMPARATOR}. Signed ΔH is H_PCM−H_L0, positive into air. Category-3 annual baseline available for {len(base)}/{len(inputs.ids)} candidates; full completion verified: {bool(completed_c3 and completed_c3['full_complete'])}. Treated loading area uses full eligible-roof coverage with η={ns['C3_BASE']['eta']} where verified geometry or results exist. Atmospheric values refer to the inherited hotspot receptor and effective vertically mixed air-layer anomaly; gross integrals use K m² s. Sensitivity spread describes tested scenarios, not a probability interval. Missing component weights, safety evidence and results remain missing. Numerical regeneration/cycling do not prove real freezing, retained latent capacity, field durability or structural safety. The PCM workbook and cached population remain DEVELOPMENT until the final reviewed workbook is frozen and required simulations repeated. No economics, health benefit, lifecycle assessment or commercial cool-roof comparison was calculated.\n"
    (output / 'category4_methodology_limitations.txt').write_text(text)
    ready = bool(completed_c3 and completed_c3['full_complete']) and assessment.decision.ne('INDETERMINATE').all() and missing_table.empty
    status = 'COMPLETE_CONDITIONAL_ASSESSMENT_DEVELOPMENT' if ready else 'PARTIAL_FEASIBILITY_WITH_INCOMPLETE_DEPLOYMENT_ASSESSMENT'
    handoff = {'status': status, 'CATEGORY4_READY': bool(ready), 'candidate_ids': list(inputs.ids), 'upstream_identity': inputs.identity, 'source_identity': ns['CUMULATIVE_SOURCE_IDENTITY'], 'category3_receipt': {k: completed_c3[k] for k in ['receipt_path', 'receipt_sha256', 'status', 'full_complete']} if completed_c3 else None, 'result_use': 'DEVELOPMENT', 'comparator': C4_COMPARATOR, 'config': config, 'artifacts': {name: c4_hash(output / name) for name in [*tables, 'category4_methodology_limitations.txt']}}
    ns['c3_write_json'](output / 'category4_handoff.json', handoff)
    ns.update(CATEGORY4_READY=bool(ready), CATEGORY4_HANDOFF=handoff, CATEGORY4_PARTIAL_STATUS=status, CATEGORY4_FEASIBILITY=feasibility, CATEGORY4_ASSESSMENT=assessment, CATEGORY4_PAPER_TABLE=paper)
    return handoff

# Executed stage lineage: original cells 29, 31, 33, 45, 47, 57, 59, 63, 71, 77, 79, 83, 84, 86, 88, 90, 92, 94
def run_preliminary():
    """Explicit original acquisition, weather, screening and plotting stages. Requires authorised Earth Engine access."""
    global AIR_ROOF_SCREEN_LOWER_UNROUNDED_C, ANNUAL_WEATHER, AOI, BUILT_FRACTION, COMMON_EE_PROJECTION, COMMON_GRID, DELHI_BOUNDARY, DELHI_MASK, DELHI_PROJECTED_SHAPE, ERA5_COVERAGE, ERA5_OVERLAP_LEDGER, ERA5_RAW, ERA5_SCREEN_METADATA, ERA5_SHORTWAVE_CORRECTION_LEDGER, ERA5_SHORTWAVE_QC_SOURCE, ERA5_SHORTWAVE_QC_SUMMARY, ERA5_SOURCE_METADATA, FUTURE_INPUT_REQUIREMENTS, HANDOFF_OBJECTS, HOTSPOT_MASK, HOTSPOT_REGION_AUDIT, HOTSPOT_SUMMARY, HOTSPOT_WEATHER_LOCATION, LANDSAT_AVAILABILITY, LANDSAT_AVAILABILITY_ROWS, LANDSAT_OUTPUTS, LST_INDEX_RELATIONSHIPS, LST_POSTMONSOON_C, LST_PREMONSOON_C, LST_ROOF_SCREEN_UPPER_UNROUNDED_C, NDBI_POSTMONSOON, NDBI_PREMONSOON, NDVI_POSTMONSOON, NDVI_PREMONSOON, OVERLAP_SCREEN_TEST_RECEIPT, PATH_UTILITIES, PCM_ALL, PCM_EXCLUDED, PCM_NEEDS_DATA, PCM_PRELIM_CANDIDATES, PCM_PRELIM_CANDIDATE_IDS, PCM_RUNNABLE, PCM_SCREEN_DERIVATION, PCM_SCREEN_LOCAL_DAILY_AUDIT, PCM_TABLE_SUMMARY, PCM_TEMP_SCREEN, PCM_USER_VERIFICATION, PCM_UTILITIES, PCM_WORKBOOK_INFO, PRELIM_EE_SESSION, PRELIM_ENVIRONMENT, PRELIM_HANDOFF, PRELIM_PROVENANCE, PRELIM_READY, PROVENANCE_TABLE, READINESS_TABLE, ROOF_SCREEN_LOWER_C, ROOF_SCREEN_UPPER_C, SEASONAL_SPATIAL_SUMMARY, SEASONAL_SPATIAL_SUMMARY_ROWS, SECTION0_HANDOFF, SECTION0_READY, SPATIAL_UTILITIES, TIME_UTILITIES, VALIDATION_UTILITIES, _pcm_frame, _pcm_parts, _pcm_record, a, all_lst, arr, array, arrays, availability, ax, axes, axis, b, bins, bmeta, boundary_cfg, boundary_feature, boundary_features, bpath, c, cmap, coefficient, color, cols_grid, column, d, daily_path, dataset_id, e, ee, end, f, feature_count, field, fig, filename, fn, gcfg, geographic_geometry, geographic_shape, ghsl_asset, ghsl_asset_id, ghsl_check, ghsl_fraction_image, ghsl_info, ghsl_native_range, gmeta, gpath, handle, handoff_raster_values, hmeta, idx, im, index, k, key, ledger_qc_units, limits, local_daily_checks, local_daily_meta, lst, lst_limits, lst_screen_values, manifest, manifest_path, n, name, native_cell_area_m2, native_crs, native_crs_definition, native_projection, native_transform, npz_path, obj, original, persisted_daily, plt, probe, projector, raw_shortwave_before_qc, raw_weather_path, raw_weather_sha256_before_qc, ready, record, relative, reloaded_materials, repeated_audit, repeated_candidates, repeated_correction_ledger, repeated_coverage, repeated_daily, repeated_daily_meta, repeated_ids, repeated_ledger, repeated_location, repeated_lower, repeated_qc_source, repeated_qc_summary, repeated_screen, repeated_summary, repeated_weather, requests, required_sources, reselected, restored, restored_daily, rmeta, row, rows, rows_grid, screen_cfg, season, season_outputs, smeta, source_qc_units, start, status, support, units, valid, value, values, weather_checks, wmeta, wpath, wraw_meta, xx, y, yy
    OVERLAP_SCREEN_TEST_RECEIPT = run_overlap_screen_tests()
    PCM_ALL, _pcm_parts, PCM_WORKBOOK_INFO = load_section0_pcm(CONFIG, PROVENANCE_REGISTRY)
    PCM_RUNNABLE = _pcm_parts['runnable']
    PCM_NEEDS_DATA = _pcm_parts['needs_data']
    PCM_EXCLUDED = _pcm_parts['excluded']
    del _pcm_parts
    PCM_TABLE_SUMMARY = pd.DataFrame({'table': ['PCM_ALL', 'PCM_RUNNABLE', 'PCM_NEEDS_DATA', 'PCM_EXCLUDED'], 'rows': [len(PCM_ALL), len(PCM_RUNNABLE), len(PCM_NEEDS_DATA), len(PCM_EXCLUDED)], 'meaning': ['Every original workbook row', 'Schema only; no Delhi selection', 'Missing/invalid data or unsupported phase law', 'Explicitly excluded; retained for audit'], 'source_quality': ['NOT CHECKED'] * 4})
    display(PCM_TABLE_SUMMARY)
    display(PCM_RUNNABLE[['pcm_id', 'material_name', 'section0_readiness', 'source_quality_status']])
    PCM_USER_VERIFICATION = dict(PCM_WORKBOOK_REVIEW)
    PCM_WORKBOOK_INFO['user_verification'] = dict(PCM_USER_VERIFICATION)
    for _pcm_frame in (PCM_ALL, PCM_RUNNABLE, PCM_NEEDS_DATA, PCM_EXCLUDED):
        _pcm_frame.attrs['user_verification'] = dict(PCM_USER_VERIFICATION)
    _pcm_record = PROVENANCE_REGISTRY.get('pcm_workbook')
    PROVENANCE_REGISTRY.register('pcm_workbook', source=_pcm_record['source'], planned_source=_pcm_record['planned_source'], availability='PRESENT', data_nature='REAL', scope_kind='material', required_in_section0=True, units=PCM_UNITS, expected_units=PCM_UNITS, native_crs=_pcm_record['native_crs'], native_resolution=_pcm_record['native_resolution'], file_info={k: _pcm_record[k] for k in ['file_path', 'sha256', 'file_size_bytes', 'file_modified_at_UTC']}, quality_status=_pcm_record['quality_status'], source_quality_status='PASS', checks=_pcm_record['checks'] + [{'name': 'user_material_value_verification', 'result': 'PASS', 'detail': PCM_USER_VERIFICATION['basis'] + '; attached workbook SHA-256 ' + PCM_USER_VERIFICATION['workbook_sha256']}], notes='Material values checked and verified by the user for the exact attached workbook. Software schema/model compatibility checks remain separate.', replace=True)
    print('PCM material-value verification:', PCM_USER_VERIFICATION['status'])
    FUTURE_INPUT_REQUIREMENTS = (('study_boundary', 'Reviewed Delhi administrative/study-domain vector boundary', 'Geometry with declared native CRS and source version', 'Required before the common grid extent is defined.'), ('era5_hourly', 'ECMWF/ERA5_LAND/HOURLY or a reviewed equivalent supplied file', {'T_air': 'degC after explicit conversion', 'u10': 'm/s', 'v10': 'm/s', 'solar_down': 'W/m2 interval mean', 'longwave_down': 'W/m2 interval mean'}, 'Full local-year coverage and source-native time/accumulation convention must be verified; not retrieved.'), ('landsat_observations', 'Landsat 8/9 Collection 2 Level 2; scenes and QA to be reviewed', 'LST degC, reflectance dimensionless after sourced scaling; native QA flags retained', 'Scene dates, masks, compositing method and native resolution are unselected; no composites created.'), ('ghsl_built_surface', 'JRC/GHSL/P2023A/GHS_BUILT_S or reviewed release/epoch', 'Native built surface m2; any derived fraction requires correct source-cell area', 'Release, observation epoch and native cell area require review; not a measured roof survey.'), ('roof_area_and_eligibility', 'Measured/sourced roof plan area and PCM installation eligibility', 'm2 and/or dimensionless fractions with documented denominator', 'Real georeferenced layers needed; GHSL must not silently substitute for surveyed roof eligibility.'), ('canopy_wind', 'Measured or independently justified canopy-height wind with height and processing metadata', {'Ux': 'm/s', 'Uy': 'm/s', 'height': 'm'}, 'Coverage and wind basis unverified.'), ('building_geometry', 'Sourced building/canopy geometry for the selected domain', 'm and georeferenced geometry', 'No universal building height is inserted.'), ('air_exchange_parameter_evidence', 'Literature/measurement/calibration records for future air-transport parameters', 'Parameter-specific SI units', 'Values, uncertainty and representativeness remain unset.'), ('roof_material_and_boundary_evidence', 'Sourced reference roof/assembly material properties and boundary conditions', 'Material-specific SI units', 'No thermal stack or bottom-boundary model is assumed here.'), ('cool_roof_optical_evidence', 'Sourced initial and aged coating optical measurements', 'Dimensionless reflectance/emissivity and dated state metadata', 'No optical or ageing values are inserted.'), ('pcm_retention_and_lifetime_evidence', 'Material/system-specific cycle retention, durability and service-life evidence', 'Fractions, cycles and years as explicitly reported', 'No degradation curve or lifetime is invented.'), ('assembly_mass_and_structural_evidence', 'Sourced installation mass and structural allowance records', 'kg/m2 with assembly scope', 'No structural feasibility decision is made.'), ('safety_and_compatibility_evidence', 'Source records for fire, toxicity, corrosion and leakage behaviour', 'Evidence/test classifications with conditions', 'No generic safety pass is assigned.'), ('cost_and_maintenance_evidence', 'Dated local quotations/records for supply, installation, maintenance, replacement and disposal', 'Currency with price year and denominator', 'No costs, totals or economics are calculated.'), ('economic_method_inputs', 'Explicit sourced financial assumptions and study horizon for any future economic section', 'Years, fractions and currency conventions', 'No discount rate, cost year or lifetime is assumed.'), ('independent_validation_observations', 'Field/CFD/benchmark observations with methods and uncertainty', 'Measurement-specific units', 'No empirical or CFD validation has been performed.'))
    register_future_inputs(PROVENANCE_REGISTRY, CONFIG)
    PROVENANCE_TABLE = PROVENANCE_REGISTRY.to_frame()
    READINESS_TABLE = make_section0_readiness(PROVENANCE_REGISTRY, PCM_TABLE_SUMMARY)
    SECTION0_READY = bool((READINESS_TABLE.loc[READINESS_TABLE.required_now == 'YES', 'section0_status'] == 'LOADED / AUDITED').all())
    if not SECTION0_READY:
        raise InputValidationError('Section 0 is not ready: a currently required input is missing.')
    PATH_UTILITIES = {fn.__name__: fn for fn in (resolve_path, sha256_file, require_file, prepare_output_directories)}
    TIME_UTILITIES = {fn.__name__: fn for fn in (to_utc_index, to_local_index, expected_hourly_intervals, annual_hourly_intervals, validate_hourly_coverage)}
    SPATIAL_UTILITIES = {fn.__name__: fn for fn in (GridSpec, validate_grid_spec, define_common_grid, raster_grid_from_file, validate_grid_alignment, validate_raster_grid)}
    VALIDATION_UTILITIES = {fn.__name__: fn for fn in (InputValidationError, is_unknown, require_fields, require_columns, require_finite_numeric, validate_config, require_registered_input)}
    PCM_UTILITIES = {fn.__name__: fn for fn in (clean_pcm_properties, load_pcm_properties, inspect_pcm_workbook, audit_cleaned_pcm, partition_pcm_tables, load_section0_pcm)}
    SECTION0_HANDOFF = {'CONFIG': CONFIG, 'PROVENANCE_REGISTRY': PROVENANCE_REGISTRY, 'PROVENANCE_TABLE': PROVENANCE_TABLE, 'PCM_ALL': PCM_ALL, 'PCM_RUNNABLE': PCM_RUNNABLE, 'PCM_NEEDS_DATA': PCM_NEEDS_DATA, 'PCM_EXCLUDED': PCM_EXCLUDED, 'PCM_UNITS': PCM_UNITS, 'PCM_WORKBOOK_INFO': PCM_WORKBOOK_INFO, 'PCM_TABLE_SUMMARY': PCM_TABLE_SUMMARY, 'ANNUAL_HOURLY_INTERVALS': ANNUAL_HOURLY_INTERVALS, 'EXPECTED_ANNUAL_HOURS': EXPECTED_ANNUAL_HOURS, 'COMMON_GRID': COMMON_GRID, 'PATH_UTILITIES': PATH_UTILITIES, 'TIME_UTILITIES': TIME_UTILITIES, 'SPATIAL_UTILITIES': SPATIAL_UTILITIES, 'VALIDATION_UTILITIES': VALIDATION_UTILITIES, 'PCM_UTILITIES': PCM_UTILITIES, 'READINESS_TABLE': READINESS_TABLE, 'SECTION0_READY': SECTION0_READY, 'ENVIRONMENT_VERSIONS': ENVIRONMENT_VERSIONS, 'REUSED_REPOSITORY_CODE': REUSED_REPOSITORY_CODE, 'FUTURE_INPUT_REQUIREMENTS': FUTURE_INPUT_REQUIREMENTS}
    display(pd.DataFrame([{'in_memory_object': name, 'type': type(obj).__name__} for name, obj in SECTION0_HANDOFF.items()]))
    PRELIM_EE_SESSION = None
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        import ee
        import requests
        import matplotlib.pyplot as plt
        from shapely.geometry import shape, mapping
        from shapely.ops import transform as transform_geometry
        from pyproj import Transformer
        from rasterio.features import geometry_mask
        from rasterio.io import MemoryFile
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        PRELIM_EE_SESSION = None
        require_prelim(bool(str(PRELIM_CONFIG['ee_project']).strip()), 'Set PRELIM_CONFIG["ee_project"] explicitly.')
        try:
            if IN_COLAB:
                ee.Authenticate(auth_mode=PRELIM_CONFIG['ee_auth_mode_colab'])
            else:
                ee.Authenticate()
            ee.Initialize(project=PRELIM_CONFIG['ee_project'])
            probe = ee.data.getAsset(PRELIM_CONFIG['boundary']['collection'])
            require_prelim(probe.get('type') in {'TABLE', 'FEATURE_COLLECTION'}, 'Boundary asset has unexpected type.')
            PRELIM_EE_SESSION = {'project': PRELIM_CONFIG['ee_project'], 'boundary_asset_probe': probe, 'initialized_at_UTC': datetime.now(timezone.utc).isoformat()}
        except Exception as error:
            raise RuntimeError('Earth Engine authentication/initialization failed. Run this cell in Colab with an authorized project and account. No scientific handoff is available.') from error
        PRELIM_RUN_DIR.mkdir(parents=True, exist_ok=False)
        PRELIM_ENVIRONMENT = {name: importlib.metadata.version(name) for name in ['earthengine-api', 'requests', 'shapely', 'matplotlib', 'numpy', 'pandas', 'rasterio', 'pyproj']}
        save_json_artifact('preliminary_config.json', {'global': CONFIG, 'preliminary': PRELIM_CONFIG, 'environment': PRELIM_ENVIRONMENT, 'run_id': PRELIM_RUN_ID}, artifact_metadata('Section-0 global policy + visible Section-1 configuration', 'configuration', status='CONFIGURATION_ONLY'))
        print('Earth Engine initialized; new output directory:', PRELIM_RUN_DIR)
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        DELHI_BOUNDARY = None
        COMMON_GRID = None
        require_prelim(PRELIM_EE_SESSION is not None, 'Authenticate Earth Engine before boundary retrieval.')
        boundary_cfg = PRELIM_CONFIG['boundary']
        try:
            boundary_features = ee.FeatureCollection(boundary_cfg['collection'])
            for key, value in boundary_cfg['filters'].items():
                boundary_features = boundary_features.filter(ee.Filter.eq(key, value))
            feature_count = int(boundary_features.size().getInfo())
            require_prelim(feature_count == boundary_cfg['required_feature_count'], f'Delhi boundary filter returned {feature_count} features; expected exactly one. Review filters explicitly.')
            boundary_feature = boundary_features.first().getInfo()
            AOI = boundary_features.geometry()
            geographic_geometry = AOI.getInfo()
        except Exception as error:
            raise RuntimeError('Real GAUL Delhi boundary retrieval failed; no inferred bounding box is available.') from error
        geographic_shape = shape(geographic_geometry)
        require_prelim(geographic_shape.geom_type in {'Polygon', 'MultiPolygon'} and (not geographic_shape.is_empty) and geographic_shape.is_valid, 'Retrieved Delhi geometry is empty/invalid/nonpolygonal. Review the source.')
        projector = Transformer.from_crs(boundary_cfg['geographic_crs'], CONFIG['spatial']['crs'], always_xy=True)
        DELHI_PROJECTED_SHAPE = transform_geometry(projector.transform, geographic_shape)
        require_prelim(DELHI_PROJECTED_SHAPE.is_valid and DELHI_PROJECTED_SHAPE.area > 0, 'Projected Delhi boundary is invalid or has zero area.')
        COMMON_GRID = define_common_grid(DELHI_PROJECTED_SHAPE.bounds, CONFIG)
        DELHI_MASK = geometry_mask([mapping(DELHI_PROJECTED_SHAPE)], out_shape=COMMON_GRID.shape, transform=COMMON_GRID.transform, invert=True, all_touched=False)
        require_prelim(DELHI_MASK.any(), 'No common-grid cell centers lie inside the retrieved boundary.')
        DELHI_BOUNDARY = {'source_collection': boundary_cfg['collection'], 'source_version': boundary_cfg['version'], 'source_asset_metadata': PRELIM_EE_SESSION['boundary_asset_probe'], 'selection': boundary_cfg['filters'], 'feature_count': feature_count, 'feature_properties': boundary_feature['properties'], 'geometry': geographic_geometry, 'crs': boundary_cfg['geographic_crs'], 'bounds_lonlat': list(geographic_shape.bounds), 'projected_geometry': mapping(DELHI_PROJECTED_SHAPE), 'projected_crs': COMMON_GRID.crs, 'projected_bounds_xy_m': list(DELHI_PROJECTED_SHAPE.bounds), 'projected_boundary_area_m2': DELHI_PROJECTED_SHAPE.area, 'boundary_membership': boundary_cfg['raster_membership'], 'data_nature': 'REAL'}
        bmeta = artifact_metadata(boundary_cfg['collection'], 'geometry', window=['2015-01-01', '2016-01-01'], grid=COMMON_GRID)
        bpath = save_json_artifact('delhi_boundary.json', DELHI_BOUNDARY, bmeta)
        save_json_artifact('delhi_boundary.geojson', {'type': 'Feature', 'geometry': geographic_geometry, 'properties': boundary_feature['properties']}, bmeta)
        save_json_artifact('common_grid.json', grid_metadata(COMMON_GRID), bmeta)
        register_prelim_source('study_boundary', boundary_cfg['collection'] + '; India/Delhi ADM1', 'geometry EPSG:4326', ['2015-01-01', '2016-01-01'], [{'name': 'real_service_geometry', 'result': 'PASS', 'detail': 'Authenticated service response; exact country/state filter; one valid polygon feature.'}], native_crs=boundary_cfg['geographic_crs'], native_resolution='VECTOR', file_info=require_file(bpath, label='Delhi boundary'))
        PRELIM_SOURCE_METADATA['study_boundary'] = DELHI_BOUNDARY
        COMMON_EE_PROJECTION = ee.Projection(COMMON_GRID.crs, list(tuple(COMMON_GRID.transform)[:6]))
        display(pd.DataFrame([grid_metadata(COMMON_GRID)]))
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        LANDSAT_OUTPUTS = {}
        LANDSAT_AVAILABILITY_ROWS = []
        for season in PRELIM_CONFIG['landsat']['windows']:
            season_outputs, availability = retrieve_landsat_season(season)
            LANDSAT_OUTPUTS.update(season_outputs)
            LANDSAT_AVAILABILITY_ROWS.extend(availability)
        LST_PREMONSOON_C = LANDSAT_OUTPUTS['LST_PREMONSOON_C']
        LST_POSTMONSOON_C = LANDSAT_OUTPUTS['LST_POSTMONSOON_C']
        NDVI_PREMONSOON = LANDSAT_OUTPUTS['NDVI_PREMONSOON']
        NDVI_POSTMONSOON = LANDSAT_OUTPUTS['NDVI_POSTMONSOON']
        NDBI_PREMONSOON = LANDSAT_OUTPUTS['NDBI_PREMONSOON']
        NDBI_POSTMONSOON = LANDSAT_OUTPUTS['NDBI_POSTMONSOON']
        LANDSAT_AVAILABILITY = pd.DataFrame(LANDSAT_AVAILABILITY_ROWS)
        save_table_artifact('landsat_availability.csv', LANDSAT_AVAILABILITY, artifact_metadata('; '.join(PRELIM_CONFIG['landsat']['collections']), 'coverage and scene counts', grid=COMMON_GRID))
        display(LANDSAT_AVAILABILITY)
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        BUILT_FRACTION = None
        gcfg = PRELIM_CONFIG['ghsl']
        ghsl_asset_id = f"{gcfg['collection']}/{gcfg['epoch']}"
        try:
            ghsl_asset = ee.Image(ghsl_asset_id)
            ghsl_info = ghsl_asset.getInfo()
            require_prelim(gcfg['band'] in ghsl_asset.bandNames().getInfo(), 'Required GHSL built_surface band missing.')
            native_projection = ghsl_asset.select(gcfg['band']).projection().getInfo()
        except Exception as error:
            raise RuntimeError('Real GHSL source retrieval failed; BUILT_FRACTION is unavailable.') from error
        native_crs_definition = native_projection.get('crs') or native_projection.get('wkt')
        require_prelim(native_crs_definition is not None, f"GHSL projection metadata contains neither 'crs' nor 'wkt': {native_projection}")
        native_transform = native_projection.get('transform')
        require_prelim(native_transform is not None and len(native_transform) == 6, f'GHSL projection metadata lacks a valid 6-element affine transform: {native_projection}')
        native_crs = CRS.from_user_input(native_crs_definition)
        require_prelim(native_crs.is_projected and all((np.isclose(axis.unit_conversion_factor, 1.0) for axis in native_crs.axis_info[:2])), 'GHSL native grid must be projected in metres for this area conversion.')
        a, b, c, d, e, f = native_transform
        native_cell_area_m2 = abs(a * e - b * d)
        require_prelim(np.isclose(native_cell_area_m2, gcfg['catalog_cell_area_m2'], rtol=0, atol=0.01), f'GHSL native area {native_cell_area_m2} disagrees with catalog cell area; review, do not guess denominator.')
        ghsl_native_range = ghsl_asset.select(gcfg['band']).reduceRegion(reducer=ee.Reducer.minMax(), geometry=AOI, crs=ee.Projection(native_crs_definition, native_transform), scale=float(np.sqrt(native_cell_area_m2)), maxPixels=PRELIM_CONFIG['weather']['max_pixels'], bestEffort=False).getInfo()
        require_prelim(ghsl_native_range.get(gcfg['band'] + '_min') is not None and ghsl_native_range.get(gcfg['band'] + '_max') is not None and (0 <= ghsl_native_range[gcfg['band'] + '_min'] <= ghsl_native_range[gcfg['band'] + '_max'] <= native_cell_area_m2), 'Native GHSL built_surface violates source cell area bounds; no automatic clipping.')
        ghsl_fraction_image = ghsl_asset.select(gcfg['band']).divide(native_cell_area_m2).rename('BUILT_FRACTION').reduceResolution(reducer=ee.Reducer.mean(), bestEffort=False, maxPixels=PRELIM_CONFIG['landsat']['reduce_resolution_max_pixels']).reproject(COMMON_EE_PROJECTION).clip(AOI)
        BUILT_FRACTION = download_common_band(ghsl_fraction_image, 'BUILT_FRACTION')
        ghsl_check = validate_field(BUILT_FRACTION, 'BUILT_FRACTION', lo=0, hi=1)
        gmeta = artifact_metadata(ghsl_asset_id, 'dimensionless built-surface fraction', window=[f"{gcfg['epoch']}-01-01", f"{gcfg['epoch'] + 1}-01-01"], grid=COMMON_GRID)
        gmeta.update(native_projection=native_projection, native_cell_area_m2=native_cell_area_m2, native_source_range=ghsl_native_range, source_unit=gcfg['source_unit'], aggregation=gcfg['spatial_method'], source_asset=ghsl_info)
        gpath = save_real_raster('BUILT_FRACTION', BUILT_FRACTION, gmeta)
        register_prelim_source('ghsl_built_surface', ghsl_asset_id + ' built_surface', 'dimensionless after m2/cell-area conversion', gmeta['data_window'], [ghsl_check], native_crs=native_crs_definition, native_resolution=[abs(a), abs(e)], file_info=gpath, notes='2020 morphology snapshot retained; built surface is not surveyed eligible roof area.')
        PRELIM_SOURCE_METADATA['ghsl_built_surface'] = gmeta
        PRELIM_CHECKS.append(ghsl_check)
        print('GHSL native area:', native_cell_area_m2, 'm²; valid built-fraction cells:', np.isfinite(BUILT_FRACTION).sum())
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        HOTSPOT_MASK = None
        HOTSPOT_MASK, HOTSPOT_SUMMARY, HOTSPOT_REGION_AUDIT = select_persistent_hotspot(LST_PREMONSOON_C, LST_POSTMONSOON_C, BUILT_FRACTION, DELHI_MASK, COMMON_GRID, PRELIM_CONFIG['hotspot'])
        hmeta = artifact_metadata('Derived solely from real seasonal Landsat LST and GHSL BUILT_FRACTION', 'classification and area m2', grid=COMMON_GRID)
        save_json_artifact('hotspot_summary.json', HOTSPOT_SUMMARY, hmeta)
        save_table_artifact('hotspot_region_audit.csv', HOTSPOT_REGION_AUDIT, hmeta)
        display(pd.DataFrame([HOTSPOT_SUMMARY]))
        display(HOTSPOT_REGION_AUDIT.head(24))
        require_prelim(HOTSPOT_MASK is not None, 'No hotspot meets the unchanged thresholds; failure summary persisted. No handoff.')
        HOTSPOT_MASK.setflags(write=False)
        save_real_raster('HOTSPOT_MASK', HOTSPOT_MASK.astype(np.float64), {**hmeta, 'units': 'binary membership (0/1)'})
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        ANNUAL_WEATHER = None
        HOTSPOT_WEATHER_LOCATION = derive_hotspot_weather_location(HOTSPOT_MASK, BUILT_FRACTION, COMMON_GRID)
        require_prelim(PRELIM_EE_SESSION is not None, 'Annual ERA5 retrieval requires real EE access.')
        try:
            ERA5_RAW, ERA5_SOURCE_METADATA = retrieve_ee_hourly_weather(HOTSPOT_WEATHER_LOCATION)
        except Exception as error:
            raise RuntimeError('Annual real ERA5-Land acquisition failed; ANNUAL_WEATHER is unavailable. Run in authorized Colab.') from error
        wraw_meta = artifact_metadata(PRELIM_CONFIG['weather']['collection'], 'K; m/s; hourly energy J/m2', window=[ERA5_SOURCE_METADATA['fetch_start_UTC'], ERA5_SOURCE_METADATA['fetch_end_UTC_exclusive']], status='ACQUIRED_REAL_SOURCE_CONVERSION_PENDING')
        raw_weather_path = save_table_artifact('era5_hourly_raw.csv', ERA5_RAW, {**wraw_meta, **ERA5_SOURCE_METADATA})
        save_json_artifact('hotspot_weather_location.json', HOTSPOT_WEATHER_LOCATION, {**wraw_meta, **ERA5_SOURCE_METADATA})
        raw_weather_sha256_before_qc = sha256_file(raw_weather_path)
        raw_shortwave_before_qc = ERA5_RAW[PRELIM_CONFIG['weather']['bands'][3]].to_numpy(copy=True)
        ANNUAL_WEATHER, ERA5_OVERLAP_LEDGER, ERA5_COVERAGE, ERA5_SHORTWAVE_QC_SOURCE, ERA5_SHORTWAVE_CORRECTION_LEDGER, ERA5_SHORTWAVE_QC_SUMMARY = prepare_annual_real_weather(ERA5_RAW)
        require_prelim(np.array_equal(raw_shortwave_before_qc, ERA5_RAW[PRELIM_CONFIG['weather']['bands'][3]].to_numpy()) and sha256_file(raw_weather_path) == raw_weather_sha256_before_qc, 'Physical QC altered acquired raw shortwave values or the raw acquisition artifact.')
        require_prelim(ANNUAL_WEATHER.attrs['spatial_support'] == ERA5_SOURCE_METADATA['spatial_support'], 'Canonical weather lost its hotspot native-cell support.')
        wmeta = artifact_metadata(PRELIM_CONFIG['weather']['collection'], ANNUAL_WEATHER.attrs['units'])
        wmeta.update(source_metadata=ERA5_SOURCE_METADATA, spatial_support=ANNUAL_WEATHER.attrs['spatial_support'], raw_source_file='era5_hourly_raw.csv', raw_source_sha256=sha256_file(raw_weather_path), intended_downstream_use='Category 1 selected-hotspot weather', coverage=ERA5_COVERAGE, time_semantics=ANNUAL_WEATHER.attrs, shortwave_physical_qc=deepcopy(ERA5_SHORTWAVE_QC_SUMMARY), solar_negative_hours=int(ANNUAL_WEATHER.shortwave_negative_flag.sum()), minimum_shortwave_Wm2=float(ANNUAL_WEATHER.G_Wm2.min()))
        wpath = save_table_artifact('annual_weather.csv', ANNUAL_WEATHER, wmeta)
        verify_cached_weather_table(wpath, ANNUAL_WEATHER)
        source_qc_units = {'raw_shortwave_energy_Jm2': 'J/m2 source UTC interval', 'model_shortwave_energy_Jm2': 'J/m2 source UTC interval', 'shortwave_qc_correction_Jm2': 'J/m2 source UTC interval', 'shortwave_qc_applied_flag': 'boolean source-hour QC flag', 'shortwave_qc_status': 'source-hour QC classification'}
        ledger_qc_units = {**source_qc_units, 'source_energy_weight': 'fraction of source UTC interval', 'overlap_seconds': 's', 'raw_shortwave_contribution_Jm2': 'J/m2 target contribution', 'model_shortwave_contribution_Jm2': 'J/m2 target contribution', 'shortwave_correction_contribution_Jm2': 'J/m2 target correction contribution'}
        save_table_artifact('era5_radiation_overlap_ledger.csv', ERA5_OVERLAP_LEDGER, {**wmeta, 'units': ledger_qc_units})
        save_table_artifact('era5_shortwave_source_qc.csv', ERA5_SHORTWAVE_QC_SOURCE, {**wmeta, 'units': source_qc_units})
        save_table_artifact('era5_shortwave_correction_ledger.csv', ERA5_SHORTWAVE_CORRECTION_LEDGER, {**wmeta, 'units': ledger_qc_units})
        save_json_artifact('era5_shortwave_qc_summary.json', ERA5_SHORTWAVE_QC_SUMMARY, {**wmeta, 'units': 'J/m2 source/annual energy as named; W/m2 canonical interval-mean minimum'})
        weather_checks = [{'name': 'provider_shortwave_physical_qc', 'result': 'PASS', 'detail': 'Raw increments preserved; only [-5,0) J/m2 mapped to zero before overlap integration; below -5 hard-fails; correction ledgers reconcile; canonical G>=0.'}, {'name': 'hotspot_native_cell_support', 'result': 'PASS', 'detail': 'Reproducible point inside selected hotspot; exactly one unchanged native ERA5 cell each hour; projection and scale retained.'}, {'name': 'annual_hourly_coverage_and_physical_units', 'result': 'PASS', 'detail': f'{len(ANNUAL_WEATHER)} exact intervals; raw native hours checked; finite quantities; real band identities.'}, {'name': 'radiation_overlap_equivalence', 'result': 'PASS', 'detail': 'Loop integral equals independently constructed source/target overlap ledger; J/m2 divided by 3600 seconds.'}]
        register_prelim_source('era5_hourly', PRELIM_CONFIG['weather']['collection'], ANNUAL_WEATHER.attrs['units'], [CONFIG['study_start_local'], CONFIG['study_end_local_exclusive']], weather_checks, native_crs=ERA5_SOURCE_METADATA['native_projection']['crs'], native_resolution={'catalog_nominal_m': PRELIM_CONFIG['weather']['catalog_nominal_resolution_m'], 'actual_nominal_scale_m': ERA5_SOURCE_METADATA['spatial_support']['native_nominal_scale_m'], 'native_transform': ERA5_SOURCE_METADATA['native_projection']['transform']}, file_info=require_file(wpath, label='Annual real ERA5 weather'), notes=json.dumps(json_ready({'spatial_support': ERA5_SOURCE_METADATA['spatial_support'], 'raw_source_file': 'era5_hourly_raw.csv', 'radiation': 'hourly energy ends at source validity time; explicit shortwave quantization/deaccumulation physical QC before overlap integration', 'shortwave_physical_qc': ERA5_SHORTWAVE_QC_SUMMARY})))
        PRELIM_SOURCE_METADATA['era5_hourly'] = wmeta
        PRELIM_CHECKS.extend(weather_checks)
        print(json.dumps(json_ready(ERA5_SOURCE_METADATA['spatial_support']), indent=2))
        display(pd.DataFrame([ERA5_COVERAGE]))
        display(ANNUAL_WEATHER.iloc[[0, -1]])
        print('Provider shortwave physical-QC correction summary:')
        print(json.dumps(json_ready(ERA5_SHORTWAVE_QC_SUMMARY), indent=2))
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        PCM_TEMP_SCREEN = None
        PCM_PRELIM_CANDIDATES = None
        PCM_PRELIM_CANDIDATE_IDS = None
        try:
            AIR_ROOF_SCREEN_LOWER_UNROUNDED_C, PCM_SCREEN_LOCAL_DAILY_AUDIT, ERA5_SCREEN_METADATA = derive_local_day_roof_screen_air(ANNUAL_WEATHER, ERA5_RAW)
        except Exception as error:
            raise RuntimeError('Complete local-day hotspot hourly air input unavailable. No missing-hour/day repair or manual bounds are permitted.') from error
        screen_cfg = PRELIM_CONFIG['pcm_screen']
        lst_screen_values = LST_PREMONSOON_C[DELHI_MASK & np.isfinite(LST_PREMONSOON_C)]
        require_prelim(lst_screen_values.size > 0, 'No real LST pixels for the roof-screen 95th percentile.')
        LST_ROOF_SCREEN_UPPER_UNROUNDED_C = float(np.percentile(lst_screen_values, screen_cfg['lst_percentile'], method=screen_cfg['quantile_method']))
        ROOF_SCREEN_LOWER_C = round(AIR_ROOF_SCREEN_LOWER_UNROUNDED_C, screen_cfg['round_decimals'])
        ROOF_SCREEN_UPPER_C = round(LST_ROOF_SCREEN_UPPER_UNROUNDED_C, screen_cfg['round_decimals'])
        require_file(CONFIG['pcm_workbook_path'], label='Current PCM workbook', expected_sha256=CONFIG['pcm_expected_sha256'])
        PCM_TEMP_SCREEN, PCM_PRELIM_CANDIDATES, PCM_PRELIM_CANDIDATE_IDS = screen_workbook_roof(PCM_ALL, PCM_RUNNABLE, ROOF_SCREEN_LOWER_C, ROOF_SCREEN_UPPER_C)
        PCM_SCREEN_DERIVATION = {'method_origin': 'Inclusive melting-interval overlap with hotspot local-day air maxima and premonsoon surface-temperature percentile', 'lower_equation': 'round(mean_local_days(max_hourly_hotspot(T_air_C per Asia/Kolkata date)), 1)', 'local_daily_audit_file': 'pcm_screen_local_daily_audit.csv', 'upper_equation': 'round(percentile(valid Delhi pre-monsoon common-grid LST_C, 95, linear), 1)', 'pass_equation': 'Tm_high_C >= lower_C AND Tm_low_C <= upper_C; PCM_RUNNABLE only', 'screen_method': CUMULATIVE_SCREEN_METHOD, 'boundary_contact_passes': True, 'purpose': 'thermal relevance only; annual activation/saturation/utilisation/regeneration belong to Category 1', 'candidate_population_final': False, 'result_use': 'DEVELOPMENT', 'daily_air': ERA5_SCREEN_METADATA, 'lower_unrounded_C': AIR_ROOF_SCREEN_LOWER_UNROUNDED_C, 'upper_unrounded_C': LST_ROOF_SCREEN_UPPER_UNROUNDED_C, 'lower_C': ROOF_SCREEN_LOWER_C, 'upper_C': ROOF_SCREEN_UPPER_C, 'lst_sample_count': int(lst_screen_values.size), 'lst_percentile': screen_cfg['lst_percentile'], 'quantile_method': screen_cfg['quantile_method'], 'round_decimals': screen_cfg['round_decimals'], 'workbook_sha256': CONFIG['pcm_expected_sha256'], 'evaluated_runnable_rows': len(PCM_RUNNABLE), 'passing_rows': len(PCM_PRELIM_CANDIDATES), 'candidate_order': 'original verified workbook order; no ranking', 'material_source_quality': 'NOT CHECKED; unchanged Section-0 status', 'review_status': screen_cfg['review_status'], 'limitations': 'coarse hotspot hourly-sampled local daily air maxima plus Delhi overpass composite LST roof proxy; not measured roof bounds or a regeneration result'}
        smeta = artifact_metadata(PRELIM_CONFIG['weather']['collection'] + ' hotspot hourly local daily Tmax; real Landsat; pinned user workbook', 'degC; pass/fail; material properties retain Section-0 units', window=PRELIM_CONFIG['landsat']['windows']['PREMONSOON'], status='PRELIMINARY_SCREEN_REVIEW_REQUIRED')
        smeta.update(screen_method=CUMULATIVE_SCREEN_METHOD, pass_equation=PCM_SCREEN_DERIVATION['pass_equation'], workbook_sha256=CONFIG['pcm_expected_sha256'], candidate_population_final=False, result_use='DEVELOPMENT')
        smeta['candidate_table_schema'] = {'columns': list(PCM_PRELIM_CANDIDATES.columns), 'numeric_columns': list(PCM_PRELIM_CANDIDATES.select_dtypes(include='number').columns), 'boolean_columns': list(PCM_PRELIM_CANDIDATES.select_dtypes(include='bool').columns), 'null_rows': {name: np.flatnonzero(PCM_PRELIM_CANDIDATES[name].isna().to_numpy()).tolist() for name in PCM_PRELIM_CANDIDATES.columns}}
        local_daily_meta = {**smeta, **ERA5_SCREEN_METADATA, 'raw_source_sha256': sha256_file(PRELIM_RUN_DIR / 'era5_hourly_raw.csv'), 'annual_source_sha256': sha256_file(PRELIM_RUN_DIR / 'annual_weather.csv'), 'units': {'local_daily_Tmax_C': 'degC', 'contributing_hourly_observations': 'count'}}
        daily_path = save_table_artifact('pcm_screen_local_daily_audit.csv', PCM_SCREEN_LOCAL_DAILY_AUDIT, local_daily_meta)
        restored_daily = pd.read_csv(daily_path)
        pd.testing.assert_frame_equal(restored_daily, PCM_SCREEN_LOCAL_DAILY_AUDIT, check_exact=False, rtol=1e-12, atol=1e-12)
        save_json_artifact('pcm_screen_derivation.json', PCM_SCREEN_DERIVATION, smeta)
        save_table_artifact('pcm_temperature_screen.csv', PCM_TEMP_SCREEN, smeta)
        save_table_artifact('pcm_prelim_candidates.csv', PCM_PRELIM_CANDIDATES, smeta)
        save_json_artifact('pcm_prelim_candidate_ids.json', PCM_PRELIM_CANDIDATE_IDS, smeta)
        local_daily_checks = [{'name': 'hotspot_local_daily_Tmax_complete', 'result': 'PASS', 'detail': f'{len(PCM_SCREEN_LOCAL_DAILY_AUDIT)} complete Asia/Kolkata days; exactly 24 acquired hourly observations per day; no interpolation.'}]
        register_prelim_source('pcm_screen_local_daily', PRELIM_CONFIG['weather']['collection'] + ' derived hotspot local daily Tmax', 'degC; hourly observation counts', ERA5_SCREEN_METADATA['window'], local_daily_checks, native_crs=ERA5_SCREEN_METADATA['spatial_support']['native_projection']['crs'], native_resolution=PROVENANCE_REGISTRY.get('era5_hourly')['native_resolution'], file_info=require_file(daily_path, label='Hotspot local daily air audit'), notes=json.dumps(json_ready(local_daily_meta)))
        PRELIM_SOURCE_METADATA['pcm_screen_local_daily'] = local_daily_meta
        PRELIM_CHECKS.extend(local_daily_checks)
        display(PCM_SCREEN_LOCAL_DAILY_AUDIT)
        display(pd.DataFrame([{'lower_unrounded_C': AIR_ROOF_SCREEN_LOWER_UNROUNDED_C, 'upper_unrounded_C': LST_ROOF_SCREEN_UPPER_UNROUNDED_C, 'screen_lower_C': ROOF_SCREEN_LOWER_C, 'screen_upper_C': ROOF_SCREEN_UPPER_C, 'daily_air_days': ERA5_SCREEN_METADATA['observed_day_count'], 'LST_pixels': len(lst_screen_values), 'passing_runnable_rows': len(PCM_PRELIM_CANDIDATES)}]))
        display(PCM_TEMP_SCREEN[['pcm_id', 'Tm_low_C', 'Tm_high_C', 'screen_lower_C', 'screen_upper_C', 'screen_status', 'reason']])
        display(PCM_PRELIM_CANDIDATES[['pcm_id', 'material_name', 'Tm_low_C', 'Tm_high_C', 'source_quality_status']])
        if not PCM_PRELIM_CANDIDATE_IDS:
            print('No schema-ready workbook material passes this derived screen. Empty candidates are preserved honestly for review.')
        PRELIM_CHECKS.append({'name': 'inclusive_overlap_eight_boundary_cases', 'result': 'PASS', 'detail': 'All eight isolated overlap cases pass; zero-width boundary contact passes; no ranking.'})
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:

        def compute_index_relationships():
            rows = []
            for season in ['PREMONSOON', 'POSTMONSOON']:
                lst = LANDSAT_OUTPUTS[f'LST_{season}_C']
                for index in ['NDVI', 'NDBI']:
                    idx = LANDSAT_OUTPUTS[f'{index}_{season}']
                    valid = DELHI_MASK & np.isfinite(lst) & np.isfinite(idx)
                    n = int(valid.sum())
                    status = 'COMPUTED'
                    if n < 2 or np.std(lst[valid]) == 0 or np.std(idx[valid]) == 0:
                        coefficient = np.nan
                        status = 'UNDEFINED: fewer than 2 overlapping cells or zero variance'
                    else:
                        coefficient = float(np.corrcoef(lst[valid], idx[valid])[0, 1])
                        require_prelim(np.isfinite(coefficient) and abs(coefficient) <= 1 + 1e-12, 'Invalid Pearson correlation.')
                    rows.append({'season': season, 'index': index, 'sample_count': n, 'excluded_boundary_cells': int(DELHI_MASK.sum()) - n, 'method': 'Pearson r; ddof-independent normalized covariance', 'coefficient': coefficient, 'missing_handling': 'pairwise finite overlap within real Delhi mask; no imputation', 'status': status})
            return pd.DataFrame(rows)
        LST_INDEX_RELATIONSHIPS = compute_index_relationships()
        SEASONAL_SPATIAL_SUMMARY_ROWS = []
        for name in ['LST_PREMONSOON_C', 'LST_POSTMONSOON_C', 'NDVI_PREMONSOON', 'NDVI_POSTMONSOON', 'NDBI_PREMONSOON', 'NDBI_POSTMONSOON']:
            arr = LANDSAT_OUTPUTS[name]
            values = arr[DELHI_MASK & np.isfinite(arr)]
            SEASONAL_SPATIAL_SUMMARY_ROWS.append({'field': name, 'valid_cells': len(values), 'mean': float(np.mean(values)), 'std_population': float(np.std(values)), 'min': float(values.min()), 'p95': float(np.percentile(values, 95)), 'max': float(values.max()), 'units': 'degC' if name.startswith('LST') else 'dimensionless'})
        values = BUILT_FRACTION[DELHI_MASK & np.isfinite(BUILT_FRACTION)]
        SEASONAL_SPATIAL_SUMMARY_ROWS.append({'field': 'BUILT_FRACTION', 'valid_cells': len(values), 'mean': float(values.mean()), 'std_population': float(values.std()), 'min': float(values.min()), 'p95': float(np.percentile(values, 95)), 'max': float(values.max()), 'units': 'dimensionless'})
        SEASONAL_SPATIAL_SUMMARY = pd.DataFrame(SEASONAL_SPATIAL_SUMMARY_ROWS)
        rmeta = artifact_metadata('Real Landsat seasonal composites and GHSL', 'descriptive statistics and correlations', grid=COMMON_GRID)
        save_table_artifact('lst_index_relationships.csv', LST_INDEX_RELATIONSHIPS, rmeta)
        save_table_artifact('seasonal_spatial_summary.csv', SEASONAL_SPATIAL_SUMMARY, rmeta)
        display(SEASONAL_SPATIAL_SUMMARY)
        display(LST_INDEX_RELATIONSHIPS)
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        all_lst = np.concatenate([LST_PREMONSOON_C[np.isfinite(LST_PREMONSOON_C)], LST_POSTMONSOON_C[np.isfinite(LST_POSTMONSOON_C)]])
        lst_limits = [float(all_lst.min()), float(all_lst.max())]
        for field, units, cmap, limits, arrays in [('LST', 'Surface temperature (°C)', 'inferno', lst_limits, [LST_PREMONSOON_C, LST_POSTMONSOON_C]), ('NDVI', 'NDVI (dimensionless)', 'BrBG', [-1, 1], [NDVI_PREMONSOON, NDVI_POSTMONSOON]), ('NDBI', 'NDBI (dimensionless)', 'BrBG_r', [-1, 1], [NDBI_PREMONSOON, NDBI_POSTMONSOON])]:
            fig, axes = plt.subplots(1, 2, figsize=PRELIM_CONFIG['figures']['map_size'], layout='constrained')
            for ax, arr, season in zip(axes, arrays, ['PREMONSOON', 'POSTMONSOON']):
                start, end = PRELIM_CONFIG['landsat']['windows'][season]
                im = draw_prelim_map(ax, arr, f'{season.title()} {field}\n{start} to {end} (end excluded)', cmap, limits)
            fig.colorbar(im, ax=axes, shrink=0.75, label=units)
            save_prelim_figure(fig, field.lower() + '_seasonal_maps.png', rmeta)
        fig, axes = plt.subplots(1, 2, figsize=PRELIM_CONFIG['figures']['map_size'], layout='constrained')
        im = draw_prelim_map(axes[0], BUILT_FRACTION, 'GHSL built fraction · 2020 epoch', 'cividis', [0, 1])
        fig.colorbar(im, ax=axes[0], shrink=0.75, label='Built surface / source cell area (dimensionless)')
        im = draw_prelim_map(axes[1], LST_PREMONSOON_C, 'Selected persistent-hot built-up region', 'inferno', lst_limits)
        rows_grid, cols_grid = np.indices(COMMON_GRID.shape)
        xx = (COMMON_GRID.transform.c + (cols_grid + 0.5) * COMMON_GRID.transform.a) / 1000
        yy = (COMMON_GRID.transform.f + (rows_grid + 0.5) * COMMON_GRID.transform.e) / 1000
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        axes[1].contour(xx, yy, HOTSPOT_MASK.astype(float), levels=[0.5], colors=['#00c5e7'], linewidths=1.5)
        from matplotlib.lines import Line2D
        axes[1].legend(handles=[Line2D([0], [0], color='#00c5e7', label='Selected region boundary')], loc='lower right', fontsize=8)
        fig.colorbar(im, ax=axes[1], shrink=0.75, label='Pre-monsoon LST (°C)')
        fig.suptitle(f"Selected physical area {HOTSPOT_SUMMARY['physical_area_m2'] / 1000000.0:.3f} km²; built-weighted area {HOTSPOT_SUMMARY['built_weighted_area_m2'] / 1000000.0:.3f} km²")
        save_prelim_figure(fig, 'built_fraction_selected_hotspot.png', rmeta)
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        for season in ['PREMONSOON', 'POSTMONSOON']:
            fig, axes = plt.subplots(1, 2, figsize=PRELIM_CONFIG['figures']['relationship_size'], layout='constrained')
            for ax, index in zip(axes, ['NDVI', 'NDBI']):
                lst, idx = (LANDSAT_OUTPUTS[f'LST_{season}_C'], LANDSAT_OUTPUTS[f'{index}_{season}'])
                valid = DELHI_MASK & np.isfinite(lst) & np.isfinite(idx)
                row = LST_INDEX_RELATIONSHIPS.loc[(LST_INDEX_RELATIONSHIPS['season'] == season) & (LST_INDEX_RELATIONSHIPS['index'] == index)].iloc[0]
                if int(valid.sum()) < 2:
                    ax.text(0.5, 0.5, 'Insufficient valid pairs', ha='center', transform=ax.transAxes)
                else:
                    bins = ax.hexbin(idx[valid], lst[valid], gridsize=PRELIM_CONFIG['figures']['hexbin_gridsize'], mincnt=1, cmap='cividis')
                    fig.colorbar(bins, ax=ax, label='Valid cell pairs per bin')
                coefficient = 'undefined' if not np.isfinite(row.coefficient) else f'{row.coefficient:+.3f}'
                ax.set_title(f'{season.title()} LST vs {index}\nPearson r={coefficient}; n={int(row.sample_count):,}')
                ax.set_xlabel(f'{index} (dimensionless)')
                ax.set_ylabel('Surface temperature (°C)')
            save_prelim_figure(fig, 'lst_index_relationships_' + season.lower() + '.png', rmeta)
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        fig, ax = plt.subplots(figsize=(10, 5), layout='constrained')
        ready = PCM_TEMP_SCREEN.loc[PCM_TEMP_SCREEN.screen_status.isin(['PASS', 'FAIL'])]
        ax.axvspan(ROOF_SCREEN_LOWER_C, ROOF_SCREEN_UPPER_C, color='#f4c46c', alpha=0.35, label=f'Plausible thermal window: {ROOF_SCREEN_LOWER_C:.1f}–{ROOF_SCREEN_UPPER_C:.1f} °C')
        for y, row in enumerate(ready.itertuples()):
            color = '#206095' if row.screen_status == 'PASS' else '#888888'
            ax.plot([row.Tm_low_C, row.Tm_high_C], [y, y], color=color, linewidth=3, marker='o', markersize=4)
            ax.text(row.Tm_high_C, y, ' ' + row.screen_status, va='center', fontsize=9)
        ax.set_yticks(range(len(ready)), ready.pcm_id.tolist())
        ax.set_xlabel('Melting range and derived screen (°C)')
        ax.set_title('Inclusive overlap screen · thermal relevance only · workbook order\nLower: hotspot mean local daily Tmax; upper: Delhi pre-monsoon LST 95th percentile')
        ax.legend(loc='best', fontsize=9)
        ax.grid(axis='x', color='#dddddd', linewidth=0.6)
        save_prelim_figure(fig, 'pcm_preliminary_temperature_screen.png', smeta)
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        required_sources = ['study_boundary', 'landsat_premonsoon', 'landsat_postmonsoon', 'ghsl_built_surface', 'era5_hourly', 'pcm_screen_local_daily']
        for dataset_id in required_sources:
            require_registered_input(dataset_id, PROVENANCE_REGISTRY, CONFIG)
        handoff_raster_values = {name: globals()[name] for name in ['LST_PREMONSOON_C', 'LST_POSTMONSOON_C', 'NDVI_PREMONSOON', 'NDVI_POSTMONSOON', 'NDBI_PREMONSOON', 'NDBI_POSTMONSOON', 'BUILT_FRACTION']}
        handoff_raster_values.update({name: array for name, array in LANDSAT_OUTPUTS.items() if '_OBS_COUNT_' in name})
        handoff_raster_values['HOTSPOT_MASK'] = HOTSPOT_MASK.astype(np.float64)
        for name, record in PRELIM_RASTER_RECORDS.items():
            require_file(record['file_info']['file_path'], label=name, expected_sha256=record['file_info']['sha256'])
            validate_grid_alignment(raster_grid_from_file(record['file_info']['file_path']), COMMON_GRID, config=CONFIG)
            require_prelim(name in handoff_raster_values, f'Unexpected raster record: {name}')
            values = handoff_raster_values[name]
            require_prelim(numeric_array_hash(values) == record['numeric_sha256'], f'{name}: in-memory numeric content changed.')
            require_prelim(record['data_nature'] == 'REAL' and record['grid']['shape'] == list(COMMON_GRID.shape), f'{name}: missing real-source/grid metadata.')
        require_prelim(HOTSPOT_MASK.dtype == np.dtype(bool) and (not (HOTSPOT_MASK & ~DELHI_MASK).any()), 'Hotspot must be a boolean mask confined to actual Delhi cells.')
        reselected, repeated_summary, repeated_audit = select_persistent_hotspot(LST_PREMONSOON_C, LST_POSTMONSOON_C, BUILT_FRACTION, DELHI_MASK, COMMON_GRID, PRELIM_CONFIG['hotspot'])
        require_prelim(np.array_equal(reselected, HOTSPOT_MASK) and repeated_summary == HOTSPOT_SUMMARY, 'Hotspot mask or thresholds changed after deterministic selection.')
        repeated_location = derive_hotspot_weather_location(HOTSPOT_MASK, BUILT_FRACTION, COMMON_GRID)
        require_prelim(repeated_location == HOTSPOT_WEATHER_LOCATION == ERA5_SOURCE_METADATA['spatial_support']['hotspot_location'], 'Canonical weather extraction point no longer matches the selected hotspot.')
        require_prelim(ANNUAL_WEATHER.attrs['spatial_support'] == ERA5_SOURCE_METADATA['spatial_support'] and ANNUAL_WEATHER.attrs['spatial_support']['weather_resampled_to_common_grid'] is False, 'Canonical weather lacks native ERA5 support or incorrectly claims common-grid weather.')
        repeated_weather, repeated_ledger, repeated_coverage, repeated_qc_source, repeated_correction_ledger, repeated_qc_summary = prepare_annual_real_weather(ERA5_RAW)
        pd.testing.assert_frame_equal(ERA5_SHORTWAVE_QC_SOURCE, repeated_qc_source, check_exact=True)
        pd.testing.assert_frame_equal(ERA5_SHORTWAVE_CORRECTION_LEDGER, repeated_correction_ledger, check_exact=True)
        require_prelim(ERA5_SHORTWAVE_QC_SUMMARY == repeated_qc_summary and (ANNUAL_WEATHER.G_Wm2 >= 0.0).all(), 'Shortwave QC summary changed or canonical shortwave is negative.')
        require_prelim(sha256_file(raw_weather_path) == raw_weather_sha256_before_qc and np.array_equal(raw_shortwave_before_qc, ERA5_RAW[PRELIM_CONFIG['weather']['bands'][3]].to_numpy()), 'Raw acquisition identity changed after physical QC.')
        for filename, original in [('era5_shortwave_source_qc.csv', ERA5_SHORTWAVE_QC_SOURCE), ('era5_shortwave_correction_ledger.csv', ERA5_SHORTWAVE_CORRECTION_LEDGER)]:
            restored = pd.read_csv(PRELIM_RUN_DIR / filename, float_precision='round_trip')
            for column in [c for c in original.columns if c.endswith('_UTC')]:
                restored[column] = pd.to_datetime(restored[column], utc=True).dt.as_unit('ns')
            pd.testing.assert_frame_equal(restored, original, check_exact=True)
        require_prelim(json.loads((PRELIM_RUN_DIR / 'era5_shortwave_qc_summary.json').read_text()) == json_ready(ERA5_SHORTWAVE_QC_SUMMARY), 'Persisted QC summary differs from checked in-memory summary.')
        pd.testing.assert_frame_equal(ERA5_OVERLAP_LEDGER, repeated_ledger, check_exact=True)
        repeated_lower, repeated_daily, repeated_daily_meta = derive_local_day_roof_screen_air(ANNUAL_WEATHER, ERA5_RAW)
        pd.testing.assert_frame_equal(PCM_SCREEN_LOCAL_DAILY_AUDIT, repeated_daily, check_exact=True)
        require_prelim(repeated_daily_meta == ERA5_SCREEN_METADATA and repeated_lower == AIR_ROOF_SCREEN_LOWER_UNROUNDED_C, 'Local daily Tmax audit or full-precision lower statistic changed.')
        pd.testing.assert_frame_equal(ANNUAL_WEATHER, repeated_weather, check_exact=True)
        require_prelim(ANNUAL_WEATHER.attrs == repeated_weather.attrs, 'Weather units/time metadata changed.')
        reloaded_materials = audit_cleaned_pcm(load_pcm_properties(CONFIG['pcm_workbook_path'], sheet_name=CONFIG['pcm_sheet_name']))
        pd.testing.assert_frame_equal(PCM_ALL, reloaded_materials, check_exact=True)
        repeated_screen, repeated_candidates, repeated_ids = screen_workbook_roof(PCM_ALL, PCM_RUNNABLE, ROOF_SCREEN_LOWER_C, ROOF_SCREEN_UPPER_C)
        pd.testing.assert_frame_equal(PCM_TEMP_SCREEN, repeated_screen, check_exact=True)
        pd.testing.assert_frame_equal(PCM_PRELIM_CANDIDATES, repeated_candidates, check_exact=True)
        require_prelim(PCM_PRELIM_CANDIDATE_IDS == repeated_ids, 'PCM candidate IDs changed.')
        require_prelim(ROOF_SCREEN_LOWER_C == round(ERA5_SCREEN_METADATA['unrounded_lower_C'], screen_cfg['round_decimals']) and ROOF_SCREEN_UPPER_C == round(float(np.percentile(lst_screen_values, screen_cfg['lst_percentile'], method=screen_cfg['quantile_method'])), screen_cfg['round_decimals']), 'PCM screen bounds cannot be reconstructed from the explicit real-data statistics.')
        for relative, record in PRELIM_ARTIFACTS.items():
            require_file(PRELIM_RUN_DIR / relative, label=relative, expected_sha256=record['sha256'])
        for filename, value in [('delhi_boundary.json', DELHI_BOUNDARY), ('common_grid.json', grid_metadata(COMMON_GRID)), ('hotspot_summary.json', HOTSPOT_SUMMARY), ('hotspot_weather_location.json', HOTSPOT_WEATHER_LOCATION), ('pcm_screen_derivation.json', PCM_SCREEN_DERIVATION)]:
            require_prelim(json.loads((PRELIM_RUN_DIR / filename).read_text()) == json_ready(value), f'{filename}: in-memory metadata differ from persisted source/derivation.')
        persisted_daily = pd.read_csv(PRELIM_RUN_DIR / 'pcm_screen_local_daily_audit.csv')
        pd.testing.assert_frame_equal(persisted_daily, PCM_SCREEN_LOCAL_DAILY_AUDIT, check_exact=False, rtol=1e-12, atol=1e-12)
        for filename, support in [('annual_weather.csv.metadata.json', ANNUAL_WEATHER.attrs['spatial_support']), ('era5_hourly_raw.csv.metadata.json', ERA5_SOURCE_METADATA['spatial_support']), ('pcm_screen_local_daily_audit.csv.metadata.json', ERA5_SCREEN_METADATA['spatial_support'])]:
            require_prelim(json.loads((PRELIM_RUN_DIR / filename).read_text())['spatial_support'] == json_ready(support), f'{filename}: persisted hotspot native-cell support differs.')
        PRELIM_CHECKS.extend([{'name': 'shortwave_physical_qc_reproducibility_and_raw_preservation', 'result': 'PASS', 'detail': 'Source/target flags, correction magnitudes, annual energy and both ledgers reproduced; raw acquisition file and values unchanged; all canonical G>=0.'}, {'name': 'hotspot_weather_and_local_daily_reproducibility', 'result': 'PASS', 'detail': 'Native hotspot location, source-linked radiation ledger, complete local-day audit and unrounded lower bound reproduced.'}, {'name': 'whole_grid_and_numeric_identity', 'result': 'PASS', 'detail': 'All raster files, affines, CRS, dimensions and numeric hashes match.'}, {'name': 'hotspot_reproducibility', 'result': 'PASS', 'detail': 'Same rule reproduces the mask, areas, counts and thresholds.'}, {'name': 'PCM_screen_and_workbook_integrity', 'result': 'PASS', 'detail': 'Pinned workbook unchanged; all rows audited; derived bounds/candidate IDs reproduced.'}, {'name': 'provenance_completeness', 'result': 'PASS', 'detail': 'All required environmental inputs have authenticated source records and passing checks.'}])
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_PROVENANCE = PROVENANCE_REGISTRY.to_frame().copy(deep=True)
        save_real_raster('DELHI_MASK', DELHI_MASK.astype(float), artifact_metadata(PRELIM_CONFIG['boundary']['collection'], 'binary cell-center boundary membership', window=['2015-01-01', '2016-01-01'], grid=COMMON_GRID))
        save_json_artifact('source_metadata.json', PRELIM_SOURCE_METADATA, artifact_metadata('Authenticated EE source responses + processing metadata', 'source metadata'))
        save_json_artifact('provenance_registry.json', PRELIM_PROVENANCE.to_dict('records'), artifact_metadata('Section-0 registry updated with actual Preliminary retrievals', 'registry snapshot', status='REAL_ENVIRONMENT_VALIDATED; MATERIAL_SOURCE_QUALITY_NOT_CHECKED'))
        save_table_artifact('preliminary_validation.csv', pd.DataFrame(PRELIM_CHECKS), artifact_metadata('Executed source/grid/time/selection/screen checks', 'validation results'))
        npz_path = PRELIM_RUN_DIR / 'preliminary_numeric_layers.npz'
        with npz_path.open('xb') as handle:
            np.savez_compressed(handle, **LANDSAT_OUTPUTS, BUILT_FRACTION=BUILT_FRACTION, HOTSPOT_MASK=HOTSPOT_MASK, DELHI_MASK=DELHI_MASK)
        track_artifact(npz_path, artifact_metadata('Real Landsat/GHSL/GAUL and deterministic hotspot selection', 'per-layer units in GeoTIFF sidecars', grid=COMMON_GRID))
        print('Preliminary validation passed; no PCM solver has run.')
        display(pd.DataFrame(PRELIM_CHECKS))
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        PRELIM_HANDOFF = None
        PRELIM_READY = False
        HANDOFF_OBJECTS = {'COMMON_GRID': COMMON_GRID, 'DELHI_BOUNDARY': DELHI_BOUNDARY, 'LST_PREMONSOON_C': LST_PREMONSOON_C, 'LST_POSTMONSOON_C': LST_POSTMONSOON_C, 'NDVI_PREMONSOON': NDVI_PREMONSOON, 'NDVI_POSTMONSOON': NDVI_POSTMONSOON, 'NDBI_PREMONSOON': NDBI_PREMONSOON, 'NDBI_POSTMONSOON': NDBI_POSTMONSOON, 'BUILT_FRACTION': BUILT_FRACTION, 'HOTSPOT_MASK': HOTSPOT_MASK, 'HOTSPOT_SUMMARY': HOTSPOT_SUMMARY, 'ERA5_SHORTWAVE_QC_SOURCE': ERA5_SHORTWAVE_QC_SOURCE, 'ERA5_SHORTWAVE_CORRECTION_LEDGER': ERA5_SHORTWAVE_CORRECTION_LEDGER, 'ERA5_SHORTWAVE_QC_SUMMARY': deepcopy(ERA5_SHORTWAVE_QC_SUMMARY), 'ANNUAL_WEATHER': ANNUAL_WEATHER, 'HOTSPOT_WEATHER_LOCATION': deepcopy(HOTSPOT_WEATHER_LOCATION), 'PCM_SCREEN_LOCAL_DAILY_AUDIT': PCM_SCREEN_LOCAL_DAILY_AUDIT, 'PCM_TEMP_SCREEN': PCM_TEMP_SCREEN, 'PCM_PRELIM_CANDIDATES': PCM_PRELIM_CANDIDATES, 'PCM_PRELIM_CANDIDATE_IDS': PCM_PRELIM_CANDIDATE_IDS, 'PRELIM_PROVENANCE': PRELIM_PROVENANCE, 'PCM_SCREEN_DERIVATION': PCM_SCREEN_DERIVATION, 'PRELIM_CONFIG': deepcopy(PRELIM_CONFIG), 'PRELIM_ARTIFACT_DIR': str(PRELIM_RUN_DIR), 'PRELIM_RUN_ID': PRELIM_RUN_ID, 'CUMULATIVE_RUN_ID': CUMULATIVE_RUN_ID, 'CUMULATIVE_SOURCE_IDENTITY': deepcopy(CUMULATIVE_SOURCE_IDENTITY), 'PCM_SCREEN_METHOD': CUMULATIVE_SCREEN_METHOD, 'PRELIM_REVIEW_STATUS': PRELIM_CONFIG['pcm_screen']['review_status']}
        manifest = {'cumulative_run_id': CUMULATIVE_RUN_ID, 'cumulative_source_identity': deepcopy(CUMULATIVE_SOURCE_IDENTITY), 'screen_method': CUMULATIVE_SCREEN_METHOD, 'screen_derivation': deepcopy(PCM_SCREEN_DERIVATION), 'candidate_population_final': False, 'result_use': 'DEVELOPMENT', 'run_id': PRELIM_RUN_ID, 'status': 'PRELIMINARY_COMPLETE_REVIEW_REQUIRED', 'scientific_source_policy': 'REAL_ONLY; NO_SUBSTITUTION; EXPLICIT_PROVIDER_SHORTWAVE_PHYSICAL_QC', 'source_quality_limit': 'PCM workbook identity/schema checked; physical citations remain NOT CHECKED', 'global_config': json_ready(CONFIG), 'preliminary_config': json_ready(PRELIM_CONFIG), 'grid': grid_metadata(COMMON_GRID), 'handoff_objects': list(HANDOFF_OBJECTS), 'annual_weather_contract': {'canonical_object': 'ANNUAL_WEATHER', 'intended_use': 'Category 1 selected hotspot', 'file': 'annual_weather.csv', 'raw_source_file': 'era5_hourly_raw.csv', 'spatial_support': deepcopy(ANNUAL_WEATHER.attrs['spatial_support']), 'raw_source_sha256': raw_weather_sha256_before_qc, 'shortwave_physical_qc': deepcopy(ERA5_SHORTWAVE_QC_SUMMARY), 'radiation_overlap_ledger_file': 'era5_radiation_overlap_ledger.csv', 'time_semantics': deepcopy(ANNUAL_WEATHER.attrs)}, 'pcm_screen_lower_contract': {'source': 'ANNUAL_WEATHER.T_air_C', 'method': PRELIM_CONFIG['pcm_screen']['lower_statistic'], 'timezone': PRELIM_CONFIG['pcm_screen']['air_timezone'], 'audit_file': 'pcm_screen_local_daily_audit.csv', 'derivation_file': 'pcm_screen_derivation.json', 'coverage': deepcopy(ERA5_SCREEN_METADATA)}, 'artifacts': deepcopy(PRELIM_ARTIFACTS), 'validation': PRELIM_CHECKS, 'source_registry_file': 'provenance_registry.json', 'candidate_review_required': True}
        manifest_path = PRELIM_RUN_DIR / 'preliminary_manifest.json'
        with manifest_path.open('x', encoding='utf-8') as handle:
            json.dump(json_ready(manifest), handle, indent=2, allow_nan=False)
        for relative, record in manifest['artifacts'].items():
            require_file(PRELIM_RUN_DIR / relative, label=relative, expected_sha256=record['sha256'])
        PRELIM_HANDOFF = HANDOFF_OBJECTS
        PRELIM_READY = True
        display(pd.DataFrame([{'object': key, 'type': type(value).__name__} for key, value in PRELIM_HANDOFF.items()]))
        print('PRELIM_HANDOFF ready; overlap lineage published. Category 1 follows the verified in-memory handoff.')
        print('Manifest:', manifest_path)

# Executed stage lineage: original cells 134, 136, 138
def run_rooftop():
    """Explicit original annual rooftop and matched-control stages. Computationally expensive; never called during cleanup."""
    global C1_COMPLETED_RUN, C1_FIGURES, C1_FIGURE_END_IST, C1_FIGURE_PCM_ID, C1_FIGURE_START_IST, C1_SCIENTIFIC_STATUS, C1_VALIDATED_INPUTS, C1_WORKBOOK_INSPECTION, CATEGORY1_HANDOFF, CATEGORY1_READY, figure
    C1_WORKBOOK_INSPECTION = c1_inspect_workbook(CATEGORY1_IO)
    display(pd.DataFrame([C1_WORKBOOK_INSPECTION['receipt']]))
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        CATEGORY1_HANDOFF = None
        CATEGORY1_READY = False
        C1_SCIENTIFIC_STATUS = 'INPUT_VALIDATION'
        try:
            C1_VALIDATED_INPUTS = c1_load_inputs(globals(), CATEGORY1_IO)
            C1_SCIENTIFIC_STATUS = 'RUNNING_REAL_ANNUAL'
            C1_COMPLETED_RUN = c1_run_category1(C1_VALIDATED_INPUTS)
            if CATEGORY1_IO['persist_real_outputs']:
                c1_persist_category1(C1_COMPLETED_RUN, C1_VALIDATED_INPUTS, CATEGORY1_IO['output_dir'])
            CATEGORY1_HANDOFF = C1_COMPLETED_RUN
            CATEGORY1_READY = True
            C1_SCIENTIFIC_STATUS = 'COMPLETE_REAL_ANNUAL'
        except Exception as error:
            CATEGORY1_HANDOFF = None
            CATEGORY1_READY = False
            C1_SCIENTIFIC_STATUS = 'STOPPED: ' + str(error)
            print('CATEGORY 1 SCIENTIFIC RUN STOPPED:', error)
            raise
    if not CACHE_FIRST and RUN_FRESH_UPSTREAM:
        if not CATEGORY1_READY:
            raise Category1InputError('No validated real annual handoff exists.')
        display(CATEGORY1_HANDOFF['annual_summary'])
        C1_FIGURE_PCM_ID = CATEGORY1_HANDOFF['candidate_ids'][0]
        C1_FIGURE_START_IST = '2025-05-15'
        C1_FIGURE_END_IST = '2025-05-18'
        C1_FIGURES = {'annual_liquid_fraction': c1_plot_liquid_fraction(CATEGORY1_HANDOFF['by_pcm'], CATEGORY1_CONFIG), 'next_morning_reset': c1_plot_morning_reset(CATEGORY1_HANDOFF['by_pcm'], CATEGORY1_CONFIG), 'physical_mechanism': c1_plot_mechanism(CATEGORY1_HANDOFF['by_pcm'][C1_FIGURE_PCM_ID], C1_FIGURE_START_IST, C1_FIGURE_END_IST), 'candidate_comparison': c1_plot_comparison(CATEGORY1_HANDOFF['annual_summary'])}
        for figure in C1_FIGURES.values():
            display(figure)
            plt.close(figure)
        print("Category-3 input: CATEGORY1_HANDOFF['by_pcm'][pcm_id]['delta_H']; roof→air W/m²")
        print('Artifacts:', CATEGORY1_HANDOFF['artifact_directory'])


def run_atmospheric_baselines():
    """Explicit three-material annual baseline plus original numerical acceptance.

    Uses completed in-memory Preliminary and Category 1 results. This expensive
    entry point is never called by import or --verify-saved.
    """
    global C3_INPUTS, C3_ROOFS, C3_BASELINE_RESULTS
    C3_INPUTS = c3_validate_handoff(globals())
    C3_ROOFS = c3_acquire_roofs(C3_INPUTS, globals(), C3_EXECUTION)
    settings = {**C3_EXECUTION, 'mode': 'REAL_BASELINE_ANNUAL'}
    C3_BASELINE_RESULTS = c3_run_baselines(C3_INPUTS, C3_ROOFS, globals(), settings, 6000)
    return C3_BASELINE_RESULTS



CUMULATIVE_SOURCE_IDENTITY = {
    'source_name': 'src/research_model.py',
    'source_content_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    'executable_source_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    'fingerprint_policy': 'Current source fingerprint; saved run identities are recorded separately.'
}

def verify_saved():
    """Inspect existing numerical outputs without starting a solver."""
    from postprocessing.saved_metrics import verify_saved_results
    return verify_saved_results(REPOSITORY_ROOT)

def main():
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--verify-saved', action='store_true', help='Check saved inputs, forcing and protected maxima; no simulation.')
    args = parser.parse_args()
    if args.verify_saved:
        print(json.dumps(verify_saved(), indent=2))
    else:
        parser.print_help()

if __name__ == '__main__':
    main()

